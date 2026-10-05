"""Decompression-bomb guard for zip-container document uploads.

``.docx``/``.xlsx``/``.pptx``/``.odt``/``.epub`` files are zip
containers of XML parts, while the upload route's caps bound only the
*compressed* request body. A modest container can therefore declare —
or, with lying headers, secretly contain — gigabytes of XML and OOM
the indexing worker: the same class as the OpenAlex partition caps and
the PDF extraction ceilings.

What the guard checks, before any parser sees the bytes:

0. **Anchoring**: the upload starts with a zip local file header (or,
   for an empty archive, the end-of-central-directory record) and the
   first entry's local header sits at offset 0. zipfile locates the
   central directory from the *end* of the file and tolerates data in
   front of the archive, so without this a legacy OLE ``.doc``/``.ppt``
   with a small zip appended would pass as ``.docx``/``.pptx`` — and
   unstructured, which sniffs the leading bytes, would then route it to
   its own unbounded soffice conversion.
1. **Declarations** (central directory only, no decompression): a
   declared central-directory byte span of at most ``max_entries``
   kibibytes, read from the end-of-central-directory (and ZIP64)
   records before ``zipfile`` walks them — zipfile materialises every
   record before the entry ceiling can count them, and the EOCD's
   record count wraps at 65,535 and is ignored by zipfile, so only
   the declared span bounds the walk up front; at
   most ``MAX_ENTRIES`` entries; only STORED or DEFLATE compression
   (CPython's zipfile decompresses BZIP2/LZMA/ZSTD chunks with no
   output limit); no encrypted entries; a STORED entry's compressed
   size equals its declared size and a DEFLATE entry's is within a
   small slack of it (a full read copies the whole compressed range);
   no two entries share a local header; every member name is
   canonical (no empty, ``.`` or ``..`` segments, so no ``//``, ``/./``
   or leading ``/``; no backslash) and unique, because readers
   normalise such names onto one path (zipfile itself truncates a name
   at NUL, so NUL-truncated names that collide are caught as
   duplicates); no entry declares more than
   the largest per-entry ceiling; and at most ``MAX_CONTAINER_BYTES``
   declared across all entries.
2. **Honesty of the declarations**: every entry's local header must
   agree with its central-directory record (method, flags, CRC, both
   sizes, name; with a data descriptor, the descriptor must repeat the
   central CRC and sizes), because pandoc's zip reader follows the
   local header where zipfile follows the central directory. Every
   DEFLATE entry is then inflated once, in bounded steps (a few MiB of
   memory, at most the declared total of work), and refused unless the
   stream ends exactly at the end of its compressed range having
   produced exactly the declared size with the declared CRC; a STORED
   entry's size is fixed by its equality check in step 1 and its CRC is
   checked too. zipfile only
   truncates to the declared size *after* inflating a chunk, and a
   full ``read()`` inflates an entry's whole compressed range in one
   step, so without this a DEFLATE entry declaring 100 bytes could
   still materialize gigabytes inside a parser.
3. **Per-kind ceilings** on the verified sizes: a parsed part (XML or
   any non-media entry) at most ``MAX_ENTRY_BYTES`` and parsed parts at
   most ``MAX_TOTAL_BYTES`` in total; a binary media entry at most
   ``MAX_MEDIA_ENTRY_BYTES`` — media ceilings apply to
   ``.docx``/``.pptx``/``.xlsx`` only (``MEDIA_CEILING_EXTENSIONS``).
   Media classification is described below.
4. **Sheet fan-out** (``.xlsx`` only): openpyxl converts a sheet part
   once per ``<sheet>`` that resolves to it, so a workbook whose sheets
   resolve to the same part more than once is refused. The sheet list
   comes from openpyxl's own ``ExcelReader`` / ``WorkbookParser`` (the
   steps and options ``load_workbook`` uses under pandas), not from a
   re-implementation, so the guard cannot disagree with the loader
   about tags, attributes, the workbook's location, relationship
   resolution (including ``TargetMode="External"`` targets, which
   openpyxl keeps unresolved) or XML decoding. Sheets must also have
   distinct relationships parts as computed by openpyxl's
   ``get_rels_path``, which openpyxl parses once per sheet. Only
   ``[Content_Types].xml``, the workbook part and its relationships are
   parsed, each first held to ``MAX_XLSX_SKELETON_PART_BYTES`` so this
   extra full-tree parse stays cheap; any openpyxl failure there is
   refused as malformed. A comment or processing instruction below the
   root's children of one of these parts or of ``xl/styles.xml`` is
   refused (``_refuse_xlsx_nested_node``): openpyxl builds every child
   node of a list such as ``fonts`` or ``bookViews`` into an entry, while
   the guard counts elements. openpyxl also builds an object from every
   entry of every list of the styles and workbook parts and from every
   child it knows of every other element (keeping the last of a
   repeated one), so each list of the two parts is capped
   (``_xlsx_list_cap``; the workbook's defined names as
   ``_check_xlsx_defined_names`` holds them), as are the elements under
   each root (``MAX_XLSX_TOP_LEVEL_ELEMENTS``), the elements under
   ``dxfs``, the gradient stops in all and the elements of one font,
   fill, border or cell-format entry (``MAX_XLSX_STYLE_ENTRY_ELEMENTS``),
   and an element that is not a list may not repeat a child's local
   name (``_refuse_xlsx_repeated_child``). openpyxl parses the print
   titles and print areas among the defined names with regular
   expressions in time quadratic in their length, on every load, so
   each defined name is held to ``MAX_XLSX_DEFINED_NAME_CHARS`` and
   the print names together to ``MAX_XLSX_PRINT_NAME_WORK``. A workbook
   listing more than ``MAX_XLSX_SHEETS`` sheets is refused: openpyxl
   and pandas do work per sheet that grows with the members and with
   the sheets; so is one with a sheet name over
   ``MAX_XLSX_SHEET_NAME_CHARS`` characters or two sheets of one name
   (lowercased). Before openpyxl parses any part, and before any part
   the guard parses, each part openpyxl's read-only load parses (found
   as it finds them: ``[Content_Types].xml``, the workbook part and its
   relationships, then the shared strings the manifest names, the core
   and custom document properties, the styles part, and every sheet's
   resolved part and relationships part; ``_xlsx_loaded_parts``) is
   streamed once with the parser that load uses for it, lxml or expat
   (``[Content_Types].xml`` with both, for msoffcrypto's minidom), and
   no other member is (``_check_xlsx_namespaces``): an encoding other
   than UTF-8 (``_require_utf8_xml``; the parsers detect and decode
   others differently, and the guard counts bytes), a namespace URI over ``MAX_XLSX_NAMESPACE_URI_CHARS`` characters,
   which lxml and expat copy into the name of every element under it,
   a DTD, nesting deeper than ``MAX_XLSX_XML_DEPTH``, a start tag of
   more than ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` attributes or more than
   ``MAX_XLSX_PART_NAMES`` distinct names (state both parsers keep per
   open element, attribute and name) is refused, as is an lxml part
   that lxml fails to parse (openpyxl's load fails on it too); an expat
   part is read up to its first syntax error, where the loader's parse
   ends. The
   styles part ``xl/styles.xml``, which openpyxl loads on every read,
   is streamed once (``_check_xlsx_styles``): a number format code over
   ``MAX_XLSX_NUMBER_FORMAT_CHARS`` characters, a cell-format list of
   more than ``MAX_XLSX_CELL_FORMATS`` entries, a gradient fill of more
   than ``MAX_XLSX_GRADIENT_STOPS`` stops, more than
   ``MAX_XLSX_STYLE_LIST_ENTRIES`` fonts, fills or borders, a numeric
   field outside 32 bits, more than ``MAX_XLSX_STYLE_HASH_CHAIN``
   distinct style objects of one hash in an indexed list openpyxl
   builds (each entry is rebuilt with openpyxl's classes to compute
   it), or a format code, cell-format
   id or named-style field given as a child element (which openpyxl
   reads in place of the attribute) is refused, as is a stylesheet
   whose date tests over the cell formats' codes exceed
   ``MAX_XLSX_NUMBER_FORMAT_WORK`` or whose named styles, each priced
   at a fixed cost plus the custom formats and the largest font, fill
   and border it can bind, exceed ``MAX_XLSX_NAMED_STYLE_WORK``
   (openpyxl repeats both per cell format or named style).
5. **Slide and section fan-out** (``.pptx``/``.docx``): a presentation
   listing one slide part under several ``<p:sldId>`` entries is
   refused, as is one whose slide placeholders would cost python-pptx's
   uncached layout and master lookups more than
   ``MAX_PPTX_PLACEHOLDER_WORK`` (each slide's placeholders times the
   shapes, nodes and relationships one inherited position read walks;
   every slide is streamed once for this, as is every slide layout and
   slide master a slide with placeholders inherits from). Every
   relationships part of a presentation or document is streamed and,
   like every other part the guard streams or parses from a
   presentation or document, first held to UTF-8 and to start tags of
   at most ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` attributes, counted in its
   bytes before lxml reads them (``_CountedXmlStream``; lxml's
   ``iterparse`` builds a start tag's element in full, with no limit,
   before any per-element ceiling sees it), then held to ``MAX_PART_RELATIONSHIPS`` entries (comments and processing
   instructions included; an entry with child elements, or any element
   with more than ``MAX_RELS_ELEMENT_ATTRIBUTES`` attributes, is
   refused) and all of them together to
   ``MAX_PACKAGE_RELATIONSHIPS``, and a document's internal
   relationships times its members to ``MAX_DOCX_RELATIONSHIP_WALK``
   (python-docx scans a list of the parts it has visited once per
   relationship as it opens the package). The parts the guard reaches through
   relationships (the main part; a presentation's slides and the
   layouts and masters it prices; a document's styles, settings,
   header and footer parts) must carry the content type, resolved from
   ``[Content_Types].xml`` as the libraries resolve it, that the
   library builds such a part from: python-pptx and python-docx pick a
   part's class by content type, and a slide or document part standing
   in for a layout, styles or settings part makes the library follow
   that part's own relationship to one the guard never priced (see
   ``_require_content_type``). A document is refused when it has more than
   ``MAX_DOCX_SECTIONS`` ``w:sectPr`` elements or more than
   ``MAX_DOCX_PAGE_BREAKS`` ``w:lastRenderedPageBreak`` elements, when its
   settings or styles part is over
   ``MAX_OPC_SKELETON_PART_BYTES`` or has more than
   ``MAX_DOCX_SETTINGS_CHILDREN`` or ``MAX_DOCX_STYLES`` top-level
   children, when a table cell spans, or a row skips, more than
   ``MAX_DOCX_GRID_SPAN`` grid columns, when a table row has more than
   ``MAX_DOCX_ROW_CELLS`` cells or spans more than
   ``MAX_DOCX_ROW_GRID_COLUMNS`` columns, when a table row has more
   than ``MAX_DOCX_ROW_OTHER_CHILDREN`` other children or a table more
   than ``MAX_DOCX_TABLE_ROW_GAP`` children between (or before) two of
   its rows, when an element of
   ``document.xml`` carries more than ``MAX_DOCX_ELEMENT_ATTRIBUTES``
   attributes or sees more than ``MAX_DOCX_NAMESPACES_IN_SCOPE``
   namespace declarations in scope, when a body-level paragraph holds
   more than ``MAX_DOCX_PARAGRAPH_NAMESPACES`` declarations, or when a
   declaration's URI or prefix is longer than
   ``MAX_DOCX_NAMESPACE_URI_CHARS`` or ``MAX_DOCX_NAMESPACE_PREFIX_CHARS``
   (abnormal input no writer produces, refused outright rather than
   priced), or when python-docx's and
   unstructured's repeated work, estimated from one streaming pass over
   ``document.xml``, exceeds a ceiling: the section lookup over the
   body (``MAX_DOCX_SECTION_BLOCK_WORK``, ``MAX_DOCX_SECTION_WORK``),
   the header/footer parts its sections reference, counted per
   reference with the node-set unions each read evaluates
   (``MAX_DOCX_HEADER_FOOTER_BYTES``), the style lookups per
   paragraph (``MAX_DOCX_STYLE_WORK``), the splitting of paragraphs at
   rendered page breaks (``MAX_DOCX_PAGE_BREAK_WORK``), the text
   extraction from runs and table cells, once per spanned column, with
   each cell's paragraphs and tables, the node-set unions python-docx
   evaluates over a cell's, a paragraph's and a run's children (and,
   in a document with no sections, the body's; and unstructured's
   page-break union over the whole body and its paragraph-text union,
   whose third branch reaches runs in inline drawings) and
   python-docx's steps up vertical merges (``MAX_DOCX_RUN_WORK``), and
   the bytes copied per page break,
   spanned column and continued merged cell
   (``MAX_DOCX_COPIED_BYTES``). These are the fan-outs that turn one
   bounded part or element into many times its work or memory; each
   ceiling holds its fan-out to a measured order of magnitude (about
   45 s, or a few hundred MB, at the slowest layout found), and the
   seven share one budget: the work's shares of their ceilings may sum
   to at most 1 (see ``_check_pptx_slide_fan_out`` and
   ``_check_docx_section_fan_out``). They do not bound a document's
   total parse time or memory (see "Known gaps").
6. **Package type** (``.docx``/``.pptx``/``.xlsx``, and the
   ``mimetype`` of ``.odt``/``.epub``): langchain's Word and PowerPoint
   loaders ask unstructured's ``detect_filetype`` for the type (when
   python-magic imports) and send a file it calls DOC or PPT to soffice
   with no timeout, whatever the extension. An OOXML package is
   therefore refused when it has a root ``mimetype`` member, when its
   member names would not make the detector name its extension, when
   an ``.xlsx`` has no ``xl/workbook.xml`` (pandas would pick another
   engine), or when a ``.docx``/``.pptx`` has no main document part as
   python-docx/python-pptx resolve it; an ``.odt``/``.epub`` whose
   ``mimetype`` member names another type is refused too (see
   ``_check_package_type``).

Containers the stdlib cannot list or read, and spreadsheet packages
openpyxl cannot resolve, are refused as ``DecompressionBombError``
(fail closed) for the failure types caught in
``validate_zip_container`` and ``_check_xlsx_sheet_fan_out``.

What this bounds, and what it does not:

- **In-process Python parsers** (python-docx, python-pptx, openpyxl,
  unstructured for ``.docx``/``.xlsx``/``.pptx``): each *read* of an
  entry now yields at most its (verified) declared size. It does not
  bound parser memory or parser time. For an XML parse tree alone the
  cost is a multiple of the XML size (around 10x measured for
  text-heavy ``.docx``, more for many small elements), which is why the
  parsed-part ceilings sit well under a worker's memory; but what the
  loaders build *from* that tree can be far larger than the XML and is
  not bounded here (see "Known gaps" below). python-docx/-pptx parse
  with lxml ``resolve_entities=False``, openpyxl uses its safe parser,
  and libxml2/expat enforce their own amplification limits, so entities
  are not expanded; but an unexpanded entity reference stays in
  python-docx's tree as a node its per-section XPath walks, and the
  streaming count of that work (step 5) sees no event for it. Every XML
  part the guard parses or streams (OPC relationships,
  ``presentation.xml``, the slides and the slide layouts and masters
  their placeholders inherit from, ``document.xml``, the ``.docx`` settings part
  and the header/footer parts its sections reference, and the ``.xlsx`` parts
  openpyxl's reader parses for the sheet check and ``xl/styles.xml``) is
  therefore refused if it carries a
  document type declaration, which is the only way to declare an
  entity; OPC (ECMA-376 Part 2) forbids DTDs in the parts it defines
  (content types, relationships, core properties), and Office writes
  none in the XML parts it generates (an SVG image it stores may carry
  one, but no loader parses images). The declaration is refused in the
  part's bytes, before any parser the guard drives receives it
  (``_XmlProlog``): libxml2 reads a whole internal subset before the
  root element's first event, and the attribute defaults it declares
  add attributes to start tags whose bytes show none. Parts the guard
  does not read are not checked. How *often* a parser reads an entry is bounded only
  where a fan-out is known and closed: shared sheet parts and
  per-format stylesheet work in ``.xlsx`` (step 4), repeated slides and per-section, per-paragraph, per-page-
  break and per-spanned-column work and header/footer repeats in
  ``.pptx``/``.docx`` (step 5), and openpyxl's per-reference drawing
  reads (below).
- **Binary media** (images, audio/video, fonts, embedded packages):
  python-pptx and python-docx load each package part once, however
  many times it is referenced, and keep it as an opaque blob — about
  1x its size (measured for one and for eight references to the same
  image). openpyxl would instead re-read a chartsheet drawing's image
  once per reference and keep every copy; the ``.xlsx`` path disables
  openpyxl's drawing reader (see ``openpyxl_hardening``; in the
  read-only mode pandas uses, chartsheet drawings are the only
  drawings openpyxl reads), so ``.xlsx`` media is not read at all.
  Only an entry of an OOXML container whose first bytes carry a known
  binary signature earns the media ceiling — never a name: part names
  are attacker-chosen and packages dispatch parts by content type.
  Each signature is anchored so that no XML parser the in-process
  loaders use (lxml/expat) accepts it as a document start. EPUB and
  ODT entries never earn it: pandoc reads an EPUB chapter by its
  declared media type, not its bytes. All but two
  (ISO BMFF and EMF) are fixed prefixes that start with neither
  ``<``, whitespace nor a complete BOM (JPEG's ``FF D8`` shares only
  its first byte with the ``FF FE`` UTF-16/32 BOM). ISO BMFF requires
  two leading zero bytes, a big-endian box size of at least 8 and
  ``ftyp`` (unparseable as XML in any encoding: ``ftyp`` read as UCS-4
  is above U+10FFFF); EMF requires record type 1 (``01 00 00 00``) and
  `` EMF`` at offset 40. As a second line, any other entry whose head
  decodes to markup (``<`` after an optional BOM and whitespace, in
  UTF-8/16/32 with or without a BOM) stays a parsed part whatever else
  it matches. A
  markup-looking name also always keeps an entry under the parsed-part
  ceilings.
- **pandoc** (``.odt``/``.epub``, see ``utilities/pandoc_conversion``)
  has its own zip reader, which locates entries through the central
  directory but takes each entry's method and compressed size from its
  local header; step 2's header agreement makes it inflate the streams
  verified here. Reference fan-out inside pandoc's readers is not
  checked; the pandoc child is bounded by its own heap cap and timeout.

Known gaps: pre-existing parser behaviour this guard does **not**
bound. It bounds container and entry sizes, compression, the honesty
of the declarations, and only the specific fan-outs listed above.

- **.xlsx grid padding.** Memory follows the *cell references*, not the
  XML size. pandas' openpyxl reader (``_OpenpyxlReader.get_sheet_data``,
  which unstructured's ``partition_xlsx`` reaches through
  ``pd.read_excel``) pads every row to the widest row
  (``data_row + (max_width - len(data_row)) * empty_cell``), so a
  few-KB sheet with cells ``XFD1`` and ``A300`` materialises
  300 x 16384 cells (about +121 MiB RSS measured), and ``XFD1`` with
  ``A1048576`` would need hundreds of GiB. The per-part ceilings do
  not limit this. Likewise, a shared string is stored once in the
  package but referenced by index from any number of cells, and
  unstructured renders each sheet with ``DataFrame.to_html``, which
  writes the string out once per referencing cell, so a long shared
  string referenced by many cells multiplies into the HTML.
- **Repeated references inside pandoc's readers** (for example an EPUB
  spine listing one chapter many times) are not counted; pandoc is
  bounded only by its heap cap and timeout.
- **.docx total parse time and memory.** The ceilings in step 5
  bound the known fan-outs, not the document: python-docx's and
  unstructured's cost stays linear in the element count and in the
  text, with a large constant. A trivial paragraph costs about
  0.6-2 ms (5,000 one-letter paragraphs under python-docx's template
  took 11 s); python-docx-built paragraphs of ordinary text 3-9 ms
  (11,300 one-run paragraphs in one section took 31 s; 5,600
  paragraphs of seven formatted runs beside 300 styles, 49 s); text
  that triggers unstructured's per-element classification 35-65 ms
  (a US city/state/ZIP regex backtracks for ~40 ms on text starting
  with 80 or more letters without a space or comma, and langdetect
  runs once per element at ~25 ms when the whole text is not one
  language), so an accepted document of a few thousand such
  paragraphs runs for minutes; and one 3 MB paragraph took 43 s and
  ~300 MB. Nothing here bounds that; only a wall-clock and memory
  limit on the in-process parse would.
- **.docx length refusals.** The ceilings share one budget and price
  each step at the slowest layout measured, so they also refuse long
  ordinary documents that would partition in under a minute. Measured
  on documents python-docx writes: a single section of more than
  about 11,300 one-run paragraphs; a Word-like document (300 styles,
  five sections, a table) of more than about 5,250-5,850 paragraphs of
  5-10 formatted runs; more than about 580 one-paragraph sections, or
  180 ten-paragraph ones (a mail merge to "individual documents").
  Long theses and books can exceed these. Vertical merges are priced
  the same way: one column merged down more than about 545 rows is
  refused (300 such rows took 8 s to partition), as is every cell of
  a 64-column table continuing down 27 rows (26 such rows took 22 s),
  or down 26 with cell widths set.
- **.pptx length refusals.** ``MAX_PPTX_PLACEHOLDER_WORK`` counts every
  slide placeholder as inheriting its position and prices each read at
  the slowest layout measured, so it also refuses long ordinary decks:
  about 2,600 title-and-content slides under Office's default theme,
  fewer under templates with more shapes per layout and master.
- **Parse-tree memory of an accepted part.** The guard's own lxml
  walks hold every open element with its attributes, and its parses of
  the skeleton parts (``_skeleton_xml``) the whole tree, about 200
  bytes per attribute, as the loaders' own parses of the same part do.
  ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` bounds one start tag, not how many
  such tags are open at once: a slide, header or footer of nested
  elements each carrying close to that many attributes costs about
  14 MB per open level (eight levels, a 5.1 MB slide, held 118 MB in
  the guard and 122 MB in python-pptx's load; a 128 MiB slide about
  3 GB). ``document.xml`` and relationships parts are not affected: an
  element over ``MAX_DOCX_ELEMENT_ATTRIBUTES`` or
  ``MAX_RELS_ELEMENT_ATTRIBUTES`` attributes is refused as it starts.
- **In-process parsers have no timeout.** python-docx, python-pptx,
  openpyxl, pandas and unstructured run inside the indexing worker with
  no wall clock, so super-linear behaviour other than the fan-outs
  listed above can pin a worker; only the soffice and pandoc
  subprocesses carry timeouts.
"""

from __future__ import annotations

import functools
import hashlib
import io
import posixpath
import re
import struct
import zipfile
import zlib
from typing import Any, Callable, NamedTuple, Optional, cast

from loguru import logger

#: Extensions whose loaders parse zip containers of XML parts.
#: (.doc/.ppt arrive as legacy OLE binaries and are converted to
#: containers by the loader, so the guard does not apply to their
#: raw upload bytes.)
ZIP_CONTAINER_EXTENSIONS = frozenset(
    {".docx", ".xlsx", ".odt", ".pptx", ".epub"}
)

#: Per-entry uncompressed ceiling. A text-heavy 200-page .docx keeps
#: its main XML part in single-digit MB; an XML part near this ceiling
#: already costs the in-process parsers on the order of 1-2+ GB of
#: memory (XML parse trees run ~10-20x the XML size), so larger parts
#: are refused rather than risk an OOM of the worker. This ceiling does
#: not bound what a loader builds from the tree (for example pandas'
#: .xlsx grid padding; see "Known gaps" in the module docstring).
MAX_ENTRY_BYTES = 128 * 1024 * 1024

#: Total uncompressed ceiling across all *parsed* (non-media) entries.
#: Parsers that walk every part (openpyxl across worksheets,
#: python-pptx across slides) pay the parse-tree multiplier on the
#: total, so the total is kept to a few entry ceilings.
MAX_TOTAL_BYTES = 512 * 1024 * 1024

#: Per-entry ceiling for binary media entries (see module docstring).
#: The packages that read media load each part once at ~1x, so a
#: single large video or high-resolution image in a deck stays
#: accepted, while a media "bomb" is still bounded.
MAX_MEDIA_ENTRY_BYTES = 512 * 1024 * 1024

#: Ceiling across *all* entries, media included. python-docx and
#: python-pptx hold every part's blob at once, so this bounds the
#: resident media (at ~1x) plus the parsed parts' raw bytes; it admits
#: a media-heavy deck of several hundred MB. It also bounds the
#: guard's own verification work (every DEFLATE entry is inflated
#: once; a STORED entry's size is fixed by its equality check).
MAX_CONTAINER_BYTES = 1024 * 1024 * 1024

#: Entry-count ceiling: bounds central-directory walks and per-part
#: parser work; real containers carry dozens of parts.
MAX_ENTRIES = 10_000


#: Names that are parsed as markup whatever their leading bytes say.
_MARKUP_SUFFIXES = (
    ".xml",
    ".rels",
    ".xhtml",
    ".html",
    ".htm",
    ".opf",
    ".ncx",
    ".svg",
)

#: Leading-byte signatures of binary media no document parser reads as
#: markup (none of them starts with ``<``).
_MEDIA_SIGNATURES = (
    b"\x89PNG\r\n\x1a\n",  # PNG
    b"\xff\xd8\xff",  # JPEG
    b"GIF87a",
    b"GIF89a",
    b"II*\x00",  # TIFF little-endian
    b"MM\x00*",  # TIFF big-endian
    b"BM",  # BMP
    b"RIFF",  # WAV / AVI / WebP
    b"\xd7\xcd\xc6\x9a",  # placeable WMF
    b"\x1f\x8b\x08",  # gzip (EMZ/WMZ)
    b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1",  # OLE compound file
    b"PK\x03\x04",  # embedded package
    b"%PDF-",
    b"\x30\x26\xb2\x75\x8e\x66\xcf\x11",  # ASF (WMV/WMA)
    b"OggS",
    b"fLaC",
    b"ID3",  # MP3
    b"\x1a\x45\xdf\xa3",  # Matroska / WebM
    b"wOFF",
    b"wOF2",
    b"OTTO",
    b"\x00\x01\x00\x00",  # TrueType
)


class DecompressionBombError(ValueError):
    """A zip-container document declares an unreasonable uncompressed size."""


#: zipfile's local file header layout (``zipfile.structFileHeader``).
_LOCAL_HEADER = struct.Struct("<4s2B4HL2L2H")
_LOCAL_HEADER_MAGIC = b"PK\x03\x04"
#: End-of-central-directory signature: the first bytes of an empty
#: archive (which zipfile, and so the guard, accepts).
_EMPTY_ARCHIVE_MAGIC = b"PK\x05\x06"
#: A central-directory record is 46 fixed bytes plus its name and
#: extra fields (both bounded by other checks once read); 1 KiB per
#: allowed entry is a generous honest ceiling for the walk.
_CENTRAL_DIR_BYTES_PER_ENTRY = 1024
#: The function ``ZipFile._RealGetContents`` itself calls to pick the
#: end records (``endrec = _EndRecData(fp)``), then walks
#: ``endrec[_ECD_SIZE]`` bytes of central directory. Calling the
#: running interpreter's own copy, rather than a re-implementation,
#: is what makes the guard bound the very span zipfile then walks:
#: the selection is subtle (the fixed 22-byte record at the end is
#: taken before any backwards search, the ZIP64 record is followed
#: through its locator's offset and may sit behind "extensible data",
#: and that ZIP64 logic changed across 3.12 patch releases), and any
#: divergence lets a writer show the guard one record and zipfile
#: another. Private but present unchanged in name and return shape in
#: every CPython the project supports (``requires-python`` < 3.15).
_zipfile_end_record = zipfile._EndRecData  # type: ignore[attr-defined]
#: Index of the central-directory size in that record
#: (``zipfile._ECD_SIZE``).
_ECD_SIZE = 5

#: General-purpose flag bits the local/central agreement check reads.
_FLAG_DEFLATE_LEVEL_BITS = 0x0006  # compression-level hint, ignored
_FLAG_DATA_DESCRIPTOR = 0x0008
_FLAG_UTF8 = 0x0800
#: Data-descriptor signature (optional before the descriptor fields).
_DATA_DESCRIPTOR_MAGIC = b"PK\x07\x08"
#: A 32-bit size of this value defers to the ZIP64 extra field.
_ZIP64_LIMIT = 0xFFFFFFFF

#: The only compression methods accepted: those the bounded inflation
#: below can verify (zlib takes an output limit per step). zipfile
#: decompresses BZIP2/LZMA/ZSTD chunks with no output limit at all, so
#: a tiny entry of those methods can materialize gigabytes in any
#: reader, this guard included. Office/ODF/EPUB writers only emit
#: STORED and DEFLATE.
_ALLOWED_METHODS = frozenset({zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED})


def _max_compress_size(file_size: int) -> int:
    """Largest compressed size accepted for a DEFLATE entry.

    Incompressible data costs ~5 bytes per 64 KiB stored block plus
    framing. A full read copies an entry's whole compressed range into
    memory, so that range must not exceed the declared size by more
    than this slack.
    """
    return file_size + (file_size >> 10) + 1024


#: Inflation step sizes for the verification pass (memory stays at a
#: few MiB whatever the entry size).
_INFLATE_INPUT_CHUNK = 64 * 1024
_INFLATE_OUTPUT_CHUNK = 1024 * 1024


def _local_entry_range(
    content: memoryview, info: zipfile.ZipInfo
) -> tuple[int, int]:
    """Byte range of *info*'s compressed data, after requiring the
    entry's local header to agree with its central-directory record.

    zipfile reads an entry through the central directory (offset,
    method, sizes) and only takes the name/extra lengths from the
    local header, but pandoc's zip reader takes the method and
    compressed size from the *local* header. An entry whose two
    headers disagree would therefore be verified here as one stream
    and inflated by pandoc as another — for instance a small honest
    central record over a local header declaring a much larger
    stream. So the method, the general-purpose flags (bits 1-2, the
    DEFLATE level hint, excepted), the CRC, both sizes and the name
    must agree. With a data descriptor (flag bit 3), the local CRC and
    sizes may each be zero (the usual streaming form) or the central
    value, and the descriptor that follows the data must repeat the
    central CRC and sizes exactly; the local data range is then the
    central one by construction.
    """
    start = info.header_offset
    header = bytes(content[start : start + _LOCAL_HEADER.size])
    if len(header) != _LOCAL_HEADER.size:
        raise zipfile.BadZipFile("truncated local header")
    (
        magic,
        _version,
        _os,
        flags,
        method,
        _time,
        _date,
        crc,
        compress_size,
        file_size,
        name_len,
        extra_len,
    ) = _LOCAL_HEADER.unpack(header)
    if magic != _LOCAL_HEADER_MAGIC:
        raise zipfile.BadZipFile("bad local header magic")
    name_start = start + _LOCAL_HEADER.size
    name = bytes(content[name_start : name_start + name_len])
    extra = bytes(
        content[name_start + name_len : name_start + name_len + extra_len]
    )
    if len(name) != name_len or len(extra) != extra_len:
        raise zipfile.BadZipFile("truncated local header")

    zip64 = _ZIP64_LIMIT in (compress_size, file_size)
    if zip64:
        compress_size, file_size = _local_zip64_sizes(extra)

    if method != info.compress_type:
        raise DecompressionBombError(
            "zip entry's local header disagrees with the central directory"
        )
    if (flags ^ info.flag_bits) & ~_FLAG_DEFLATE_LEVEL_BITS:
        raise DecompressionBombError(
            "zip entry's local header disagrees with the central directory"
        )
    encoding = "utf-8" if flags & _FLAG_UTF8 else "cp437"
    if name != info.orig_filename.encode(encoding):
        raise DecompressionBombError(
            "zip entry's local header disagrees with the central directory"
        )
    declared = (crc, compress_size, file_size)
    central = (info.CRC, info.compress_size, info.file_size)
    if flags & _FLAG_DATA_DESCRIPTOR:
        agrees = all(
            local in (0, value) for local, value in zip(declared, central)
        )
    else:
        agrees = declared == central
    if not agrees:
        raise DecompressionBombError(
            "zip entry's local header disagrees with the central directory"
        )

    data_start = name_start + name_len + extra_len
    data_end = data_start + info.compress_size
    if data_end > len(content):
        raise zipfile.BadZipFile("entry data runs past the archive")
    if flags & _FLAG_DATA_DESCRIPTOR:
        _check_data_descriptor(content, data_end, info, zip64=zip64)
    return data_start, data_end


def _local_zip64_sizes(extra: bytes) -> tuple[int, int]:
    """(compressed, uncompressed) sizes from a local ZIP64 extra field.

    A local header that sets either 32-bit size to 0xFFFFFFFF must
    carry both 64-bit sizes (uncompressed first) in its ZIP64 field.
    """
    pos = 0
    while pos + 4 <= len(extra):
        tag, size = struct.unpack_from("<2H", extra, pos)
        body = extra[pos + 4 : pos + 4 + size]
        if tag == 0x0001:
            if len(body) < 16:
                break
            file_size, compress_size = struct.unpack_from("<2Q", body)
            return compress_size, file_size
        pos += 4 + size
    raise zipfile.BadZipFile("local header lacks its ZIP64 sizes")


def _check_data_descriptor(
    content: memoryview, offset: int, info: zipfile.ZipInfo, *, zip64: bool
) -> None:
    """Require the data descriptor at *offset* to repeat the central
    CRC and sizes (signature optional; 8-byte sizes after a ZIP64
    local header, 4-byte otherwise)."""
    size_format = "<L2Q" if zip64 else "<3L"
    record = struct.Struct(size_format)
    tail = bytes(content[offset : offset + 4 + record.size])
    candidates = [tail[: record.size]]
    if tail[:4] == _DATA_DESCRIPTOR_MAGIC:
        candidates.append(tail[4:])
    expected = (info.CRC, info.compress_size, info.file_size)
    if not any(
        len(fields) == record.size and record.unpack(fields) == expected
        for fields in candidates
    ):
        raise DecompressionBombError(
            "zip entry's data descriptor disagrees with the central directory"
        )


def _inflate_head(content: memoryview, info: zipfile.ZipInfo) -> bytes:
    """Verify *info*'s data is exactly what it declares, and return its
    first 64 bytes.

    CPython's zipfile truncates an entry to its declared size only
    *after* inflating each read chunk, and a full ``read()`` inflates
    the entry's whole compressed range in one step, so a DEFLATE entry
    that declares 100 bytes can still materialize gigabytes inside a
    parser. Inflating here in bounded steps, and refusing any entry
    whose real output exceeds its declaration, is what makes the
    declared sizes an actual bound on parser reads.

    The stream must also *end* exactly where it is declared to: the
    inflater has to reach the end of the DEFLATE stream at the last
    byte of the compressed range (no truncation, no trailing bytes),
    having produced exactly the declared size with the declared CRC.
    A reader that does not stop at the central directory's compressed
    size (pandoc's, which follows the local header) would otherwise
    keep inflating past the range this check saw.
    """
    data_start, data_end = _local_entry_range(content, info)
    if info.compress_type == zipfile.ZIP_STORED:
        crc = 0
        for pos in range(data_start, data_end, _INFLATE_OUTPUT_CHUNK):
            crc = zlib.crc32(
                content[pos : min(pos + _INFLATE_OUTPUT_CHUNK, data_end)], crc
            )
        if crc != info.CRC:
            raise DecompressionBombError("zip entry fails its CRC check")
        return bytes(content[data_start : min(data_end, data_start + 64)])

    if data_start == data_end and info.file_size == 0 and info.CRC == 0:
        return b""  # an empty DEFLATE entry some writers emit bare

    inflater = zlib.decompressobj(-zlib.MAX_WBITS)
    produced = 0
    crc = 0
    head = b""
    pos = data_start
    pending = b""
    while not inflater.eof:
        if not pending and pos < data_end:
            pending = bytes(
                content[pos : min(pos + _INFLATE_INPUT_CHUNK, data_end)]
            )
            pos += len(pending)
        # Once the range is consumed, ``pending`` is b"" and this drains
        # output zlib still holds: a step that hit its output cap can
        # consume the last input byte before emitting all its output
        # (and before reaching the end of the stream).
        out = inflater.decompress(pending, _INFLATE_OUTPUT_CHUNK)
        pending = inflater.unconsumed_tail
        if not out and not pending and pos >= data_end and not inflater.eof:
            break  # range consumed and drained before the stream ended
        produced += len(out)
        if produced > info.file_size:
            raise DecompressionBombError(
                "zip entry inflates past its declared size"
            )
        crc = zlib.crc32(out, crc)
        if len(head) < 64:
            head += out[: 64 - len(head)]
    if not inflater.eof:
        raise DecompressionBombError("zip entry's DEFLATE stream is truncated")
    if pending or inflater.unused_data or pos != data_end:
        raise DecompressionBombError(
            "zip entry's DEFLATE stream ends before its compressed range"
        )
    if produced != info.file_size:
        raise DecompressionBombError(
            "zip entry inflates to less than its declared size"
        )
    if crc != info.CRC:
        raise DecompressionBombError("zip entry fails its CRC check")
    return head


#: Containers whose entries may earn the media ceiling. Their parsers
#: (python-docx/-pptx over lxml, openpyxl over lxml/expat) never parse
#: a signature-headed media part as XML. EPUB/ODT go to pandoc, whose
#: readers may still parse such an entry (an EPUB chapter is read by
#: its manifest media type, not its bytes), so every EPUB/ODT entry
#: stays under the parsed-part ceilings.
MEDIA_CEILING_EXTENSIONS = frozenset({".docx", ".pptx", ".xlsx"})


def _is_canonical_member_name(name: str) -> bool:
    """True for a plain relative member name.

    Refuses empty, ``.`` and ``..`` segments (so ``//``, ``/./``, a
    leading ``/`` and traversal) and backslashes. Such names make
    distinct members that readers normalise onto one path — openpyxl's
    ``get_rels_path`` collapses ``a//b.xml`` and ``a/b.xml`` onto one
    ``_rels`` part, for instance — and no Office/ODF/EPUB writer emits
    them. A single trailing ``/`` (a directory entry) is allowed. NUL
    never reaches here: zipfile truncates names at NUL when it reads
    the central directory, and names that then collide are refused by
    the duplicate-name check.
    """
    if not name or "\\" in name:
        return False
    body = name[:-1] if name.endswith("/") else name
    return all(segment not in ("", ".", "..") for segment in body.split("/"))


def _is_media_head(info: zipfile.ZipInfo, head: bytes) -> bool:
    """True when *info* is binary media by content (see module docstring)."""
    if info.is_dir():
        return False
    if info.filename.lower().endswith(_MARKUP_SUFFIXES):
        return False
    # ISO BMFF (MP4 / MOV / M4A / HEIC): a big-endian box size of at
    # least 8 (so the head starts with two zero bytes) then "ftyp".
    # Checked before the markup sniff: such a head cannot parse as XML
    # in any encoding ("ftyp" read as UCS-4 is above U+10FFFF, and two
    # leading zero bytes are a NUL in UTF-16), while a real MP4 whose
    # ftyp box is 60 bytes long starts 00 00 00 3C, which the UTF-32BE
    # decode below would otherwise read as "<".
    if (
        head[:2] == b"\x00\x00"
        and head[4:8] == b"ftyp"
        and int.from_bytes(head[:4], "big") >= 8
    ):
        return True
    if _looks_like_markup(head):
        return False
    if head.startswith(_MEDIA_SIGNATURES):
        return True
    # EMF: EMR_HEADER record type 1 with the " EMF" signature at 40.
    return head[:4] == b"\x01\x00\x00\x00" and head[40:44] == b" EMF"


#: Encodings an XML parser can auto-detect a document start in.
_MARKUP_ENCODINGS = (
    "utf-8",
    "utf-16-le",
    "utf-16-be",
    "utf-32-le",
    "utf-32-be",
)


def _looks_like_markup(head: bytes) -> bool:
    """True when *head* could begin an XML document.

    Conservative: a match only ever keeps an entry under the parsed-part
    ceilings. Undecodable bytes become U+FFFD, never ``<``.
    """
    for encoding in _MARKUP_ENCODINGS:
        text = head.decode(encoding, errors="replace")
        if text.lstrip("\ufeff \t\r\n").startswith("<"):
            return True
    return False


#: unstructured's ``_ZipFileDetector`` rules, in its order: the first
#: member-name pattern any member matches (``re.match``) names the
#: package type; with no match it reads a root ``mimetype`` member as
#: the type.
_SNIFFED_ZIP_TYPES = (
    (".docx", re.compile(r"word/document.*\.xml$")),
    (".xlsx", re.compile(r"xl/workbook.*\.xml$")),
    (".pptx", re.compile(r"ppt/presentation.*\.xml$")),
)

#: The member pandas' ``inspect_excel_format`` needs (names lowercased)
#: to read a workbook with openpyxl rather than another engine.
_PANDAS_XLSX_MEMBER = "xl/workbook.xml"

#: The root ``mimetype`` member of an ODF or EPUB package, and the media
#: type it must name for each extension when present.
_MIMETYPE_MEMBER = "mimetype"
_PACKAGE_MIMETYPES = {
    ".odt": "application/vnd.oasis.opendocument.text",
    ".epub": "application/epub+zip",
}

#: Longest ``mimetype`` member read (the media types above are at most
#: 39 bytes).
_MAX_MIMETYPE_MEMBER_BYTES = 256


def _check_package_type(archive: zipfile.ZipFile, ext: str) -> None:
    """Refuse a package whose member layout makes a content sniffer pick
    a different parser than its extension does.

    langchain's Word and PowerPoint loaders call unstructured's
    ``detect_filetype`` (whenever python-magic imports) and route a file
    it calls DOC or PPT to ``partition_doc``/``partition_ppt``, which run
    soffice with no timeout and partition its output without this
    guard. For a zip the detector returns DOCX, XLSX or PPTX from the
    first of its member-name patterns that matches
    (``_SNIFFED_ZIP_TYPES``), and otherwise trusts a root ``mimetype``
    member's content: a ``.docx`` with no ``word/document*.xml`` and a
    ``mimetype`` of ``application/msword`` went to soffice. An OOXML
    package (``.docx``/``.xlsx``/``.pptx``) is therefore refused when it
    has a ``mimetype`` member (OPC packages have none) or when the
    detector's patterns would not name its own extension, and an
    ``.xlsx`` also when pandas would not pick openpyxl for it (no
    ``xl/workbook.xml``; pandas would pick pyxlsb or odfpy instead). An
    ``.odt`` or ``.epub`` (converted by pandoc with an explicit input
    format, never sniffed on the upload path) is refused when its
    ``mimetype`` member, if present, names another type. Called after
    the entries are verified, so the ``mimetype`` read is bounded."""
    names = archive.namelist()
    if ext in _PACKAGE_MIMETYPES:
        if _MIMETYPE_MEMBER not in names:
            return
        info = archive.getinfo(_MIMETYPE_MEMBER)
        if info.file_size > _MAX_MIMETYPE_MEMBER_BYTES:
            raise DecompressionBombError("package mimetype member is too long")
        try:
            declared: Optional[str] = archive.read(info).decode("utf-8").strip()
        except UnicodeDecodeError:
            declared = None
        if declared != _PACKAGE_MIMETYPES[ext]:
            raise DecompressionBombError(
                "package mimetype does not match its extension"
            )
        return
    if _MIMETYPE_MEMBER in names:
        raise DecompressionBombError(
            "office package has a mimetype member, which a content sniffer "
            "reads as its type"
        )
    sniffed = next(
        (
            kind
            for kind, pattern in _SNIFFED_ZIP_TYPES
            if any(pattern.match(name) for name in names)
        ),
        None,
    )
    if sniffed != ext:
        raise DecompressionBombError(
            "office package's part names do not match its extension"
        )
    if ext == ".xlsx" and _PANDAS_XLSX_MEMBER not in {
        name.lower() for name in names
    }:
        raise DecompressionBombError(
            "spreadsheet package has no xl/workbook.xml part"
        )


def _declared_central_directory_size(content: bytes) -> int:
    """Read the central-directory size zipfile will walk, before
    ``zipfile`` is asked to walk the records.

    The end records are selected by zipfile's own ``_EndRecData`` — the
    same call ``ZipFile`` makes — so the size returned is exactly the
    one zipfile then reads (``endrec[_ECD_SIZE]``, which the ZIP64
    record overrides whenever zipfile accepts one). It reads only the
    end records and walks none. Fails closed: when zipfile would find
    no usable end record, or rejects the one it finds, the container
    is refused here rather than the bound standing aside.
    """
    try:
        endrec = _zipfile_end_record(io.BytesIO(content))
    except (zipfile.BadZipFile, OSError, struct.error) as exc:
        raise DecompressionBombError("malformed zip container") from exc
    if not endrec:
        raise DecompressionBombError("malformed zip container")
    return int(endrec[_ECD_SIZE])


def _check_central_directory_span(content: bytes, *, max_entries: int) -> None:
    """Refuse a central directory whose declared byte span alone
    exceeds the entry-count-scaled bound.

    ``zipfile`` materialises every central-directory record before the
    entry-count ceiling can count them (a hand-packed record flood a
    fraction of the upload cap in size buys hundreds of MiB of RSS),
    and the EOCD's record count wraps at 65,535 — zipfile ignores it —
    so only the declared byte span bounds that walk up front. A
    container whose end records zipfile cannot read is refused too.
    """
    cd_size = _declared_central_directory_size(content)
    limit = max_entries * _CENTRAL_DIR_BYTES_PER_ENTRY
    if cd_size > limit:
        raise DecompressionBombError(
            f"zip container's central directory spans {cd_size} bytes, "
            f"exceeding the {limit}-byte guard limit"
        )


def validate_zip_container(
    content: bytes,
    extension: str,
    *,
    max_entry_bytes: int = MAX_ENTRY_BYTES,
    max_total_bytes: int = MAX_TOTAL_BYTES,
    max_media_entry_bytes: int = MAX_MEDIA_ENTRY_BYTES,
    max_container_bytes: int = MAX_CONTAINER_BYTES,
    max_entries: int = MAX_ENTRIES,
    zip_extensions: frozenset[str] = ZIP_CONTAINER_EXTENSIONS,
) -> None:
    """Refuse zip-container documents whose declared expansion is hostile.

    Raises ``DecompressionBombError`` (a ``ValueError``) when the
    container declares a central-directory byte span over
    ``max_entries`` kibibytes (checked before ``zipfile`` materialises
    the records); declares more entries than ``max_entries``; a parsed
    part over ``max_entry_bytes`` or parsed parts over
    ``max_total_bytes`` in total; a binary media entry over
    ``max_media_entry_bytes``; more than ``max_container_bytes``
    across all entries; or when it is not a readable zip at all; or
    when its layout lets a content sniffer pick another parser than
    the extension does (``_check_package_type``). Non-container
    extensions are returned untouched, and rejection happens before
    any parser sees the bytes.
    """
    ext = (
        extension.lower()
        if extension.startswith(".")
        else f".{extension.lower()}"
    )
    if ext not in zip_extensions:
        return

    # zipfile finds the archive from the end of the file, so anything
    # may precede it; content sniffers (unstructured) read the start.
    if not content.startswith((_LOCAL_HEADER_MAGIC, _EMPTY_ARCHIVE_MAGIC)):
        raise DecompressionBombError(
            "zip container does not start with a zip signature"
        )

    # zipfile materialises every central-directory record before the
    # entry-count ceiling can refuse them, so the span the end records
    # declare is bounded first — whatever their (wrapped, ignored)
    # record counts claim.
    _check_central_directory_span(content, max_entries=max_entries)

    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            parsed_total, container_total, entries = _check_entries(
                archive,
                memoryview(content),
                max_entry_bytes=max_entry_bytes,
                max_total_bytes=max_total_bytes,
                max_media_entry_bytes=max_media_entry_bytes,
                max_container_bytes=max_container_bytes,
                max_entries=max_entries,
                media_allowed=ext in MEDIA_CEILING_EXTENSIONS,
            )
            _check_package_type(archive, ext)
            if ext == ".xlsx":
                _check_xlsx_sheet_fan_out(content)
            elif ext == ".pptx":
                _check_pptx_slide_fan_out(content)
            elif ext == ".docx":
                _check_docx_section_fan_out(content)
    except DecompressionBombError:
        raise
    except (
        zipfile.BadZipFile,
        # A container the stdlib cannot even list or open (undecodable
        # UTF-8-flagged names, unsupported versions or compression,
        # encrypted entries, truncated records, corrupt deflate data)
        # is refused too: fail closed, never pass it on.
        ValueError,
        NotImplementedError,
        RuntimeError,
        EOFError,
        OSError,
        struct.error,
        zlib.error,
    ) as exc:
        raise DecompressionBombError("malformed zip container") from exc

    # Deliberately logs no user-controlled content: entry count and
    # declared sizes only.
    logger.debug(
        "zip container accepted: {} entries, {} parsed / {} total "
        "declared bytes",
        entries,
        parsed_total,
        container_total,
    )


def _check_entries(
    archive: zipfile.ZipFile,
    content: memoryview,
    *,
    max_entry_bytes: int,
    max_total_bytes: int,
    max_media_entry_bytes: int,
    max_container_bytes: int,
    max_entries: int,
    media_allowed: bool,
) -> tuple[int, int, int]:
    infos = archive.infolist()
    if len(infos) > max_entries:
        raise DecompressionBombError(
            f"zip container declares {len(infos)} entries "
            f"(ceiling {max_entries})"
        )

    # Data in front of the archive shifts every header offset (zipfile
    # compensates); a genuine container's first entry starts at 0.
    if infos and min(info.header_offset for info in infos) != 0:
        raise DecompressionBombError(
            "zip container has data before its first entry"
        )

    # Pass 1: declaration-only checks, before any entry data is touched.
    largest_entry_ceiling = max(max_entry_bytes, max_media_entry_bytes)
    header_offsets: set[int] = set()
    names: set[str] = set()
    container_total = 0
    for info in infos:
        if not _is_canonical_member_name(info.filename):
            raise DecompressionBombError(
                "zip container has a non-canonical member name"
            )
        if info.filename in names:
            raise DecompressionBombError(
                "zip container has two members with the same name"
            )
        names.add(info.filename)
        if info.header_offset in header_offsets:
            raise DecompressionBombError(
                "zip container has entries sharing one local header"
            )
        header_offsets.add(info.header_offset)
        if info.file_size > largest_entry_ceiling:
            raise DecompressionBombError(
                f"zip entry declares {info.file_size} bytes "
                f"(ceiling {largest_entry_ceiling})"
            )
        if info.compress_type not in _ALLOWED_METHODS:
            raise DecompressionBombError(
                "zip container uses an unsupported compression method"
            )
        if info.flag_bits & 0x1:
            raise DecompressionBombError("zip container has encrypted entries")
        if info.compress_size > _max_compress_size(info.file_size) or (
            info.compress_type == zipfile.ZIP_STORED
            and info.compress_size != info.file_size
        ):
            raise DecompressionBombError(
                "zip entry's compressed size does not match its declared size"
            )
        container_total += info.file_size
        if container_total > max_container_bytes:
            raise DecompressionBombError(
                f"zip container declares over {max_container_bytes} bytes "
                "total uncompressed"
            )

    # Pass 2: verify each declaration by bounded inflation (total work
    # bounded by pass 1's container total), then classify and apply the
    # per-kind ceilings.
    parsed_total = 0
    for info in infos:
        head = _inflate_head(content, info)
        if media_allowed and _is_media_head(info, head):
            # Binary media is held to its own per-entry ceiling and
            # counts only toward the whole-container total.
            if info.file_size > max_media_entry_bytes:
                raise DecompressionBombError(
                    f"zip media entry {info.filename!r} declares "
                    f"{info.file_size} bytes "
                    f"(ceiling {max_media_entry_bytes})"
                )
            continue
        if info.file_size > max_entry_bytes:
            raise DecompressionBombError(
                f"zip entry {info.filename!r} declares "
                f"{info.file_size} bytes (ceiling {max_entry_bytes})"
            )
        parsed_total += info.file_size
        if parsed_total > max_total_bytes:
            raise DecompressionBombError(
                f"zip container declares over {max_total_bytes} bytes "
                "of parsed parts uncompressed"
            )
    return parsed_total, container_total, len(infos)


#: The parser families openpyxl's read path (and msoffcrypto, which
#: unstructured runs first) parses ``.xlsx`` parts with: lxml
#: (``openpyxl.xml.functions.fromstring``) for the package skeleton, the
#: styles, the document properties, chartsheets and every relationships
#: part, and expat (``xml.etree``/defusedxml ``iterparse``) for
#: worksheets and shared strings; msoffcrypto reads
#: ``[Content_Types].xml`` with minidom, expat too.
_LXML = "lxml"
_EXPAT = "expat"


class _XlsxScanTarget:
    """What one parser of ``_scan_xlsx_part`` reports to: an lxml parser
    target, or the handlers of an expat parser (``expat_start``), whose
    element and attribute names pyexpat interns in *interned*, so they
    are counted from there, after each chunk. It refuses a namespace URI
    over ``MAX_XLSX_NAMESPACE_URI_CHARS`` characters, a DTD, nesting
    deeper than ``MAX_XLSX_XML_DEPTH`` and more than
    ``MAX_XLSX_PART_NAMES`` distinct names, and builds nothing."""

    def __init__(self, interned: Optional[dict] = None):
        self.depth = 0
        self.names: set[str] = set()
        self.interned = interned

    def start_ns(self, prefix: Optional[str], uri: Optional[str]) -> None:
        if uri is not None and len(uri) > MAX_XLSX_NAMESPACE_URI_CHARS:
            raise DecompressionBombError(
                "spreadsheet part declares a namespace URI of more "
                f"than {MAX_XLSX_NAMESPACE_URI_CHARS} characters"
            )
        self.names.add(f"xmlns:{prefix or ''}")
        self._count_names()

    def doctype(self, *_declaration) -> None:
        raise DecompressionBombError("document package part declares a DTD")

    def start(self, tag: str, attributes: dict) -> None:
        self.names.add(tag)
        if attributes:
            self.names.update(attributes)
        self._count_names()
        self.expat_start(tag, attributes)

    def expat_start(self, _tag: str, _attributes) -> None:
        self.depth += 1
        if self.depth > MAX_XLSX_XML_DEPTH:
            raise DecompressionBombError(
                "spreadsheet part nests elements more than "
                f"{MAX_XLSX_XML_DEPTH} deep"
            )

    def _count_names(self) -> None:
        interned = len(self.interned) if self.interned is not None else 0
        if len(self.names) + interned > MAX_XLSX_PART_NAMES:
            raise DecompressionBombError(
                "spreadsheet part uses more than "
                f"{MAX_XLSX_PART_NAMES} distinct names"
            )

    def end(self, _tag) -> None:
        self.depth -= 1

    def close(self) -> None:
        return None


def _check_xlsx_namespaces(
    archive: zipfile.ZipFile,
    parts: dict[str, frozenset[str]],
    scanned: set[tuple[str, str]],
) -> None:
    """Stream each of *parts* (member name -> parser families) once,
    with the parsers of the families the loaders parse it with, before
    any of them parses it (pairs in *scanned* are skipped and new ones
    added). A missing member is skipped: openpyxl skips it too or fails
    to load.

    Refused, in every member scanned: an encoding other than UTF-8
    (``_require_utf8_xml``), a start tag that may carry more than
    ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` attributes (``_XmlTagSigns``) and a
    DTD (``_XmlProlog``; see ``_refuse_dtd``: it could spell a
    declaration through an entity), all checked before either parser
    sees the bytes (``_CountedXmlStream``; the lxml push parser buffered
    a whole internal subset before its doctype callback ran); then, as
    ``_XlsxScanTarget`` receives them, a namespace URI over
    ``MAX_XLSX_NAMESPACE_URI_CHARS`` characters (lxml and expat copy it
    into the name of every element under it), a DTD again (a second
    line), nesting deeper than ``MAX_XLSX_XML_DEPTH`` and more than
    ``MAX_XLSX_PART_NAMES`` distinct names: both parsers keep state for
    every open element (expat about 145 bytes) and every distinct name
    for the whole parse, and neither, as the scan drives it, sets a
    limit. No tree is built. An lxml member is fed to an lxml parser
    with the options of openpyxl's parser and is refused when that
    parse fails: openpyxl's load of a part libxml2 rejects fails too,
    and its parse from a bytes buffer decodes UTF-32 with a byte-order
    mark, which this push parser rejects at the first byte (the
    encoding check before it refuses it first). An expat member is fed to an expat parser
    with namespace processing, as ``xml.etree`` creates it, until its
    first syntax error, where the loader's parse ends too. Python runs
    once per chunk, per declaration and per element start and end."""
    for name, families in parts.items():
        todo = frozenset(f for f in families if (name, f) not in scanned)
        if not todo:
            continue
        scanned.update((name, family) for family in todo)
        try:
            info = archive.getinfo(name)
        except KeyError:
            continue
        _scan_xlsx_part(archive, info, todo)


def _scan_xlsx_part(
    archive: zipfile.ZipFile, info: zipfile.ZipInfo, families: frozenset[str]
) -> None:
    """Scan one member for ``_check_xlsx_namespaces``."""
    from lxml import etree

    # nosemgrep: python.lang.security.use-defused-xml.use-defused-xml, reason: bounded OOXML scan refuses DTD and entity declarations
    from xml.parsers import expat

    # Each live parser: (feed, finish, the errors that end its parse,
    # whether such an error refuses the member).
    live: list[
        tuple[
            Callable[..., Any],
            Callable[..., Any],
            tuple[type[Exception], ...],
            bool,
        ]
    ] = []
    if _LXML in families:
        # openpyxl's parser: XMLParser(resolve_entities=False), whose
        # other options (no network, no huge tree) are lxml's defaults.
        lxml_parser = etree.XMLParser(
            target=_XlsxScanTarget(), resolve_entities=False, no_network=True
        )
        live.append(
            (lxml_parser.feed, lxml_parser.close, (etree.XMLSyntaxError,), True)
        )
    if _EXPAT in families:
        # The separator xml.etree uses; any turns namespace processing
        # on, which decides what expat rejects (an unbound prefix).
        # pyexpat interns every element and attribute name it passes
        # the handlers in this dict (as expat keeps its own entry for
        # each), so its size is the distinct names read so far.
        interned: dict = {}
        expat_parser = expat.ParserCreate(
            namespace_separator="}", intern=interned
        )
        expat_target = _XlsxScanTarget(interned)
        expat_parser.StartNamespaceDeclHandler = expat_target.start_ns
        expat_parser.StartDoctypeDeclHandler = expat_target.doctype
        expat_parser.ordered_attributes = True  # a list costs less
        expat_parser.StartElementHandler = expat_target.expat_start
        expat_parser.EndElementHandler = expat_target.end

        # pyexpat reports an encoding it cannot decode as a LookupError
        # or ValueError, where xml.etree's parse fails too.
        def expat_feed(chunk: bytes, final: bool = False) -> None:
            try:
                expat_parser.Parse(chunk, final)
            finally:
                # At most one chunk's names past the ceiling.
                expat_target._count_names()

        live.append(
            (
                expat_feed,
                functools.partial(expat_feed, b"", True),
                (expat.ExpatError, LookupError, ValueError),
                False,
            )
        )
    with archive.open(info) as member:
        stream = _CountedXmlStream(member)
        while live and (chunk := stream.read(_XML_SCAN_CHUNK)):
            for parser in list(live):
                feed, _finish, errors, refuse = parser
                if not _xlsx_scan_step(
                    functools.partial(feed, chunk), errors, refuse
                ):
                    live.remove(parser)
    for _feed, finish, errors, refuse in live:
        _xlsx_scan_step(finish, errors, refuse)


def _xlsx_scan_step(step, errors: tuple, refuse: bool) -> bool:
    """Run ``step()`` for ``_scan_xlsx_part``; False when it raises one
    of *errors*, where that parser's parse of the member ends, or a
    refusal when *refuse* is set. A refusal (a ``ValueError`` too)
    propagates."""
    try:
        step()
    except DecompressionBombError:
        raise
    except errors as exc:
        if refuse:
            raise DecompressionBombError(
                "malformed spreadsheet package: a part openpyxl parses "
                "with lxml is not XML lxml reads"
            ) from exc
        return False
    return True


#: Chunk size of the namespace scan and of ``_CountedXmlStream``; at
#: most ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` bytes (see ``_XmlTagSigns``).
_XML_SCAN_CHUNK = 64 * 1024

#: Ceiling on the element nesting depth of a scanned ``.xlsx`` part.
#: expat, which openpyxl streams worksheets and shared strings with
#: (and msoffcrypto's minidom ``[Content_Types].xml``), keeps about 145
#: bytes per open element and sets no limit, so 128 MiB of ``<a>``, a
#: 130 KB file, would hold about 6 GB in expat alone (and
#: xml.etree builds a tree element per level); libxml2's push parser
#: driving a target, as the scan drives it, keeps about 36 bytes per
#: level and sets no limit either. libxml2 building a tree, as
#: openpyxl's parse of the other parts does, refuses more than 256
#: levels; this is the same ceiling. Office writers nest worksheet
#: elements about ten deep.
MAX_XLSX_XML_DEPTH = 256

#: Ceiling on the distinct names (element names, attribute names and
#: declared namespace prefixes, an element or attribute name counted
#: with its namespace URI) in a scanned ``.xlsx`` part. expat keeps an
#: entry for every distinct element and attribute name and every
#: prefix for the whole parse (about 140 bytes per element name), as
#: libxml2 keeps its name dictionary, so 1,000,000 distinct element
#: names, 8.9 MB of a worksheet, held about 140 MB in the scan (a
#: 128 MiB part about 2 GB). SpreadsheetML, its extensions and the
#: package parts define a few hundred names.
MAX_XLSX_PART_NAMES = 10_000

#: Ceiling on the ``=`` characters between two ``<`` in an XML part
#: the guard streams or counts (every ``.xlsx`` part the namespace scan
#: covers, and each ``.docx``/``.pptx`` part and relationships part the
#: guard parses): an upper bound on the attributes of a start tag,
#: since an attribute value cannot hold a ``<``. It is counted in the
#: part's bytes, which ``_require_utf8_xml`` holds to UTF-8 (no
#: multi-byte UTF-8 sequence holds either byte), before any parser
#: receives them. An attribute without ``=`` is a syntax error, and the
#: only other source of attributes, the defaults of a DTD's
#: attribute-list declarations (``xmlns:`` ones included), is refused
#: in the same bytes before then (``_XmlProlog``), so no parser builds
#: a start tag of more attributes.
#: expat keeps about 108 bytes per attribute of the tag it is parsing
#: and sets no limit (one tag of 1,000,000 attributes, 11 MB, held
#: about 175 MB in the scan). libxml2 refuses a start tag over 10 MB
#: up front only when it parses from a buffer (``fromstring``): its
#: push parser, as the scan drives it, buffers the whole tag and
#: reports that limit only once the tag is complete (4,000,000
#: attributes, 47 MB, held 551 MB first), and lxml's ``iterparse``,
#: which the ``.docx``/``.pptx`` walkers and the relationships check
#: use, builds such an element in full (2,000,000 attributes, 23 MB,
#: held 406 MB; one of 1,000,000 cost the guard about 200 MB in a
#: ``.docx``, ``.pptx`` or relationships part). Office writers put a
#: few dozen attributes on an element; this also admits a text node of
#: 32,767 ``=``, the most a cell holds in Excel.
MAX_XML_TAG_ATTRIBUTE_SIGNS = 65_536

#: An XML declaration as XML 1.0 (section 2.8) spells it, in bytes: the
#: version, then optionally the encoding name (group 1 or 2) and the
#: standalone flag, separated by XML whitespace.
_XML_DECLARATION = re.compile(
    rb"<\?xml[ \t\r\n]+version[ \t\r\n]*=[ \t\r\n]*"
    rb"(?:\"1\.[0-9]+\"|'1\.[0-9]+')"
    rb"(?:[ \t\r\n]+encoding[ \t\r\n]*=[ \t\r\n]*"
    rb"(?:\"([A-Za-z][A-Za-z0-9._-]*)\"|'([A-Za-z][A-Za-z0-9._-]*)'))?"
    rb"(?:[ \t\r\n]+standalone[ \t\r\n]*=[ \t\r\n]*"
    rb"(?:\"(?:yes|no)\"|'(?:yes|no)'))?"
    rb"[ \t\r\n]*\?>"
)


def _require_utf8_xml(head: bytes) -> None:
    """Refuse an XML part starting with *head* (its first
    ``_XML_SCAN_CHUNK`` bytes, or all of a shorter part) unless both
    parsers read it as UTF-8: an optional UTF-8 byte-order mark, then
    ``<`` or XML whitespace, no NUL byte anywhere in *head*, and either
    no XML declaration or one, complete within *head*, that names no
    encoding or names UTF-8 (in any case).

    The guard counts a part's ``<`` and ``=`` in bytes
    (``_XmlTagSigns``), which holds only for an ASCII-compatible
    encoding both parsers agree on, and they do not agree on UTF-16:
    expat reads a part whose first or second byte is zero as UTF-16
    without a byte-order mark (`` \\x00<\\x00`` as UTF-16LE), where
    libxml2 sees UTF-8 and fails, and libxml2 switches to UTF-16 after
    an ASCII declaration naming ``UTF-16LE``, which expat rejects. In
    UTF-16 a character whose low byte is ``3C`` (U+013C) is a ``<`` to
    a byte count, so a start tag of millions of attributes passed it.
    OPC (ECMA-376 Part 2) allows XML parts in UTF-8 or UTF-16, but
    every producer surveyed writes UTF-8 (Excel, Word and PowerPoint;
    LibreOffice's serializer, which writes nothing else; openpyxl,
    XlsxWriter, python-docx, python-pptx and SheetJS), as do all 2,238
    XML members of the 88 local Office files checked; the only UTF-16
    package found is a reader's test fixture (PhpSpreadsheet's). The NUL and
    first-byte tests also refuse UTF-32 and EBCDIC (``4C 6F A7 94``),
    which libxml2 detects and expat does not, and the declaration test
    every other declared encoding."""
    body = head[3:] if head.startswith(b"\xef\xbb\xbf") else head
    if b"\x00" in head or (body and body[:1] not in b"< \t\r\n"):
        raise DecompressionBombError(
            "document package part is not encoded in UTF-8"
        )
    if body.startswith(b"<?xml") and body[5:6] in (b" ", b"\t", b"\r", b"\n"):
        declaration = _XML_DECLARATION.match(body)
        if declaration is None:
            raise DecompressionBombError(
                "document package part's XML declaration is malformed "
                "or too long"
            )
        name = declaration.group(1) or declaration.group(2)
        if name is not None and name.lower() != b"utf-8":
            raise DecompressionBombError(
                "document package part is not encoded in UTF-8"
            )


class _XmlTagSigns:
    """Count the ``=`` between two ``<`` of a UTF-8 part
    (``_require_utf8_xml``) streamed in chunks of at most
    ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` bytes, and refuse more than that.
    Bytes are counted: no multi-byte UTF-8 sequence holds either byte.
    Only the run open at a chunk's end is carried: a run inside one
    chunk is shorter than the ceiling."""

    def __init__(self) -> None:
        self._open = 0

    def feed(self, chunk: bytes) -> None:
        first = chunk.find(b"<")
        if first < 0:
            self._open += chunk.count(b"=")
        else:
            self._open += chunk.count(b"=", 0, first)
            self._check()
            self._open = chunk.count(b"=", chunk.rfind(b"<"))
        self._check()

    def _check(self) -> None:
        if self._open > MAX_XML_TAG_ATTRIBUTE_SIGNS:
            raise DecompressionBombError(
                "document package part has a start tag that may carry "
                f"more than {MAX_XML_TAG_ATTRIBUTE_SIGNS} attributes"
            )


class _XmlProlog:
    """Refuse a markup declaration in the prolog of a UTF-8 part
    (``_require_utf8_xml``) streamed in chunks: any ``<!`` before the
    root element's start tag that does not open a comment (``<!--``).

    A document type declaration, the only place an ``<!ENTITY``,
    ``<!ATTLIST``, ``<!ELEMENT`` or ``<!NOTATION`` can stand, is legal
    only there; in content, or before the root with no ``<!DOCTYPE``,
    lxml and expat fail on ``<!`` that opens neither a comment nor a
    CDATA section without reading a declaration. Refusing it in the
    bytes keeps the internal subset from every parser: libxml2 reads
    (and its push parser first buffers) the whole subset before the
    root's first event, where ``_refuse_dtd`` looks, building every
    declaration (one million attribute declarations, a 17.9 MB part,
    held 424 MB), and attribute defaults it declares, ``xmlns:`` ones
    included, are added to the start tags they name, attributes the
    ``=`` count (``_XmlTagSigns``) never sees. Comments and processing
    instructions before the root are skipped as the parsers end them
    (``-->``, ``?>``); ``<!-`` without a second ``-``, or ``--``
    inside a comment, which both parsers reject, is refused.
    Scanning stops at the first ``<`` that opens neither, the root
    element (or markup both parsers reject), so ``<!`` in content,
    comments and CDATA sections is never examined."""

    def __init__(self) -> None:
        self._state = "prolog"  # or "dash", "comment", "pi", "done"
        self._carry = b""

    def feed(self, chunk: bytes) -> None:
        if self._state == "done":
            return
        data = self._carry + chunk
        self._carry = b""
        at = 0
        while True:
            if self._state == "pi":
                end = data.find(b"?>", at)
                if end < 0:
                    if data.endswith(b"?") and len(data) - 1 >= at:
                        self._carry = b"?"
                    return
                at = end + 2
                self._state = "prolog"
            elif self._state == "dash":  # after "<!-"
                if at >= len(data):
                    return
                if data[at : at + 1] != b"-":
                    raise DecompressionBombError(
                        "document package part has a malformed comment "
                        "before its root element"
                    )
                at += 1
                self._state = "comment"
            elif self._state == "comment":
                end = data.find(b"--", at)
                if end < 0:
                    if data.endswith(b"-") and len(data) - 1 >= at:
                        self._carry = b"-"
                    return
                if end + 2 >= len(data):
                    self._carry = data[end:]
                    return
                if data[end + 2 : end + 3] != b">":
                    raise DecompressionBombError(
                        "document package part has a malformed comment "
                        "before its root element"
                    )
                at = end + 3
                self._state = "prolog"
            else:
                start = data.find(b"<", at)
                if start < 0:
                    return
                if len(data) - start < 3:
                    # "<" or "<!" (no declaration yet) reach the parser
                    # and are decided with the next chunk.
                    self._carry = data[start:]
                    return
                after = data[start + 1 : start + 2]
                if after == b"?":
                    self._state = "pi"
                    at = start + 2
                elif after != b"!":
                    self._state = "done"
                    return
                elif data[start + 2 : start + 3] == b"-":
                    self._state = "dash"
                    at = start + 3
                else:
                    raise DecompressionBombError(
                        "document package part declares a DTD"
                    )


class _CountedXmlStream:
    """A binary stream over *stream* (a member's stream or a
    ``BytesIO``) that an lxml ``iterparse`` reads in place of it: the
    part is held to ``_require_utf8_xml`` before its first byte is
    returned, and every chunk is counted by ``_XmlTagSigns`` and
    scanned by ``_XmlProlog`` before it is returned, so the parser
    never receives a start tag of more than
    ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` attributes or more of a document
    type declaration than the ``<!`` that opens it (when a chunk ends
    there; the next chunk is refused). A read returns at most
    ``_XML_SCAN_CHUNK`` bytes."""

    def __init__(self, stream) -> None:
        self._stream = stream
        self._signs: Optional[_XmlTagSigns] = None
        self._prolog = _XmlProlog()
        self._pending = b""

    def read(self, size: Optional[int] = -1) -> bytes:
        if self._signs is None:
            head = self._stream.read(_XML_SCAN_CHUNK)
            _require_utf8_xml(head)
            self._signs = _XmlTagSigns()
            self._signs.feed(head)
            self._prolog.feed(head)
            self._pending = head
        if size is None or size < 0 or size > _XML_SCAN_CHUNK:
            size = _XML_SCAN_CHUNK
        if self._pending:
            chunk = self._pending[:size]
            self._pending = self._pending[size:]
            return chunk
        chunk = cast(bytes, self._stream.read(size))
        self._signs.feed(chunk)
        self._prolog.feed(chunk)
        return chunk


def _read_counted_xml(archive: zipfile.ZipFile, info: zipfile.ZipInfo) -> bytes:
    """Member *info*, a whole part about to be parsed from a buffer,
    read through ``_CountedXmlStream``: held to ``_require_utf8_xml``,
    counted by ``_XmlTagSigns`` and scanned by ``_XmlProlog`` chunk by
    chunk as it is read, so a refused part is refused before the rest
    of it is read."""
    with archive.open(info) as member:
        stream = _CountedXmlStream(member)
        return b"".join(
            iter(functools.partial(stream.read, _XML_SCAN_CHUNK), b"")
        )


def _check_xlsx_sheet_fan_out(content: bytes) -> None:
    """Refuse a workbook listing the same sheet part more than once.

    openpyxl (via pandas) reads and converts a sheet part once per
    ``<sheet>`` that resolves to it, so N ``<sheet>`` entries sharing
    one part turn one bounded part into N times the parsing work and
    extracted text. Genuine workbooks never share a sheet part.

    The sheet list is computed by openpyxl itself — the same
    ``ExcelReader`` steps ``load_workbook`` runs (``read_manifest``,
    ``read_workbook``, ``WorkbookParser.find_sheets``) with the options
    pandas passes (``read_only``, ``data_only``, ``keep_links=False``,
    no VBA) — so the guard cannot disagree with the loader about which
    part each sheet names. Those steps parse only ``[Content_Types].xml``,
    the workbook part and the workbook's relationships (external links
    are not read with ``keep_links=False``); no worksheet, drawing or
    shared-string part is parsed. Each of those parts is first held to
    ``MAX_XLSX_SKELETON_PART_BYTES`` (16 MiB), scanned
    (``_check_xlsx_namespaces``, ``[Content_Types].xml`` with expat as
    well, for msoffcrypto) and checked before openpyxl parses it: a DTD
    is refused, and so, after one lxml parse
    of the part, is a comment or processing instruction below its root's
    children (``_refuse_xlsx_nested_node``; openpyxl would build one
    into a list entry), and the workbook part's lists and elements are
    held as ``_check_xlsx_workbook_tree`` holds them. The cost is two
    extra parses of at most that
    much XML per part, the same parse the loader performs right after.
    Once the sheet list passes (``_check_xlsx_sheet_list``), every other
    part openpyxl's load parses (``_xlsx_loaded_parts``) is scanned the
    same way, each with the parser the load uses for it, and the styles
    part is then streamed once (``_check_xlsx_styles``).

    ``rel.target`` is compared exactly: ``read_worksheets`` looks it up
    by exact name among the archive's members, so only identical
    strings load the same member. Every ``<sheet>``'s name, with or
    without a relationship, is held to ``MAX_XLSX_SHEET_NAME_CHARS``,
    and no two may be equal once lowercased.
    """
    try:
        import openpyxl  # noqa: F401
    except ImportError:  # no openpyxl, no openpyxl fan-out
        return

    try:
        # Imported inside the fail-closed block: if an openpyxl upgrade
        # moves these, every .xlsx is refused (and the control tests
        # fail) rather than the check silently switching off.
        from openpyxl.packaging.relationship import get_rels_path
        from openpyxl.reader.excel import ExcelReader, _find_workbook_part
        from openpyxl.xml.constants import ARC_CONTENT_TYPES

        reader = ExcelReader(
            io.BytesIO(content),
            read_only=True,
            keep_vba=False,
            data_only=True,
            keep_links=False,
        )
        with reader.archive:
            archive = reader.archive
            scanned: set[tuple[str, str]] = set()
            # The parts parsed here are package skeleton, a few KB in
            # real workbooks; a tighter ceiling than the general parsed
            # part keeps this extra full-tree parse cheap. Each is
            # scanned (_check_xlsx_namespaces) before it is parsed, and
            # [Content_Types].xml for msoffcrypto's minidom too.
            _require_skeleton_size(
                archive, ARC_CONTENT_TYPES, scanned, (_LXML, _EXPAT)
            )
            reader.read_manifest()
            workbook = _find_workbook_part(reader.package).PartName[1:]
            _require_skeleton_size(archive, workbook, scanned, workbook=True)
            _require_skeleton_size(archive, get_rels_path(workbook), scanned)
            reader.read_workbook()
            sheets = list(reader.parser.find_sheets())
            _check_xlsx_sheet_list(
                [rel.target for _sheet, rel in sheets],
                [sheet.name or "" for sheet in reader.parser.sheets],
            )
            _check_xlsx_namespaces(
                archive, _xlsx_loaded_parts(reader, sheets), scanned
            )
            _check_xlsx_styles(archive)
    except DecompressionBombError:
        raise
    except Exception as exc:  # openpyxl's failure modes are open-ended
        raise DecompressionBombError("malformed spreadsheet package") from exc


def _check_xlsx_sheet_list(targets: list[str], names: list[str]) -> None:
    """Refuse the sheet list ``_check_xlsx_sheet_fan_out`` read (each
    sheet's resolved part and every sheet's name) for the reasons its
    docstring gives."""
    from openpyxl.packaging.relationship import get_rels_path

    if len(targets) > MAX_XLSX_SHEETS:
        raise DecompressionBombError(
            f"spreadsheet lists more than {MAX_XLSX_SHEETS} sheets"
        )
    if any(len(name) > MAX_XLSX_SHEET_NAME_CHARS for name in names):
        raise DecompressionBombError(
            "spreadsheet sheet name is longer than "
            f"{MAX_XLSX_SHEET_NAME_CHARS} characters"
        )
    folded = [name.lower() for name in names]
    if len(folded) != len(set(folded)):
        raise DecompressionBombError("spreadsheet lists two sheets of one name")
    if len(targets) != len(set(targets)):
        raise DecompressionBombError(
            "spreadsheet lists the same sheet part more than once"
        )
    # openpyxl also parses each sheet's own relationships part, found
    # via get_rels_path — which collapses targets that differ only in
    # repeated slashes onto one path — so those must be unique too.
    rels_paths = [get_rels_path(target) for target in targets]
    if len(rels_paths) != len(set(rels_paths)):
        raise DecompressionBombError(
            "spreadsheet sheets share one relationships part"
        )


def _xlsx_loaded_parts(reader, sheets) -> dict[str, frozenset[str]]:
    """The parts openpyxl's read-only load (``ExcelReader.read`` with
    the options pandas passes) parses after the workbook skeleton, each
    with its parser family, found as ``read`` finds them: the shared
    strings part the manifest names (``read_strings``, expat), the core
    and custom document properties and the styles part when present
    (``read_properties``, ``read_custom``, ``apply_stylesheet``, lxml),
    and for every sheet whose resolved part exists
    (``read_worksheets``) its relationships part when present
    (``get_dependents``, lxml) and the part itself: a chartsheet's
    (``"chartsheet"`` in the relationship type) with lxml
    (``read_chartsheet``), a worksheet's with expat, which pandas
    streams it with. ``read_theme`` reads the theme's bytes without
    parsing them, and drawings are never read
    (``openpyxl_hardening``), so neither is listed, nor is any part
    found only by its name or folder (images, unreferenced members)."""
    from openpyxl.packaging.relationship import get_rels_path
    from openpyxl.xml.constants import (
        ARC_CORE,
        ARC_CUSTOM,
        ARC_STYLE,
        SHARED_STRINGS,
    )

    valid = set(reader.valid_files)
    parts: dict[str, set[str]] = {}

    def add(name: str, family: str) -> None:
        if name in valid:
            parts.setdefault(name, set()).add(family)

    strings = reader.package.find(SHARED_STRINGS)
    if strings is not None:
        add(strings.PartName[1:], _EXPAT)
    for name in (ARC_CORE, ARC_CUSTOM, ARC_STYLE):
        add(name, _LXML)
    for _sheet, rel in sheets:
        if rel.target not in valid:
            continue
        add(get_rels_path(rel.target), _LXML)
        add(rel.target, _LXML if "chartsheet" in rel.Type else _EXPAT)
    return {name: frozenset(families) for name, families in parts.items()}


#: Ceiling for the package-skeleton parts the fan-out check parses in
#: full ([Content_Types].xml, the workbook part and its relationships).
#: Typical workbooks keep them in kilobytes, and thousands of sheets
#: stay well under 1 MiB, but the workbook part grows with defined
#: names: a workbook with more than roughly 170k defined names (about
#: 100 bytes each) exceeds 16 MiB and is refused. At the general
#: 128 MiB parsed-part ceiling, one extra full-tree parse would cost
#: gigabytes.
MAX_XLSX_SKELETON_PART_BYTES = 16 * 1024 * 1024

#: Ceiling on the sheets a workbook lists (``<sheet>`` entries with a
#: relationship id, as openpyxl's ``find_sheets`` yields them). openpyxl
#: checks each sheet's target against the archive's member list, a
#: Python list (``rel.target not in self.valid_files``), before it skips
#: a sheet whose part is missing, and pandas looks every sheet up by
#: name in the list of all sheet names, then in the workbook's sheets
#: (``raise_if_bad_sheet_by_name`` and ``Workbook.__getitem__``): sheets
#: times members, and sheets squared. 10,000 empty sheets, a 2.5 MB file
#: within the entry ceiling, took 40 s in ``pandas.read_excel`` before
#: the change, and 100,000 sheets without parts next to 9,000 other
#: members, a 2.2 MB file, 7.6 s in ``openpyxl.load_workbook`` alone
#: (a 16 MiB workbook part lists about 370,000). 2,000 empty sheets take
#: about 2 s in ``pandas.read_excel``, and 2,000 one-cell sheets about
#: 5 s in unstructured's ``partition_xlsx`` (500, 1.2 s). Workbooks
#: rarely have more than a few hundred sheets.
MAX_XLSX_SHEETS = 2_000

#: Ceiling on the characters of a sheet's name (``<sheet name>``).
#: openpyxl names every chartsheet it reads after comparing its name,
#: lowercased, with every sheet name before it, and pandas looks every
#: worksheet up by name in lists of all of them, so long names multiply
#: the per-sheet work ``MAX_XLSX_SHEETS`` bounds: 2,000 sheets of
#: 8,000-character names, a 554 KB file, took 6.4 s in
#: ``pandas.read_excel`` (of 31 characters, about 2 s). Excel limits a
#: sheet name to 31 characters and Google Sheets to 100. Two sheets
#: whose names are equal once lowercased (``str.lower``) are refused
#: too: Excel forbids them, and openpyxl renames such a chartsheet by
#: joining every sheet name into one string and searching it with a
#: regular expression, once per chartsheet.
MAX_XLSX_SHEET_NAME_CHARS = 255

#: Ceiling on the characters of a defined name's value in the workbook
#: part (measured at the longest of a ``definedName`` entry's attribute
#: values, its text and its children's: openpyxl reads the value from
#: an ``attr_text`` or ``attr-text`` attribute or the text, and other
#: fields from attributes or children of their names). Excel limits a
#: formula, and the reference a name stands for, to 8,192 characters.
MAX_XLSX_DEFINED_NAME_CHARS = 8_192

#: Ceiling on the work openpyxl's load spends on the print titles and
#: print areas of the workbook's defined names, each counted at its
#: value's length squared plus ``XLSX_PRINT_NAME_BASE_UNITS``. On every
#: load (pandas' read-only one too, ``WorkbookParser.assign_names``)
#: openpyxl parses each sheet-local name that starts like
#: ``_xlnm.Print_Titles`` or ``_xlnm.Print_Area`` (a prefix match, so
#: one sheet can hold any number of them) with
#: ``PrintTitles.from_string`` or ``PrintArea.from_string``, whose
#: regular expressions scan for a sheet title from every position of
#: the value and back off at its end: time quadratic in its length, up
#: to about 34 ns per character squared (a value of 40,000 ``#`` took
#: 14 s and 31 s, and a 1 MB one would take hours, holding the GIL),
#: plus about 50 us per name (65,536 values of 32 characters took
#: 5.5 s). The guard counts such names whatever their sheet. This
#: ceiling admits one value at ``MAX_XLSX_DEFINED_NAME_CHARS``, which
#: takes 2.3 to 2.9 s to load, 16,476 of 45 characters (2.2 s) or
#: 31,776 of 8 (0.9 s); Excel writes one print area and one set of
#: print titles per sheet, a sheet reference per range (under 100
#: characters each).
XLSX_PRINT_NAME_BASE_UNITS = 2_048
MAX_XLSX_PRINT_NAME_WORK = (
    MAX_XLSX_DEFINED_NAME_CHARS**2 + XLSX_PRINT_NAME_BASE_UNITS
)
_XLSX_PRINT_NAMES = ("_xlnm.Print_Titles", "_xlnm.Print_Area")

#: Ceiling on the characters of a namespace URI declared in any XML
#: member of an ``.xlsx`` package. lxml builds every element and
#: attribute name it hands back as ``{URI}local``, and openpyxl (and
#: the guard) read the name of every element of every part they parse,
#: so one long URI declared once and used through a short prefix is
#: copied again for every element: 2,000 empty elements under a
#: 1,000,000-character URI in a worksheet, a 5.9 KB file, took 3.8 s
#: in ``pandas.read_excel``. Office writers declare URIs well under 100
#: characters; this is the ``.docx`` ceiling
#: (``MAX_DOCX_NAMESPACE_URI_CHARS``). The scan that checks it in
#: every part openpyxl parses (``_check_xlsx_namespaces``), with a
#: Python call per element for the depth ceiling, costs the guard about
#: 6.5 s per 100 MB of worksheet XML (pandas takes about 20 s to read it)
#: and about 8 s per 100 MB of element-dense lxml parts.
MAX_XLSX_NAMESPACE_URI_CHARS = 1_024


#: Ceiling on the characters of a custom number format code
#: (``<numFmt formatCode>``) in ``xl/styles.xml``. Excel's format
#: dialog and the ``.xls`` Format record hold a code to 255
#: characters; this allows four times that.
MAX_XLSX_NUMBER_FORMAT_CHARS = 1_024

#: Ceiling on the entries of each of the cell-format lists of
#: ``xl/styles.xml`` (``<xf>`` elements of ``cellXfs`` and of
#: ``cellStyleXfs``, ``<cellStyle>`` elements of ``cellStyles``), each
#: counted over every such list in the part. Excel's specifications
#: limit a workbook to 65,490 unique cell formats and cell styles.
MAX_XLSX_CELL_FORMATS = 65_490

#: Ceiling on the work openpyxl's stylesheet load spends on number
#: formats, which it does on every load (``apply_stylesheet``, under
#: pandas' read-only load too). ``Stylesheet._normalise_numbers`` runs
#: ``is_date_format`` and ``is_timedelta_format`` (a split and three
#: regular-expression scans) over the format code of every ``<xf>`` of
#: ``cellXfs``, so N formats sharing one long code cost N times its
#: length: one 1,000,000-character code under 2,000 formats, a 3.3 KB
#: file, took 53 s to load. One scan skips ahead past a ``[`` only at
#: the next ``]``, so a code with B unclosed ``[`` costs up to B times
#: its length. Each format is counted at its code's length times
#: ``XLSX_NUMBER_FORMAT_SCAN_WEIGHT`` plus its ``[`` (a code shared by
#: several ``numFmt`` entries at the costliest of them; a built-in
#: format at 64 characters); a unit measured about 0.5 ns for codes of
#: zeros and of brackets alike: 3,600 formats of 1,024 ``[``, counted
#: at 0.99 of this ceiling, load in 2.1 s, and 65,490 formats of 1,024
#: zeros, at 0.84, in 2.9 s. 65,490 formats of an ordinary 30-character
#: date code with one bracket count about 100 million.
MAX_XLSX_NUMBER_FORMAT_WORK = 4_000_000_000
XLSX_NUMBER_FORMAT_SCAN_WEIGHT = 50
_XLSX_BUILTIN_FORMAT_CHARS = 64

#: Ceiling on the work openpyxl spends per named style, which it does
#: on every load (under pandas' read-only load too). For every named
#: style (``<cellStyle>``, duplicates dropped) ``Stylesheet``
#: builds a ``NamedStyle`` with fresh default font, fill, border,
#: alignment and protection objects, rebuilds the dict of every custom
#: number format (``custom_formats``) when its ``cellStyleXfs`` format
#: is custom, and ``apply_stylesheet`` binds it to the workbook
#: (``NamedStyle.bind``), which hashes its font, fill, border,
#: alignment, protection and number format into the workbook's indexed
#: lists and compares each with an equal entry already listed. A hash
#: walks the whole object, so a gradient fill costs every stop, and a
#: comparison reads every string, so a font costs its name's length.
#: Each named style is counted at ``XLSX_NAMED_STYLE_BASE_UNITS`` plus
#: the ``numFmt`` entries plus, for the largest font, the largest fill
#: and the largest border of the part (whichever one it binds),
#: ``XLSX_STYLE_NODE_UNITS`` per element and one unit per
#: ``XLSX_STYLE_CHARS_PER_UNIT`` characters of attribute values and
#: text, plus the longest format code at the same rate. A unit is
#: about 130 ns (one entry of the rebuilt format dict); an element
#: (one gradient stop is two) about 1.5 us, and the fixed work about
#: 75 us. Named styles are counted as the distinct ``xfId`` values of
#: ``cellStyle`` entries within ``-n`` to ``n - 1`` for the ``n``
#: ``cellStyleXfs`` entries: openpyxl drops a style whose ``xfId``
#: repeats, indexes a negative one from the end (so ``n`` formats
#: carry up to ``2n`` named styles) and fails to load an id out of that
#: range. 8,800 named styles over 4,400 formats and 4,400 custom
#: formats took 5.3 s to load, and 1,000 named styles of a 1,000-stop
#: gradient fill about 3 s; 20,000 named styles of a minimal font, fill
#: and border, counted at 0.66 of this ceiling, load in 1.8 s, 2,000 of
#: a 256-stop gradient fill, at 0.89, in 1.6 s, and a workbook bloated
#: to 15,000 named styles of Excel's default font, fills and border
#: with 300 custom formats, at 0.84, in 2 s (20,000 such, at 1.12,
#: take 2.7 s and are refused).
MAX_XLSX_NAMED_STYLE_WORK = 20_000_000
XLSX_NAMED_STYLE_BASE_UNITS = 600
XLSX_STYLE_NODE_UNITS = 16
XLSX_STYLE_CHARS_PER_UNIT = 256

#: Ceiling on the ``<stop>`` elements of one gradient fill in
#: ``xl/styles.xml``. openpyxl hashes every stop each time a named style
#: binds the fill. Excel's cell gradient dialog sets two or three
#: stops (its "Fill Effects" offer two colours); this allows 256.
MAX_XLSX_GRADIENT_STOPS = 256

_XLSX_STYLES_PART = "xl/styles.xml"
#: The ``xl/styles.xml`` lists whose entries openpyxl binds once per
#: named style (each read as a list of all the container's children,
#: whatever their tag; comments and processing instructions are refused
#: there, see ``_refuse_xlsx_nested_node``).
_XLSX_BOUND_LISTS = ("fonts", "fills", "borders")


def _refuse_xlsx_nested_node(depth: int, part: str) -> None:
    """Refuse a comment or processing instruction of the ``.xlsx``
    *part* (named for the message) that is *depth* elements deep,
    unless it lies outside the root (*depth* 0) or directly under it
    (*depth* 1).

    openpyxl parses ``xl/styles.xml`` and the workbook skeleton parts
    with lxml, which keeps comments and processing instructions as
    child nodes, and ``NestedSequence`` builds one entry per child
    *node* of its container, whatever it is
    (``[expected_type.from_tree(el) for el in node]``): a comment among
    the ``fonts``, ``borders`` or ``dxfs`` of ``xl/styles.xml`` becomes a
    default ``Font``, ``Border`` or ``DifferentialStyle``, one among
    ``mruColors`` a default ``Color``, and one among the workbook's
    ``bookViews`` a default ``BookView`` (among ``fills``,
    ``indexedColors`` and ``sheets`` openpyxl fails to load instead).
    The guard counts elements, so 2,000,000 empty comments after one
    font, a 25 KB file whose styles part is 14 MB, were accepted and
    took openpyxl 37 s and 960 MB to load as 2,000,001 fonts, and
    2,390,000 among the ``bookViews`` of a 16 MiB workbook part took the
    guard's own openpyxl parse of it 19.6 s and 1.1 GB. Elsewhere below
    the root openpyxl skips them (``Serialisable.from_tree``, which
    reads ``Sequence`` lists such as ``cellXfs``, ``numFmts``,
    ``cellStyles`` or a gradient's stops, and single entries) or fails
    to load (in a ``fill`` before its pattern, or in a ``patternFill``).
    Directly under the root ``Serialisable.from_tree`` skips them too,
    at the cost of the tree node lxml builds for each, and the
    relationships list, read as every child node of its root, fails to
    load, which leaves openpyxl without the workbook's sheets (refused
    as malformed). Excel writes no comments or processing instructions
    in these parts, so any below the root's children is refused rather
    than counted, and every child node of a list the guard counts is
    then an element it counts."""
    if depth >= 2:
        raise DecompressionBombError(
            f"spreadsheet {part} part has a comment or processing "
            "instruction below its top level"
        )


#: Ceiling on the entries of each of the font, fill and border lists of
#: ``xl/styles.xml``, each counted over every such list in the part.
#: openpyxl builds every entry into a style object and an indexed list
#: of each list on every load, and the guard rebuilds each differently
#: spelled entry once (``MAX_XLSX_STYLE_HASH_CHAIN``). Excel's
#: specifications allow 512 fonts per workbook and 256 fill and line
#: styles; this allows 8,192 of each. A 16 MB styles part with all
#: three lists at this ceiling and 65,490 cell formats, each with its
#: own alignment and protection, loads in 5.3 s (the guard takes 10 s).
MAX_XLSX_STYLE_LIST_ENTRIES = 8_192

#: Ceilings on the other lists of ``xl/styles.xml`` and of the workbook
#: part that openpyxl builds an object from per entry, each counted over
#: every such list in the part (``_XLSX_STYLE_LISTS`` and
#: ``_XLSX_WORKBOOK_LISTS``). Excel documents 200 to 250 custom number
#: formats per workbook; the differential formats of conditional
#: formatting and tables (``dxfs``) are held to its 65,490 cell formats;
#: Excel writes a 64-colour ``indexedColors`` palette and up to ten
#: ``mruColors``; custom table styles, workbook views, external
#: references and pivot caches have no documented limit and number a
#: handful in real files. A table style has one element per table part
#: (28 kinds; counted per table style). Extensions (``<ext>`` of every
#: ``extLst``, counted over the part) number one per cell format at most
#: in Excel's output (``xfComplement``), plus a few.
MAX_XLSX_NUMBER_FORMATS = 8_192
MAX_XLSX_DIFFERENTIAL_FORMATS = 65_490
MAX_XLSX_COLOR_LIST_ENTRIES = 1_024
MAX_XLSX_TABLE_STYLES = 1_024
MAX_XLSX_TABLE_STYLE_ELEMENTS = 32
MAX_XLSX_EXTENSIONS = 131_072
MAX_XLSX_WORKBOOK_LIST_ENTRIES = 8_192

#: Ceiling on the elements below ``dxfs`` in ``xl/styles.xml`` (its
#: entries and everything in them). A differential format holds up to a
#: font, number format, fill, alignment, border and protection, each a
#: handful of elements, and openpyxl builds them all: 10,000 with every
#: one of those filled in (43 elements each, a 9 MB part) take 3.5 s to
#: load. Excel writes a few elements per differential format.
MAX_XLSX_DIFFERENTIAL_FORMAT_ELEMENTS = 500_000

#: Ceiling on the gradient ``<stop>`` elements of ``xl/styles.xml``
#: together (fills and differential formats). openpyxl builds every
#: stop and its colour, and the guard rebuilds every distinct fill:
#: 1,000 fills of 256 stops, a 14 MB part, took 4.5 s to load and the
#: guard 8 s, and the 8,192 fills allowed would take about 37 s and
#: 66 s. Excel's gradients have two or three stops.
MAX_XLSX_GRADIENT_STOPS_TOTAL = 65_536

#: Ceiling on the elements of one entry of the font, fill, border and
#: cell-format lists of ``xl/styles.xml`` (the entry and everything in
#: it, the content of an extension included). The guard holds each such
#: entry whole until it ends, to rebuild it with openpyxl's classes,
#: and then drops it; lxml frees a detached subtree in linear time, but
#: when a Python proxy of a node inside it is still alive (as
#: ``iterparse`` leaves some) and the part declares a namespace, as
#: every Office part does, lxml instead moves the subtree to a document
#: of its own in time quadratic in its nodes: one fill holding 80,000
#: elements took the guard 1.6 s to drop, 160,000 took 6.5 s. Excel
#: writes a dozen or so; a gradient fill at ``MAX_XLSX_GRADIENT_STOPS``
#: stops has 514.
MAX_XLSX_STYLE_ENTRY_ELEMENTS = 1_024

#: Ceiling on the child elements of the root of ``xl/styles.xml`` and of
#: the workbook part. openpyxl builds every child it knows into an
#: object (keeping only the last of a repeated one), so 1,800,000 empty
#: ``<calcPr/>`` in a 16 MiB workbook part took its parse 17 s (the
#: guard's own parse as long again). Excel writes about a dozen.
MAX_XLSX_TOP_LEVEL_ELEMENTS = 256

#: Range of the integers openpyxl reads from ``xl/styles.xml`` (ids,
#: colour indices, charsets and the like; the schema types them as
#: 32-bit, ``xsd:int`` or ``xsd:unsignedInt``), and the magnitude of its
#: floats (sizes, tints, gradient geometry; Excel's largest font is 409
#: points). Python hashes a number modulo ``2**61 - 1``, so integers
#: ``1 + k * (2**61 - 1)`` are distinct values of one hash: 4,000 fonts
#: differing only in such a ``charset`` took 6.6 s to load (a 51 KB
#: part) and 8,000 of them under 8,000 named styles 81 s. Within this
#: range distinct integers hash apart, which also keeps the dicts
#: openpyxl and the guard key by format and style ids linear.
XLSX_STYLE_INT_MIN = -(2**31)
XLSX_STYLE_INT_MAX = 2**32 - 1
MAX_XLSX_STYLE_FLOAT = float(2**32)

#: Ceiling on the distinct values of one hash in any of the indexed
#: lists openpyxl's stylesheet load builds (fonts, fills, borders,
#: alignments, protections, cell formats, number format codes) and in
#: its set of named-style names. Each list is a dict keyed by the
#: style object, whose hash openpyxl takes over the object's field
#: values, so distinct objects of one hash cost a comparison with each
#: other on insertion (quadratic in their number) and on every later
#: lookup, which each named style repeats for its font, fill, border,
#: alignment and protection. Bounding the numbers alone does not
#: prevent it: a float field reaches almost any hash value (``x`` and
#: ``x * 2**-61`` share one), and the hash of a tuple of free 32-bit
#: fields can be steered by a meet-in-the-middle search. Real
#: stylesheets have none; this allows the pair Python's hash folds
#: (``-1`` and ``-2``). 28,000 named styles each binding the second of
#: a pair in every list loaded in 8.2 s, against 7 s without pairs.
MAX_XLSX_STYLE_HASH_CHAIN = 2

#: The lists of ``xl/styles.xml`` under its root (matched by local
#: name) and under its ``colors``, each read as a list of entries: by
#: ``NestedSequence`` (every child node) or ``Sequence`` (every child of
#: the entry's name; ``cellXfs`` also reads ``alignment`` and
#: ``protection`` children), all counted here.
_XLSX_STYLE_LISTS = (
    "numFmts",
    "fonts",
    "fills",
    "borders",
    "cellStyleXfs",
    "cellXfs",
    "cellStyles",
    "dxfs",
    "tableStyles",
)
_XLSX_COLOR_LISTS = ("indexedColors", "mruColors")

#: The lists of the workbook part under its root, each counted over
#: every such list in the part: ``NestedSequence`` lists and the
#: ``Sequence`` lists of the smart tag, function group and web publish
#: objects. ``definedNames`` is held only by the part's 16 MiB ceiling
#: (``MAX_XLSX_SKELETON_PART_BYTES``).
_XLSX_WORKBOOK_LISTS = (
    "bookViews",
    "sheets",
    "customWorkbookViews",
    "externalReferences",
    "pivotCaches",
    "smartTagTypes",
    "functionGroups",
    "webPublishObjects",
)


def _xlsx_list_cap(name: str) -> int:
    """The ceiling on the entries of the ``.xlsx`` list *name*."""
    return {
        "numFmts": MAX_XLSX_NUMBER_FORMATS,
        "fonts": MAX_XLSX_STYLE_LIST_ENTRIES,
        "fills": MAX_XLSX_STYLE_LIST_ENTRIES,
        "borders": MAX_XLSX_STYLE_LIST_ENTRIES,
        "cellStyleXfs": MAX_XLSX_CELL_FORMATS,
        "cellXfs": MAX_XLSX_CELL_FORMATS,
        "cellStyles": MAX_XLSX_CELL_FORMATS,
        "dxfs": MAX_XLSX_DIFFERENTIAL_FORMATS,
        "tableStyles": MAX_XLSX_TABLE_STYLES,
        "indexedColors": MAX_XLSX_COLOR_LIST_ENTRIES,
        "mruColors": MAX_XLSX_COLOR_LIST_ENTRIES,
        "tableStyle": MAX_XLSX_TABLE_STYLE_ELEMENTS,
        "extLst": MAX_XLSX_EXTENSIONS,
        "sheets": MAX_XLSX_SHEETS,
    }.get(name, MAX_XLSX_WORKBOOK_LIST_ENTRIES)


def _refuse_xlsx_list_entry(part: str, name: str, count: int) -> None:
    """Refuse the ``.xlsx`` *part* once list *name* has *count* entries
    over its ceiling (``_xlsx_list_cap``)."""
    cap = _xlsx_list_cap(name)
    if count > cap:
        if name == "sheets":
            raise DecompressionBombError(
                f"spreadsheet lists more than {cap} sheets"
            )
        raise DecompressionBombError(
            f"spreadsheet {part} list more than {cap} {name} entries"
        )


def _refuse_xlsx_repeated_child(part: str, parent: str, seen: set, tag: str):
    """Refuse the ``.xlsx`` *part* when an element that is not a list
    (*parent*, whose child local names so far are *seen*) has a second
    child named *tag*. openpyxl builds every child it knows into an
    object and keeps only the last of a repeated one, so repeats cost a
    build each and are bounded by nothing else; Excel writes none."""
    if tag in seen:
        raise DecompressionBombError(
            f"spreadsheet {part} element {parent} repeats its {tag} child"
        )
    seen.add(tag)


#: The ``xl/styles.xml`` lists whose entries the guard rebuilds as
#: openpyxl's style objects (held in full until the entry ends).
_XLSX_MODELLED_LISTS = ("fonts", "fills", "borders", "cellXfs", "cellStyleXfs")


class _XlsxIndexedList:
    """The distinct values openpyxl's ``IndexedList`` of *what* holds,
    grouped by hash: ``add`` returns a value's index as
    ``IndexedList.add`` does and refuses a value that would be the
    ``MAX_XLSX_STYLE_HASH_CHAIN + 1``-th distinct one of its hash. The
    groups are keyed by hash values (ints), which reduce modulo
    ``2**61 - 1`` again, so at most eight of them share a slot here."""

    def __init__(self, what: str, initial: tuple = ()) -> None:
        self._what = what
        self._groups: dict[int, list[tuple[object, int]]] = {}
        self._size = 0
        for value in initial:
            self.add(value)

    def add(self, value) -> int:
        group = self._groups.setdefault(hash(value), [])
        for known, index in group:
            if known is value or known == value:
                return index
        if len(group) >= MAX_XLSX_STYLE_HASH_CHAIN:
            raise DecompressionBombError(
                f"spreadsheet styles hold more than "
                f"{MAX_XLSX_STYLE_HASH_CHAIN} distinct {self._what} of "
                "one hash"
            )
        group.append((value, self._size))
        self._size += 1
        return self._size - 1


@functools.lru_cache(maxsize=1)
def _xlsx_numeric_fields() -> tuple[
    frozenset[str], frozenset[str], frozenset[str]
]:
    """Names of the integer fields and of the float fields of the
    openpyxl style classes ``Stylesheet`` reads (walked from it through
    every complex field and its subclasses), and of those read from a
    ``val`` attribute of a child element of that name (``<sz val>``).
    Derived from openpyxl so an upgrade that adds a field is covered.
    No name is an integer in one class and a float in another."""
    from openpyxl.descriptors import Descriptor
    from openpyxl.descriptors.serialisable import Serialisable
    from openpyxl.styles.stylesheet import Stylesheet

    ints: set[str] = set()
    floats: set[str] = set()
    nested: set[str] = set()
    seen: set[type] = set()
    pending: list[type] = [Stylesheet]
    while pending:
        cls = pending.pop()
        if cls in seen:
            continue
        seen.add(cls)
        pending.extend(cls.__subclasses__())
        for name in dir(cls):
            desc = getattr(cls, name, None)
            if not isinstance(desc, Descriptor):
                continue
            kind = getattr(desc, "expected_type", None)
            if isinstance(kind, type) and issubclass(kind, Serialisable):
                pending.append(kind)
            elif kind is int or kind is float:
                (ints if kind is int else floats).add(name)
                if getattr(desc, "nested", False):
                    nested.add(name)
    return frozenset(ints), frozenset(floats), frozenset(nested)


def _check_xlsx_style_number(name: str, value: Optional[str]) -> None:
    """Refuse *value* of the numeric style field *name* if openpyxl would
    read it as an integer outside ``XLSX_STYLE_INT_MIN`` to
    ``XLSX_STYLE_INT_MAX`` or a float of magnitude over
    ``MAX_XLSX_STYLE_FLOAT`` (NaN included). A value that is not a
    number is left to openpyxl, which fails to load it."""
    ints, floats, _nested = _xlsx_numeric_fields()
    if value is None or (name not in ints and name not in floats):
        return
    try:
        number: float = int(value) if name in ints else float(value)
    except ValueError:
        return
    if name in ints:
        fits = XLSX_STYLE_INT_MIN <= number <= XLSX_STYLE_INT_MAX
    else:
        fits = abs(number) <= MAX_XLSX_STYLE_FLOAT
    if not fits:
        raise DecompressionBombError(
            f"spreadsheet style {name} value is out of range"
        )


def _xlsx_text_fields(cls) -> frozenset[str]:
    """Names of the fields of openpyxl style class *cls* that
    ``Serialisable.from_tree`` also reads from a child element of that
    local name (its text), overriding the attribute: every descriptor
    that is neither a converter itself nor typed by a class with its
    own ``from_tree``. Derived from openpyxl so an upgrade that adds a
    field is covered; ``CellStyle`` gives ``numFmtId``, ``fontId``,
    ``fillId``, ``borderId``, ``xfId``, ``quotePrefix``,
    ``pivotButton`` and the ``apply*`` flags (its ``alignment``,
    ``protection`` and ``extLst`` children, which Excel writes, are
    complex types), ``_NamedCellStyle`` ``name``, ``xfId``,
    ``builtinId``, ``iLevel``, ``hidden`` and ``customBuiltin`` (its
    ``extLst`` is complex)."""
    from openpyxl.descriptors import Descriptor

    fields = set()
    for name in dir(cls):
        desc = getattr(cls, name, None)
        if (
            isinstance(desc, Descriptor)
            and not hasattr(desc, "from_tree")
            and hasattr(desc, "expected_type")
            and not hasattr(desc.expected_type, "from_tree")
        ):
            fields.add(name)
    return frozenset(fields)


def _check_xlsx_styles(archive: zipfile.ZipFile) -> None:
    """Refuse a workbook whose ``xl/styles.xml`` (the fixed name
    openpyxl reads it by) declares a DTD, has a number format code over
    ``MAX_XLSX_NUMBER_FORMAT_CHARS``, more than ``MAX_XLSX_CELL_FORMATS``
    entries in a cell-format list, a gradient fill of more than
    ``MAX_XLSX_GRADIENT_STOPS`` stops, a field the guard reads given as
    a child element, more than ``MAX_XLSX_STYLE_LIST_ENTRIES`` fonts,
    fills or borders, a numeric field out of the range
    ``_check_xlsx_style_number`` allows, more than
    ``MAX_XLSX_STYLE_HASH_CHAIN`` distinct style objects of one hash in
    any indexed list openpyxl builds from it, a comment or processing
    instruction below the root's children (``_refuse_xlsx_nested_node``)
    or more than ``MAX_XML_OUTSIDE_ROOT_NODES`` outside the root, more
    than ``MAX_XLSX_TOP_LEVEL_ELEMENTS`` elements under the root, more
    entries in its lists (``_XLSX_STYLE_LISTS``, ``_XLSX_COLOR_LISTS``,
    a table style's elements and every ``extLst``'s extensions, every
    child element counting) than ``_xlsx_list_cap`` allows, more than
    ``MAX_XLSX_DIFFERENTIAL_FORMAT_ELEMENTS`` elements under ``dxfs``,
    more than ``MAX_XLSX_GRADIENT_STOPS_TOTAL`` gradient stops in all,
    an element that is not a list repeating a child's local name
    (``_refuse_xlsx_repeated_child``; the content of an extension, which
    openpyxl does not read, excepted), a font, fill, border or
    cell-format entry of more than ``MAX_XLSX_STYLE_ENTRY_ELEMENTS``
    elements, or would cost openpyxl's stylesheet load more than
    ``MAX_XLSX_NUMBER_FORMAT_WORK`` or ``MAX_XLSX_NAMED_STYLE_WORK``.
    Elements are classed as openpyxl reads them: an ``extLst`` is an
    extension list only as the child of an element that is not a list
    (as a list's entry it is that entry, and as a fill's first child a
    gradient whose stops count), and only the ``ext`` children of an
    extension list are extensions.

    The part is read as openpyxl reads it: children are matched by
    local name in any namespace (``Serialisable.from_tree``) under any
    root, and numeric ids are read with ``int``. openpyxl keeps the last
    of repeated lists; every one is counted here. ``from_tree`` also
    reads a field from a child element of the field's name, whose text
    replaces the attribute (``<numFmt numFmtId="164"><formatCode>``),
    so any element child of a ``numFmt`` is refused, as is a child of a
    cell-format ``xf`` or a ``cellStyle`` named after one of its
    attribute fields (``_xlsx_text_fields``); Excel writes neither.
    The part is streamed once, each element dropped once read (an entry
    of the font, fill, border and cell-format lists once that entry has
    been rebuilt with openpyxl's own ``from_tree``), so the guard holds
    the path to the current element, the ``cellStyle`` ids, the cell
    formats' format ids and the distinct style objects (which openpyxl
    holds too). A non-integer id, which openpyxl cannot load, is refused
    as malformed."""
    from lxml import etree
    from openpyxl.styles.alignment import Alignment
    from openpyxl.styles.borders import Border
    from openpyxl.styles.cell_style import CellStyle
    from openpyxl.styles.fills import Fill
    from openpyxl.styles.fonts import Font
    from openpyxl.styles.named_styles import _NamedCellStyle
    from openpyxl.styles.protection import Protection

    try:
        info = archive.getinfo(_XLSX_STYLES_PART)
    except KeyError:
        return
    ints, floats, nested = _xlsx_numeric_fields()
    numeric = ints | floats
    entry_classes = {"fonts": Font, "fills": Fill, "borders": Border}
    # openpyxl's indexed lists, seeded as it seeds them.
    lists = {
        "fonts": _XlsxIndexedList("fonts"),
        "fills": _XlsxIndexedList("fills"),
        "borders": _XlsxIndexedList("borders"),
    }
    alignments = _XlsxIndexedList("alignments", (Alignment(),))
    protections = _XlsxIndexedList("protections", (Protection(),))
    cell_styles = _XlsxIndexedList("cell formats")
    strings = _XlsxIndexedList("format codes or style names")
    modelled: dict[str, set[bytes]] = {
        name: set() for name in _XLSX_MODELLED_LISTS
    }
    text_fields = {
        ("cellXfs", "xf"): _xlsx_text_fields(CellStyle),
        ("cellStyleXfs", "xf"): _xlsx_text_fields(CellStyle),
        ("cellStyles", "cellStyle"): _xlsx_text_fields(_NamedCellStyle),
    }
    format_units: dict[int, int] = {}
    cell_format_ids: list[int] = []
    style_ids: set[int] = set()
    counts = {"cellXfs": 0, "cellStyleXfs": 0, "cellStyles": 0}
    entry_tags = {
        "cellXfs": "xf",
        "cellStyleXfs": "xf",
        "cellStyles": "cellStyle",
    }
    bound_units = dict.fromkeys(_XLSX_BOUND_LISTS, 0)
    number_formats = longest_code = depth = 0
    entry_nodes = entry_chars = stops = outside = top = held = 0
    all_stops = dxf_elements = 0
    extensions = [0]
    path: list[str] = []
    # Per open element: ("root"), ("set", child names so far) for an
    # element that is not a list, ("list", name, counter) for a list,
    # ("gradient", child names but stops), or ("skip") inside an
    # extension, whose content openpyxl does not read.
    frames: list[tuple] = []
    totals = {name: [0] for name in _XLSX_STYLE_LISTS + _XLSX_COLOR_LISTS}
    with archive.open(info) as stream:
        for event, element in etree.iterparse(
            _CountedXmlStream(stream),
            events=("start", "end", "comment", "pi"),
            resolve_entities=False,
            no_network=True,
        ):
            if event in ("comment", "pi"):
                _refuse_xlsx_nested_node(depth, "styles")
                if depth == 0:
                    outside = _note_outside_root_node(outside)
                # Outside the root or directly under it, where openpyxl
                # skips it: drop the siblings before it (it gets no end
                # event; the next sibling's end drops it in turn).
                parent = element.getparent()
                if parent is not None:
                    while element.getprevious() is not None:
                        del parent[0]
                continue
            if event == "end":
                tag = path[-1]
                # from_tree reads such a child's text (before its own
                # first child, if it has any) in place of the attribute.
                if (
                    tag in numeric
                    and tag not in nested
                    and (element.text or "").strip()
                ):
                    _check_xlsx_style_number(tag, element.text)
                if depth == 3 and path[1] in _XLSX_MODELLED_LISTS:
                    # An entry spelled like one already rebuilt adds
                    # nothing new (bloated workbooks repeat entries).
                    spelling = hashlib.blake2b(
                        etree.tostring(element), digest_size=16
                    ).digest()
                    if spelling not in modelled[path[1]]:
                        modelled[path[1]].add(spelling)
                        _model_xlsx_style_entry(
                            element,
                            path[1],
                            entry_classes,
                            lists,
                            alignments,
                            protections,
                            cell_styles,
                        )
                if depth >= 3 and path[1] in bound_units:
                    entry_nodes += 1
                    entry_chars += len(element.text or "") + sum(
                        len(value) for value in element.attrib.values()
                    )
                    if depth == 3:
                        units = (
                            entry_nodes * XLSX_STYLE_NODE_UNITS
                            + entry_chars // XLSX_STYLE_CHARS_PER_UNIT
                        )
                        if units > bound_units[path[1]]:
                            bound_units[path[1]] = units
                        entry_nodes = entry_chars = 0
                depth -= 1
                path.pop()
                frames.pop()
                if depth >= 3 and path[1] in _XLSX_MODELLED_LISTS:
                    continue  # rebuilt with its entry, then dropped
                element.clear()
                parent = element.getparent()
                if parent is not None:
                    while element.getprevious() is not None:
                        del parent[0]
                continue
            depth += 1
            tag = etree.QName(element).localname
            path.append(tag)
            if depth == 1:
                _refuse_dtd(element)
                frames.append(("root",))
                continue
            frame = frames[-1]
            kind = frame[0]
            if depth >= 3 and path[1] in _XLSX_MODELLED_LISTS:
                # Held whole until the entry ends, extension content
                # included (see MAX_XLSX_STYLE_ENTRY_ELEMENTS).
                held = 1 if depth == 3 else held + 1
                if held > MAX_XLSX_STYLE_ENTRY_ELEMENTS:
                    raise DecompressionBombError(
                        f"spreadsheet style {path[2]} entry has more than "
                        f"{MAX_XLSX_STYLE_ENTRY_ELEMENTS} elements"
                    )
            if kind == "skip":
                frames.append(frame)
            else:
                if kind == "root":
                    top += 1
                    if top > MAX_XLSX_TOP_LEVEL_ELEMENTS:
                        raise DecompressionBombError(
                            "spreadsheet styles part has more than "
                            f"{MAX_XLSX_TOP_LEVEL_ELEMENTS} top-level elements"
                        )
                elif kind == "list":
                    frame[2][0] += 1
                    _refuse_xlsx_list_entry("styles", frame[1], frame[2][0])
                elif kind == "gradient" and tag == "stop":
                    stops += 1
                    all_stops += 1
                    if stops > MAX_XLSX_GRADIENT_STOPS:
                        raise DecompressionBombError(
                            "spreadsheet gradient fill has more than "
                            f"{MAX_XLSX_GRADIENT_STOPS} stops"
                        )
                    if all_stops > MAX_XLSX_GRADIENT_STOPS_TOTAL:
                        raise DecompressionBombError(
                            "spreadsheet gradient fills have more than "
                            f"{MAX_XLSX_GRADIENT_STOPS_TOTAL} stops together"
                        )
                else:
                    _refuse_xlsx_repeated_child(
                        "styles", path[-2], frame[1], tag
                    )
                if path[1] == "dxfs" and depth >= 3:
                    dxf_elements += 1
                    if dxf_elements > MAX_XLSX_DIFFERENTIAL_FORMAT_ELEMENTS:
                        raise DecompressionBombError(
                            "spreadsheet differential formats have more "
                            f"than {MAX_XLSX_DIFFERENTIAL_FORMAT_ELEMENTS} "
                            "elements"
                        )
                # The kind of this element, as openpyxl reads it. It
                # reads a fill's first child, whatever its tag, as a
                # gradient unless that tag (namespace included) contains
                # "patternFill"; every child of a list as an entry
                # (``NestedSequence``, whatever its tag) or not at all
                # (``Sequence``, another tag), never as an extension
                # list; and an extension's content not at all.
                if tag == "ext" and kind == "list" and frame[1] == "extLst":
                    frames.append(("skip",))
                elif (
                    kind == "set"
                    and len(frame[1]) == 1
                    and (
                        (depth == 4 and path[1] == "fills")
                        or path[-2] == "fill"
                    )
                    and "patternFill" not in element.tag
                ):
                    stops = 0
                    frames.append(("gradient", set()))
                elif tag == "extLst" and kind != "list":
                    frames.append(("list", tag, extensions))
                elif depth == 2 and tag in totals:
                    frames.append(("list", tag, totals[tag]))
                elif depth == 3 and path[1] == "colors" and tag in totals:
                    frames.append(("list", tag, totals[tag]))
                elif depth == 3 and path[1] == "tableStyles":
                    frames.append(("list", "tableStyle", [0]))
                else:
                    frames.append(("set", set()))
            for key, value in element.attrib.items():
                _check_xlsx_style_number(key, value)
            if tag in nested:
                _check_xlsx_style_number(tag, element.get("val"))
            if depth >= 3 and path[-2] == "numFmt":
                raise DecompressionBombError(
                    "spreadsheet number format has a child element"
                )
            if depth == 4 and tag in text_fields.get((path[1], path[2]), ()):
                raise DecompressionBombError(
                    f"spreadsheet style {path[2]} gives {tag} as a child "
                    "element"
                )
            if depth != 3:
                continue
            container = path[1]
            if container == "numFmts" and tag == "numFmt":
                number_formats += 1
                code = element.get("formatCode") or ""
                strings.add(code)
                if len(code) > MAX_XLSX_NUMBER_FORMAT_CHARS:
                    raise DecompressionBombError(
                        "spreadsheet number format code is longer than "
                        f"{MAX_XLSX_NUMBER_FORMAT_CHARS} characters"
                    )
                longest_code = max(longest_code, len(code))
                format_id = int(element.get("numFmtId"))
                units = _xlsx_format_units(code)
                if units > format_units.get(format_id, 0):
                    format_units[format_id] = units
            elif entry_tags.get(container) == tag:
                counts[container] += 1
                if container == "cellXfs":
                    cell_format_ids.append(int(element.get("numFmtId", 0)))
                elif container == "cellStyles":
                    style_ids.add(int(element.get("xfId")))
                    strings.add(element.get("name"))
    builtin = _XLSX_BUILTIN_FORMAT_CHARS * XLSX_NUMBER_FORMAT_SCAN_WEIGHT
    work = sum(
        format_units.get(format_id, builtin) for format_id in cell_format_ids
    )
    if work > MAX_XLSX_NUMBER_FORMAT_WORK:
        raise DecompressionBombError(
            f"spreadsheet number formats cost {work} units in openpyxl's "
            f"stylesheet load (ceiling {MAX_XLSX_NUMBER_FORMAT_WORK})"
        )
    style_formats = counts["cellStyleXfs"]
    named_styles = sum(
        -style_formats <= style_id < style_formats for style_id in style_ids
    )
    per_style = (
        XLSX_NAMED_STYLE_BASE_UNITS
        + number_formats
        + sum(bound_units.values())
        + longest_code // XLSX_STYLE_CHARS_PER_UNIT
    )
    named = named_styles * per_style
    if named > MAX_XLSX_NAMED_STYLE_WORK:
        raise DecompressionBombError(
            f"spreadsheet named styles cost {named} units in openpyxl's "
            f"stylesheet load (ceiling {MAX_XLSX_NAMED_STYLE_WORK})"
        )


def _model_xlsx_style_entry(
    element,
    container: str,
    entry_classes: dict,
    lists: dict[str, _XlsxIndexedList],
    alignments: _XlsxIndexedList,
    protections: _XlsxIndexedList,
    cell_styles: _XlsxIndexedList,
) -> None:
    """Rebuild one entry of a styles list as openpyxl's load does and add
    it to the indexed lists it lands in: a font, fill or border entry
    (any tag; ``NestedSequence.from_tree``) to its list
    (``apply_stylesheet``); a cell format ``xf``'s alignment and
    protection to theirs and its style array to the cell formats
    (``CellStyleList._to_array``); a named style format's alignment and
    protection to theirs (``NamedStyle.bind``, which adds an absent one
    as the default already listed)."""
    if container in entry_classes:
        lists[container].add(entry_classes[container].from_tree(element))
        return
    from lxml import etree
    from openpyxl.styles.cell_style import CellStyle

    if etree.QName(element).localname != "xf":
        return

    xf = CellStyle.from_tree(element)
    alignment = (
        alignments.add(xf.alignment) if xf.alignment is not None else None
    )
    protection = (
        protections.add(xf.protection) if xf.protection is not None else None
    )
    if container == "cellXfs":
        style = xf.to_array()
        if alignment is not None:
            style.alignmentId = alignment
        if protection is not None:
            style.protectionId = protection
        cell_styles.add(style)


def _xlsx_format_units(code: str) -> int:
    """Units of ``MAX_XLSX_NUMBER_FORMAT_WORK`` one date/timedelta test
    of format *code* costs openpyxl."""
    return max(len(code), 1) * (
        XLSX_NUMBER_FORMAT_SCAN_WEIGHT + code.count("[")
    )


def _require_skeleton_size(
    archive: zipfile.ZipFile,
    name: str,
    scanned: set[tuple[str, str]],
    families: tuple[str, ...] = (_LXML,),
    *,
    workbook: bool = False,
) -> None:
    """Hold *name* to ``MAX_XLSX_SKELETON_PART_BYTES``, scan it with
    *families* (``_check_xlsx_namespaces``, which records it in
    *scanned*) and refuse it if
    it declares a DTD (refused by the scan, in its bytes) or has a
    comment or processing instruction below its root's children
    (``_refuse_xlsx_nested_node``), found by one lxml parse of the part
    before openpyxl parses it, walking only its comments and processing
    instructions. The *workbook* part's lists and elements are then held
    as ``_check_xlsx_workbook_tree`` holds them."""
    from lxml import etree

    try:
        info = archive.getinfo(name)
    except KeyError:
        return  # openpyxl reports the missing part itself
    if info.file_size > MAX_XLSX_SKELETON_PART_BYTES:
        raise DecompressionBombError(
            "spreadsheet package part is larger than "
            f"{MAX_XLSX_SKELETON_PART_BYTES} bytes"
        )
    _check_xlsx_namespaces(archive, {name: frozenset(families)}, scanned)
    with archive.open(info) as stream:
        for _event, element in etree.iterparse(
            _CountedXmlStream(stream),
            events=("start",),
            resolve_entities=False,
            no_network=True,
        ):
            _refuse_dtd(element)
            break
    root = etree.fromstring(
        archive.read(info),
        etree.XMLParser(resolve_entities=False, no_network=True),
    )
    for node in root.iter(etree.Comment, etree.ProcessingInstruction):
        _refuse_xlsx_nested_node(
            1 if node.getparent() is root else 2, "workbook skeleton"
        )
    if workbook:
        _check_xlsx_workbook_tree(root)
    del root


def _check_xlsx_workbook_tree(root) -> None:
    """Refuse a workbook part (parsed, with no comment or processing
    instruction below its root's children) whose root has more than
    ``MAX_XLSX_TOP_LEVEL_ELEMENTS`` child elements, whose lists
    (``_XLSX_WORKBOOK_LISTS``) hold more entries than ``_xlsx_list_cap``
    allows, counted over every such list, or that has an extension list
    of more than ``MAX_XLSX_EXTENSIONS`` entries, an element that is
    not a list repeating a child's local name
    (``_refuse_xlsx_repeated_child``) or defined names that
    ``_check_xlsx_defined_names`` refuses.

    This covers the elements openpyxl's ``WorkbookPackage.from_tree``
    builds objects from: the root's children, the entries of its lists
    and the children of those entries and of the root's other children
    (an extension list's extensions, whose content it does not read).
    Deeper elements it skips or reads as text, and it builds a defined
    name from each ``definedName`` of a ``definedNames`` list, its
    attributes, its text and its children's."""
    from lxml import etree

    totals = dict.fromkeys(_XLSX_WORKBOOK_LISTS + ("extLst",), 0)
    top = 0
    print_work = [0]

    def entry(element) -> None:
        if len(element) == 0:
            return
        seen: set = set()
        parent = etree.QName(element).localname
        for child in element:
            tag = etree.QName(child).localname
            _refuse_xlsx_repeated_child("workbook", parent, seen, tag)
            if tag == "extLst":
                totals[tag] += len(child)
                _refuse_xlsx_list_entry("workbook", tag, totals[tag])

    for child in root.iterchildren(etree.Element):
        top += 1
        if top > MAX_XLSX_TOP_LEVEL_ELEMENTS:
            raise DecompressionBombError(
                "spreadsheet workbook part has more than "
                f"{MAX_XLSX_TOP_LEVEL_ELEMENTS} top-level elements"
            )
        tag = etree.QName(child).localname
        if tag in totals:
            totals[tag] += len(child)
            _refuse_xlsx_list_entry("workbook", tag, totals[tag])
            if tag != "extLst":  # openpyxl does not read inside an <ext>
                for item in child:
                    entry(item)
        elif tag == "definedNames":
            _check_xlsx_defined_names(child, print_work)
        else:
            entry(child)


def _check_xlsx_defined_names(names, print_work: list[int]) -> None:
    """Refuse a ``definedNames`` list of the workbook part with a defined
    name whose value is longer than ``MAX_XLSX_DEFINED_NAME_CHARS``, or
    that brings the print names' work (*print_work*, counted over every
    such list of the part) over ``MAX_XLSX_PRINT_NAME_WORK``.

    openpyxl's ``Serialisable.from_tree`` reads a defined name's value
    (``attr_text``) from an ``attr_text`` or ``attr-text`` attribute
    (it turns ``-`` in an attribute name into ``_``), replaced by the
    element's text when there is any, and its other fields from
    attributes or from child elements of their names, so every entry is
    read here, whatever its tag, at the longest of its attribute values,
    its text and its children's. A name that starts like
    ``_xlnm.Print_Titles`` or ``_xlnm.Print_Area`` (openpyxl matches
    the reserved names as a prefix) is counted at that length squared
    plus ``XLSX_PRINT_NAME_BASE_UNITS``, whatever its sheet."""
    from lxml import etree

    for item in names.iterchildren(etree.Element):
        value = max(
            len(text or "")
            for text in (
                item.text,
                *(c.text for c in item),
                *item.attrib.values(),
            )
        )
        if value > MAX_XLSX_DEFINED_NAME_CHARS:
            raise DecompressionBombError(
                "spreadsheet defined name is longer than "
                f"{MAX_XLSX_DEFINED_NAME_CHARS} characters"
            )
        spellings = [item.get("name")] + [
            c.text for c in item if etree.QName(c).localname == "name"
        ]
        if any(
            (spelling or "").startswith(_XLSX_PRINT_NAMES)
            for spelling in spellings
        ):
            print_work[0] += value * value + XLSX_PRINT_NAME_BASE_UNITS
            if print_work[0] > MAX_XLSX_PRINT_NAME_WORK:
                raise DecompressionBombError(
                    "spreadsheet print areas and titles cost "
                    f"{print_work[0]} units in openpyxl's load (ceiling "
                    f"{MAX_XLSX_PRINT_NAME_WORK})"
                )


def _refuse_dtd(element) -> None:
    """Refuse the part *element* belongs to if it has a document type
    declaration, internal subset or not.

    A DTD is the only way to declare an entity. lxml with
    ``resolve_entities=False`` keeps each reference to a declared
    entity as an ``_Entity`` node, and with an external subset it does
    the same for undeclared ones (the subset is not loaded, so they
    are not errors); ``iterparse`` emits no event for either, so the
    guard's streaming counts would miss them. OPC (ECMA-376 Part 2)
    forbids DTDs in the parts it defines and Office writes none in the
    XML parts it generates.
    Every part the guard parses is first read through
    ``_CountedXmlStream``, whose ``_XmlProlog`` refuses the declaration
    in the bytes, before a parser reads it; this is a second line. It looks at the root's start event, by which libxml2 has
    read the whole internal subset (one million attribute declarations,
    a 17.9 MB part, held 424 MB first)."""
    docinfo = element.getroottree().docinfo
    if (
        docinfo.doctype
        or docinfo.internalDTD is not None
        or docinfo.externalDTD is not None
    ):
        raise DecompressionBombError("document package part declares a DTD")


#: Ceiling on the comments and processing instructions outside the root
#: element (before or after it) of a part the guard streams with
#: comment events and does not otherwise count them in: the ``.xlsx``
#: styles part, ``.pptx`` slides, layouts and masters, and ``.docx``
#: ``document.xml``, headers and footers. lxml's ``iterparse``
#: emits the event for one before the root in time linear in the nodes
#: already there, so N of them cost the guard N squared: 60,000 empty
#: comments before a slide's root, a 420 KB part, kept the guard busy
#: for 16.7 s (40,000, 5.6 s; after the root they cost nothing extra).
#: Office writers put none there.
MAX_XML_OUTSIDE_ROOT_NODES = 64


def _note_outside_root_node(count: int) -> int:
    """Count one more comment or processing instruction outside a
    part's root element, refusing the part past
    ``MAX_XML_OUTSIDE_ROOT_NODES``; returns the new count."""
    count += 1
    if count > MAX_XML_OUTSIDE_ROOT_NODES:
        raise DecompressionBombError(
            "document package part has more than "
            f"{MAX_XML_OUTSIDE_ROOT_NODES} comments or processing "
            "instructions outside its root element"
        )
    return count


# --- Reference fan-out in .pptx / .docx -----------------------------------
#
# python-pptx and python-docx load each package part once, but the text
# extractors walk *references*: unstructured's ``partition_pptx`` emits a
# slide once per ``<p:sldId>`` that resolves to it, and ``partition_docx``
# emits a header/footer part once per section that references it, after
# python-docx has located each section's content with XPath over the
# body-level elements and over every body-level paragraph's children
# (so N sections cost N passes over the body; see
# ``MAX_DOCX_SECTION_BLOCK_WORK`` and ``MAX_DOCX_SECTION_WORK``). One
# bounded part referenced N times is therefore N times the work and the
# extracted text, which no per-entry ceiling bounds. The same holds per
# paragraph: unstructured looks up each non-empty paragraph's style
# about four times (each a scan of the styles part and of the document
# part's relationships), reads each run's text about five times, and
# splits a paragraph at each rendered page break. python-docx's and
# unstructured's cost is not modelled exactly: the section count and
# the parts read per section or per paragraph are capped
# (``MAX_DOCX_SECTIONS``, ``MAX_PART_RELATIONSHIPS``,
# ``MAX_DOCX_SETTINGS_CHILDREN``, ``MAX_DOCX_STYLES``), and the work
# that grows with the body is counted as products of those counts,
# each priced at the slowest layout measured. The checks below parse in
# full only package skeleton parts (relationships, ``presentation.xml``,
# the ``.docx`` settings and styles parts), each held to
# ``MAX_OPC_SKELETON_PART_BYTES``; ``document.xml`` is not a skeleton
# part: it is streamed once without building a tree, and is bounded
# only by the parsed-part ceiling (``MAX_ENTRY_BYTES``, 128 MiB).

#: Ceiling for the skeleton parts the .pptx/.docx fan-out checks parse
#: in full (package and part relationships, ``presentation.xml``, the
#: ``.docx`` settings and styles parts). Real ones are kilobytes (a
#: styles part with hundreds of styles is a few hundred KB); 10,000
#: slides list in well under 1 MiB.
MAX_OPC_SKELETON_PART_BYTES = 16 * 1024 * 1024

#: Ceiling on the entries (children of the root element) of every
#: relationships part (``*.rels`` member) of a ``.docx`` or ``.pptx``
#: package, the parts the guard reads (the package's, the main
#: document's and the presentation's) and the others alike: the
#: libraries parse the relationships part of every part they reach
#: (see ``MAX_PACKAGE_RELATIONSHIPS``). python-docx finds the settings
#: and styles parts by scanning every relationship of the document part
#: in Python (uncached): unstructured does that twice per section for
#: the settings part, which this cap keeps small (10,000 relationships
#: under ``MAX_DOCX_SECTIONS`` sections measured about 3 s), and about
#: four times per non-empty paragraph for the styles part, which
#: ``MAX_DOCX_STYLE_WORK`` counts. Real documents have tens to a few
#: hundred relationships (one per image, hyperlink and header/footer
#: part, plus about a dozen fixed parts), so a document with thousands
#: of images or external hyperlinks still fits.
MAX_PART_RELATIONSHIPS = 10_000

#: Ceiling on the attributes of any element of a relationships part.
#: The OPC schema (ECMA-376 Part 2) gives ``Relationship`` four
#: (``Id``, ``Type``, ``Target``, ``TargetMode``) and ``Relationships``
#: none; namespace declarations are not attributes. Checked at each
#: element's start, so a part whose first entry carries hundreds of
#: thousands of attributes is refused after that entry rather than
#: read to its end.
MAX_RELS_ELEMENT_ATTRIBUTES = 64

#: Ceiling on the relationships of all the relationships parts of a
#: ``.docx`` or ``.pptx`` package together. python-docx and python-pptx
#: load every part reachable through relationships and parse each
#: one's relationships part, about 5-8 us per relationship (a million,
#: in 100 parts of 10,000 external hyperlinks each, a 5.9 MB file, took
#: 5.2 s to open in python-docx and 8.1 s in python-pptx), so 50,000
#: cost well under a second. Real documents have tens to hundreds; one
#: with 10,000 images and as many hyperlinks has about 20,000.
MAX_PACKAGE_RELATIONSHIPS = 50_000

#: Ceiling on python-docx's relationship walks when it opens a package:
#: its ``PackageReader._walk_phys_parts`` keeps the parts it has visited
#: in a list and, for every internal relationship of every part it
#: reaches, scans that list (``partname in visited_partnames``), and
#: ``OpcPackage.iter_rels`` (run on load to gather the image parts)
#: does the same with its own list. The guard counts the internal
#: relationships (``TargetMode`` other than ``External``) of every
#: relationships part times the archive's members (an upper bound on
#: the parts visited: each must exist), each member counted as one step
#: plus one per ``DOCX_RELATIONSHIP_WALK_NAME_BYTES_PER_STEP`` bytes of
#: its name as CPython stores it (comparing names of equal length reads
#: them: one, two or four bytes per character, by the name's widest
#: code point; ``_str_width``). A step measured 17-25 ns: 9,000 parts
#: and one part with 200,000 relationships to the last of them, a
#: 1.66 MB file with every relationships part the guard read under
#: ``MAX_PART_RELATIONSHIPS``, took 41 s to open before the change, and
#: 2,000 parts with 20,000, 1.0 s; 9,000 parts with 19,000 internal
#: relationships, counted at 0.91 of this ceiling, open in 3.0 s. Names
#: of four-byte characters cost two to four times as much per character
#: as ASCII ones and were counted per character before: 2,400 parts
#: named by 450 emoji and 10,114 internal relationships open in 2.1 s
#: and are counted at 1.01 of this ceiling (0.34 per character), and
#: 2,000 parts named by 2,000 emoji and 10,014 relationships, counted
#: at 3.3 (0.89 per character), in 9.5 s, so opening at the ceiling
#: costs about 2-5 s whatever the names. A document with 10,000 images
#: (one part and one relationship each) counts about 110 million.
MAX_DOCX_RELATIONSHIP_WALK = 200_000_000
DOCX_RELATIONSHIP_WALK_NAME_BYTES_PER_STEP = 256

#: Ceiling on the ``w:sectPr`` elements anywhere in ``document.xml``
#: (python-docx enumerates those in body-level paragraph properties and
#: at body level as sections; any other is counted too). python-docx and
#: unstructured do work once per section that this guard does not model
#: exactly: each section re-evaluates the ``_sectPrs`` XPath (whose
#: union merge compares the two kinds of ``w:sectPr`` pairwise and whose
#: returned nodes cost about 100 ns each), and unstructured reads the
#: settings part's ``evenAndOddHeaders`` twice (a scan of the document
#: part's relationships and a search of the settings root's children).
#: This cap, with ``MAX_PART_RELATIONSHIPS`` and
#: ``MAX_DOCX_SETTINGS_CHILDREN``, bounds each of those factors instead:
#: 1,000 sections packed into one paragraph's properties, or split
#: between paragraph and body level, partition in under 2 s. Word writes
#: one ``w:sectPr`` per section break plus the body's own; real
#: documents rarely have more than a few hundred (a mail merge to
#: "individual documents" writes one per record). Each section's
#: paragraphs also count toward ``MAX_DOCX_SECTION_BLOCK_WORK`` and the
#: shared budget: measured on documents python-docx writes (where each
#: section break is a paragraph of its own), more than about 580
#: one-paragraph sections, or about 180 ten-paragraph ones, are
#: refused.
MAX_DOCX_SECTIONS = 1_000

#: Ceiling on the children of the settings part's root element. The
#: part is the one the document part's settings relationship names, as
#: python-docx resolves it, and unstructured searches its root's
#: children twice per section. Word's settings.xml has well under 200.
MAX_DOCX_SETTINGS_CHILDREN = 5_000

#: Ceiling on the quadratic part of python-docx's per-section content
#: lookup: the sum, over sections, of the body-level paragraphs and
#: tables (``w:p``/``w:tbl``) from the start of the body to the
#: section's end, times the child nodes of the body (elements, comments,
#: processing instructions and text). python-docx 1.2 finds a section's
#: blocks with ``preceding-sibling::*[self::w:p | self::w:tbl]`` from the
#: section's end, which returns that section's blocks and every earlier
#: one, and libxml2 (2.14 measured) sorts that reverse-axis result into
#: document order with comparisons that each walk sibling links towards
#: the end of the body. One evaluation therefore costs up to (blocks
#: returned) x (body nodes) steps; it runs once per section. A step
#: costs about 5 ns while the body fits in the CPU caches, but a large
#: body is walked from memory, and the cost then grows with what each
#: node occupies: measured end to end in ``partition_docx`` behind 1,240
#: paragraphs, about 16 ns per step for bare elements, ~48 ns with three
#: attributes, ~70-100 ns with 8-40 attributes or twenty child elements,
#: and ~80 ns per counted step for three-attribute elements separated by
#: text. The ceiling prices every step at that slowest layout, ~100 ns,
#: so this fan-out at the ceiling measured about 45 s (1,240 paragraphs
#: before 348,000 twenty-attribute bookmarks, at 96% of it: 40 s; before
#: the change, 1,240 paragraphs and 2,000,000 three-attribute bookmarks,
#: a 234 KiB file, were accepted and took 121 s). Bodies of plain
#: paragraphs that python-docx would walk at a few ns per step are
#: priced the same. On its own the ceiling would admit a single section
#: of about 21,000 paragraphs and tables; with the shared budget (the
#: style lookup costs about as much per paragraph) the measured limits
#: are about half that (see "Known gaps" in the module docstring).
#: Text between body-level elements is counted whether or not it is
#: whitespace (python-docx's parser drops most whitespace-only text, but
#: not all), so a pretty-printed ``document.xml`` counts up to twice its
#: nodes.
MAX_DOCX_SECTION_BLOCK_WORK = 450_000_000

#: Ceiling on the linear part of python-docx's per-section work, in
#: units of one body node per section: the section count times (the
#: body's child nodes, which each section's sibling walk and
#: ``count()`` visit, plus the nodes the ``_sectPrs`` XPath visits
#: divided by ``DOCX_SECTION_XPATH_NODES_PER_UNIT``, plus the node pairs
#: its union compares divided by ``DOCX_SECTION_UNION_PAIRS_PER_UNIT``).
#: The ``_sectPrs`` XPath is
#: ``/w:document/w:body/w:p/w:pPr/w:sectPr | /w:document/w:body/w:sectPr``;
#: its child steps visit every child node (elements, comments,
#: processing instructions, text) of the document node, of
#: ``w:document``, of each ``w:body``, of each body-level ``w:p`` and of
#: each such paragraph's ``w:pPr``, at 12-33 ns per visit, and libxml2
#: merges its two branches by comparing every paragraph-level
#: ``w:sectPr`` with every body-level one (1-4 ns per pair). The weights
#: are generous rather than exact: at this ceiling 1,000 body-level
#: sections over ~9,700 body nodes took 1-2 s in ``partition_docx``.
MAX_DOCX_SECTION_WORK = 10_000_000

#: Section-XPath node visits that count as one unit of
#: ``MAX_DOCX_SECTION_WORK`` (one body-level element per section).
DOCX_SECTION_XPATH_NODES_PER_UNIT = 64

#: Node pairs compared by the section XPath's union (per evaluation)
#: that count as one unit of ``MAX_DOCX_SECTION_WORK``.
DOCX_SECTION_UNION_PAIRS_PER_UNIT = 1_000

#: Ceiling on header/footer bytes as extracted: each section's
#: header/footer reference costs one text read of the referenced part,
#: so the referenced part's size is counted once per reference. A read
#: costs up to about 3.5 us per byte (empty runs, empty paragraphs or
#: empty table cells: 20 references to a 1.2 MB header of empty runs
#: took ~73 s), so the sum is held to 12 MiB, about 45 s (39-40 s
#: measured for this fan-out at the ceiling). Real documents reference kilobyte-sized
#: headers from a handful of sections; a mail merge of 1,000 records
#: whose header and footer total 12 KB would be refused. Each read also
#: evaluates the part's node-set unions, which are not linear in its
#: size (see ``DOCX_UNION_STEPS_PER_HEADER_BYTE``): a header holding a
#: table before 8,000 empty paragraphs and 40,000 empty elements,
#: referenced by ten sections, a 37 KB file, took 35 s and was counted
#: at 0.23 of this ceiling before.
MAX_DOCX_HEADER_FOOTER_BYTES = 12 * 1024 * 1024

#: Ceiling on the children of the styles part's root element (the part
#: the document part's styles relationship names, as python-docx
#: resolves it). Word writes the styles a document uses plus its
#: defaults, typically tens to a few hundred ``w:style`` elements;
#: python-docx's own template has 164.
MAX_DOCX_STYLES = 5_000

#: Ceiling on unstructured's style lookups: the paragraphs that can have
#: text (body-level ``w:p`` with a ``w:r`` or ``w:hyperlink`` child, plus
#: one per ``w:lastRenderedPageBreak`` in them, since each page break
#: splits off a fragment that is looked up again) times (the styles
#: root's children, plus their attributes divided by
#: ``DOCX_STYLE_ATTRIBUTES_PER_ENTRY``, plus the document part's
#: relationships divided by ``DOCX_RELATIONSHIPS_PER_STYLE_ENTRY``).
#: unstructured reads ``paragraph.style`` about four times per such
#: paragraph; each read scans every relationship of the document part in
#: Python (``part_related_by``) and, for a paragraph whose style is
#: unset or does not resolve, loops over every ``w:style`` in Python
#: (``default_for``), reading two attributes of each. Measured at about
#: 15 us per unit (50 paragraphs beside 300,000 styles, an 805 KB file,
#: took 132 s before the change), so this fan-out at the ceiling is
#: about 45 s: 597 paragraphs beside 5,000 styles took 40 s, and 5,664
#: paragraphs beside 10,000 relationships, at 90% of it, 24 s. On its
#: own it would admit about 18,000 non-empty paragraphs under
#: python-docx's template (166 entries), fewer with more styles; with
#: the shared budget the measured limits are lower (see "Known gaps").
MAX_DOCX_STYLE_WORK = 3_000_000

#: Document-part relationships that count as one styles-part entry in
#: ``MAX_DOCX_STYLE_WORK`` (a relationship costs about 0.1 us per read,
#: a style about 3.5 us).
DOCX_RELATIONSHIPS_PER_STYLE_ENTRY = 32

#: Attributes of the styles root's children that count as one entry in
#: ``MAX_DOCX_STYLE_WORK``: libxml2 finds an attribute by scanning the
#: element's attributes, so styles of 50 or 250 attributes cost 1.4 or
#: 3.1 times as much per read as plain ones.
DOCX_STYLE_ATTRIBUTES_PER_ENTRY = 32

#: Ceiling on the sibling sorting of unstructured's page-break splits:
#: the sum over body-level paragraphs of K * (K + 3) * S**2, K being the
#: paragraph's ``w:lastRenderedPageBreak`` elements and S the largest
#: child-node count of the paragraph or of any element in it.
#: unstructured splits a paragraph at its first rendered page break and
#: recurses into the rest; each split evaluates python-docx XPath with
#: reverse sibling axes (``precedes_all_content`` once per break left in
#: the first run, and the fragment builders), each sorted at up to S**2
#: steps. Measured at up to ~5 ns per unit (one break after 55,000
#: elements of one run took 34-61 s; one run with 600 breaks, a 37 KB
#: file, ran for over 180 s before the change), so this fan-out at the
#: ceiling is about 45 s. Word writes about one rendered page break per
#: page, in paragraphs of tens to hundreds of runs.
MAX_DOCX_PAGE_BREAK_WORK = 9_000_000_000

#: Ceiling on the ``w:lastRenderedPageBreak`` elements anywhere in
#: ``document.xml``. unstructured asks once per document whether it has
#: any, with a four-branch union over the whole body
#: (``w:body/w:p/w:r``, ``w:body/w:p/w:hyperlink/w:r`` and the same two
#: under ``w:body/w:tbl/w:tr/w:tc``), and libxml2 merges each branch by
#: comparing every node with every node already merged, then sorts the
#: result with comparisons that walk the siblings after the breaks:
#: 20,000 breaks in one table cell after 20,000 in body paragraphs took
#: 0.8 s, quadrupling per doubling, and 10,000 breaks in each of the two
#: cell branches with 800-1,200 empty elements between consecutive ones
#: 2.3-2.5 s (~25 ns per pair), while a break in a cell costs the other
#: ceilings about one unit of ``MAX_DOCX_RUN_WORK``. Word writes one
#: rendered page break per page in each flow of text the page boundary
#: cuts (the body, and each cell of a table row the boundary crosses),
#: so a 1,000-page document of ten-column tables has about 10,000 and a
#: book of 20,000 pages this many. This ceiling bounds the merges, not
#: the sort: once a later branch's break precedes an earlier branch's,
#: the sort walks the shared parent's children after the breaks, so
#: filler there costs per break (1,000 breaks of each hyperlink and
#: plain branch in table rows before 160,000 empty table children, a
#: 42 KB file, took 2.2-2.6 s per evaluation, 640,000 of them 10.5 s,
#: and 10,000 rows of each branch alternating with empty spacer rows
#: 3.8-7.6 s). The merges' pairs and the sort's steps are both priced
#: in ``MAX_DOCX_RUN_WORK`` (``DOCX_PAGE_BREAK_UNION_PAIRS_PER_UNIT``,
#: ``DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT``): those three shapes were
#: counted at 0.02, 0.02 and 0.36 of it before, and are now counted at
#: 1.24, 4.85 and 3.06 of it and refused; the sort's price covers
#: siblings that lie far apart in memory and libxml2's binary-insertion
#: runs, which plain filler and 1,000 breaks per branch do not reach,
#: so the union at that ceiling costs about 30 s. Before the change, 96,000
#: breaks in each of the two cell branches, about 0.9 of
#: ``MAX_DOCX_RUN_WORK``, were accepted and would take ~18 s
#: (extrapolated: a 45 KB file with 20,000 in each took 0.77 s and was
#: counted at 0.19 of it). ``w:br w:type="page"`` is
#: not part of this union (it is one of a run's text branches, see
#: ``DOCX_UNION_STEPS_PER_RUN_WORK_UNIT``) and is not counted here.
MAX_DOCX_PAGE_BREAKS = 20_000

#: Node pairs compared by that union's merges (the sum, over pairs of
#: branches, of the product of their break counts) that count as one
#: unit of ``MAX_DOCX_RUN_WORK``: 2-3 ns per pair for the merges and up
#: to ~25 ns with the sort measured, priced at 30 ns.
DOCX_PAGE_BREAK_UNION_PAIRS_PER_UNIT = 1_000

#: Steps of that union's sort that count as one unit of
#: ``MAX_DOCX_RUN_WORK``, a step priced at 150 ns. With two or more
#: non-empty branches, the merged set is out of document order as soon
#: as a later branch's break precedes an earlier branch's, and libxml2
#: then compares breaks under a common ancestor by walking that
#: ancestor's children from one towards the other, to the end of them
#: when the order is reversed. The guard counts, for every element, the
#: union's breaks below it, B, times its child nodes (text included),
#: weighted by ``B + DOCX_PAGE_BREAK_SORT_INSERTION_WALKS *
#: min(B, DOCX_PAGE_BREAK_SORT_INSERTION_RUN)``, summed, and only when
#: at least two branches are non-empty (one branch is already in
#: document order). The weight covers libxml2's timsort: below 64 nodes,
#: and for each run of the merged set shorter than its minimum run
#: (32-64 nodes; the merged set is at most four ascending runs, one per
#: branch), it sorts by binary insertion, about log2(n) comparisons, so
#: up to six walks, per node: 62 breaks of one branch before one of
#: another walked the filler after them 12 times per break, against
#: about 2.5 times for 2,000 breaks. Each walk step is a dependent load
#: of the next sibling, so its cost is set by how far apart siblings
#: lie in memory, not by what they are: 3-9 ns for siblings allocated
#: back to back, and 35-45 ns, a cache miss each, once each carries 16
#: or more attributes, a few descendants or text (64 and 256 attributes
#: cost no more). Measured per weighted step, the worst shapes found
#: (62 breaks of one branch and one of another, or 31 of each of four
#: branches, before 150,000-200,000 filler elements with 16-64
#: attributes each; 1,000 breaks of one branch and one of another before
#: 100,000 filler elements with 64 attributes) cost 70-95 ns, so the
#: sort at the ceiling costs about 30 s. Priced at 30 ns per unweighted
#: step before, 1,000 breaks of each of two branches in table rows
#: before 600,000 filler elements with 16 empty attributes each, a
#: 386 KB file, were accepted at 0.83 of ``MAX_DOCX_RUN_WORK`` and took
#: 50 s for one evaluation of the union.
DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT = 200

#: Extra walks per break of the binary-insertion runs of that union's
#: sort, and the breaks they apply to (see
#: ``DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT``).
DOCX_PAGE_BREAK_SORT_INSERTION_WALKS = 4
DOCX_PAGE_BREAK_SORT_INSERTION_RUN = 64

#: Ceiling on unstructured's text extraction from body-level paragraphs
#: and tables: the sum over them of (K + 1) * (``DOCX_RUN_WORK_ITEM_WEIGHT``
#: * (``w:r``, ``w:tr`` and ``w:tc`` elements) + child nodes of the
#: ``w:r`` elements) + K * (descendant nodes) //
#: ``DOCX_FRAGMENT_NODES_PER_RUN_WORK_UNIT``, K being a paragraph's
#: rendered page breaks (0 for a table). unstructured reads each
#: fragment's text about five times, and each read of a run is an XPath
#: evaluation (~25 us) plus a Python loop over the run's children;
#: tables cost about as much per row and cell (cells counted as
#: ``_DocxSectionStats.run_work`` describes, with python-docx's steps up
#: vertical merges priced by ``DOCX_MERGE_STEP_WEIGHT``); and each
#: page-break split deep-copies the paragraph. Measured at up to ~30 us
#: per unit (1.8 million tabs in one run; 120 paragraphs of 2,479 one-character
#: runs, a 76 KB file, took 40 s, and the cost was bounded only by the
#: 128 MiB part ceiling before the change), so this fan-out at the
#: ceiling is about 45 s. On its own it would admit about 50,000
#: paragraphs of five runs; with the shared budget, far fewer (see
#: "Known gaps").
MAX_DOCX_RUN_WORK = 1_500_000

#: Units of ``MAX_DOCX_RUN_WORK`` per run, table row or table cell.
DOCX_RUN_WORK_ITEM_WEIGHT = 5

#: Descendant nodes of a split paragraph that count as one unit of
#: ``MAX_DOCX_RUN_WORK`` per page break (the deep copy and the walks of
#: each fragment: ~650 ns per node and break measured).
DOCX_FRAGMENT_NODES_PER_RUN_WORK_UNIT = 32

#: Ceiling on every ``w:gridSpan``, ``w:gridBefore`` and ``w:gridAfter``
#: value in ``document.xml``. python-docx's ``_Row.cells`` yields a cell
#: once per grid column it spans, and unstructured re-reads the cell's
#: text for each yield (and pads a row with one empty string per
#: ``gridBefore``/``gridAfter`` column), so the value is a multiplier
#: on a cell's work and on the text it retains: a 950-byte file with
#: one cell spanning 400,000 columns took 53 s, and at 2**31 - 1
#: python-docx would build a tuple of about 17 GB. Word's tables are
#: at most 63 columns wide, and the values of any writer are bounded
#: by the table grid; 1,000 leaves 16 times Word's width. A value
#: python-docx cannot read as an integer (``int()`` of the attribute,
#: as python-docx converts it) is refused too, as is a missing one or
#: one of more than ``_MAX_DOCX_GRID_VALUE_CHARS`` characters. Negative
#: values are accepted (they span and pad nothing).
MAX_DOCX_GRID_SPAN = 1_000

#: Longest ``w:val`` of a grid element the guard parses.
_MAX_DOCX_GRID_VALUE_CHARS = 32

#: Most attributes any one element of ``document.xml`` may carry. lxml
#: reads each attribute value by scanning the element's attribute list
#: again (``attrib.values()``, ``items()``), so reading them all is
#: quadratic in their number: 10,000 attributes took 0.35 s, 40,000
#: 11 s, and one paragraph of 50,000 empty attributes, a 145 KB file,
#: kept this guard itself busy for 40 s before. Real writers put at
#: most a few dozen on an element (Word's ``w:p`` carries up to seven
#: revision and paragraph ids, ``wp:anchor`` eleven; the most found in
#: Word, LibreOffice and python-docx output was twelve, and VML shapes
#: allow a few dozen), so 256 is refused outright, checked by counting
#: (linear) before any attribute is read. lxml has built the element by
#: then; one of more than ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` attributes is
#: refused before lxml receives it (``_CountedXmlStream``).
MAX_DOCX_ELEMENT_ATTRIBUTES = 256

#: Most namespace declarations in scope at any element of
#: ``document.xml`` (its own and its ancestors'). libxml2 resolves the
#: namespace of every element and attribute it copies by walking the
#: declarations in scope, so a paragraph whose elements use N
#: declarations costs N squared per copy, and python-docx copies a
#: paragraph twice per rendered page break: 8,000 declarations on the
#: ``w:document`` element used by 8,000 elements took 0.6 s per copy,
#: and the cost quadruples per doubling. Word declares about 35 on
#: ``w:document`` and drawings add a few each; 256 is refused outright.
MAX_DOCX_NAMESPACES_IN_SCOPE = 256

#: Most namespace declarations inside one body-level paragraph of
#: ``document.xml`` (on the ``w:p`` and on everything in it). lxml's
#: ``remove()``, which python-docx calls on a copy of the paragraph for
#: every rendered page break, re-homes the removed subtree's namespace
#: declarations through a cache it searches linearly, so its cost grows
#: with the square of the declarations in that subtree: 5,000 one-
#: declaration elements took 0.007 s per removal, 40,000 0.4-0.8 s, and
#: six page breaks before 40,000 of them, a 142 KB file, took 5 s to
#: partition and were accepted before (at the copied-bytes ceiling a
#: paragraph of 480,000 would cost minutes). At this cap a removal
#: costs well under a millisecond (a paragraph at the cap with 140 page
#: breaks partitioned in 1.7 s, against 1.65 s without the
#: declarations). Word declares its namespaces on ``w:document`` and up
#: to five more per inline picture (python-docx two), so 1,024 admits
#: paragraphs of 200 pictures. Tables are not split, so they are not
#: capped.
MAX_DOCX_PARAGRAPH_NAMESPACES = 1_024

#: Longest namespace URI, and longest prefix, a declaration in
#: ``document.xml`` may have, in characters; checked when the
#: declaration is announced, before any element name is read. lxml
#: builds every element and attribute name it hands back as
#: ``{URI}local``, so one long URI declared once and used through a
#: short prefix is copied again for every name read under it: a single
#: 8 MB URI used by 3,000 elements, a 44 KB file, ran this guard out of
#: a 3 GB address space within 3 s (a 1 MB one took 3.5 s), and a 1 MB
#: URI on 1,024 attribute names grew it by ~0.5 GB. With the cap, the
#: guard reads each name a bounded number of times, so a name costs it
#: at most about 1 KB per read (100,000 elements named through a URI
#: at the cap took about a third longer than through Word's own). The
#: namespaces Office writers use are well under 100
#: characters (the longest in Word's output,
#: ``http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing``,
#: has 71), and their prefixes under ten.
MAX_DOCX_NAMESPACE_URI_CHARS = 1_024
MAX_DOCX_NAMESPACE_PREFIX_CHARS = 64

#: Most ``w:tc`` children of one ``w:tr``, and most grid columns one row
#: may cover (its ``gridBefore``, the spans of its cells and its
#: ``gridAfter``). python-docx resolves a cell that continues a
#: vertical merge (``row.cells``) by computing its grid offset, which
#: reads the span of every earlier cell of its row, and then walking
#: the row above to that offset, recursively up the merge; unstructured
#: reads ``row.cells`` twice per table. Each such step therefore costs
#: time proportional to the row's width, and a merged row costs its
#: width squared: two rows of 1,000 cells, the second all continuing, a
#: 37 KB file, took 6-7 s per read of the cells (4,000 cells, minutes).
#: The steps are priced in ``MAX_DOCX_RUN_WORK`` (see
#: ``DOCX_MERGE_STEP_WEIGHT``); this cap keeps each step's price, which
#: the guard bounds linearly in the row's width, inside the range it was
#: measured over (past a few hundred cells a step grows faster than
#: linearly; a merged row of 256 cells: ~0.4 s per read, 0.84 s to
#: partition). Word's tables are at most
#: 63 columns wide, so 256 cells is four times that; 2,048 columns
#: admits a row of a cell at the ``MAX_DOCX_GRID_SPAN`` cap with that
#: much padding on both sides.
MAX_DOCX_ROW_CELLS = 256
MAX_DOCX_ROW_GRID_COLUMNS = 2_048

#: Units of ``MAX_DOCX_RUN_WORK`` per step python-docx takes to resolve
#: a cell continuing a vertical merge (``_tc_above``: two XPath
#: evaluations and the grid offset computations, about 0.1 ms per read
#: of ``row.cells`` for a one-column table), plus one per
#: ``DOCX_MERGE_CELLS_PER_UNIT`` cells whose span it reads (the cells
#: before the continuing cell in its row, and those before the cell it
#: lands on in the row above, or all of them if there is none: 2-4 us
#: each per read) and one per ``DOCX_MERGE_NODES_PER_UNIT`` nodes it
#: walks: the child and descendant nodes of the row before the cell and
#: of the row above (which those reads and the row above's cell list
#: walk); the row's children up to its first ``w:trPr`` (all of them if
#: it has none; capped by ``MAX_DOCX_ROW_OTHER_CHILDREN``) and the nodes
#: in that ``w:trPr`` (``grid_before``'s ``find`` calls, wherever the
#: ``w:trPr`` sits); and the table's children from the row above
#: through the row (``_tr_above``; capped by
#: ``MAX_DOCX_TABLE_ROW_GAP``). A cell whose merge goes up D rows takes D
#: steps, each priced at its own rows, so one column merged down R rows
#: costs R squared over two steps. The prices cover both of
#: unstructured's reads: end to end in ``partition_docx``, merged
#: tables of 1-256 columns and 2-500 rows (and rows with 256 other
#: children and 256 more between each two, or cells whose ``w:tcPr``
#: holds 2,000 children), at 0.84-0.96 of the ceiling, took 1-21 us per
#: unit counted, under the ~30 us per unit the ceiling assumes; 64
#: columns merged down 26 rows, at 0.96 of the ceiling, took 22 s.
DOCX_MERGE_STEP_WEIGHT = 10
DOCX_MERGE_CELLS_PER_UNIT = 2
DOCX_MERGE_NODES_PER_UNIT = 16

#: Most children of a table row python-docx reads (a direct ``w:tr``
#: child of a ``w:tbl``) that are not ``w:tc`` elements, and most
#: children of a ``w:tbl`` other than ``w:tr`` elements before or
#: between two of its ``w:tr`` children (``w:tblPr`` and ``w:tblGrid``
#: included); elements, comments and processing instructions count
#: (the text between them is at most one node more each). Every step
#: up a vertical merge reads ``grid_before`` of the continuing cell's
#: row, an lxml ``find`` over all the row's children when it has no
#: ``w:trPr``, and evaluates ``_tr_above``
#: (``ancestor::w:tr[1]/preceding-sibling::w:tr[1]``), which visits
#: every sibling between that row and the one above. The siblings are
#: walked, not their content: a row padded with 200,000 empty elements
#: after its cells, a 1.6 KB file, took 7.6 s per read of the cells,
#: and 26 merged rows of 64 cells with 200,000 between each pair, an
#: 81 KB file, 58 s per read (unstructured reads twice), at 0.09 and
#: 0.90 of the ceiling before; 20,000 wrapped in one ``w:sdt`` per
#: row cost nothing. In a table
#: Word, LibreOffice and python-docx write, a row's other children are
#: its ``w:trPr`` and ``w:tblPrEx`` and the children between rows are
#: none (the ``w:tblPr`` and ``w:tblGrid`` before the first row); the
#: schema also allows bookmark, comment, permission and revision range
#: markers, ``w:proofErr``, and ``w:sdt``/``w:customXml`` wrappers
#: there, one element per wrapped cell or group of rows. 256, the
#: ``MAX_DOCX_ROW_CELLS`` cap, admits Word's widest row (63 columns)
#: with every cell wrapped and room for markers; a step at both caps
#: takes ~0.12 ms per read instead of ~0.10 ms, and the siblings are
#: priced in it (see ``DOCX_MERGE_STEP_WEIGHT``). Children after a
#: table's last row, which no step walks, are not capped (a table may
#: hold all its rows in ``w:sdt`` wrappers).
MAX_DOCX_ROW_OTHER_CHILDREN = 256
MAX_DOCX_TABLE_ROW_GAP = 256

#: Descendant nodes of a table cell that count as one unit of
#: ``MAX_DOCX_RUN_WORK``, before the cell's work is multiplied by its
#: span (or added for each cell continuing its merge).
#: unstructured reads a cell's content once per grid column it is
#: yielded for (``iter_inner_content`` and ``paragraphs`` walk its
#: children, ``Paragraph.text`` each paragraph's), so content other
#: than runs is walked span times too: a cell spanning 1,000 columns
#: with 60,000 empty elements after its paragraph took ~1.7 s to
#: partition (~28 ns per node and yield) and was counted 5,005 units.
#: Real cells hold tens of nodes, which count nothing. That walk is
#: linear only for nodes python-docx skips; the cell's paragraphs and
#: tables are priced as items (``DOCX_CELL_BLOCK_ITEM_WEIGHT``) and the
#: union that selects them by ``DOCX_UNION_STEPS_PER_RUN_WORK_UNIT``.
DOCX_CELL_NODES_PER_UNIT = 256

#: Units of ``MAX_DOCX_RUN_WORK`` per ``w:p`` or ``w:tbl`` child of a
#: table cell, before the cell's work is multiplied by its span:
#: unstructured builds a python-docx object for each of them and reads
#: its text (and a paragraph's runs for emphasis) once per yield, about
#: 20 us for an empty paragraph and 7 us for an empty table measured
#: (16,000 of either in one cell, 1 against 8 yields).
DOCX_CELL_BLOCK_ITEM_WEIGHT = 1

#: Node-set unions python-docx and unstructured evaluate on each read
#: of a cell, paragraph or run: ``./w:p | ./w:tbl`` over a cell's
#: children (and a header's, and, in a document with no sections, the
#: body's), ``w:r | w:hyperlink`` over a paragraph's, and the six
#: text-element branches of ``CT_R.text`` over a run's.
#: libxml2 (2.14 measured) merges each branch's result into the
#: previous ones by comparing every node with every node already
#: there, then sorts the merged set into document order with
#: comparisons that each walk sibling links towards the end of the
#: parent; the merged set is out of order as soon as a later branch's
#: node precedes an earlier branch's, so one evaluation with M nodes
#: selected out of N child nodes costs up to M * N steps (per branch
#: merged after the first): ~2 ns a step in cache, ~60 ns from memory
#: in isolation (one ``w:tbl`` before 10,000 twenty-attribute
#: paragraphs and 40,000 forty-attribute elements: 31 s for one
#: evaluation), and up to ~115 ns per step and read end to end in
#: ``partition_docx`` (1,000 ``w:t`` before a ``w:br`` and 40,000 to
#: 60,000 elements of 40-120 attributes in one run). A step is priced
#: at 120 ns, so one unit (~30 us) is 250 steps of one evaluation; the
#: steps of a cell, paragraph or run are counted ``DOCX_UNION_READS``
#: times (then multiplied, like the rest of its work, by its span, its
#: page breaks or its merges). At 0.82-0.86 of the ceiling under the
#: 100 ns first tried, such a run took 35 s, such a paragraph (a
#: hyperlink before 1,000 runs and 60,000 forty-attribute elements)
#: 21 s and such a cell (a table before 1,000 paragraphs and 60,000
#: forty-attribute elements) 11 s. A body-level paragraph's text is
#: also read with unstructured's
#: ``w:r | w:hyperlink | w:r/descendant::wp:inline[ancestor::w:drawing][1]//w:r``,
#: whose third branch selects the runs inside an inline drawing of each
#: of the paragraph's runs: those are counted as a third branch
#: (any ``w:r`` under a ``wp:inline`` under a ``w:drawing`` in one of
#: the paragraph's ``w:r`` children), and since they are not the
#: paragraph's children, and libxml2 merges each run's share of them
#: into the earlier shares the same way, the evaluation is counted at
#: M * max(M, N) steps: a run whose inline drawing holds one run, before
#: 10,000 plain runs (a 37 KB file), took 0.37 s per evaluation, and
#: before 40,000, 8.9 s; ten runs whose drawings hold 5,000 runs each,
#: 2.9 s (counted at 0.04 and 0.17 of the ceiling before, so that
#: 0.9 of it would have bought minutes). Sorting two runs of one drawing
#: walks that drawing's nodes after the earlier one, so such an
#: evaluation also counts M times the child nodes of every node inside
#: the paragraph's runs' inline drawings (the inlines' own included):
#: one run with a hyperlink after it, whose drawing holds 1,000 runs
#: before 100,000 empty elements (a 38 KB file), took 0.31-0.39 s per
#: evaluation and was counted at 0.04 of the ceiling before (1.65 now);
#: the costliest accepted shape measured with it, 30 runs before 1.8
#: million elements in one drawing (52 KB, 0.92 of the ceiling), takes
#: 2.3 s per evaluation (~40 ns per step). Only an evaluation with at least
#: two non-empty branches is counted: one branch needs no merge and is
#: already in order (an evaluation already in order costs only the
#: merge's comparisons, but is counted the same). Before the change, a 37 KB file with one
#: cell spanning 1,000 columns that held 16,000 empty paragraphs and
#: 16,000 empty tables was accepted at 0.09 of the ceiling and took
#: ~1.5 s per yield (~25 minutes).
DOCX_UNION_STEPS_PER_RUN_WORK_UNIT = 250

#: Evaluations counted per union of a cell, paragraph or run. Measured
#: per body-level paragraph fragment: five of each run's union and four
#: of the paragraph's (its text twice, unstructured's classification
#: lookup and its hyperlink walk); per cell, span + 1 of the cell's
#: union and of its paragraphs' and 2 * span + 1 of its runs' (one or
#: two per yield and one for the table's text), which six per yield
#: covers.
DOCX_UNION_READS = 6

#: Union steps of one header/footer read that count as one byte of
#: ``MAX_DOCX_HEADER_FOOTER_BYTES`` (a byte read is priced at ~3.5 us,
#: a step at 120 ns): each reference reads the part's root union, and
#: each of its cells', paragraphs' and runs' unions, once.
DOCX_UNION_STEPS_PER_HEADER_BYTE = 29

#: Attributes up to which the guard reads an element's attribute values
#: with ``attrib.values()``, quadratic in their number but fastest for
#: few (about 0.4 us for three); above it, with an ``@*`` XPath, linear
#: with a larger fixed cost (about 60 us for 256 against 140 us).
_DOCX_ATTRIBUTE_VALUES_SCAN = 64

#: ``gridBefore``/``gridAfter`` padding columns that count as one unit
#: of ``MAX_DOCX_RUN_WORK`` (one empty string yielded and one ``<td/>``
#: emitted each: ~0.5 us and ~20 bytes measured).
DOCX_GRID_PADDING_PER_RUN_WORK_UNIT = 8

#: Ceiling on the bytes python-docx and unstructured *copy* out of
#: ``document.xml`` beyond reading it once: the text, attributes,
#: namespace declarations and nodes that one bounded paragraph or cell
#: turns into many copies, each priced at what a copy of it was
#: measured to keep.
#:
#: - A paragraph with K rendered page breaks: unstructured's
#:   ``iter_paragraph_items`` recurses once per break, and each level
#:   deep-copies the rest of the paragraph twice (python-docx's
#:   preceding and following fragments) and keeps both until the
#:   paragraph is done. Runs and hyperlinks leave the fragment they do
#:   not belong to, so their content is kept up to K + 1 times (80
#:   breaks before 3 MB of text, a 40 KB file, kept ~230 MB more than
#:   the same paragraph without them, and 20 breaks before 200,000
#:   nested empty elements ~530 MB, ~130 bytes a node); everything else
#:   stays in all 2K copies: the ``w:p`` element with its attributes
#:   and namespace declarations, its ``w:pPr``, the comments and
#:   processing instructions in it, the text after it (lxml copies an
#:   element's tail), and a declaration of each namespace it uses that
#:   an ancestor declares (libxml2 re-declares those on every copy).
#:   Each break is therefore counted as two copies of the paragraph,
#:   of its tail, and of each namespace declaration of the
#:   ``w:document`` and ``w:body`` elements whose URI an element or
#:   attribute name in the paragraph uses: 8 breaks before a 2 MB
#:   namespace URI on the ``w:p``, a 39 KB file, kept ~29 MB, and such
#:   declarations were not counted before (a URI that long is now
#:   refused outright, see ``MAX_DOCX_NAMESPACE_URI_CHARS``).
#: - A table cell spanning S grid columns: its text is extracted and
#:   kept S times (counted as S - 1 extra copies, multiplied through
#:   nested tables), each copy priced at
#:   ``DOCX_TABLE_COPY_BYTES_PER_TEXT_BYTE`` per byte of its text, since
#:   unstructured keeps each copy as a string and twice more HTML-escaped
#:   (in its row's and in the table's HTML): 20 extra copies of a 1 MB
#:   cell of ``"`` and one emoji, a 38 KB file, kept ~1 GB.
#: - A cell continuing a vertical merge: python-docx yields the cell
#:   its merge starts from, so that cell's text is kept once more per
#:   continuing row, priced as a spanned cell's copies are. The
#:   streaming pass cannot resolve which cell that is, so it counts the
#:   largest cell (times its span) of the earlier rows of the same
#:   table: 300 continuing rows under a 1 MB cell, a 38 KB file, kept
#:   ~930 MB and took 21 s, and 20 under a 1 MB cell of ``"`` and one
#:   emoji kept ~1 GB. A table whose continuing rows times its largest
#:   earlier cell exceed about 2 MiB of text (``MAX_DOCX_COPIED_BYTES``
#:   divided by ``DOCX_TABLE_COPY_BYTES_PER_TEXT_BYTE``; 200 continuing
#:   rows beside a 10 KB cell) is therefore refused, even when the cells
#:   actually merged are small. A cell continues exactly
#:   when python-docx and unstructured say so (``tc.vMerge ==
#:   "continue"``): the first ``w:vMerge`` of the cell's first
#:   ``w:tcPr`` has ``w:val="continue"`` or no ``w:val``. Any other
#:   cell, whatever other ``w:vMerge`` or ``w:tcPr`` elements it
#:   carries, is counted as an ordinary cell, once per spanned column.
#:
#: Real documents copy a few MB: Word writes about one rendered page
#: break per page, in paragraphs of a few KB, and spans and merges of
#: small cells. Element, attribute and processing-instruction names
#: are not counted: lxml's copies share them through the parser's name
#: dictionary (measured: no growth). The count is not a measurement:
#: each item is priced at about the most its copy was measured to keep
#: (an element ~130 bytes, an attribute ~275, a declaration ~110,
#: a byte of lxml text about 1; a byte of table-cell text ~52 at worst,
#: ~12 for plain ASCII), and the shapes measured kept less than they
#: were counted (20 page breaks before 200,000 nested empty elements
#: ~530 MB, counted ~1 GB; 20 before 200 elements of 256 empty
#: attributes ~310 MB, counted ~590 MB; 20 extra copies of a 1 MB cell of ``"`` and
#: one emoji ~1 GB, counted ~1.3 GB), so a document at this ceiling
#: keeps on the order of 128 MiB of copies, not more.
MAX_DOCX_COPIED_BYTES = 128 * 1024 * 1024

#: Bytes a copied element, comment, processing instruction or
#: namespace declaration counts in ``MAX_DOCX_COPIED_BYTES`` besides
#: its text, attribute values, prefix and URI: an lxml node (~130 bytes
#: per copied empty element, ~110 per copied declaration measured).
DOCX_COPIED_BYTES_PER_NODE = 128

#: Bytes an attribute counts in ``MAX_DOCX_COPIED_BYTES`` besides its
#: value: libxml2 copies it as an attribute node and a text node
#: (~275 bytes per empty attribute per copy measured, the same with or
#: without a namespace).
DOCX_COPIED_BYTES_PER_ATTRIBUTE = 288

#: Bytes a byte of table-cell text counts in ``MAX_DOCX_COPIED_BYTES``
#: per extra copy of the cell (spanned column or continued merged
#: row). unstructured's ``_convert_table_to_html`` keeps every copy as
#: a string, then HTML-escaped in its row's string and again in the
#: table's, all at once: a ``"`` (one byte of XML text) becomes
#: ``&quot;`` (six characters) and one astral character in a string
#: makes it four bytes per character, so a byte costs up to 4 + 2 x 24
#: = 52 bytes (measured ~49 per extra copy; ~12 for plain ASCII text).
DOCX_TABLE_COPY_BYTES_PER_TEXT_BYTE = 64

#: Children of python-docx's default styles part, which it creates when
#: the document part has no styles relationship; the style lookup
#: (``MAX_DOCX_STYLE_WORK``) scans at least that many entries.
DOCX_DEFAULT_STYLES_ENTRIES = 6

_RELS_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_R_ID = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
)
_RT_OFFICE_DOCUMENT = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    "officeDocument"
)
_P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_W_HDRFTR_REFS = (f"{_W_NS}headerReference", f"{_W_NS}footerReference")


def _skeleton_xml(archive: zipfile.ZipFile, name: str):
    """Parse member *name* (held to ``MAX_OPC_SKELETON_PART_BYTES``,
    and read through ``_read_counted_xml``: UTF-8, start tags of at
    most ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` attributes, no DTD) without
    entity resolution or network access; ``None`` if absent."""
    from lxml import etree

    try:
        info = archive.getinfo(name)
    except KeyError:
        return None
    if info.file_size > MAX_OPC_SKELETON_PART_BYTES:
        raise DecompressionBombError(
            "document package part is larger than "
            f"{MAX_OPC_SKELETON_PART_BYTES} bytes"
        )
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    root = etree.fromstring(_read_counted_xml(archive, info), parser=parser)
    _refuse_dtd(root)
    return root


def _part_rels(
    archive: zipfile.ZipFile,
    packuri_cls,
    partname: str,
    *,
    skip_missing_targets: bool,
) -> dict[str, tuple[str, Optional[str]]]:
    """``{rId: (reltype, target partname or None if external)}`` for
    *partname*, built the way the loading library builds it: later
    duplicates of an rId win, and (python-pptx only,
    *skip_missing_targets*) internal relationships whose target is not
    in the archive are dropped before that. A relationships part with
    more than ``MAX_PART_RELATIONSHIPS`` children is refused."""
    members = set(archive.namelist())
    root = _skeleton_xml(archive, packuri_cls(partname).rels_uri.membername)
    rels: dict[str, tuple[str, Optional[str]]] = {}
    if root is None:
        return rels
    if len(root) > MAX_PART_RELATIONSHIPS:
        raise DecompressionBombError(
            f"package relationships part has {len(root)} entries "
            f"(ceiling {MAX_PART_RELATIONSHIPS})"
        )
    base_uri = packuri_cls(partname).baseURI
    for rel in root.findall(f"{_RELS_NS}Relationship"):
        if rel.get("TargetMode", "Internal") == "External":
            target = None
        else:
            target = packuri_cls.from_rel_ref(base_uri, rel.get("Target"))
            if skip_missing_targets and target[1:] not in members:
                continue
        rels[rel.get("Id")] = (rel.get("Type"), target)
    return rels


def _check_package_relationships(
    archive: zipfile.ZipFile, *, price_walk: bool
) -> None:
    """Refuse a ``.docx``/``.pptx`` package any of whose relationships
    parts (``*.rels`` members, whichever part they belong to, reached
    or not) is over ``MAX_OPC_SKELETON_PART_BYTES``, declares a DTD or
    has more than ``MAX_PART_RELATIONSHIPS`` entries, whose
    relationships parts have more than ``MAX_PACKAGE_RELATIONSHIPS``
    entries together, or (*price_walk*, python-docx) whose internal
    relationships times its members cost python-docx's relationship
    walks more than ``MAX_DOCX_RELATIONSHIP_WALK``. Each part is
    streamed and each entry dropped once read, so at most the root and
    one entry are held: every element child of the root counts as an
    entry, internal unless its ``TargetMode`` is ``External``, as the
    libraries read it, and every comment and processing instruction
    counts as an entry too. A part with an element below an entry, or
    an element with more than ``MAX_RELS_ELEMENT_ATTRIBUTES``
    attributes, is refused when that element starts: the OPC schema
    gives ``Relationship`` four attributes and text content only, and
    without this one entry holding millions of children (16 MiB, about
    16 KB deflated, in a part no library reads) was built in full,
    costing the guard 3.1 s and ~550 MB per part. Each part is read
    through ``_CountedXmlStream``, so lxml never receives a start tag of
    more than ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` attributes: ``iterparse``
    builds one in full before its start event (1,000,000 attributes
    cost this check about 200 MB before)."""
    from lxml import etree

    total = internal = 0
    for info in archive.infolist():
        if not info.filename.lower().endswith(".rels"):
            continue
        if info.file_size > MAX_OPC_SKELETON_PART_BYTES:
            raise DecompressionBombError(
                "document package part is larger than "
                f"{MAX_OPC_SKELETON_PART_BYTES} bytes"
            )
        entries = depth = 0
        with archive.open(info) as stream:
            for event, element in etree.iterparse(
                _CountedXmlStream(stream),
                events=("start", "end", "comment", "pi"),
                resolve_entities=False,
                no_network=True,
            ):
                if event == "end":
                    depth -= 1
                    if depth == 1:
                        # Keep no tree: drop each entry once read.
                        element.clear()
                        parent = element.getparent()
                        while element.getprevious() is not None:
                            del parent[0]
                    continue
                if event == "start":
                    depth += 1
                    if depth > 2:
                        raise DecompressionBombError(
                            "package relationships entry has child elements"
                        )
                    if len(element.attrib) > MAX_RELS_ELEMENT_ATTRIBUTES:
                        raise DecompressionBombError(
                            "package relationships element has more than "
                            f"{MAX_RELS_ELEMENT_ATTRIBUTES} attributes"
                        )
                    if depth == 1:
                        _refuse_dtd(element)
                        continue
                    if element.get("TargetMode") != "External":
                        internal += 1
                # An entry, or a comment or processing instruction
                # anywhere in the part (each stays in the tree until
                # the entry it is in, or the next one, ends).
                entries += 1
                if entries > MAX_PART_RELATIONSHIPS:
                    raise DecompressionBombError(
                        "package relationships part has more than "
                        f"{MAX_PART_RELATIONSHIPS} entries"
                    )
        total += entries
        if total > MAX_PACKAGE_RELATIONSHIPS:
            raise DecompressionBombError(
                "package relationships parts have more than "
                f"{MAX_PACKAGE_RELATIONSHIPS} entries together"
            )
    if not price_walk or not internal:
        return
    names = archive.namelist()
    steps = (
        len(names)
        + sum(len(name) * _str_width(name) for name in names)
        // DOCX_RELATIONSHIP_WALK_NAME_BYTES_PER_STEP
    )
    if internal * steps > MAX_DOCX_RELATIONSHIP_WALK:
        raise DecompressionBombError(
            f"package's {internal} internal relationships over {len(names)} "
            f"members cost {internal * steps} steps in python-docx's "
            f"relationship walk (ceiling {MAX_DOCX_RELATIONSHIP_WALK})"
        )


def _str_width(text: str) -> int:
    """Bytes per character CPython stores *text* with (PEP 393): 1 when
    every code point is below 256, 2 below 65,536, else 4. Comparing
    two equal strings of equal length reads ``len * width`` bytes."""
    if text.isascii():
        return 1
    widest = max(text)
    if widest <= "\xff":
        return 1
    return 2 if widest <= "\uffff" else 4


def _main_part(
    archive: zipfile.ZipFile, packuri_cls, *, skip_missing_targets: bool
) -> Optional[str]:
    """The package's main document partname, resolved as the library
    resolves it, or ``None`` when the library cannot load the package
    at all: python-pptx and python-docx both require exactly one
    internal officeDocument relationship (``part_with_reltype`` raises
    otherwise) whose target is in the archive. Such a package is
    refused: the library cannot extract text from it, and a package
    without its main part is what a content sniffer may take for
    another type (see ``_check_package_type``)."""
    rels = _part_rels(
        archive, packuri_cls, "/", skip_missing_targets=skip_missing_targets
    )
    mains = [
        target
        for reltype, target in rels.values()
        if reltype == _RT_OFFICE_DOCUMENT
    ]
    if len(mains) != 1 or mains[0] is None:
        return None
    if mains[0][1:] not in set(archive.namelist()):
        return None
    return mains[0]


_CT_NS = "{http://schemas.openxmlformats.org/package/2006/content-types}"
_CT_PREFIX = "application/vnd.openxmlformats-officedocument."
_CT_PML_SLIDE = frozenset({_CT_PREFIX + "presentationml.slide+xml"})
_CT_PML_SLIDE_LAYOUT = frozenset(
    {_CT_PREFIX + "presentationml.slideLayout+xml"}
)
_CT_PML_SLIDE_MASTER = frozenset(
    {_CT_PREFIX + "presentationml.slideMaster+xml"}
)
#: python-pptx's ``Presentation()`` accepts these main parts only.
_CT_PML_MAINS = frozenset(
    {
        _CT_PREFIX + "presentationml.presentation.main+xml",
        "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
    }
)
#: python-docx's ``Document()`` accepts this main part only.
_CT_WML_MAINS = frozenset({_CT_PREFIX + "wordprocessingml.document.main+xml"})
_CT_WML_STYLES = frozenset({_CT_PREFIX + "wordprocessingml.styles+xml"})
_CT_WML_SETTINGS = frozenset({_CT_PREFIX + "wordprocessingml.settings+xml"})
_CT_WML_HEADER = frozenset({_CT_PREFIX + "wordprocessingml.header+xml"})
_CT_WML_FOOTER = frozenset({_CT_PREFIX + "wordprocessingml.footer+xml"})


class _ContentTypes(NamedTuple):
    """``[Content_Types].xml`` as python-pptx and python-docx read it:
    ``Override`` content types by lower-cased part name and ``Default``
    ones by lower-cased extension, a later entry replacing an earlier
    one with the same key."""

    overrides: dict[str, Optional[str]]
    defaults: dict[str, Optional[str]]


def _content_types(archive: zipfile.ZipFile) -> _ContentTypes:
    """Read the package's content types (see ``_ContentTypes``; the
    part is held to ``MAX_OPC_SKELETON_PART_BYTES`` and refused with a
    DTD). A package without the part, or with an entry missing its
    part name or extension, is refused: both libraries fail to open
    it."""
    root = _skeleton_xml(archive, "[Content_Types].xml")
    if root is None:
        raise DecompressionBombError("document package has no content types")
    overrides: dict[str, Optional[str]] = {}
    defaults: dict[str, Optional[str]] = {}
    for kind, key, table in (
        ("Override", "PartName", overrides),
        ("Default", "Extension", defaults),
    ):
        for entry in root.findall(f"{_CT_NS}{kind}"):
            name = entry.get(key)
            if name is None:
                raise DecompressionBombError(
                    "document package has a malformed content type entry"
                )
            table[name.lower()] = entry.get("ContentType")
    return _ContentTypes(overrides, defaults)


def _content_type(types: _ContentTypes, partname: str) -> Optional[str]:
    """The content type the libraries give *partname* (an absolute
    part name): its ``Override``, else the ``Default`` of its extension
    (``PackURI.ext``), else ``None``."""
    key = partname.lower()
    if key in types.overrides:
        return types.overrides[key]
    ext = posixpath.splitext(partname)[1]
    ext = ext[1:] if ext.startswith(".") else ext
    return types.defaults.get(ext.lower())


def _require_content_type(
    types: _ContentTypes, partname: str, allowed: frozenset[str], what: str
) -> None:
    """Refuse the package unless *partname*'s content type is in
    *allowed*.

    python-pptx and python-docx pick a part's class from its content
    type, not from the relationship that reaches it, and several
    classes answer the same property by following their own
    relationship of the same type: a slide whose layout relationship
    targets another slide reaches that slide's layout (``SlidePart``
    has ``slide_layout`` too), and a document part whose styles or
    settings relationship targets another document part reaches that
    part's styles or settings. The guard reads the part a relationship
    targets as the kind the relationship names, so such a chain would
    let python-pptx or python-docx read parts the guard never priced
    (a slide over a chained 300-placeholder layout was counted at 1,200
    units, the same layout reached directly at about 90,000). A part
    the library would build as another class either fails there or is
    such a chain, so nothing a document legitimately holds is lost."""
    content_type = _content_type(types, partname)
    if content_type not in allowed:
        raise DecompressionBombError(
            f"document package's {what} part has the wrong content type"
        )


#: Ceiling on python-pptx's placeholder inheritance as unstructured
#: drives it: the sum over slides of the slide's placeholders times the
#: cost of one inherited read (``_pptx_inherited_read_units``).
#: unstructured reads ``top`` and ``left`` of every shape on a slide up
#: to three times each (to sort the shapes, and twice more to skip
#: off-slide ones); a slide placeholder with no position of its own
#: inherits it, and each such read rebuilds the slide part's
#: relationships by type, finds the slide layout's placeholder of the
#: same ``idx`` by building a proxy for each placeholder of the layout's
#: shape tree in turn (two to three XPath evaluations each, every one
#: walking the shape's children, its first child's and that child's
#: ``p:nvPr``'s), and, when that one has no position either, does the
#: same over the slide master's shape tree. Nothing is cached, so a
#: slide of P placeholders over a layout of L shapes costs about
#: 6 * P * L proxy builds: before the change, 100 placeholders over a
#: 100-placeholder layout, a 30 KB file, took 5.2 s and 200 over 200,
#: 18.6 s, quadrupling per doubling. A unit measured 400-430 us end to
#: end in ``partition_pptx`` (the 200-over-200 slide is 46,000 units,
#: 0.51 of the ceiling), so this fan-out at the ceiling is about 40 s.
#: Office's default theme, as python-pptx and pandoc ship it, has 5-10
#: shapes per layout and 7 on the master, and its title-and-content
#: slide two placeholders without a position: 34 units per slide, so
#: about 2,600 such slides fit; templates with tens of shapes per
#: layout and master fit proportionally fewer (see "Known gaps").
MAX_PPTX_PLACEHOLDER_WORK = 90_000

#: Nodes below a layout's or master's shapes (``_PptxShapeTree.nodes``)
#: that count as one unit of ``MAX_PPTX_PLACEHOLDER_WORK``: each of the
#: six reads walks them at ~50 ns per node (a unit is ~420 us), so one
#: counted twice for the text before it is priced at about ten times
#: that.
PPTX_PLACEHOLDER_NODES_PER_UNIT = 256

#: Relationships of the slide and layout parts that count as one unit of
#: ``MAX_PPTX_PLACEHOLDER_WORK``: python-pptx rebuilds the part's
#: relationships by type on every read (100 placeholders on a slide of
#: 10,000 relationships took 3.7 s more, ~0.6 us per relationship and
#: read), so one is priced at about 3.5 times that.
PPTX_PLACEHOLDER_RELATIONSHIPS_PER_UNIT = 32

_RT_SLIDE_LAYOUT = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    "slideLayout"
)
_RT_SLIDE_MASTER = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    "slideMaster"
)
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_P_CSLD = f"{_P_NS}cSld"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_P_SPTREE = f"{_P_NS}spTree"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_P_NVPR = f"{_P_NS}nvPr"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_P_PH = f"{_P_NS}ph"


class _PptxShapeTree(NamedTuple):
    """What ``_pptx_shape_tree`` counts in one slide-like part."""

    #: Children of ``p:cSld/p:spTree`` (shapes and anything else but
    #: text).
    shapes: int
    #: Their children, grandchildren and great-grandchildren (elements,
    #: comments and processing instructions; python-pptx's
    #: ``./*[1]/p:nvPr/p:ph`` walks some of them), each counted twice
    #: for the text node that may precede it.
    nodes: int
    #: Children holding a ``p:ph`` in a ``p:nvPr`` of one of their own
    #: children: python-pptx's placeholders (it looks in the first
    #: child only, so this can only overcount).
    placeholders: int


def _check_pptx_slide_fan_out(content: bytes) -> None:
    """Refuse a presentation listing the same slide part more than once,
    or whose placeholder inheritance costs more than
    ``MAX_PPTX_PLACEHOLDER_WORK``.

    python-pptx yields one slide per ``<p:sldId>`` (first ``p:sldIdLst``
    of the presentation part), resolved through that part's
    relationships, so N entries resolving to one slide part make
    unstructured extract that slide N times. Genuine presentations list
    each slide part once. A ``<p:sldId>`` without a resolvable internal
    relationship is refused too: python-pptx fails on it, so nothing
    indexable is lost, and the check never has to guess. Each slide, and
    each slide layout a slide with placeholders relates to and each
    slide master such a layout relates to, is then streamed once (see
    ``_pptx_shape_tree``). Every relationships part of the package is
    held to ``MAX_PART_RELATIONSHIPS`` entries, and all of them together
    to ``MAX_PACKAGE_RELATIONSHIPS`` (``_check_package_relationships``).
    """
    try:
        import pptx  # noqa: F401
    except ImportError:  # no python-pptx, no .pptx loader
        return
    try:
        # Inside the fail-closed block, as for .xlsx: if an upgrade
        # moves the helper, uploads are refused, not left unchecked.
        from pptx.opc.packuri import PackURI

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            _check_package_relationships(archive, price_walk=False)
            types = _content_types(archive)
            targets = _pptx_slide_targets(archive, PackURI, types)
            repeated = len(targets) != len(set(targets))
            work = (
                0
                if repeated
                else _pptx_placeholder_work(
                    archive, PackURI, targets, MAX_PPTX_PLACEHOLDER_WORK, types
                )
            )
    except DecompressionBombError:
        raise
    except Exception as exc:  # lxml/zipfile failure modes are open-ended
        raise DecompressionBombError("malformed presentation package") from exc
    if repeated:
        raise DecompressionBombError(
            "presentation lists the same slide part more than once"
        )
    if work > MAX_PPTX_PLACEHOLDER_WORK:
        raise DecompressionBombError(
            "presentation's placeholders cost more than "
            f"{MAX_PPTX_PLACEHOLDER_WORK} units of python-pptx's layout and "
            "master lookups"
        )


def _pptx_placeholder_work(
    archive: zipfile.ZipFile,
    packuri_cls,
    slides: list[str],
    limit: int,
    content_types: Optional[_ContentTypes] = None,
) -> int:
    """The sum over *slides* of each slide's placeholders times the
    units of one inherited read (see ``MAX_PPTX_PLACEHOLDER_WORK``),
    returned as soon as it exceeds *limit* (so that no more parts are
    streamed than needed to refuse). Every placeholder is counted as
    inheriting (whether or not it has a position of its own), every
    slide layout a slide relates to (python-pptx fails on more than
    one) and every master such a layout relates to are counted at the
    costliest, and each part is streamed at most once. Each such layout
    and master must have the content type python-pptx builds it from
    (see ``_require_content_type``)."""
    if content_types is None:
        content_types = _content_types(archive)
    types = content_types
    members = set(archive.namelist())
    trees: dict[str, _PptxShapeTree] = {}
    rels: dict[str, dict[str, tuple[str, Optional[str]]]] = {}

    def tree(partname: str) -> _PptxShapeTree:
        if partname not in trees:
            trees[partname] = _pptx_shape_tree(archive, partname[1:])
        return trees[partname]

    def related(
        partname: str, reltype: str, allowed: frozenset[str], what: str
    ) -> list[str]:
        if partname not in rels:
            rels[partname] = _part_rels(
                archive, packuri_cls, partname, skip_missing_targets=True
            )
        targets = [
            target
            for kind, target in rels[partname].values()
            if kind == reltype and target is not None and target[1:] in members
        ]
        for target in targets:
            _require_content_type(types, target, allowed, what)
        return targets

    work = 0
    for slide in slides:
        placeholders = tree(slide).placeholders
        if not placeholders:
            continue
        read = 0
        for layout in related(
            slide, _RT_SLIDE_LAYOUT, _CT_PML_SLIDE_LAYOUT, "slide layout"
        ):
            masters = related(
                layout, _RT_SLIDE_MASTER, _CT_PML_SLIDE_MASTER, "slide master"
            )
            relationships = len(rels[slide]) + len(rels[layout])
            candidates: list[Optional[str]] = list(masters) or [None]
            for master in candidates:
                read = max(
                    read,
                    _pptx_inherited_read_units(
                        tree(layout),
                        _PptxShapeTree(0, 0, 0)
                        if master is None
                        else tree(master),
                        relationships,
                    ),
                )
        work += placeholders * read
        if work > limit:
            break
    return work


def _pptx_inherited_read_units(
    layout: _PptxShapeTree, master: _PptxShapeTree, relationships: int
) -> int:
    """Units of ``MAX_PPTX_PLACEHOLDER_WORK`` for one slide placeholder:
    the shapes of its layout and master, their nodes per
    ``PPTX_PLACEHOLDER_NODES_PER_UNIT`` and the slide's and layout's
    *relationships* per ``PPTX_PLACEHOLDER_RELATIONSHIPS_PER_UNIT``
    (each rounded up)."""
    return (
        layout.shapes
        + master.shapes
        - (-(layout.nodes + master.nodes) // PPTX_PLACEHOLDER_NODES_PER_UNIT)
        - (-relationships // PPTX_PLACEHOLDER_RELATIONSHIPS_PER_UNIT)
    )


def _pptx_shape_tree(archive: zipfile.ZipFile, name: str) -> _PptxShapeTree:
    """Count member *name*'s shape tree (see ``_PptxShapeTree``) in one
    streaming pass, with no tree kept. A DTD, a namespace URI or
    prefix over the caps ``document.xml`` is held to, or more than
    ``MAX_XML_OUTSIDE_ROOT_NODES`` comments and processing instructions
    outside the root, is refused. An
    element's name is read only down to the depth of ``p:ph``, so the
    pass is linear in the part's size."""
    from lxml import etree

    shapes = nodes = placeholders = outside = 0
    # Names of the open elements down to p:ph's depth (7), and the depth.
    path: list[str] = []
    depth = 0
    in_tree = has_ph = False
    with archive.open(name) as stream:
        for event, element in etree.iterparse(
            _CountedXmlStream(stream),
            events=("start-ns", "start", "end", "comment", "pi"),
            resolve_entities=False,
            no_network=True,
        ):
            if event == "start-ns":
                prefix, uri = element
                if (
                    len(uri) > MAX_DOCX_NAMESPACE_URI_CHARS
                    or len(prefix) > MAX_DOCX_NAMESPACE_PREFIX_CHARS
                ):
                    raise DecompressionBombError(
                        "presentation part declares a namespace URI of more "
                        f"than {MAX_DOCX_NAMESPACE_URI_CHARS} or a prefix of "
                        f"more than {MAX_DOCX_NAMESPACE_PREFIX_CHARS} "
                        "characters"
                    )
                continue
            if event == "end":
                if depth <= 7:
                    path.pop()
                    if depth == 3:
                        in_tree = False
                depth -= 1
                element.clear(keep_tail=False)
            elif event == "start":
                depth += 1
                if depth == 1:
                    _refuse_dtd(element)
                if depth <= 7:
                    tag = element.tag
                    path.append(tag)
                    if depth == 3:
                        in_tree = path[1] == _P_CSLD and tag == _P_SPTREE
                    elif in_tree and depth == 4:
                        shapes += 1
                        has_ph = False
                    elif in_tree:
                        nodes += 2
                        if (
                            depth == 7
                            and tag == _P_PH
                            and path[5] == _P_NVPR
                            and not has_ph
                        ):
                            has_ph = True
                            placeholders += 1
            elif depth == 0:
                outside = _note_outside_root_node(outside)
            elif in_tree and depth == 3:
                # A comment or processing instruction among the shapes,
                # which python-pptx iterates over too.
                shapes += 1
            elif in_tree and 5 <= depth + 1 <= 7:
                # A comment or processing instruction among the nodes.
                nodes += 2
            if event != "start":
                # Drop what has been read (comments and processing
                # instructions get no end event).
                parent = element.getparent()
                if parent is not None:
                    while element.getprevious() is not None:
                        del parent[0]
    return _PptxShapeTree(shapes, nodes, placeholders)


def _pptx_slide_targets(
    archive: zipfile.ZipFile,
    packuri_cls,
    content_types: Optional[_ContentTypes] = None,
) -> list[str]:
    """The slide partname of every ``<p:sldId>``, in order. The main
    part and each slide must have the content type python-pptx builds
    them from (see ``_require_content_type``)."""
    main = _main_part(archive, packuri_cls, skip_missing_targets=True)
    if main is None:
        raise DecompressionBombError(
            "presentation package has no main document part"
        )
    if content_types is None:
        content_types = _content_types(archive)
    _require_content_type(content_types, main, _CT_PML_MAINS, "main")
    root = _skeleton_xml(archive, main[1:])
    if root is None:  # pragma: no cover - _main_part checked membership
        return []
    rels = _part_rels(archive, packuri_cls, main, skip_missing_targets=True)
    sld_id_lst = root.find(f"{_P_NS}sldIdLst")
    if sld_id_lst is None:
        return []
    targets = []
    for sld_id in sld_id_lst.findall(f"{_P_NS}sldId"):
        resolved = rels.get(sld_id.get(_R_ID))
        if resolved is None or resolved[1] is None:
            raise DecompressionBombError(
                "presentation lists a slide it cannot resolve"
            )
        _require_content_type(
            content_types, resolved[1], _CT_PML_SLIDE, "slide"
        )
        targets.append(resolved[1])
    return targets


def _check_docx_section_fan_out(content: bytes) -> None:
    """Bound python-docx's and unstructured's per-section,
    per-paragraph, per-page-break and per-spanned-column work and
    header/footer repeats (the known fan-outs; not the document's total
    parse cost, see "Known gaps" in the module docstring).

    Refused: more than ``MAX_DOCX_SECTIONS`` ``w:sectPr`` elements or
    ``MAX_DOCX_PAGE_BREAKS`` ``w:lastRenderedPageBreak`` elements
    anywhere in ``document.xml``; a relationships part with
    more than ``MAX_PART_RELATIONSHIPS`` entries, relationships parts
    with more than ``MAX_PACKAGE_RELATIONSHIPS`` together, or internal
    relationships times members over ``MAX_DOCX_RELATIONSHIP_WALK``
    (``_check_package_relationships``); a settings or styles
    part over ``MAX_OPC_SKELETON_PART_BYTES``, with a DTD, or with more
    than ``MAX_DOCX_SETTINGS_CHILDREN`` or ``MAX_DOCX_STYLES`` children
    of its root; a ``w:gridSpan``, ``w:gridBefore`` or ``w:gridAfter``
    over ``MAX_DOCX_GRID_SPAN`` or that python-docx cannot read; an
    element with more than ``MAX_DOCX_ELEMENT_ATTRIBUTES`` attributes
    or more than ``MAX_DOCX_NAMESPACES_IN_SCOPE`` namespace declarations
    in scope; a body-level paragraph with more than
    ``MAX_DOCX_PARAGRAPH_NAMESPACES`` declarations; a declaration whose
    URI or prefix is longer than ``MAX_DOCX_NAMESPACE_URI_CHARS`` or
    ``MAX_DOCX_NAMESPACE_PREFIX_CHARS``; a table row of more than
    ``MAX_DOCX_ROW_CELLS`` cells or ``MAX_DOCX_ROW_GRID_COLUMNS``
    columns; the
    sum over sections of the body-level blocks up to
    each section's end, times the body's child nodes, over
    ``MAX_DOCX_SECTION_BLOCK_WORK``; the section count times the linear
    per-section work (the body's child nodes, plus the nodes the
    ``_sectPrs`` XPath visits divided by
    ``DOCX_SECTION_XPATH_NODES_PER_UNIT``, plus the node pairs its union
    compares divided by ``DOCX_SECTION_UNION_PAIRS_PER_UNIT``) over
    ``MAX_DOCX_SECTION_WORK``; the summed size of the parts the
    sections' ``headerReference``/``footerReference`` children point at
    (counted once per reference, with the union steps of each read)
    over ``MAX_DOCX_HEADER_FOOTER_BYTES``;
    the style lookups over ``MAX_DOCX_STYLE_WORK``; the page-break
    splitting over ``MAX_DOCX_PAGE_BREAK_WORK``; the text extraction
    from runs and table cells (once per spanned column) over
    ``MAX_DOCX_RUN_WORK`` (with each cell's paragraphs and tables and
    the node-set unions of cells, paragraphs and runs, see
    ``DOCX_UNION_STEPS_PER_RUN_WORK_UNIT``, and the merges and sort of
    unstructured's page-break union, see
    ``DOCX_PAGE_BREAK_UNION_PAIRS_PER_UNIT`` and
    ``DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT``); the bytes copied per page
    break, spanned column and continued merged cell over
    ``MAX_DOCX_COPIED_BYTES``;
    and, since each ceiling is the work one kind alone may cost, any
    document whose seven works' shares of their ceilings sum to more
    than 1.
    Sections are the ``w:sectPr`` elements python-docx enumerates
    (``w:body/w:p/w:pPr/w:sectPr`` and ``w:body/w:sectPr``). A
    reference that does not resolve to a part in the archive is
    refused, as python-docx fails on it. ``document.xml`` is streamed
    with no tree kept, so this costs one parse pass over it.

    The section, relationship, settings and styles caps bound the
    per-section and per-paragraph work python-docx and unstructured do
    that is not counted (see ``MAX_DOCX_SECTIONS``); it is bounded by
    those caps, not modelled.
    """
    try:
        import docx  # noqa: F401
    except ImportError:  # no python-docx, no .docx loader
        return
    try:
        from docx.opc.packuri import PackURI

        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            _check_package_relationships(archive, price_walk=True)
            stats = _docx_section_stats(archive, PackURI)
    except DecompressionBombError:
        raise
    except Exception as exc:  # lxml/zipfile failure modes are open-ended
        raise DecompressionBombError(
            "malformed word-processing package"
        ) from exc

    if stats.block_work > MAX_DOCX_SECTION_BLOCK_WORK:
        raise DecompressionBombError(
            f"document's {stats.sections} sections cost {stats.block_work} "
            "block-by-node steps in python-docx's section lookup (ceiling "
            f"{MAX_DOCX_SECTION_BLOCK_WORK})"
        )
    sections = stats.sections
    work = (
        sections * stats.body_nodes
        + sections * stats.xpath_nodes // DOCX_SECTION_XPATH_NODES_PER_UNIT
        + sections * stats.union_pairs // DOCX_SECTION_UNION_PAIRS_PER_UNIT
    )
    if work > MAX_DOCX_SECTION_WORK:
        raise DecompressionBombError(
            f"document has {sections} sections over {stats.body_nodes} "
            f"body nodes, {stats.xpath_nodes} section-lookup nodes and "
            f"{stats.union_pairs} section-lookup pairs (work {work}, ceiling "
            f"{MAX_DOCX_SECTION_WORK})"
        )
    if stats.referenced_bytes > MAX_DOCX_HEADER_FOOTER_BYTES:
        raise DecompressionBombError(
            "document's header/footer references total over "
            f"{MAX_DOCX_HEADER_FOOTER_BYTES} bytes"
        )
    style_work = stats.styled_paragraphs * (
        stats.style_entries
        + stats.style_attributes // DOCX_STYLE_ATTRIBUTES_PER_ENTRY
        + stats.relationships // DOCX_RELATIONSHIPS_PER_STYLE_ENTRY
    )
    if style_work > MAX_DOCX_STYLE_WORK:
        raise DecompressionBombError(
            f"document's {stats.styled_paragraphs} paragraphs cost "
            f"{style_work} steps in python-docx's style lookup over "
            f"{stats.style_entries} styles part entries and "
            f"{stats.relationships} relationships (ceiling "
            f"{MAX_DOCX_STYLE_WORK})"
        )
    if stats.page_break_work > MAX_DOCX_PAGE_BREAK_WORK:
        raise DecompressionBombError(
            "document's rendered page breaks cost "
            f"{stats.page_break_work} steps in python-docx's paragraph "
            f"splitting (ceiling {MAX_DOCX_PAGE_BREAK_WORK})"
        )
    if stats.run_work > MAX_DOCX_RUN_WORK:
        raise DecompressionBombError(
            f"document's runs and table cells cost {stats.run_work} units "
            f"of python-docx text extraction (ceiling {MAX_DOCX_RUN_WORK})"
        )
    if stats.copied_bytes > MAX_DOCX_COPIED_BYTES:
        raise DecompressionBombError(
            f"document's page breaks, spanned and merged cells copy "
            f"{stats.copied_bytes} bytes in python-docx and unstructured "
            f"(ceiling {MAX_DOCX_COPIED_BYTES})"
        )
    # Each ceiling above is the work one kind alone may cost; together
    # they share one budget, so a document near several ceilings at
    # once costs no more than one at a single ceiling.
    shares = (
        stats.block_work / MAX_DOCX_SECTION_BLOCK_WORK,
        work / MAX_DOCX_SECTION_WORK,
        stats.referenced_bytes / MAX_DOCX_HEADER_FOOTER_BYTES,
        style_work / MAX_DOCX_STYLE_WORK,
        stats.page_break_work / MAX_DOCX_PAGE_BREAK_WORK,
        stats.run_work / MAX_DOCX_RUN_WORK,
        stats.copied_bytes / MAX_DOCX_COPIED_BYTES,
    )
    if sum(shares) > 1:
        raise DecompressionBombError(
            "document's python-docx work takes "
            f"{sum(shares):.2f} of the combined budget (section lookup, "
            "section work, header/footer bytes, style lookup, page-break "
            "splitting, text extraction and copied bytes, each as a share "
            "of its ceiling)"
        )


# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_BODY = f"{_W_NS}body"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_P = f"{_W_NS}p"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_PPR = f"{_W_NS}pPr"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_SECTPR = f"{_W_NS}sectPr"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_TBL = f"{_W_NS}tbl"
_W_BLOCKS = (_W_P, _W_TBL)
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_R = f"{_W_NS}r"
_W_RUN_CONTAINERS = (_W_R, f"{_W_NS}hyperlink")
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_TR = f"{_W_NS}tr"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_TC = f"{_W_NS}tc"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_TRPR = f"{_W_NS}trPr"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_TCPR = f"{_W_NS}tcPr"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_GRID_SPAN = f"{_W_NS}gridSpan"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_GRID_BEFORE = f"{_W_NS}gridBefore"
_W_GRID_VALUES = (_W_GRID_SPAN, _W_GRID_BEFORE, f"{_W_NS}gridAfter")
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_VMERGE = f"{_W_NS}vMerge"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_VAL = f"{_W_NS}val"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_LAST_RENDERED_PAGE_BREAK = f"{_W_NS}lastRenderedPageBreak"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_HYPERLINK = f"{_W_NS}hyperlink"
# nosemgrep: semgrep.rules.sql-string-concatenation, reason: OOXML qualified name constant; no SQL sink
_W_DRAWING = f"{_W_NS}drawing"
_WP_INLINE = (
    "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    "inline"
)
#: The paths of unstructured's page-break union branches below the
#: root, by depth of the ``w:lastRenderedPageBreak`` (see
#: ``MAX_DOCX_PAGE_BREAKS``).
_W_PAGE_BREAK_BRANCHES = {
    5: (_W_BODY, _W_P, _W_R),
    6: (_W_BODY, _W_P, _W_HYPERLINK, _W_R),
    8: (_W_BODY, _W_TBL, _W_TR, _W_TC, _W_P, _W_R),
    9: (_W_BODY, _W_TBL, _W_TR, _W_TC, _W_P, _W_HYPERLINK, _W_R),
}
#: The branches of ``CT_R.text``'s union, one bit each.
_W_RUN_TEXT_BRANCHES = {
    f"{_W_NS}{name}": 1 << bit
    for bit, name in enumerate(
        ("br", "cr", "noBreakHyphen", "ptab", "t", "tab")
    )
}
#: Elements whose children python-docx selects with a union.
_W_UNION_PARENTS = frozenset({_W_TC, _W_P, _W_R})
_RT_SETTINGS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    "settings"
)
_RT_STYLES = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"
)


class _DocxSectionStats(NamedTuple):
    """What ``_docx_section_stats`` counts in one streaming pass."""

    #: ``w:sectPr`` elements anywhere in ``document.xml``.
    sect_prs: int
    #: The sections python-docx enumerates (N1 + N2 below).
    sections: int
    #: Child nodes of the body: elements, comments, processing
    #: instructions and (non-empty) text.
    body_nodes: int
    #: Node visits of one ``_sectPrs`` XPath evaluation.
    xpath_nodes: int
    #: N1 * N2: paragraph-level times body-level section ``w:sectPr``.
    union_pairs: int
    #: Sum over sections of the body-level ``w:p``/``w:tbl`` up to the
    #: section's end, times ``body_nodes``.
    block_work: int
    #: Header/footer bytes, counted once per reference, plus the union
    #: steps of each read (``DOCX_UNION_STEPS_PER_HEADER_BYTE``).
    referenced_bytes: int
    #: Body-level paragraphs with a ``w:r`` or ``w:hyperlink`` child
    #: (the only ones whose text can be non-empty), plus the
    #: ``w:lastRenderedPageBreak`` elements inside them: an upper bound
    #: on the paragraph fragments whose style unstructured looks up.
    styled_paragraphs: int
    #: Sum over body-level paragraphs of K * (K + 3) * S**2, K being the
    #: paragraph's ``w:lastRenderedPageBreak`` elements and S the largest
    #: child-node count of the paragraph or of any element in it.
    page_break_work: int
    #: Sum over body-level paragraphs and tables of (K + 1) *
    #: (``DOCX_RUN_WORK_ITEM_WEIGHT`` * (``w:r``, ``w:tr`` and ``w:tc``
    #: elements in it) + child nodes of its ``w:r`` elements) + K * (its
    #: descendant nodes) // ``DOCX_FRAGMENT_NODES_PER_RUN_WORK_UNIT``
    #: (K = 0 for a table). Within a table, a cell's work (with one
    #: unit per ``DOCX_CELL_NODES_PER_UNIT`` of its descendant nodes) is
    #: counted once per grid column it spans (through nested tables), a
    #: cell continuing a vertical merge instead counts the largest cell
    #: of the earlier rows, and a row counts its
    #: ``gridBefore``/``gridAfter`` columns
    #: (``DOCX_GRID_PADDING_PER_RUN_WORK_UNIT`` per unit) and
    #: python-docx's steps up the merges its cells continue (see
    #: ``_docx_merge_steps``). A cell also counts its ``w:p`` and
    #: ``w:tbl`` children (``DOCX_CELL_BLOCK_ITEM_WEIGHT``), and every
    #: cell, paragraph and run the steps of its union
    #: (``DOCX_UNION_STEPS_PER_RUN_WORK_UNIT``; a body-level paragraph's
    #: with the runs in its runs' inline drawings as a third branch,
    #: and the nodes in those drawings), multiplied with the rest of its
    #: work; without sections, the body's union is added once, and the
    #: pairs merged and the steps sorted by unstructured's page-break
    #: union are added once (``DOCX_PAGE_BREAK_UNION_PAIRS_PER_UNIT``,
    #: ``DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT``).
    run_work: int
    #: Bytes copied beyond one read (see ``MAX_DOCX_COPIED_BYTES``):
    #: per body-level paragraph, 2K times its bytes, its tail and the
    #: ancestors' namespace declarations of the URIs it uses; per
    #: spanned cell, (span - 1) times its text bytes; per continuing
    #: merged cell, the text bytes of the largest cell of the earlier
    #: rows of its table; cell text times
    #: ``DOCX_TABLE_COPY_BYTES_PER_TEXT_BYTE``.
    copied_bytes: int
    #: Children of the styles part's root, at least
    #: ``DOCX_DEFAULT_STYLES_ENTRIES`` (python-docx's default part).
    style_entries: int
    #: Attributes of those children.
    style_attributes: int
    #: Relationships of the document part.
    relationships: int


def _docx_section_stats(archive: zipfile.ZipFile, packuri_cls):
    """Section statistics of the main document part (see
    ``_DocxSectionStats``), from one streaming pass, after checking the
    relationships, settings and styles parts.

    The union pairs are N1 * N2, N1 being the ``w:body/w:p/w:pPr/w:sectPr``
    elements and N2 the ``w:body/w:sectPr`` elements: libxml2 compares
    every node of one union branch's result with every node of the
    other's to drop duplicates.

    The block work is, summed over sections, the body-level
    ``w:p``/``w:tbl`` elements from the start of the body through the
    paragraph holding a paragraph-level ``w:sectPr`` (or before a
    body-level one), which is what python-docx's ``preceding-sibling``
    lookup returns for that section, times the body's child nodes,
    which bound the sibling walk of each comparison sorting them (see
    ``MAX_DOCX_SECTION_BLOCK_WORK``). Blocks and nodes are counted
    across all ``w:body`` elements, which can only overcount. Text
    nodes are counted from the text before each body child and after
    the last one, so ``element.clear`` keeps tails until then.

    The node visits are those one evaluation of python-docx's
    ``/w:document/w:body/w:p/w:pPr/w:sectPr | /w:document/w:body/w:sectPr``
    makes: every element, comment and processing instruction that is a
    child of the document node, of the root element, or of a ``w:body``
    child of it (each counted twice, since both union branches walk
    these levels), or of a body-level ``w:p`` or of that paragraph's
    ``w:pPr`` (counted once). Text nodes are not counted; each sits
    beside a counted node or alone in its parent. Entity reference
    nodes, for which ``iterparse`` emits no event, cannot occur: a
    ``document.xml`` with a DTD is refused in its bytes, before lxml
    reads it (``_CountedXmlStream``).
    Counting does not check the root's tag, which only adds nodes the
    XPath would not visit.

    Per body-level paragraph and table (python-docx's block items; the
    content of ``w:sdt`` and other body-level elements is not read), the
    child nodes of every element in it are counted as python-docx's tree
    has them, text included: elements, comments and processing
    instructions, the text before each and the text after the last.
    From those come the page-break and run work (see
    ``MAX_DOCX_PAGE_BREAK_WORK`` and ``MAX_DOCX_RUN_WORK``) and the
    paragraphs whose style is looked up (``MAX_DOCX_STYLE_WORK``).
    ``w:lastRenderedPageBreak`` elements are counted at any depth of a
    paragraph and ``w:r`` elements at any depth of a block, though
    python-docx reads only some of them, which can only overcount. The
    bytes of each block (``MAX_DOCX_COPIED_BYTES``) are its text, its
    attribute values, its namespace declarations (prefix and URI, from
    the ``start-ns`` events), ``DOCX_COPIED_BYTES_PER_NODE`` per node
    and per declaration and ``DOCX_COPIED_BYTES_PER_ATTRIBUTE`` per
    attribute; the extra copies of a table cell count its text bytes
    times ``DOCX_TABLE_COPY_BYTES_PER_TEXT_BYTE``; a paragraph's copies also count the text after
    it and the declarations on the ``w:document`` and ``w:body``
    elements of the namespace URIs its element and attribute names
    use. A ``w:tc`` anywhere in a block weighs its content by its
    ``w:gridSpan`` (the largest under any of its ``w:tcPr``, where
    python-docx reads the first ``w:gridSpan`` of the first
    ``w:tcPr``: never less), unless it continues a vertical merge as
    python-docx reads it (the first ``w:vMerge`` of its first
    ``w:tcPr``, with ``w:val`` absent or ``continue``): such a cell
    adds the largest earlier cell of its table instead, which the cell
    python-docx actually yields cannot exceed, and is not itself a
    candidate for that largest cell. A row python-docx reads (a direct
    ``w:tr`` child of a table) also counts python-docx's steps up the
    vertical merges its cells continue (``_docx_merge_steps``), from the
    spans and ``gridBefore`` as python-docx reads them.

    Streaming stops with a refusal as soon as ``MAX_DOCX_SECTIONS`` or
    ``MAX_DOCX_PAGE_BREAKS`` is exceeded, at the first grid value over ``MAX_DOCX_GRID_SPAN``, at
    the first element with more than ``MAX_DOCX_ELEMENT_ATTRIBUTES``
    attributes (counted before any is read), at the first with more
    than ``MAX_DOCX_NAMESPACES_IN_SCOPE`` namespace declarations in
    scope, at the first namespace declaration whose URI or prefix is
    over its cap (before any name under it is read), at the first
    paragraph over ``MAX_DOCX_PARAGRAPH_NAMESPACES``, at the first row
    cell over ``MAX_DOCX_ROW_CELLS``, at the first other row child over
    ``MAX_DOCX_ROW_OTHER_CHILDREN``, at the first row after more than
    ``MAX_DOCX_TABLE_ROW_GAP`` other table children and at the end of
    the first row over ``MAX_DOCX_ROW_GRID_COLUMNS``.

    The pass reads each element's name (lxml builds it, namespace URI
    included, on every read) at most twice and each attribute name at
    most once, and each text, attribute value and comment at most
    twice, so its time is linear in the size of ``document.xml``, a
    name costing up to ``MAX_DOCX_NAMESPACE_URI_CHARS`` more per read
    than its own length; the state it keeps is bounded by the depth (at
    most 256, libxml2's limit), the namespace declarations in scope
    (``pending_ns`` is refused as soon as it would exceed them) and the
    open tables' last rows."""
    from lxml import etree

    main = _main_part(archive, packuri_cls, skip_missing_targets=False)
    if main is None:
        raise DecompressionBombError(
            "word-processing package has no main document part"
        )
    types = _content_types(archive)
    _require_content_type(types, main, _CT_WML_MAINS, "main")
    rels = _part_rels(archive, packuri_cls, main, skip_missing_targets=False)
    _docx_skeleton_entries(
        archive,
        rels,
        _RT_SETTINGS,
        MAX_DOCX_SETTINGS_CHILDREN,
        "settings",
        types,
        _CT_WML_SETTINGS,
    )
    styles, style_attributes = _docx_skeleton_entries(
        archive,
        rels,
        _RT_STYLES,
        MAX_DOCX_STYLES,
        "styles",
        types,
        _CT_WML_STYLES,
    )
    sizes = {info.filename: info.file_size for info in archive.infolist()}
    body_nodes = xpath_nodes = referenced_bytes = outside = 0
    # Blocks seen so far, and their sum over sections.
    sect_prs = blocks = section_blocks = 0
    # w:sectPr in body-level paragraph properties, and body-level ones.
    paragraph_sections = body_sections = 0
    styled_paragraphs = page_break_work = run_work = copied_bytes = 0
    path: list[str] = []
    # Child nodes seen so far of each element in ``path``.
    children: list[int] = []
    # The body-level paragraph or table being read, if any (0 if none,
    # else 3, its depth): whether it is a paragraph with a w:r or
    # w:hyperlink child, its page breaks, descendant nodes, largest
    # child-node count, bytes (nodes, attribute values and text, see
    # ``MAX_DOCX_COPIED_BYTES``) and text bytes as read once.
    block_depth = 0
    is_paragraph = has_run = False
    page_breaks = block_nodes = widest = block_bytes = block_text = 0
    # Run-work units and text bytes of the block as unstructured reads
    # ``row.cells`` (python-docx shares a cell's objects between its
    # yields, so only its text is copied):
    # frames[0] for the block, one more per w:tc open in it, which is
    # folded into its parent's frame, times its span, when it ends.
    frames: list[list[int]] = []
    # The w:tc, w:tbl and w:tr elements open in the block.
    cells: list[_OpenCell] = []
    tables: list[_OpenTable] = []
    rows: list[_OpenRow] = []
    # Namespace declarations in the body-level paragraph being read.
    paragraph_ns = 0
    # Union branch counts of each element in ``path`` whose children
    # python-docx selects with a union (see
    # ``DOCX_UNION_STEPS_PER_RUN_WORK_UNIT``), None for the others: a
    # cell's w:p and w:tbl, a paragraph's w:r and w:hyperlink, a run's
    # text branches (as a bit set) and their count.
    unions: list[Optional[list[int]]] = []
    # The body's w:p and w:tbl children (its union is read once when the
    # document has no sections).
    body_paragraphs = body_tables = 0
    # Header/footer references per part.
    references: dict[str, int] = {}
    # w:lastRenderedPageBreak elements anywhere, and those in each branch
    # of unstructured's page-break union, by depth (see
    # MAX_DOCX_PAGE_BREAKS).
    all_breaks = 0
    break_branches = dict.fromkeys(_W_PAGE_BREAK_BRANCHES, 0)
    # The union's breaks below each element in ``path``, and the sum
    # over elements of those breaks (weighted for the sort's binary
    # insertion) times the element's child nodes (see
    # DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT).
    breaks_below: list[int] = []
    break_sort_steps = 0
    # The w:drawing elements, and the wp:inline elements below one, open
    # in a w:r child of the body-level paragraph being read: the runs
    # under such an inline are the third branch of unstructured's
    # paragraph union (see DOCX_UNION_STEPS_PER_RUN_WORK_UNIT).
    drawings = inlines = 0

    def visits(depth: int) -> int:
        """Section-XPath visits of a node at *depth* (1 = a child of the
        document node), given the ancestors in ``path[:depth - 1]``."""
        if depth <= 2:
            return 2
        if path[1] != _W_BODY:
            return 0
        if depth == 3:
            return 2
        if path[2] != _W_P:
            return 0
        if depth == 4:
            return 1
        if depth == 5 and path[3] == _W_PPR:
            return 1
        return 0

    def text_before(node) -> Optional[str]:
        """The text node that precedes *node* among its siblings, if
        any: python-docx keeps text that is not whitespace, and some
        that is."""
        previous = node.getprevious()
        text: Optional[str] = (
            node.getparent().text if previous is None else previous.tail
        )
        return text

    # The namespace declarations announced by ``start-ns`` events for
    # the next element, as (URI, bytes), and the bytes of those on the
    # w:document and w:body elements by URI: libxml2 re-declares one on
    # each copy of a paragraph that uses its URI.
    pending_ns: list[tuple[str, int]] = []
    ancestor_ns: dict[str, int] = {}
    # The namespace URIs of ``ancestor_ns`` that element and attribute
    # names in the body-level paragraph being read use (at most one
    # entry per declaration there, so this stays small however many
    # names use them).
    used_uris: set[str] = set()
    # Copies kept of the text after the last body-level paragraph,
    # added when that text is seen.
    tail_copies = 0

    attribute_values = etree.XPath("@*", smart_strings=False)

    def node_bytes(node, is_element: bool) -> int:
        """What copying *node* (an element, comment or processing
        instruction) costs, without its children or the text around
        them: a node, and per attribute its node and its value."""
        size = DOCX_COPIED_BYTES_PER_NODE
        if is_element:
            # At most MAX_DOCX_ELEMENT_ATTRIBUTES, checked at its start.
            count = len(node.attrib)
            if count:
                size += count * DOCX_COPIED_BYTES_PER_ATTRIBUTE
                values = (
                    node.attrib.values()
                    if count <= _DOCX_ATTRIBUTE_VALUES_SCAN
                    else attribute_values(node)
                )
                for value in values:
                    size += _text_bytes(value)
        else:
            size += _text_bytes(node.text)
        return size

    # Depth of the section-defining w:sectPr being read, if any:
    # python-docx reads header/footer references from its direct
    # children only.
    section_depth = 0
    # Namespace declarations in scope at each element of ``path``.
    in_scope: list[int] = []
    with archive.open(main[1:]) as stream:
        for event, element in etree.iterparse(
            _CountedXmlStream(stream),
            events=("start-ns", "start", "end", "comment", "pi"),
            resolve_entities=False,
            no_network=True,
        ):
            if event == "start-ns":
                prefix, uri = element
                # Before any name under it is read (see
                # MAX_DOCX_NAMESPACE_URI_CHARS).
                if (
                    len(uri) > MAX_DOCX_NAMESPACE_URI_CHARS
                    or len(prefix) > MAX_DOCX_NAMESPACE_PREFIX_CHARS
                ):
                    raise DecompressionBombError(
                        "document declares a namespace URI of more than "
                        f"{MAX_DOCX_NAMESPACE_URI_CHARS} or a prefix of more "
                        f"than {MAX_DOCX_NAMESPACE_PREFIX_CHARS} characters"
                    )
                pending_ns.append(
                    (
                        uri,
                        DOCX_COPIED_BYTES_PER_NODE
                        + _text_bytes(prefix)
                        + _text_bytes(uri),
                    )
                )
                # The element they belong to is checked at its start;
                # refusing here already keeps pending_ns as small.
                scope = in_scope[-1] if in_scope else 0
                if scope + len(pending_ns) > MAX_DOCX_NAMESPACES_IN_SCOPE:
                    raise DecompressionBombError(
                        "document has more than "
                        f"{MAX_DOCX_NAMESPACES_IN_SCOPE} namespace declarations "
                        "in scope"
                    )
                continue
            if event in ("comment", "pi"):
                if not path:
                    # Outside the root, where python-docx never looks.
                    outside = _note_outside_root_node(outside)
                    continue
                xpath_nodes += visits(len(path) + 1)
                before = text_before(element)
                if children:
                    children[-1] += 1 + bool(before)
                if len(path) == 2 and path[1] == _W_BODY:
                    body_nodes += 1 + bool(before)
                    if tail_copies:
                        copied_bytes += tail_copies * _text_bytes(before)
                        tail_copies = 0
                if block_depth:
                    text = _text_bytes(before)
                    frames[-1][1] += text
                    block_text += text
                    block_bytes += text + node_bytes(element, False)
                    _note_docx_table_child(len(path) + 1, None, tables, rows)
                # Comments and PIs get no end event, so drop the
                # siblings before them here too.
                parent = element.getparent()
                if parent is not None:
                    while element.getprevious() is not None:
                        del parent[0]
                continue
            # Each element's name is read once per event: lxml builds it
            # afresh, URI included, on every read.
            tag = element.tag
            if event == "end":
                if len(path) == section_depth:
                    section_depth = 0
                # Inside an inline drawing of one of the paragraph's runs
                # (the inline itself included): its child nodes are
                # walked when the third branch's runs are sorted.
                in_inline = bool(inlines)
                if drawings and len(path) > block_depth + 1:
                    # The same test as at the element's start: the
                    # drawings open above it have not changed.
                    if tag == _W_DRAWING:
                        drawings -= 1
                    elif tag == _WP_INLINE:
                        inlines -= 1
                # The text after the element's last child (or all of
                # its text, if it has none).
                last = element[-1] if len(element) else None
                trailing = element.text if last is None else last.tail
                if len(path) == 2 and tag == _W_BODY:
                    body_nodes += bool(trailing)
                    if tail_copies:
                        copied_bytes += tail_copies * _text_bytes(trailing)
                        tail_copies = 0
                own = children.pop() + bool(trailing)
                counter = unions.pop()
                below = breaks_below.pop()
                if below:
                    break_sort_steps += own * (
                        below
                        + DOCX_PAGE_BREAK_SORT_INSERTION_WALKS
                        * min(below, DOCX_PAGE_BREAK_SORT_INSERTION_RUN)
                    )
                    if breaks_below:
                        breaks_below[-1] += below
                if in_inline and block_depth:
                    paragraph_union = unions[block_depth - 1]
                    if paragraph_union is not None:
                        paragraph_union[3] += own
                if block_depth:
                    text = _text_bytes(trailing)
                    frames[-1][1] += text
                    block_text += text
                    block_bytes += text
                    block_nodes += own
                    widest = max(widest, own)
                    extra = 0
                    if counter is not None:
                        extra = _union_units(_union_steps(tag, counter, own))
                        if tag == _W_TC:
                            extra += DOCX_CELL_BLOCK_ITEM_WEIGHT * (
                                counter[0] + counter[1]
                            )
                    _fold_docx_text_work(
                        tag,
                        len(path),
                        own,
                        frames,
                        cells,
                        tables,
                        rows,
                        block_nodes,
                        extra,
                    )
                    if len(path) == block_depth:
                        block_depth = 0
                        units, weighted_text = frames.pop()
                        if has_run:
                            styled_paragraphs += 1 + page_breaks
                        page_break_work += (
                            page_breaks * (page_breaks + 3) * widest * widest
                        )
                        run_work += (page_breaks + 1) * units + (
                            page_breaks
                            * block_nodes
                            // DOCX_FRAGMENT_NODES_PER_RUN_WORK_UNIT
                        )
                        # Extra copies of table cells (spanned columns,
                        # continued merges), as unstructured's HTML
                        # keeps them; two fragments per break, each a
                        # deep copy of the paragraph (its tail included)
                        # that re-declares the ancestors' namespaces it
                        # uses.
                        copied_bytes += DOCX_TABLE_COPY_BYTES_PER_TEXT_BYTE * (
                            weighted_text - block_text
                        ) + (
                            2
                            * page_breaks
                            * (
                                block_bytes
                                + sum(ancestor_ns[uri] for uri in used_uris)
                            )
                        )
                        tail_copies = 2 * page_breaks
                path.pop()
                in_scope.pop()
                # Nothing is read after an element's start event but
                # the text before the next sibling, so the tree is
                # dropped as it is built, tails kept until then.
                element.clear(keep_tail=True)
                parent = element.getparent()
                if parent is not None:
                    while element.getprevious() is not None:
                        del parent[0]
                continue
            # Counted, not read: reading an element's attribute values
            # is quadratic in their number.
            if len(element.attrib) > MAX_DOCX_ELEMENT_ATTRIBUTES:
                raise DecompressionBombError(
                    "document has an element with more than "
                    f"{MAX_DOCX_ELEMENT_ATTRIBUTES} attributes"
                )
            declared = len(pending_ns)
            in_scope.append((in_scope[-1] if in_scope else 0) + declared)
            if in_scope[-1] > MAX_DOCX_NAMESPACES_IN_SCOPE:
                raise DecompressionBombError(
                    "document has more than "
                    f"{MAX_DOCX_NAMESPACES_IN_SCOPE} namespace declarations "
                    "in scope"
                )
            before = None
            if not path:
                _refuse_dtd(element)
            else:
                before = text_before(element)
                children[-1] += 1 + bool(before)
            path.append(tag)
            children.append(0)
            unions.append(None)
            breaks_below.append(0)
            depth = len(path)
            xpath_nodes += visits(depth)
            if tag == _W_LAST_RENDERED_PAGE_BREAK:
                all_breaks += 1
                if all_breaks > MAX_DOCX_PAGE_BREAKS:
                    raise DecompressionBombError(
                        "document has more than "
                        f"{MAX_DOCX_PAGE_BREAKS} rendered page breaks"
                    )
                branch = _W_PAGE_BREAK_BRANCHES.get(depth)
                if branch is not None and tuple(path[1:-1]) == branch:
                    break_branches[depth] += 1
                    breaks_below[-1] = 1
            ns_bytes = 0
            if pending_ns:
                for uri, size in pending_ns:
                    ns_bytes += size
                    if depth <= 2:
                        ancestor_ns[uri] = ancestor_ns.get(uri, 0) + size
                del pending_ns[:]
            if tag == _W_SECTPR:
                sect_prs += 1
                if sect_prs > MAX_DOCX_SECTIONS:
                    raise DecompressionBombError(
                        "document has more than "
                        f"{MAX_DOCX_SECTIONS} sections (w:sectPr elements)"
                    )
            elif tag in _W_GRID_VALUES:
                _note_docx_grid_value(element, tag, path, cells, rows)
            elif tag == _W_TCPR:
                # python-docx reads a cell's first w:tcPr only...
                if cells and cells[-1].depth == depth - 1:
                    cells[-1].tc_pr = 1 if cells[-1].tc_pr == 0 else 2
            elif tag == _W_TRPR:
                # ...and a row's first w:trPr.
                if rows and rows[-1].depth == depth - 1:
                    row = rows[-1]
                    row.tr_pr = 1 if row.tr_pr == 0 else 2
                    if row.tr_pr == 1:
                        row.tr_pr_nodes = -block_nodes
                        # The children find() visits up to it.
                        row.tr_pr_index = children[-2]
            elif tag == _W_VMERGE:
                # The first w:vMerge of the cell's first w:tcPr, whose
                # missing w:val means "continue" (``tc.vMerge ==
                # "continue"``).
                if (
                    cells
                    and cells[-1].depth == depth - 2
                    and path[-2] == _W_TCPR
                    and cells[-1].tc_pr == 1
                    and not cells[-1].merge_read
                ):
                    cells[-1].merge_read = True
                    cells[-1].continues = (
                        element.get(_W_VAL, "continue") == "continue"
                    )
            if depth == 3 and path[1] == _W_BODY:
                body_nodes += 1 + bool(before)
                body_paragraphs += tag == _W_P
                body_tables += tag == _W_TBL
                if tail_copies:
                    copied_bytes += tail_copies * _text_bytes(before)
                    tail_copies = 0
                if tag in _W_BLOCKS:
                    blocks += 1
                    block_depth = depth
                    is_paragraph = tag == _W_P
                    has_run = False
                    page_breaks = block_nodes = widest = 0
                    drawings = inlines = 0
                    block_bytes = block_text = paragraph_ns = 0
                    used_uris.clear()
                    frames.append([0, 0])
                    del cells[:], tables[:], rows[:]
                    # The text before the block is not part of it.
                    before = None
            elif block_depth and is_paragraph:
                if depth == 4 and tag in _W_RUN_CONTAINERS:
                    has_run = True
                elif tag == _W_LAST_RENDERED_PAGE_BREAK:
                    page_breaks += 1
            if block_depth:
                parent_union = unions[-2]
                if parent_union is not None:
                    _note_union_child(path[-2], tag, parent_union)
                if tag in _W_UNION_PARENTS:
                    unions[-1] = [0, 0, 0, 0]
                if (
                    is_paragraph
                    and depth > block_depth + 1
                    and path[block_depth] == _W_R
                ):
                    if tag == _W_DRAWING:
                        drawings += 1
                    elif tag == _WP_INLINE:
                        if drawings:
                            inlines += 1
                    elif tag == _W_R and inlines:
                        # A run of the paragraph union's third branch.
                        paragraph_union = unions[block_depth - 1]
                        if paragraph_union is not None:
                            paragraph_union[2] += 1
                # Before the element is itself recorded as an open
                # table, row or cell.
                _note_docx_table_child(depth, tag, tables, rows)
                if tag == _W_TC:
                    frames.append([0, 0])
                    nodes_before = 0
                    if rows and rows[-1].depth == depth - 1:
                        row = rows[-1]
                        row.count += 1
                        if row.count > MAX_DOCX_ROW_CELLS:
                            raise DecompressionBombError(
                                "document has a table row of more than "
                                f"{MAX_DOCX_ROW_CELLS} cells"
                            )
                        # The row's nodes before this cell: its earlier
                        # children (text included) and their content.
                        nodes_before = (
                            block_nodes - row.nodes + children[-2] - 1
                        )
                    cells.append(_OpenCell(depth, nodes_before, block_nodes))
                elif tag == _W_TBL:
                    tables.append(_OpenTable(depth))
                elif tag == _W_TR:
                    # Rows python-docx reads: direct children of a table.
                    table_row = bool(tables) and tables[-1].depth == depth - 1
                    rows.append(_OpenRow(depth, table_row, block_nodes))
                    if table_row:
                        # The table's children before this row and after
                        # the last one python-docx reads (text, this row
                        # and, for the first, w:tblPr and w:tblGrid
                        # included), which _tr_above walks.
                        table = tables[-1]
                        rows[-1].gap = children[-2] - table.row_children
                        table.row_children = children[-2]
                text = _text_bytes(before)
                frames[-1][1] += text
                block_text += text
                block_bytes += text + node_bytes(element, True) + ns_bytes
                if is_paragraph:
                    if declared:
                        paragraph_ns += declared
                        if paragraph_ns > MAX_DOCX_PARAGRAPH_NAMESPACES:
                            raise DecompressionBombError(
                                "document has a paragraph with more than "
                                f"{MAX_DOCX_PARAGRAPH_NAMESPACES} namespace "
                                "declarations"
                            )
                    if ancestor_ns:
                        _note_used_uris(
                            tag, element.keys(), ancestor_ns, used_uris
                        )
            if tag == _W_SECTPR and (
                (depth == 3 and path[1] == _W_BODY)
                or (depth == 5 and path[1:4] == [_W_BODY, _W_P, _W_PPR])
            ):
                if depth == 3:
                    body_sections += 1
                else:
                    paragraph_sections += 1
                # Through the enclosing paragraph (already counted), or
                # the blocks before a body-level w:sectPr.
                section_blocks += blocks
                section_depth = depth
            elif (
                section_depth
                and depth == section_depth + 1
                and tag in _W_HDRFTR_REFS
            ):
                resolved = rels.get(element.get(_R_ID))
                if (
                    resolved is None
                    or resolved[1] is None
                    or resolved[1][1:] not in sizes
                ):
                    raise DecompressionBombError(
                        "document references a header or footer it "
                        "cannot resolve"
                    )
                if tag == _W_HDRFTR_REFS[0]:
                    _require_content_type(
                        types, resolved[1], _CT_WML_HEADER, "header"
                    )
                else:
                    _require_content_type(
                        types, resolved[1], _CT_WML_FOOTER, "footer"
                    )
                referenced_bytes += sizes[resolved[1][1:]]
                references[resolved[1][1:]] = (
                    references.get(resolved[1][1:], 0) + 1
                )
    if (
        not paragraph_sections + body_sections
        and body_paragraphs
        and body_tables
    ):
        # Without sections, unstructured reads the body's own union.
        run_work += _union_units(blocks * body_nodes, reads=1)
    # The merges of unstructured's page-break union (see
    # MAX_DOCX_PAGE_BREAKS): each branch against the branches before it.
    merged = break_pairs = 0
    for count in break_branches.values():
        break_pairs += merged * count
        merged += count
    if break_pairs:
        run_work += -(-break_pairs // DOCX_PAGE_BREAK_UNION_PAIRS_PER_UNIT)
        # Two non-empty branches (else there are no pairs): the merged
        # set may be out of document order, and sorting it walks the
        # siblings after the breaks.
        run_work += -(-break_sort_steps // DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT)
    if referenced_bytes <= MAX_DOCX_HEADER_FOOTER_BYTES:
        # Each reference reads its part's unions once (see
        # DOCX_UNION_STEPS_PER_HEADER_BYTE); a sum already over the
        # ceiling is refused without reading them.
        for name, count in references.items():
            referenced_bytes = min(
                referenced_bytes
                + _docx_part_union_steps(archive, name)
                * count
                // DOCX_UNION_STEPS_PER_HEADER_BYTE,
                _SATURATED,
            )
    return _DocxSectionStats(
        sect_prs=sect_prs,
        sections=paragraph_sections + body_sections,
        body_nodes=body_nodes,
        xpath_nodes=xpath_nodes,
        union_pairs=paragraph_sections * body_sections,
        block_work=section_blocks * body_nodes,
        referenced_bytes=referenced_bytes,
        styled_paragraphs=styled_paragraphs,
        page_break_work=page_break_work,
        run_work=run_work,
        copied_bytes=copied_bytes,
        style_entries=max(styles, DOCX_DEFAULT_STYLES_ENTRIES),
        style_attributes=style_attributes,
        relationships=len(rels),
    )


#: Work and byte counts saturate here: nested spans multiply, and a
#: count past every ceiling needs no more digits.
_SATURATED = 1 << 64


def _note_union_child(parent: str, tag: str, counter: list[int]) -> None:
    """Count *tag*, a child of *parent* (a cell, a paragraph, a run or a
    header/footer root), in *counter*, its parent's union branches (see
    ``DOCX_UNION_STEPS_PER_RUN_WORK_UNIT``)."""
    if parent == _W_P:
        if tag == _W_R:
            counter[0] += 1
        elif tag == _W_HYPERLINK:
            counter[1] += 1
    elif parent == _W_R:
        bit = _W_RUN_TEXT_BRANCHES.get(tag)
        if bit:
            counter[0] |= bit
            counter[1] += 1
    elif tag == _W_P:
        counter[0] += 1
    elif tag == _W_TBL:
        counter[1] += 1


def _union_steps(tag: str, counter: list[int], own: int) -> int:
    """Steps of one evaluation of the union over the *own* child nodes
    of a *tag* element whose branches *counter* counted: per branch
    merged after the first non-empty one, the nodes selected times the
    child nodes (or times the nodes selected, when a branch selects
    nodes that are not children: a body-level paragraph's runs in inline
    drawings, whose sort also walks the nodes in those drawings,
    ``counter[3]``), and nothing with fewer than two non-empty
    branches."""
    if tag == _W_R:
        branches = bin(counter[0]).count("1")
        selected = counter[1]
    else:
        branches = sum(1 for count in counter[:3] if count)
        selected = sum(counter[:3])
    if branches < 2:
        return 0
    steps = (branches - 1) * selected * max(own, selected)
    if len(counter) > 3 and counter[2]:
        # A body-level paragraph's runs in inline drawings: sorting
        # them walks the nodes of those drawings (``counter[3]``).
        steps += selected * counter[3]
    return steps


def _union_units(steps: int, reads: int = DOCX_UNION_READS) -> int:
    """Units of ``MAX_DOCX_RUN_WORK`` for *steps* union steps evaluated
    *reads* times (rounded up)."""
    if not steps:
        return 0
    return min(
        -(-steps * reads // DOCX_UNION_STEPS_PER_RUN_WORK_UNIT), _SATURATED
    )


def _docx_part_union_steps(archive: zipfile.ZipFile, name: str) -> int:
    """Union steps of one read of header/footer part *name*: its root's
    ``./w:p | ./w:tbl``, and the union of every cell, paragraph and run
    in it (see ``DOCX_UNION_STEPS_PER_RUN_WORK_UNIT``), each evaluated
    once per read. Streamed with no tree kept; a DTD, or a namespace
    URI or prefix over its cap, is refused as in ``document.xml``, and
    so are more than ``MAX_XML_OUTSIDE_ROOT_NODES`` comments and
    processing instructions outside the root."""
    from lxml import etree

    steps = outside = 0
    path: list[str] = []
    children: list[int] = []
    unions: list[Optional[list[int]]] = []
    with archive.open(name) as stream:
        for event, element in etree.iterparse(
            _CountedXmlStream(stream),
            events=("start-ns", "start", "end", "comment", "pi"),
            resolve_entities=False,
            no_network=True,
        ):
            if event == "start-ns":
                prefix, uri = element
                if (
                    len(uri) > MAX_DOCX_NAMESPACE_URI_CHARS
                    or len(prefix) > MAX_DOCX_NAMESPACE_PREFIX_CHARS
                ):
                    raise DecompressionBombError(
                        "header or footer declares a namespace URI of more "
                        f"than {MAX_DOCX_NAMESPACE_URI_CHARS} or a prefix of "
                        f"more than {MAX_DOCX_NAMESPACE_PREFIX_CHARS} "
                        "characters"
                    )
                continue
            if event == "end":
                last = element[-1] if len(element) else None
                trailing = element.text if last is None else last.tail
                own = children.pop() + bool(trailing)
                counter = unions.pop()
                tag = path.pop()
                if counter is not None:
                    steps += _union_steps(tag, counter, own)
            elif not path:
                if event == "start":
                    _refuse_dtd(element)
                    path.append(element.tag)
                    children.append(0)
                    unions.append([0, 0])
                else:
                    # A comment or processing instruction outside the
                    # root.
                    outside = _note_outside_root_node(outside)
                continue
            else:
                previous = element.getprevious()
                text = (
                    element.getparent().text
                    if previous is None
                    else previous.tail
                )
                children[-1] += 1 + bool(text)
                if event == "start":
                    tag = element.tag
                    if unions[-1] is not None:
                        _note_union_child(path[-1], tag, unions[-1])
                    path.append(tag)
                    children.append(0)
                    unions.append([0, 0] if tag in _W_UNION_PARENTS else None)
                    continue
            # An ended element, comment or processing instruction: drop
            # the siblings before it (tails are kept until then).
            if event == "end":
                element.clear(keep_tail=True)
            parent = element.getparent()
            if parent is not None:
                while element.getprevious() is not None:
                    del parent[0]
    return steps


def _note_used_uris(
    tag: str, attributes: list[str], declared: dict[str, int], used: set[str]
) -> None:
    """Add to *used* the namespace URIs in *declared* that the element
    name *tag* and the *attributes* names (Clark notation) use."""
    for name in (tag, *attributes):
        if name[:1] == "{":
            uri = name[1 : name.index("}")]
            if uri in declared:
                used.add(uri)


class _OpenCell:
    """A ``w:tc`` open in the block being read."""

    __slots__ = (
        "depth",
        "span",
        "read_span",
        "continues",
        "tc_pr",
        "merge_read",
        "span_read",
        "nodes_before",
        "start_nodes",
    )

    def __init__(self, depth: int, nodes_before: int, start_nodes: int) -> None:
        self.depth = depth
        # The largest w:gridSpan under any of its w:tcPr (the weight of
        # its work, never less than python-docx's), and the span
        # python-docx reads: the first w:gridSpan of its first w:tcPr.
        self.span = 1
        self.read_span = 1
        # Whether it continues a vertical merge as python-docx reads it.
        self.continues = False
        # 0: no w:tcPr seen yet, 1: in its first, 2: past its first;
        # and whether that first one's w:vMerge and w:gridSpan were read.
        self.tc_pr = 0
        self.merge_read = False
        self.span_read = False
        # Child and descendant nodes of its row before it.
        self.nodes_before = nodes_before
        # The block's node count when the cell started.
        self.start_nodes = start_nodes


class _OpenRow:
    """A ``w:tr`` open in the block being read."""

    __slots__ = (
        "depth",
        "padding",
        "cells",
        "count",
        "columns",
        "tr_pr",
        "grid_before",
        "before_read",
        "nodes",
        "others",
        "gap",
        "tr_pr_nodes",
        "tr_pr_index",
    )

    def __init__(self, depth: int, table_row: bool, nodes: int) -> None:
        self.depth = depth
        # Its gridBefore + gridAfter columns (any w:trPr, negative ones
        # pad nothing).
        self.padding = 0
        # For a row python-docx reads (a direct child of a table), per
        # direct w:tc: (the span python-docx reads, whether it continues
        # a vertical merge, the row's nodes before it).
        self.cells: Optional[list[tuple[int, bool, int]]] = (
            [] if table_row else None
        )
        # Direct w:tc children, and the columns they span (largest
        # spans, at least one each).
        self.count = 0
        self.columns = 0
        # 0: no w:trPr seen yet, 1: in its first, 2: past its first; and
        # python-docx's grid_before (first w:gridBefore of the first
        # w:trPr) once read.
        self.tr_pr = 0
        self.grid_before = 0
        self.before_read = False
        # The block's node count when the row started.
        self.nodes = nodes
        # Children other than w:tc elements (elements, comments and
        # processing instructions; see MAX_DOCX_ROW_OTHER_CHILDREN).
        self.others = 0
        # For a row python-docx reads, its table's child nodes from the
        # last such row before it (excluded) through it (included).
        self.gap = 0
        # The nodes in its first w:trPr (where grid_before is looked
        # up), once that has ended, and its child nodes up to and
        # including that w:trPr (0 if it has none).
        self.tr_pr_nodes = 0
        self.tr_pr_index = 0


class _OpenTable:
    """A ``w:tbl`` open in the block being read."""

    __slots__ = (
        "depth",
        "largest_units",
        "largest_text",
        "above",
        "others",
        "row_children",
    )

    def __init__(self, depth: int) -> None:
        self.depth = depth
        # The largest cell so far that does not continue a merge (its
        # run-work units and text bytes, times its span).
        self.largest_units = 0
        self.largest_text = 0
        # The last row python-docx reads, as ``tc_at_grid_offset``
        # searches it: {grid offset: (index of the cell found there,
        # merge-step units of that cell if it continues a merge, else
        # None)}, its cell count and its nodes.
        self.above: Optional[
            tuple[dict[int, tuple[int, Optional[int]]], int, int]
        ] = None
        # Children other than w:tr elements since the last w:tr child
        # (elements, comments and processing instructions; see
        # MAX_DOCX_TABLE_ROW_GAP), and its child nodes (text included)
        # through the last w:tr child python-docx reads.
        self.others = 0
        self.row_children = 0


def _note_docx_table_child(
    depth: int,
    tag: Optional[str],
    tables: list[_OpenTable],
    rows: list[_OpenRow],
) -> None:
    """Count a node starting at *depth* (an element named *tag*, or a
    comment or processing instruction if None) among the children of
    the open table or row it belongs to, refusing a row python-docx
    reads with more than ``MAX_DOCX_ROW_OTHER_CHILDREN`` children other
    than ``w:tc`` and a ``w:tr`` child of a table after more than
    ``MAX_DOCX_TABLE_ROW_GAP`` other children since the last."""
    if tables and tables[-1].depth == depth - 1:
        table = tables[-1]
        if tag != _W_TR:
            table.others += 1
        elif table.others > MAX_DOCX_TABLE_ROW_GAP:
            raise DecompressionBombError(
                "document has a table with more than "
                f"{MAX_DOCX_TABLE_ROW_GAP} children before or between "
                "its rows"
            )
        else:
            table.others = 0
    elif rows and rows[-1].depth == depth - 1 and tag != _W_TC:
        row = rows[-1]
        if row.cells is not None:
            row.others += 1
            if row.others > MAX_DOCX_ROW_OTHER_CHILDREN:
                raise DecompressionBombError(
                    "document has a table row with more than "
                    f"{MAX_DOCX_ROW_OTHER_CHILDREN} children that are not "
                    "cells"
                )


def _docx_merge_steps(
    row: _OpenRow,
    row_cells: list[tuple[int, bool, int]],
    table: _OpenTable,
    nodes: int,
    row_children: int,
) -> int:
    """Run-work units of python-docx's steps up vertical merges for
    *row_cells*, the cells of *row* (which ends with the block's node
    count at *nodes* and has *row_children* child nodes), and record
    *row* on *table* as the row above the next.

    A cell continuing a merge at grid offset O (``grid_before`` plus
    the spans, as python-docx reads them, of the cells before it) takes
    one step to the row above, where ``tc_at_grid_offset`` returns the
    first cell starting exactly at O unless an earlier one starts after
    O (it stops there); if that cell continues too, its own steps
    follow. Each step is priced at the cells and nodes it walks (see
    ``DOCX_MERGE_STEP_WEIGHT``): the row's nodes before the cell, its
    children up to its first ``w:trPr`` (all of them if it has none)
    and the nodes in that ``w:trPr`` (``grid_before``'s two ``find``
    calls; nodes before the cell may be counted twice), the
    table's children from the row above through this one
    (``_tr_above``'s sibling walk) and the row above's nodes.
    Without a row above, or a cell at O, python-docx raises after the
    step, which is priced the same."""
    walked = (row.tr_pr_index or row_children) + row.tr_pr_nodes + row.gap
    above = table.above
    found: dict[int, tuple[int, Optional[int]]] = {}
    highest: Optional[int] = None
    offset = row.grid_before
    total = 0
    for index, (span, continues, before) in enumerate(row_cells):
        steps: Optional[int] = None
        if continues:
            if above is None:
                steps = _merge_step(index, before + walked)
            else:
                cells_above, count_above, nodes_above = above
                hit = cells_above.get(offset)
                if hit is None:
                    steps = _merge_step(
                        index + count_above, before + walked + nodes_above
                    )
                else:
                    steps = _merge_step(
                        index + hit[0], before + walked + nodes_above
                    )
                    if hit[1] is not None:
                        steps += hit[1]
            steps = min(steps, _SATURATED)
            total = min(total + steps, _SATURATED)
        if highest is None or offset > highest:
            found[offset] = (index, steps)
            highest = offset
        offset += span
    table.above = (found, len(row_cells), nodes - row.nodes)
    return total


def _merge_step(cells: int, nodes: int) -> int:
    """Run-work units of one step up a vertical merge that reads the
    spans of *cells* cells and walks *nodes* nodes."""
    return (
        DOCX_MERGE_STEP_WEIGHT
        + cells // DOCX_MERGE_CELLS_PER_UNIT
        + nodes // DOCX_MERGE_NODES_PER_UNIT
    )


def _text_bytes(text: Optional[str]) -> int:
    """UTF-8 size of *text* (0 for ``None``), as lxml stores it."""
    if not text:
        return 0
    return len(text) if text.isascii() else len(text.encode("utf-8"))


def _note_docx_grid_value(element, tag, path, cells, rows) -> None:
    """Refuse a ``w:gridSpan``/``w:gridBefore``/``w:gridAfter`` whose
    value python-docx cannot read or that exceeds
    ``MAX_DOCX_GRID_SPAN`` (anywhere in ``document.xml``), and record
    it on the open cell or row it belongs to: a ``w:gridSpan`` child of
    a cell's ``w:tcPr`` (the largest, if several, and the one
    python-docx reads), and ``gridBefore`` / ``gridAfter`` children of
    a row's ``w:trPr`` (negative ones pad nothing; python-docx's
    ``grid_before`` is the first ``gridBefore`` of the first
    ``w:trPr``)."""
    value = element.get(_W_VAL)
    if value is None:
        raise DecompressionBombError(
            "document has a table grid value python-docx cannot read"
        )
    if len(value) > _MAX_DOCX_GRID_VALUE_CHARS:
        # int() of a long digit string is not linear in its length.
        raise DecompressionBombError(
            "document has a table grid value of more than "
            f"{_MAX_DOCX_GRID_VALUE_CHARS} characters"
        )
    try:
        # python-docx converts the attribute with int() itself.
        number = int(value)
    except ValueError:
        raise DecompressionBombError(
            "document has a table grid value python-docx cannot read"
        ) from None
    if number > MAX_DOCX_GRID_SPAN:
        raise DecompressionBombError(
            f"document has a table cell spanning or skipping {number} grid "
            f"columns (ceiling {MAX_DOCX_GRID_SPAN})"
        )
    depth = len(path)
    if tag == _W_GRID_SPAN:
        if cells and cells[-1].depth == depth - 2 and path[-2] == _W_TCPR:
            cell = cells[-1]
            cell.span = max(cell.span, number)
            if cell.tc_pr == 1 and not cell.span_read:
                cell.span_read = True
                cell.read_span = number
    elif rows and rows[-1].depth == depth - 2 and path[-2] == _W_TRPR:
        row = rows[-1]
        row.padding += max(number, 0)
        if tag == _W_GRID_BEFORE and row.tr_pr == 1 and not row.before_read:
            row.before_read = True
            row.grid_before = number


def _fold_docx_text_work(
    tag, depth, own, frames, cells, tables, rows, nodes, extra=0
) -> None:
    """Add the element that ends at *depth* (with *own* child nodes;
    the block has *nodes* nodes so far) to the run-work units and text
    bytes of ``frames``: ``DOCX_RUN_WORK_ITEM_WEIGHT`` units for a run
    (plus its children), row or cell (plus one per
    ``DOCX_CELL_NODES_PER_UNIT`` of its descendant nodes); a row's
    ``gridBefore`` and ``gridAfter`` columns and python-docx's steps up
    the vertical merges its cells continue (``_docx_merge_steps``); and
    an ending
    cell's whole frame folded into its parent's, times its span, or,
    for a cell continuing a vertical merge, plus the largest cell of its
    table so far (python-docx yields the merge's first cell in its place
    when unstructured reads ``row.cells``). A row spanning more than
    ``MAX_DOCX_ROW_GRID_COLUMNS`` columns is refused."""
    units = text = 0
    if tag == _W_TC and cells and cells[-1].depth == depth:
        cell = cells.pop()
        units, text = frames.pop()
        # Its descendant nodes, read once per yield (see
        # DOCX_CELL_NODES_PER_UNIT).
        units += (
            DOCX_RUN_WORK_ITEM_WEIGHT
            + (nodes - cell.start_nodes) // DOCX_CELL_NODES_PER_UNIT
            + extra
        )
        extra = 0
        if rows and rows[-1].depth == depth - 1:
            row = rows[-1]
            row.columns += max(cell.span, 1)
            if row.cells is not None:
                row.cells.append(
                    (cell.read_span, cell.continues, cell.nodes_before)
                )
        if cell.continues:
            if tables:
                units += tables[-1].largest_units
                text += tables[-1].largest_text
        else:
            span = max(cell.span, 1)
            units = min(units * span, _SATURATED)
            text = min(text * span, _SATURATED)
            if tables:
                table = tables[-1]
                table.largest_units = max(table.largest_units, units)
                table.largest_text = max(table.largest_text, text)
    elif tag == _W_R:
        units = DOCX_RUN_WORK_ITEM_WEIGHT + own
    elif tag == _W_TR:
        units = DOCX_RUN_WORK_ITEM_WEIGHT
        if rows and rows[-1].depth == depth:
            row = rows.pop()
            if row.padding + row.columns > MAX_DOCX_ROW_GRID_COLUMNS:
                raise DecompressionBombError(
                    "document has a table row spanning more than "
                    f"{MAX_DOCX_ROW_GRID_COLUMNS} grid columns"
                )
            units += -(-row.padding // DOCX_GRID_PADDING_PER_RUN_WORK_UNIT)
            # A row python-docx reads is a direct child of tables[-1].
            if row.cells is not None and tables:
                units += _docx_merge_steps(
                    row, row.cells, tables[-1], nodes, own
                )
    elif tag == _W_TRPR and rows and rows[-1].depth == depth - 1:
        row = rows[-1]
        if row.tr_pr == 1:
            # Its first w:trPr ends: the nodes in it.
            row.tr_pr_nodes += nodes
    elif tag == _W_TBL and tables and tables[-1].depth == depth:
        tables.pop()
    frame = frames[-1]
    frame[0] = min(frame[0] + units + extra, _SATURATED)
    frame[1] = min(frame[1] + text, _SATURATED)


def _docx_skeleton_entries(
    archive: zipfile.ZipFile,
    rels: dict[str, tuple[str, Optional[str]]],
    reltype: str,
    ceiling: int,
    what: str,
    types: _ContentTypes,
    allowed: frozenset[str],
) -> tuple[int, int]:
    """Hold the document part's *reltype* target (the settings or the
    styles part) to the skeleton treatment, and return the children of
    its root and their attributes ((0, 0) if there is none).

    python-docx resolves both parts through the document part's
    relationship of that type (and makes a small default one when there
    is none). unstructured reads ``settings.odd_and_even_pages_header_footer``
    twice per section, each a search of the settings root's children,
    and a paragraph's style about four times per paragraph, each a
    Python loop over the styles root's ``w:style`` children that reads
    two attributes of each (see ``MAX_DOCX_STYLE_WORK``). Every internal
    relationship of *reltype* whose target is in the archive is parsed
    (``MAX_OPC_SKELETON_PART_BYTES``, DTD refused) and refused if its
    root has more than *ceiling* children (elements, comments and
    processing instructions). More than one such relationship makes
    python-docx fail; checking each, and returning the largest, is
    simply conservative. Each must have the content type in *allowed*
    (see ``_require_content_type``): python-docx builds a document part
    from a styles or settings target typed as one, whose own styles or
    settings it then reads instead."""
    largest = (0, 0)
    for rel_type, target in rels.values():
        if rel_type != reltype or target is None:
            continue
        root = _skeleton_xml(archive, target[1:])
        if root is None:
            continue
        _require_content_type(types, target, allowed, what)
        if len(root) > ceiling:
            raise DecompressionBombError(
                f"document {what} part has {len(root)} entries "
                f"(ceiling {ceiling})"
            )
        attributes = sum(
            len(child.attrib) for child in root if isinstance(child.tag, str)
        )
        largest = (
            max(largest[0], len(root)),
            max(largest[1], attributes),
        )
    return largest

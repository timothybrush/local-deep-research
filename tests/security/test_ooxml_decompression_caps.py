"""Decompression-bomb ceilings for zip-container uploads.

The rag upload route indexes arbitrary document types through
``document_loaders``; ``.docx``/``.xlsx``/``.odt`` are zip containers
of XML that the parsers (python-docx, openpyxl, unstructured) read
entry-whole into memory (``archive.read(...)``). Upload caps bound the
COMPRESSED body, not the uncompressed side: a modest ``.xlsx`` can
decompress to gigabytes of ``sharedStrings`` and OOM the indexing
worker — the same class as the OpenAlex partition caps and the PDF
extraction ceilings.

The guard checks the central directory's declarations (entry count,
compression method — STORED/DEFLATE only, no encryption, compressed
size consistent with the declared size, declared total), then inflates
every entry in bounded steps and refuses any that produces more than
it declares: zipfile truncates to the declared size only *after*
inflating, so an unverified declaration bounds nothing. Verified sizes
then face the per-kind ceilings. For the in-process Python parsers
that bounds the bytes of each *read* (not their memory, which runs a
multiple of it). pandoc's own zip reader is bounded separately (see
test_pandoc_conversion_bounds.py).
"""

from __future__ import annotations

import io
import math
import posixpath
import random
import re
import struct
import time
import zipfile
import zlib
import contextlib
import copy
from contextlib import contextmanager
from unittest.mock import patch

import pytest


def _zip_bytes(entries: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries.items():
            zf.writestr(name, data)
    return buf.getvalue()


def _guard():
    from local_deep_research.document_loaders import zip_container_guard

    return zip_container_guard


def _ooxml_package(entries: dict[str, bytes], ext: str) -> dict[str, bytes]:
    """*entries* plus, after them, what a .docx or .pptx needs for the
    guard to accept its layout: a package relationship to a main part
    that exists (and so a part name unstructured's sniffer reads as the
    extension) of the content type the library loads it as. Entries
    already present are kept."""
    if ext == ".docx":
        main = "word/document.xml"
        body = f"<w:document xmlns:w='{_W_MAIN}'><w:body/></w:document>"
        main_ct = "wordprocessingml.document.main+xml"
    else:
        main = "ppt/presentation.xml"
        body = (
            "<p:presentation xmlns:p='http://schemas.openxmlformats.org/"
            "presentationml/2006/main'/>"
        )
        main_ct = "presentationml.presentation.main+xml"
    skeleton = {
        "[Content_Types].xml": (
            "<Types xmlns='http://schemas.openxmlformats.org/package/2006/"
            "content-types'><Override PartName='/"
            + main
            + "' ContentType='application/vnd.openxmlformats-officedocument."
            + main_ct
            + "'/></Types>"
        ).encode(),
        "_rels/.rels": _rels(("rId1", "officeDocument", main)).encode(),
        main: body.encode(),
    }
    return {
        **entries,
        **{k: v for k, v in skeleton.items() if k not in entries},
    }


class TestOoxmlDecompressionCaps:
    """Oversized zip-container uploads must be refused pre-parse."""

    def test_oversized_declared_entry_is_refused(self):
        guard = _guard()

        bomb = _zip_bytes({"xl/sharedStrings.xml": b"\x00" * 5_000_000})

        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(
                bomb, ".xlsx", max_entry_bytes=1_000_000
            )

    def test_oversized_declared_total_is_refused(self):
        guard = _guard()

        bomb = _zip_bytes(
            {
                "xl/worksheets/sheet1.xml": b"\x00" * 900_000,
                "xl/worksheets/sheet2.xml": b"\x00" * 900_000,
            }
        )

        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(
                bomb, ".xlsx", max_total_bytes=1_000_000
            )

    def test_entry_flood_is_refused(self):
        guard = _guard()

        flood = _zip_bytes({f"r/{i}.xml": b"x" for i in range(50)})

        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(flood, ".docx", max_entries=10)

    def test_load_from_bytes_applies_the_guard(self):
        """The upload funnel refuses the bomb before any parser runs.

        Uses the entry-count ceiling so the default knobs are exercised
        without materializing a half-gigabyte entry in the test.
        """
        from local_deep_research.document_loaders.bytes_loader import (
            load_from_bytes,
        )

        bomb = _zip_bytes({f"word/part{i}.xml": b"x" for i in range(10_001)})

        with pytest.raises(ValueError) as excinfo:
            load_from_bytes(bomb, ".docx", "hostile.docx")

        assert "DecompressionBomb" in type(excinfo.value).__name__ or (
            "entries" in str(excinfo.value).lower()
        ), f"bomb reached the parser layer unguarded ({excinfo.value!r})"

    def test_non_container_extensions_bypass_the_guard(self):
        """Plain-text and PDF bytes never hit the zip check."""
        guard = _guard()

        guard.validate_zip_container(b"not a zip at all", ".txt")
        guard.validate_zip_container(b"%PDF-not-a-zip", ".pdf")

    @pytest.mark.parametrize("ext", [".pptx", ".epub"])
    def test_every_registry_zip_container_is_guarded(self, ext):
        """The guard's extension set must cover every zip-backed loader."""
        guard = _guard()

        bomb = _zip_bytes({"content/part.xml": b"\x00" * 5_000_000})

        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(bomb, ext, max_entry_bytes=1_000_000)


def _central_dir_offset(data: bytes) -> int:
    """Offset of the first central-directory record."""
    return data.index(b"PK\x01\x02")


class TestCentralDirectorySpanIsBoundedPreParse:
    """zipfile materialises every central-directory record before the
    entry-count ceiling can count them, so the span the archive's end
    records declare is refused first — no records need be packed."""

    def test_huge_declared_span_is_refused_pre_parse(self):
        """A stub container whose EOCD claims a 500 MB central
        directory is refused on the declaration alone."""
        guard = _guard()
        data = bytearray(_zip_bytes({"word/document.xml": b"<w/>"}))
        struct.pack_into(
            "<I", data, data.rfind(b"PK\x05\x06") + 12, 500_000_000
        )

        with pytest.raises(
            guard.DecompressionBombError, match="central directory spans"
        ):
            guard.validate_zip_container(bytes(data), ".docx")

    def test_scaled_ceiling_is_used(self):
        guard = _guard()
        data = bytearray(_zip_bytes({"word/document.xml": b"<w/>"}))
        struct.pack_into("<I", data, data.rfind(b"PK\x05\x06") + 12, 11_001)

        with pytest.raises(
            guard.DecompressionBombError, match="central directory spans"
        ):
            guard.validate_zip_container(bytes(data), ".docx", max_entries=10)

    def test_zip64_declared_span_is_refused_pre_parse(self):
        """A ZIP64 end record declaring the big span is refused too,
        with the 32-bit EOCD declaring none: the ZIP64 record zipfile
        follows governs."""
        guard = _guard()
        data = _record_flood(40, cd_size_in_eocd=0, zip64=True)

        with _zipfile_must_not_be_constructed(guard):
            with pytest.raises(guard.DecompressionBombError, match="1880"):
                guard.validate_zip_container(data, ".docx", max_entries=1)

    def test_honest_small_span_falls_through(self):
        """An honest container's central directory is tiny; the
        pre-parse bound must not fire and the guard still accepts it."""
        guard = _guard()
        package = _zip_bytes(_ooxml_package({}, ".docx"))

        guard.validate_zip_container(package, ".docx")

    def test_empty_archive_falls_through(self):
        # The EOCD of an empty archive declares a zero-byte span.
        _guard().validate_zip_container(_zip_bytes({}), ".epub")

    def test_absent_eocd_is_refused_pre_parse(self):
        """No end record zipfile can find: the bound fails closed itself
        rather than standing aside for ZipFile."""
        guard = _guard()

        with _zipfile_must_not_be_constructed(guard):
            with pytest.raises(guard.DecompressionBombError, match="malformed"):
                guard.validate_zip_container(b"PK\x03\x04 truncated", ".docx")

    def test_end_record_zipfile_rejects_is_refused_pre_parse(self):
        """A ZIP64 record zipfile refuses as corrupt (its offset plus
        size miss the locator's offset) is refused before ZipFile."""
        guard = _guard()
        data = bytearray(_record_flood(2, cd_size_in_eocd=0, zip64=True))
        zip64_at = data.rfind(b"PK\x06\x06")
        struct.pack_into("<Q", data, zip64_at + 48, 7)

        with _zipfile_must_not_be_constructed(guard):
            with pytest.raises(guard.DecompressionBombError, match="malformed"):
                guard.validate_zip_container(bytes(data), ".docx")


def _record_flood(
    count: int,
    *,
    cd_size_in_eocd: int | None = None,
    eocd_counts: bytes | None = None,
    eocd_offset: int | None = None,
    zip64: bool = False,
    zip64_extensible: bytes = b"",
    comment: bytes = b"",
) -> bytes:
    """A stub local header, then *count* 47-byte central-directory
    records, then hand-packed end records — the shape of a record
    flood, kept tiny (the tests lower ``max_entries`` instead)."""
    local = b"PK\x03\x04" + b"\x00" * 26
    record = b"PK\x01\x02" + struct.pack("<6H3L5H2L", *([0] * 16)) + b"a"
    record = record[:28] + struct.pack("<H", 1) + record[30:]
    cd = record * count
    cd_offset = len(local)
    out = local + cd
    if zip64:
        zip64_at = len(out)
        out += (
            b"PK\x06\x06"
            + struct.pack(
                "<Q2H2I4Q",
                44 + len(zip64_extensible),
                45,
                45,
                0,
                0,
                count,
                count,
                len(cd),
                cd_offset,
            )
            + zip64_extensible
            + b"PK\x06\x07"
            + struct.pack("<IQI", 0, zip64_at, 1)
        )
    counts = eocd_counts or struct.pack("<2H", count, count)
    out += (
        b"PK\x05\x06"
        + struct.pack("<2H", 0, 0)
        + counts
        + struct.pack(
            "<2LH",
            len(cd) if cd_size_in_eocd is None else cd_size_in_eocd,
            cd_offset if eocd_offset is None else eocd_offset,
            len(comment),
        )
        + comment
    )
    return out


@contextmanager
def _zipfile_must_not_be_constructed(guard):
    """The bound must refuse before zipfile walks any record."""
    with patch.object(
        guard.zipfile,
        "ZipFile",
        side_effect=AssertionError("ZipFile was constructed"),
    ) as zip_file:
        yield
    assert not zip_file.called, "ZipFile walked the central directory"


class TestEndRecordSelectionMatchesZipfile:
    """The span is read from the end records zipfile itself selects;
    a layout where a naive backwards search for the signatures picks a
    different record (or none) must not slip past the bound. Each
    archive holds 40 records (1880 bytes) against a 1 KiB limit."""

    def test_signature_in_eocd_entry_count_fields_is_refused(self):
        """zipfile takes the fixed 22-byte record at the end; the
        signature written into its (ignored) entry counts must not
        hide the record from the bound."""
        guard = _guard()
        data = _record_flood(40, eocd_counts=b"PK\x05\x06")
        assert zipfile._EndRecData(io.BytesIO(data))[5] == 1880

        with _zipfile_must_not_be_constructed(guard):
            with pytest.raises(guard.DecompressionBombError, match="1880"):
                guard.validate_zip_container(data, ".docx", max_entries=1)

    def test_signature_in_eocd_offset_field_is_refused(self):
        """The offset 0x06054b50 spells the signature; zipfile ignores
        the offset when it locates the directory."""
        guard = _guard()
        data = _record_flood(40, eocd_offset=0x06054B50)
        assert zipfile._EndRecData(io.BytesIO(data))[5] == 1880

        with _zipfile_must_not_be_constructed(guard):
            with pytest.raises(guard.DecompressionBombError, match="1880"):
                guard.validate_zip_container(data, ".docx", max_entries=1)

    def test_fake_zip64_record_in_comment_is_refused(self):
        """A zeroed ZIP64 record in the archive comment declares no
        span; zipfile follows the locator to the real one."""
        guard = _guard()
        fake = b"PK\x06\x06" + b"\x00" * 52
        data = _record_flood(40, cd_size_in_eocd=0, zip64=True, comment=fake)
        assert zipfile._EndRecData(io.BytesIO(data))[5] == 1880

        with _zipfile_must_not_be_constructed(guard):
            with pytest.raises(guard.DecompressionBombError, match="1880"):
                guard.validate_zip_container(data, ".docx", max_entries=1)

    def test_zip64_record_behind_extensible_data_is_refused(self):
        """The ZIP64 record sits more than 64 KiB before its locator
        (zip64 extensible data); zipfile follows the locator's offset
        to it, so the bound must too."""
        guard = _guard()
        data = _record_flood(
            40,
            cd_size_in_eocd=0,
            zip64=True,
            zip64_extensible=b"\x00" * 70_000,
        )
        if zipfile._EndRecData(io.BytesIO(data))[5] != 1880:
            pytest.skip("this zipfile does not follow the ZIP64 locator")

        with _zipfile_must_not_be_constructed(guard):
            with pytest.raises(guard.DecompressionBombError, match="1880"):
                guard.validate_zip_container(data, ".docx", max_entries=1)

    def test_layouts_are_walked_by_zipfile_when_unbounded(self):
        """Control: without the bound, zipfile really walks all 40
        records of each layout — the refusals above are not an
        artefact of malformed archives."""
        layouts = [
            _record_flood(40, eocd_counts=b"PK\x05\x06"),
            _record_flood(40, eocd_offset=0x06054B50),
            _record_flood(
                40,
                cd_size_in_eocd=0,
                zip64=True,
                comment=b"PK\x06\x06" + b"\x00" * 52,
            ),
        ]
        for data in layouts:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                assert len(archive.infolist()) == 40


class TestMalformedContainersFailClosed:
    """Anything the stdlib cannot even list is refused as the guard's
    own typed error — never passed on, never a stray exception."""

    def test_not_a_zip_is_refused(self):
        guard = _guard()

        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(b"PK\x03\x04 truncated", ".docx")

    def test_truncated_central_directory_is_refused(self):
        guard = _guard()
        data = _zip_bytes({"word/document.xml": b"<w/>"})

        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(data[: len(data) - 30], ".docx")

    def test_undecodable_utf8_entry_name_is_refused(self):
        """Flag bit 11 promises UTF-8 names; invalid bytes make
        ``infolist()`` raise UnicodeDecodeError."""
        guard = _guard()
        data = bytearray(_zip_bytes({"word/abc.xml": b"<w/>"}))
        cd = _central_dir_offset(data)
        flags = struct.unpack_from("<H", data, cd + 8)[0] | 0x800
        struct.pack_into("<H", data, cd + 8, flags)
        name_at = cd + 46
        data[name_at + 5 : name_at + 8] = b"\xff\xfe\xfd"

        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(bytes(data), ".docx")

    def test_unsupported_zip_version_is_refused(self):
        """An extract version above what zipfile supports makes
        ``infolist()`` raise NotImplementedError."""
        guard = _guard()
        data = bytearray(_zip_bytes({"word/document.xml": b"<w/>"}))
        cd = _central_dir_offset(data)
        struct.pack_into("<B", data, cd + 6, 0xFF)

        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(bytes(data), ".xlsx")


_OLE_HEADER = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504


class TestContainerIsAnchoredAtOffsetZero:
    """zipfile finds the central directory from the end of the file and
    tolerates data in front of the archive; unstructured sniffs the
    leading bytes. A legacy OLE file with a zip appended must not pass
    as a container (it would be routed to an unbounded soffice)."""

    @pytest.mark.parametrize("ext", [".docx", ".pptx"])
    def test_ole_file_with_appended_zip_is_refused(self, ext):
        guard = _guard()
        data = _OLE_HEADER + _zip_bytes({"word/document.xml": b"<w/>"})
        # Premise: zipfile alone accepts it.
        assert zipfile.ZipFile(io.BytesIO(data)).namelist()

        with pytest.raises(guard.DecompressionBombError, match="signature"):
            guard.validate_zip_container(data, ext)

    def test_prepended_junk_is_refused(self):
        guard = _guard()
        data = b"junk" * 16 + _zip_bytes({"word/document.xml": b"<w/>"})

        with pytest.raises(guard.DecompressionBombError, match="signature"):
            guard.validate_zip_container(data, ".docx")

    def test_prepended_data_behind_a_zip_signature_is_refused(self):
        """Junk that itself starts with a local-header signature still
        shifts the first entry off offset 0."""
        guard = _guard()
        data = (
            b"PK\x03\x04"
            + b"\x00" * 60
            + _zip_bytes({"word/document.xml": b"<w/>"})
        )
        assert zipfile.ZipFile(io.BytesIO(data)).namelist()

        with pytest.raises(guard.DecompressionBombError, match="before its"):
            guard.validate_zip_container(data, ".docx")

    def test_normal_docx_is_accepted(self):
        docx = pytest.importorskip("docx")

        buf = io.BytesIO()
        docx.Document().save(buf)

        _guard().validate_zip_container(buf.getvalue(), ".docx")

    @pytest.mark.parametrize("ext", [".odt", ".epub"])
    def test_empty_archive_is_still_accepted(self, ext):
        # An empty .docx/.pptx/.xlsx has no main part and is refused for
        # that (TestPackageTypeMatchesTheExtension); the anchoring check
        # itself accepts the end-of-central-directory record.
        _guard().validate_zip_container(_zip_bytes({}), ext)


_PNG = b"\x89PNG\r\n\x1a\n"

#: Scaled-down ceilings so the tests never materialize real sizes.
_SCALED = {
    "max_entry_bytes": 1_000_000,
    "max_total_bytes": 2_000_000,
    "max_media_entry_bytes": 4_000_000,
    "max_container_bytes": 7_000_000,
}


class TestMediaEntries:
    """Binary media is parsed by nothing (~1x memory), so a large image
    or video must not be refused by the parsed-part ceilings — while a
    large XML part still is, whatever it is named."""

    def test_large_media_entry_is_accepted(self):
        guard = _guard()
        doc = _zip_bytes(
            _ooxml_package(
                {
                    "ppt/slides/slide1.xml": b"<p:sld/>",
                    "ppt/media/image1.png": _PNG + b"\x00" * 3_000_000,
                },
                ".pptx",
            )
        )

        guard.validate_zip_container(doc, ".pptx", **_SCALED)

    def test_large_xml_entry_is_still_refused(self):
        guard = _guard()
        doc = _zip_bytes({"ppt/slides/slide1.xml": b"<a/>" * 750_000})

        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(doc, ".pptx", **_SCALED)

    def test_media_name_does_not_earn_the_media_ceiling(self):
        """Names are attacker-chosen and parts dispatch by content
        type: XML stored under ``media/*.png`` is still a parsed part."""
        guard = _guard()
        doc = _zip_bytes({"ppt/media/image1.png": b"<a/>" * 750_000})

        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(doc, ".pptx", **_SCALED)

    def test_markup_name_keeps_parsed_ceiling_despite_media_bytes(self):
        guard = _guard()
        doc = _zip_bytes({"word/document.xml": _PNG + b"\x00" * 3_000_000})

        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(doc, ".docx", **_SCALED)

    def test_media_entry_over_its_own_ceiling_is_refused(self):
        guard = _guard()
        doc = _zip_bytes(
            {"word/media/v.mp4": b"\x00\x00\x00\x18ftyp" + b"\x00" * 5_000_000}
        )

        with pytest.raises(
            guard.DecompressionBombError, match="ceiling 4000000"
        ):
            guard.validate_zip_container(doc, ".docx", **_SCALED)

    def test_media_does_not_count_toward_the_parsed_total(self):
        guard = _guard()
        doc = _zip_bytes(
            _ooxml_package(
                {
                    "ppt/media/a.png": _PNG + b"\x00" * 3_000_000,
                    "ppt/slides/slide1.xml": b"<a/>" * 200_000,
                    "ppt/slides/slide2.xml": b"<a/>" * 200_000,
                },
                ".pptx",
            )
        )

        guard.validate_zip_container(doc, ".pptx", **_SCALED)

    def test_container_total_still_holds_with_media(self):
        guard = _guard()
        doc = _zip_bytes(
            {
                "ppt/media/a.png": _PNG + b"\x00" * 3_000_000,
                "ppt/media/b.png": _PNG + b"\x00" * 3_000_000,
                "ppt/media/c.png": _PNG + b"\x00" * 3_000_000,
            }
        )

        with pytest.raises(guard.DecompressionBombError, match="total"):
            guard.validate_zip_container(doc, ".pptx", **_SCALED)

    def test_corrupt_media_stream_fails_closed(self):
        """Sniffing decompresses a few bytes; a corrupt deflate stream
        must surface as the guard's typed rejection."""
        guard = _guard()
        data = bytearray(
            _zip_bytes({"ppt/media/a.png": _PNG + b"\x00" * 200_000})
        )
        name_len, extra_len = struct.unpack_from("<HH", data, 26)
        start = 30 + name_len + extra_len
        data[start : start + 8] = b"\xff" * 8

        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(bytes(data), ".pptx", **_SCALED)

    def test_real_pptx_with_large_picture_is_accepted(self):
        pptx = pytest.importorskip("pptx")
        pil = pytest.importorskip("PIL.Image")
        from pptx.util import Inches

        tiny = io.BytesIO()
        pil.new("RGB", (4, 4)).save(tiny, "PNG")
        # PNG readers ignore trailing bytes: pad the picture past the
        # scaled parsed-part ceiling.
        picture = io.BytesIO(tiny.getvalue() + b"\x00" * 2_000_000)
        deck = pptx.Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.add_picture(picture, Inches(1), Inches(1))
        buf = io.BytesIO()
        deck.save(buf)

        _guard().validate_zip_container(buf.getvalue(), ".pptx", **_SCALED)


def _set_declared_size(data: bytes, name: str, size: int) -> bytes:
    """Rewrite *name*'s uncompressed size in both headers (a lie)."""
    out = bytearray(data)
    encoded = name.encode()
    local = 0
    while True:
        local = out.index(b"PK\x03\x04", local)
        name_len = struct.unpack_from("<H", out, local + 26)[0]
        if bytes(out[local + 30 : local + 30 + name_len]) == encoded:
            struct.pack_into("<I", out, local + 22, size)
            break
        local += 4
    central = 0
    while True:
        central = out.index(b"PK\x01\x02", central)
        name_len = struct.unpack_from("<H", out, central + 28)[0]
        if bytes(out[central + 46 : central + 46 + name_len]) == encoded:
            struct.pack_into("<I", out, central + 24, size)
            return bytes(out)
        central += 4


def _method_zip(method: int, payload: bytes) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("word/document.xml", b"<w/>")
        zf.writestr("word/media/blob.bin", payload, compress_type=method)
    return buf.getvalue()


def _optional_methods():
    methods = [
        pytest.param(zipfile.ZIP_BZIP2, id="bzip2"),
        pytest.param(zipfile.ZIP_LZMA, id="lzma"),
    ]
    zstd = getattr(zipfile, "ZIP_ZSTANDARD", None)
    if zstd is not None:
        methods.append(pytest.param(zstd, id="zstd"))
    return methods


class TestCompressionMethodAndHonesty:
    """zipfile decompresses BZIP2/LZMA/ZSTD with no output limit, and
    truncates DEFLATE to the declared size only after inflating: the
    guard allows only STORED/DEFLATE and verifies every declaration.

    Payloads are 64 KiB of zeros, so a regression is detected by the
    assertion, never by exhausting memory."""

    @pytest.mark.parametrize("method", _optional_methods())
    def test_non_deflate_methods_are_refused(self, method):
        guard = _guard()
        try:
            data = _method_zip(method, b"\x00" * 65536)
        except (RuntimeError, NotImplementedError):
            pytest.skip("this Python cannot write that method")
        # Lie about the size, as an attacker would.
        data = _set_declared_size(data, "word/media/blob.bin", 100)

        with pytest.raises(
            guard.DecompressionBombError, match="compression method"
        ):
            guard.validate_zip_container(data, ".docx")

    def test_deflate_entry_inflating_past_its_declaration_is_refused(self):
        guard = _guard()
        data = _zip_bytes({"word/document.xml": b"\x00" * 65536})
        data = _set_declared_size(data, "word/document.xml", 1000)

        with pytest.raises(guard.DecompressionBombError, match="inflates past"):
            guard.validate_zip_container(data, ".docx")

    def test_stored_entry_with_mismatched_sizes_is_refused(self):
        guard = _guard()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:
            zf.writestr("word/document.xml", b"x" * 5000)
        # Within the DEFLATE slack, so only the STORED rule catches it.
        data = _set_declared_size(buf.getvalue(), "word/document.xml", 4500)

        with pytest.raises(
            guard.DecompressionBombError, match="compressed size"
        ):
            guard.validate_zip_container(data, ".docx")

    def test_compressed_range_far_beyond_declaration_is_refused(self):
        """A full read copies the whole compressed range; it must stay
        within a small slack of the declared size."""
        guard = _guard()
        noise = random.Random(0).randbytes(20_000)  # incompressible
        data = _zip_bytes({"word/document.xml": noise})
        data = _set_declared_size(data, "word/document.xml", 10)

        with pytest.raises(
            guard.DecompressionBombError, match="compressed size"
        ):
            guard.validate_zip_container(data, ".docx")

    def test_encrypted_entries_are_refused(self):
        guard = _guard()
        data = bytearray(_zip_bytes({"word/document.xml": b"<w/>"}))
        cd = _central_dir_offset(bytes(data))
        flags = struct.unpack_from("<H", data, cd + 8)[0] | 0x1
        struct.pack_into("<H", data, cd + 8, flags)

        with pytest.raises(guard.DecompressionBombError, match="encrypted"):
            guard.validate_zip_container(bytes(data), ".docx")

    def test_verification_inflates_in_bounded_steps(self):
        guard = _guard()
        calls: list = []
        real = zlib.decompressobj

        class _Recording:
            def __init__(self, *args):
                self._inner = real(*args)

            def decompress(self, data, max_length=0):
                calls.append(max_length)
                return self._inner.decompress(data, max_length)

            def __getattr__(self, name):
                return getattr(self._inner, name)

        data = _zip_bytes({"content.xml": b"<a/>" * 500_000})
        with patch.object(guard.zlib, "decompressobj", _Recording):
            guard.validate_zip_container(data, ".odt")

        assert calls, "entries were not verified"
        assert all(0 < n <= 4 * 1024 * 1024 for n in calls), calls

    def test_honest_deflate_and_stored_entries_pass(self):
        guard = _guard()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("a.xml", b"<a/>" * 10_000, zipfile.ZIP_DEFLATED)
            zf.writestr("b.xml", b"<b/>" * 10_000, zipfile.ZIP_STORED)
            zf.writestr("c.bin", bytes(range(256)) * 100, zipfile.ZIP_DEFLATED)

        guard.validate_zip_container(buf.getvalue(), ".odt")


class TestOoxmlCapDefaults:
    """Structural bounds: generous for real documents, hard ceilings."""

    def test_defaults_hold_realistic_margins(self):
        guard = _guard()

        # Real office documents stay far below these; bombs from a
        # 3 GB-compressed upload cap cannot expand past them.
        # Python XML parse trees cost ~10-20x the XML they parse, so the
        # entry ceiling must keep that product within a worker's memory.
        # (It does not bound what loaders build from the tree, such as
        # pandas' .xlsx grid padding; see the guard's "Known gaps".)
        assert guard.MAX_ENTRY_BYTES <= 128 * 1024 * 1024
        assert guard.MAX_TOTAL_BYTES <= 512 * 1024 * 1024
        assert guard.MAX_ENTRY_BYTES <= guard.MAX_TOTAL_BYTES
        assert guard.MAX_ENTRIES <= 10_000

    def test_media_defaults_hold_realistic_margins(self):
        """Media costs ~1x (measured), so it may be larger than a parsed
        part — but the whole container stays bounded."""
        guard = _guard()

        assert (
            guard.MAX_ENTRY_BYTES
            < guard.MAX_MEDIA_ENTRY_BYTES
            <= 512 * 1024 * 1024
        )
        assert (
            guard.MAX_TOTAL_BYTES
            <= guard.MAX_CONTAINER_BYTES
            <= 1024 * 1024 * 1024
        )
        assert guard.MAX_MEDIA_ENTRY_BYTES <= guard.MAX_CONTAINER_BYTES

    def test_benign_real_documents_pass_with_defaults(self):
        """A genuine docx/xlsx sails through the default ceilings."""
        docx = pytest.importorskip("docx")
        from local_deep_research.document_loaders.bytes_loader import (
            extract_text_from_bytes,
        )

        buf = io.BytesIO()
        document = docx.Document()
        document.add_paragraph("hello ooxml guard")
        document.save(buf)

        text = extract_text_from_bytes(buf.getvalue(), ".docx", "ok.docx")
        assert text is not None and "hello ooxml guard" in text


_XLSX_CT = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Default Extension="png" ContentType="image/png"/>'
    '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
    '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
    '<Override PartName="/xl/chartsheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.chartsheet+xml"/>'
    '<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
    "</Types>"
)
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_RELS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _rels(*rels: tuple[str, str, str]) -> str:
    body = "".join(
        f'<Relationship Id="{rid}" Type="{_REL_NS}/{kind}" Target="{target}"/>'
        for rid, kind, target in rels
    )
    return f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{_PKG_RELS}">{body}</Relationships>'


def _xlsx_with_image_references(references: int, image: bytes) -> bytes:
    """A chartsheet whose drawing references ONE image N times."""
    anchor = (
        '<xdr:absoluteAnchor><xdr:pos x="0" y="0"/><xdr:ext cx="1" cy="1"/>'
        '<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{i}" name="p{i}"/><xdr:cNvPicPr/>'
        '</xdr:nvPicPr><xdr:blipFill><a:blip r:embed="rId1"/></xdr:blipFill>'
        "<xdr:spPr/></xdr:pic><xdr:clientData/></xdr:absoluteAnchor>"
    )
    parts = {
        "[Content_Types].xml": _XLSX_CT,
        "_rels/.rels": _rels(("rId1", "officeDocument", "xl/workbook.xml")),
        "xl/workbook.xml": (
            '<?xml version="1.0" encoding="UTF-8"?><workbook xmlns='
            '"http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'xmlns:r="{_REL_NS}"><sheets>'
            '<sheet name="S1" sheetId="1" r:id="rId1"/>'
            '<sheet name="C1" sheetId="2" r:id="rId2"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": _rels(
            ("rId1", "worksheet", "worksheets/sheet1.xml"),
            ("rId2", "chartsheet", "chartsheets/sheet2.xml"),
        ),
        "xl/worksheets/sheet1.xml": (
            '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns='
            '"http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>hello'
            "</t></is></c></row></sheetData></worksheet>"
        ),
        "xl/chartsheets/sheet2.xml": (
            '<?xml version="1.0" encoding="UTF-8"?><chartsheet xmlns='
            '"http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'xmlns:r="{_REL_NS}"><sheetViews><sheetView workbookViewId="0"/>'
            '</sheetViews><drawing r:id="rId1"/></chartsheet>'
        ),
        "xl/chartsheets/_rels/sheet2.xml.rels": _rels(
            ("rId1", "drawing", "../drawings/drawing1.xml")
        ),
        "xl/drawings/drawing1.xml": (
            '<?xml version="1.0" encoding="UTF-8"?><xdr:wsDr xmlns:xdr='
            '"http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            f'xmlns:r="{_REL_NS}">'
            + "".join(anchor.format(i=i + 2) for i in range(references))
            + "</xdr:wsDr>"
        ),
        "xl/drawings/_rels/drawing1.xml.rels": _rels(
            ("rId1", "image", "../media/image1.png")
        ),
    }
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, text in parts.items():
            zf.writestr(name, text)
        zf.writestr("xl/media/image1.png", image)
    return buf.getvalue()


def _small_png() -> bytes:
    pil = pytest.importorskip("PIL.Image")
    out = io.BytesIO()
    pil.new("RGB", (2, 2)).save(out, "PNG")
    return out.getvalue() + b"\x00" * 65536


@contextmanager
def _counting_reads(member: str):
    reads: list = []
    real = zipfile.ZipFile.read

    def spy(self, name, *args, **kwargs):
        target = name.filename if isinstance(name, zipfile.ZipInfo) else name
        if target == member:
            reads.append(target)
        return real(self, name, *args, **kwargs)

    with patch.object(zipfile.ZipFile, "read", spy):
        yield reads


class TestReferenceFanOut:
    """One entry referenced N times must not be read N times.

    openpyxl reads a drawing image once per reference and keeps every
    copy, so a ~100 KB .xlsx referencing one large image N times costs
    N times its size — no per-entry ceiling bounds that. The .xlsx path
    disables openpyxl's drawing reader. python-pptx/python-docx load
    each part once, so they are pinned here too."""

    REFS = 8

    def test_premise_openpyxl_reads_an_image_once_per_reference(self):
        """Control: with openpyxl's own hook, reads scale with refs."""
        openpyxl = pytest.importorskip("openpyxl")
        import openpyxl.reader.excel as excel_reader
        from openpyxl.reader.drawings import find_images

        data = _xlsx_with_image_references(self.REFS, _small_png())
        with (
            patch.object(excel_reader, "find_images", find_images),
            _counting_reads("xl/media/image1.png") as reads,
        ):
            openpyxl.load_workbook(io.BytesIO(data), read_only=True)

        assert len(reads) == self.REFS

    def test_xlsx_upload_never_reads_drawing_images(self):
        pytest.importorskip("openpyxl")
        import openpyxl.reader.excel as excel_reader
        from openpyxl.reader.drawings import find_images

        from local_deep_research.document_loaders import (
            extract_text_from_bytes,
        )
        from local_deep_research.document_loaders.loader_registry import (
            LOADER_REGISTRY,
        )

        if ".xlsx" not in LOADER_REGISTRY:
            pytest.skip(".xlsx not registered in this environment")
        data = _xlsx_with_image_references(self.REFS, _small_png())
        # Start from openpyxl's own hook so the upload path must apply
        # the hardening itself.
        with (
            patch.object(excel_reader, "find_images", find_images),
            _counting_reads("xl/media/image1.png") as reads,
        ):
            text = extract_text_from_bytes(data, ".xlsx", "amp.xlsx")

        assert text is not None and "hello" in text
        assert reads == [], f"image read {len(reads)}x for {self.REFS} refs"

    def test_xlsx_upload_rejects_when_drawing_hook_disappears(
        self, monkeypatch
    ):
        """An openpyxl upgrade must not restore unbounded drawing reads."""
        pytest.importorskip("openpyxl")
        import openpyxl.reader.excel as excel_reader

        from local_deep_research.document_loaders.bytes_loader import (
            load_from_bytes,
        )
        from local_deep_research.document_loaders.zip_container_guard import (
            DecompressionBombError,
        )

        data = _xlsx_with_image_references(1, _small_png())
        monkeypatch.delattr(excel_reader, "find_images")
        with pytest.raises(
            DecompressionBombError, match="drawing guard unavailable"
        ):
            load_from_bytes(data, ".xlsx", "untrusted.xlsx")

    def test_openpyxl_still_exposes_the_hook(self):
        """Pin the openpyxl drawing hook used by this hardening."""
        pytest.importorskip("openpyxl")
        import openpyxl.reader.excel as excel_reader

        from local_deep_research.document_loaders import openpyxl_hardening

        assert hasattr(excel_reader, "find_images")
        assert openpyxl_hardening.disable_openpyxl_drawing_reads() is True

    @pytest.mark.parametrize("ext", [".pptx", ".docx"])
    def test_pptx_and_docx_read_a_shared_image_once(self, ext):
        from local_deep_research.document_loaders import (
            extract_text_from_bytes,
        )

        png = _small_png()
        buf = io.BytesIO()
        if ext == ".pptx":
            pptx = pytest.importorskip("pptx")
            from pptx.util import Inches

            deck = pptx.Presentation()
            for i in range(self.REFS):
                slide = deck.slides.add_slide(deck.slide_layouts[5])
                slide.shapes.title.text = f"slide {i}"
                slide.shapes.add_picture(io.BytesIO(png), Inches(1), Inches(1))
            deck.save(buf)
            member = "ppt/media/image1.png"
        else:
            docx = pytest.importorskip("docx")
            document = docx.Document()
            for i in range(self.REFS):
                document.add_paragraph(f"para {i}")
                document.add_picture(io.BytesIO(png))
            document.save(buf)
            member = "word/media/image1.png"

        with _counting_reads(member) as reads:
            text = extract_text_from_bytes(buf.getvalue(), ext, f"x{ext}")

        assert text is not None
        assert len(reads) <= 1, f"{member} read {len(reads)}x"


_SHEET_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_WORKBOOK_CT = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
)


def _external(rels_xml: str, rel_ids: frozenset[str]) -> str:
    """Mark the given relationship ids ``TargetMode="External"``."""
    for rel_id in rel_ids:
        rels_xml = rels_xml.replace(
            f'Id="{rel_id}" ', f'Id="{rel_id}" TargetMode="External" '
        )
    return rels_xml


def _xlsx_with_sheets(
    sheet_ids: list[str],
    rels: list[tuple[str, str]],
    workbook: str = "xl/workbook.xml",
    override: bool = True,
    external: frozenset[str] = frozenset(),
    workbook_decl: str = '<?xml version="1.0" encoding="UTF-8"?>',
) -> bytes:
    """Workbook whose ``<sheet>`` r:ids and relationship targets are
    given explicitly; ``xl/worksheets/sheet{1,2}.xml`` exist."""
    sheets = "".join(
        f'<sheet name="S{i}" sheetId="{i + 1}" r:id="{rid}"/>'
        for i, rid in enumerate(sheet_ids)
    )
    worksheet = (
        '<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{ns}">'
        '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>{text}</t>'
        "</is></c></row></sheetData></worksheet>"
    )
    parts = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8"?><Types xmlns='
            '"http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/'
            'vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            + (
                f'<Override PartName="/{workbook}" ContentType="{_WORKBOOK_CT}"/>'
                if override
                else ""
            )
            + "</Types>"
        ),
        "_rels/.rels": _rels(("rId1", "officeDocument", workbook)),
        workbook: (
            f'{workbook_decl}<workbook xmlns="{_SHEET_MAIN}" '
            f'xmlns:r="{_REL_NS}"><sheets>{sheets}</sheets></workbook>'
        ),
        posixpath.join(
            posixpath.dirname(workbook),
            "_rels",
            posixpath.basename(workbook) + ".rels",
        ): _external(
            _rels(*((rid, "worksheet", target) for rid, target in rels)),
            external,
        ),
        "xl/worksheets/sheet1.xml": worksheet.format(
            ns=_SHEET_MAIN, text="one"
        ),
        "xl/worksheets/sheet2.xml": worksheet.format(
            ns=_SHEET_MAIN, text="two"
        ),
    }
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, text in parts.items():
            zf.writestr(name, text)
    return buf.getvalue()


class TestSheetFanOut:
    """openpyxl converts a worksheet once per ``<sheet>`` that resolves
    to it: N entries sharing one part multiply parsing work and text by
    N. Genuine workbooks never share a sheet part, so it is refused."""

    def test_repeated_rid_is_refused(self):
        guard = _guard()
        data = _xlsx_with_sheets(
            ["rId1"] * 4, [("rId1", "worksheets/sheet1.xml")]
        )

        with pytest.raises(guard.DecompressionBombError, match="same sheet"):
            guard.validate_zip_container(data, ".xlsx")

    @pytest.mark.parametrize(
        "second_target",
        [
            "worksheets/sheet1.xml",
            "/xl/worksheets/sheet1.xml",
            "./worksheets/../worksheets/sheet1.xml",
        ],
    )
    def test_distinct_rels_to_one_target_are_refused(self, second_target):
        """openpyxl resolves each relationship target itself (relative,
        ``/``-rooted, ``./..`` forms); two that resolve to the same
        member name are one part read twice."""
        guard = _guard()
        data = _xlsx_with_sheets(
            ["rId1", "rId2"],
            [("rId1", "worksheets/sheet1.xml"), ("rId2", second_target)],
        )

        with pytest.raises(guard.DecompressionBombError, match="same sheet"):
            guard.validate_zip_container(data, ".xlsx")

    def test_workbook_located_through_content_types_is_checked(self):
        """openpyxl finds the workbook by content type, so a workbook at
        a non-default path is checked too."""
        guard = _guard()
        # With an xl/workbook.xml that openpyxl does not read, which
        # pandas needs to pick openpyxl at all.
        data = _rewrite_zip(
            _xlsx_with_sheets(
                ["rId1"] * 3,
                [("rId1", "worksheets/sheet1.xml")],
                workbook="xl/book.xml",
            ),
            add={"xl/workbook.xml": b"<unread/>"},
        )

        with pytest.raises(guard.DecompressionBombError, match="same sheet"):
            guard.validate_zip_container(data, ".xlsx")

    def test_workbook_without_a_workbook_override_is_refused(self):
        """openpyxl cannot locate this workbook (no workbook content
        type), so the guard cannot resolve its sheets: fail closed."""
        guard = _guard()
        data = _xlsx_with_sheets(
            ["rId1"] * 3, [("rId1", "worksheets/sheet1.xml")], override=False
        )

        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(data, ".xlsx")

    def test_repeated_relationship_id_is_refused(self):
        guard = _guard()
        data = _xlsx_with_sheets(
            ["rId1", "rId2"],
            [
                ("rId1", "worksheets/sheet1.xml"),
                ("rId1", "worksheets/sheet2.xml"),
                ("rId2", "worksheets/sheet2.xml"),
            ],
        )

        # openpyxl keeps the last duplicate Id, so both sheets resolve
        # to sheet2.xml.
        with pytest.raises(guard.DecompressionBombError, match="same sheet"):
            guard.validate_zip_container(data, ".xlsx")

    def test_distinct_sheets_are_accepted(self):
        guard = _guard()
        data = _xlsx_with_sheets(
            ["rId1", "rId2"],
            [
                ("rId1", "worksheets/sheet1.xml"),
                ("rId2", "worksheets/sheet2.xml"),
            ],
        )

        guard.validate_zip_container(data, ".xlsx")

    def test_upload_path_refuses_fan_out_and_keeps_real_workbooks(self):
        openpyxl = pytest.importorskip("openpyxl")
        from local_deep_research.document_loaders import (
            extract_text_from_bytes,
        )
        from local_deep_research.document_loaders.loader_registry import (
            LOADER_REGISTRY,
        )

        if ".xlsx" not in LOADER_REGISTRY:
            pytest.skip(".xlsx not registered in this environment")

        fan_out = _xlsx_with_sheets(
            ["rId1"] * 4, [("rId1", "worksheets/sheet1.xml")]
        )
        assert extract_text_from_bytes(fan_out, ".xlsx", "fan.xlsx") is None

        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "alpha"
        for i in range(3):
            workbook.create_sheet(f"S{i}")["A1"] = f"beta{i}"
        buf = io.BytesIO()
        workbook.save(buf)
        text = extract_text_from_bytes(buf.getvalue(), ".xlsx", "ok.xlsx")
        assert text is not None
        assert "alpha" in text and "beta2" in text


class TestSheetFanOutExternalAndEncoding:
    def test_external_rel_naming_the_internal_part_is_refused(self):
        """openpyxl keeps an External target unresolved and still loads
        it when the raw string names a package part."""
        guard = _guard()
        data = _xlsx_with_sheets(
            ["rId1", "rId2"],
            [
                ("rId1", "worksheets/sheet1.xml"),
                ("rId2", "xl/worksheets/sheet1.xml"),
            ],
            external=frozenset({"rId2"}),
        )

        with pytest.raises(guard.DecompressionBombError, match="same sheet"):
            guard.validate_zip_container(data, ".xlsx")

    def test_bogus_xml_encoding_fails_closed(self):
        guard = _guard()
        data = _xlsx_with_sheets(
            ["rId1"],
            [("rId1", "worksheets/sheet1.xml")],
            workbook_decl='<?xml version="1.0" encoding="no-such-codec"?>',
        )

        with pytest.raises(
            guard.DecompressionBombError, match="not encoded in UTF-8"
        ):
            guard.validate_zip_container(data, ".xlsx")


def _document_zip(name: str, body: bytes) -> bytes:
    return _zip_bytes(
        _ooxml_package(
            {"word/_rels/document.xml.rels": b"<r/>", name: body}, ".docx"
        )
    )


class TestMediaClassifierIsAnchored:
    """Only bytes no XML parser accepts as a document start may earn
    the media ceiling."""

    MARKUP = b"<w:document>" + b"<w:p>payload</w:p>" * 200_000

    @pytest.mark.parametrize(
        "head",
        [
            pytest.param(b"<!--ftyp-->", id="comment-ftyp"),
            pytest.param(b"  \n<?xml version='1.0'?>", id="whitespace-decl"),
            pytest.param(b"\xef\xbb\xbf<", id="utf8-bom"),
            pytest.param("<!--x-->".encode("utf-16-le"), id="utf16le-no-bom"),
            pytest.param("﻿<a".encode("utf-16-be"), id="utf16be-bom"),
            pytest.param("<a".encode("utf-32-le"), id="utf32le-no-bom"),
        ],
    )
    def test_markup_head_stays_a_parsed_part(self, head):
        guard = _guard()
        data = _document_zip("word/document.bin", head + self.MARKUP)

        with pytest.raises(guard.DecompressionBombError, match="ceiling"):
            guard.validate_zip_container(data, ".docx", **_SCALED)

    def test_ftyp_needs_a_plausible_box_size(self):
        guard = _guard()
        head = b"\x00\x00\x00\x04ftypisom"  # box size 4 < 8
        data = _document_zip("word/media/v.mp4", head + b"\x00" * 3_000_000)

        with pytest.raises(guard.DecompressionBombError, match="ceiling"):
            guard.validate_zip_container(data, ".docx", **_SCALED)

    def test_ftyp_without_a_box_shape_is_not_media(self):
        """``ftyp`` at offset 4 alone is not enough: the head must start
        with the two zero bytes of a plausible box size."""
        guard = _guard()
        data = _document_zip(
            "word/media/v.mp4", b"ABCDftypisom" + b"\x00" * 3_000_000
        )

        with pytest.raises(guard.DecompressionBombError, match="ceiling"):
            guard.validate_zip_container(data, ".docx", **_SCALED)

    def test_mp4_with_a_60_byte_ftyp_box_is_media(self):
        """00 00 00 3C reads as "<" in UTF-32BE, but "ftyp" as UCS-4 is
        above U+10FFFF, so no XML parser accepts this head: it must keep
        the media ceiling."""
        guard = _guard()
        head = b"\x00\x00\x00\x3cftypisom"
        data = _document_zip("word/media/v.mp4", head + b"\x00" * 3_000_000)

        guard.validate_zip_container(data, ".docx", **_SCALED)

    def test_real_mp4_head_is_still_media(self):
        guard = _guard()
        head = b"\x00\x00\x00\x18ftypisom"
        data = _document_zip("word/media/v.mp4", head + b"\x00" * 3_000_000)

        guard.validate_zip_container(data, ".docx", **_SCALED)


class TestDeclarationChecksBeforeInflation:
    def test_oversized_declaration_is_refused_without_inflating(self):
        guard = _guard()
        data = _zip_bytes({"word/media/a.png": _PNG + b"\x00" * 5_000_000})

        def boom(*args, **kwargs):
            raise AssertionError("entry inflated before its size was checked")

        with patch.object(guard, "_inflate_head", boom):
            with pytest.raises(guard.DecompressionBombError, match="ceiling"):
                guard.validate_zip_container(data, ".docx", **_SCALED)

    def test_entries_sharing_a_local_header_are_refused(self):
        guard = _guard()
        data = bytearray(
            _zip_bytes({"word/a.xml": b"<a/>", "word/b.xml": b"<b/>"})
        )
        first = data.index(b"PK\x01\x02")
        second = data.index(b"PK\x01\x02", first + 4)
        offset = struct.unpack_from("<I", data, first + 42)[0]
        struct.pack_into("<I", data, second + 42, offset)

        with pytest.raises(guard.DecompressionBombError, match="local header"):
            guard.validate_zip_container(bytes(data), ".docx")

    def test_local_extra_field_is_honoured(self):
        """The data offset uses the LOCAL header's extra length, which
        may differ from the central directory's (as zipfile reads it)."""
        guard = _guard()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            info = zipfile.ZipInfo("word/document.xml")
            info.compress_type = zipfile.ZIP_DEFLATED
            info.extra = struct.pack("<HH", 0xCAFE, 12) + b"x" * 12
            zf.writestr(info, b"<w:document>" + b"<w:p/>" * 5000)
        data = bytearray(buf.getvalue())
        # Drop the extra field from the central record only: 16 bytes in
        # the local header, 0 in the central directory.
        central = data.index(b"PK\x01\x02")
        name_len, extra_len = struct.unpack_from("<HH", data, central + 28)
        assert extra_len == 16
        struct.pack_into("<H", data, central + 30, 0)
        start = central + 46 + name_len
        del data[start : start + extra_len]
        eocd = data.rindex(b"PK\x05\x06")
        size = struct.unpack_from("<I", data, eocd + 12)[0]
        struct.pack_into("<I", data, eocd + 12, size - extra_len)
        assert (
            zipfile.ZipFile(io.BytesIO(bytes(data)))
            .read("word/document.xml")
            .startswith(b"<w:document>")
        )

        # .odt: a lone .docx part without its package is refused for
        # its layout (TestPackageTypeMatchesTheExtension).
        guard.validate_zip_container(bytes(data), ".odt")


class TestHardeningIdempotence:
    def test_reloaded_hook_is_patched_again(self):
        """Identity, not a flag: a restored original hook (as after a
        reload of openpyxl) is replaced again."""
        pytest.importorskip("openpyxl")
        import openpyxl.reader.excel as excel_reader
        from openpyxl.reader.drawings import find_images

        from local_deep_research.document_loaders import openpyxl_hardening

        openpyxl_hardening.disable_openpyxl_drawing_reads()
        with patch.object(excel_reader, "find_images", find_images):
            assert openpyxl_hardening.disable_openpyxl_drawing_reads() is True
            assert excel_reader.find_images is openpyxl_hardening._no_drawings


def _openpyxl_shaped_xlsx(
    sheets_xml: str,
    rels_xml: str,
    workbook_entry: str = "xl/workbook.xml",
    workbook_partname: str = "/xl/workbook.xml",
    rels_prefix: bytes = b"",
    extra: dict | None = None,
) -> bytes:
    """Hand-built workbook with full control over tags and paths."""
    parts = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8"?><Types xmlns='
            '"http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/'
            'vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            f'<Override PartName="{workbook_partname}" '
            f'ContentType="{_WORKBOOK_CT}"/></Types>'
        ).encode(),
        "_rels/.rels": (
            f'<Relationships xmlns="{_PKG_RELS}"><Relationship Id="rId1" '
            f'Type="{_REL_NS}/officeDocument" Target="{workbook_entry}"/>'
            "</Relationships>"
        ).encode(),
        workbook_entry: (
            f'<workbook xmlns="{_SHEET_MAIN}" xmlns:r="{_REL_NS}">'
            f"<sheets>{sheets_xml}</sheets></workbook>"
        ).encode(),
        posixpath.join(
            posixpath.dirname(workbook_entry),
            "_rels",
            posixpath.basename(workbook_entry) + ".rels",
        ): rels_prefix
        + f'<Relationships xmlns="{_PKG_RELS}">{rels_xml}</Relationships>'.encode(),
    }
    sheet = (
        f'<worksheet xmlns="{_SHEET_MAIN}"><sheetData><row r="1"><c r="A1" '
        't="inlineStr"><is><t>{}</t></is></c></row></sheetData></worksheet>'
    )
    parts["xl/worksheets/sheet1.xml"] = sheet.format("one").encode()
    parts["xl/worksheets/sheet2.xml"] = sheet.format("two").encode()
    parts.update(extra or {})
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, body in parts.items():
            zf.writestr(name, body)
    return buf.getvalue()


def _ws_rel(rel_id: str, target: str, tag: str = "Relationship") -> str:
    return (
        f'<{tag} Id="{rel_id}" Type="{_REL_NS}/worksheet" Target="{target}"/>'
    )


def _sheet(n: int, rel_id: str, attr: str = "r:id", tag: str = "sheet") -> str:
    return f'<{tag} name="s{n}" sheetId="{n}" {attr}="{rel_id}"/>'


class TestSheetFanOutFollowsOpenpyxl:
    """Shapes where a re-implemented parser and openpyxl disagreed: the
    guard now asks openpyxl itself which part each sheet names."""

    @staticmethod
    def _assert_refused(data: bytes) -> None:
        guard = _guard()
        with pytest.raises(guard.DecompressionBombError, match="same sheet"):
            guard.validate_zip_container(data, ".xlsx")

    def test_relationship_child_with_any_tag(self):
        pytest.importorskip("openpyxl")
        self._assert_refused(
            _openpyxl_shaped_xlsx(
                _sheet(1, "rId1") + _sheet(2, "rId2"),
                _ws_rel("rId1", "worksheets/sheet1.xml")
                + _ws_rel("rId2", "worksheets/sheet1.xml", tag="Foo"),
            )
        )

    def test_sheet_child_with_any_tag(self):
        pytest.importorskip("openpyxl")
        self._assert_refused(
            _openpyxl_shaped_xlsx(
                _sheet(1, "rId1") + _sheet(2, "rId1", tag="foo"),
                _ws_rel("rId1", "worksheets/sheet1.xml"),
            )
        )

    def test_plain_id_attribute(self):
        pytest.importorskip("openpyxl")
        self._assert_refused(
            _openpyxl_shaped_xlsx(
                "".join(_sheet(i, "rId1", attr="id") for i in range(1, 5)),
                _ws_rel("rId1", "worksheets/sheet1.xml"),
            )
        )

    def test_workbook_partname_without_leading_slash(self):
        """openpyxl derives the part as ``PartName[1:]``."""
        pytest.importorskip("openpyxl")
        self._assert_refused(
            _openpyxl_shaped_xlsx(
                _sheet(1, "rId1") + _sheet(2, "rId1"),
                _ws_rel("rId1", "/xl/worksheets/sheet1.xml"),
                workbook_entry="l/workbook.xml",
                workbook_partname="xl/workbook.xml",
                # Read by pandas' engine choice only, not by openpyxl.
                extra={"xl/workbook.xml": b"<unread/>"},
            )
        )

    def test_bom_and_declared_encoding_disagree(self):
        """UTF-8 BOM with a Latin-1 declaration: the relationships part
        is refused for its declared encoding before openpyxl parses it
        (an encoding other than UTF-8 is refused in every part openpyxl
        parses)."""
        pytest.importorskip("openpyxl")
        guard = _guard()
        with pytest.raises(
            guard.DecompressionBombError, match="not encoded in UTF-8"
        ):
            guard.validate_zip_container(
                self._bom_and_latin1_workbook(), ".xlsx"
            )

    @staticmethod
    def _bom_and_latin1_workbook() -> bytes:
        return _openpyxl_shaped_xlsx(
            _sheet(1, "rId1") + _sheet(2, "rId2"),
            _ws_rel("rId1", "worksheets/&#233;.xml")
            + _ws_rel("rId2", "worksheets/\u00e9.xml"),
            rels_prefix=(
                b'\xef\xbb\xbf<?xml version="1.0" encoding="ISO-8859-1"?>'
            ),
            extra={
                "xl/worksheets/\u00e9.xml": (
                    f'<worksheet xmlns="{_SHEET_MAIN}"><sheetData/></worksheet>'
                ).encode()
            },
        )

    def test_openpyxl_failure_is_refused_as_malformed(self):
        pytest.importorskip("openpyxl")
        guard = _guard()
        data = _openpyxl_shaped_xlsx(
            _sheet(1, "rId9"), _ws_rel("rId1", "worksheets/sheet1.xml")
        )

        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(data, ".xlsx")


class TestSheetFanOutCheckReadsOnlyThePackageSkeleton:
    def test_no_worksheet_drawing_or_string_part_is_opened(self):
        """The fan-out check parses [Content_Types].xml, the workbook
        and its relationships, and streams the styles part — nothing a
        sheet could amplify. (The namespace scan before it streams every
        member once, building no tree; see TestXlsxReadPathStringWork.)"""
        openpyxl = pytest.importorskip("openpyxl")
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "alpha"
        workbook.create_sheet("S1")["A1"] = "beta"
        buf = io.BytesIO()
        workbook.save(buf)
        guard = _guard()

        opened: list = []
        real_open = zipfile.ZipFile.open

        def spy(self, name, *args, **kwargs):
            opened.append(
                name.filename if isinstance(name, zipfile.ZipInfo) else name
            )
            return real_open(self, name, *args, **kwargs)

        # The guard's own verification inflates entries from the raw
        # bytes, so every ZipFile.open during validation is openpyxl's.
        with (
            patch.object(zipfile.ZipFile, "open", spy),
            patch.object(guard, "_check_xlsx_namespaces", lambda *_a: None),
        ):
            guard.validate_zip_container(buf.getvalue(), ".xlsx")

        assert set(opened) <= {
            "[Content_Types].xml",
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
            "xl/styles.xml",
        }, opened


class TestCanonicalMemberNames:
    """Member names that readers normalise onto another path are
    refused before anything is inflated; real writers never emit them."""

    @pytest.mark.parametrize(
        "name",
        [
            pytest.param("xl//worksheets/sheet3.xml", id="double-slash"),
            pytest.param("xl/./worksheets/sheet3.xml", id="dot-segment"),
            pytest.param("xl/../xl/worksheets/sheet3.xml", id="dotdot"),
            pytest.param("/xl/worksheets/sheet3.xml", id="leading-slash"),
            pytest.param("xl\\worksheets\\sheet3.xml", id="backslash"),
        ],
    )
    def test_non_canonical_name_is_refused(self, name):
        guard = _guard()
        data = _zip_bytes({"word/document.xml": b"<w/>", name: b"<a/>"})

        with pytest.raises(guard.DecompressionBombError, match="non-canonical"):
            guard.validate_zip_container(data, ".docx")

    def test_duplicate_member_name_is_refused(self):
        guard = _guard()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("word/document.xml", b"<w/>")
            with pytest.warns(UserWarning):
                zf.writestr("word/document.xml", b"<w2/>")

        with pytest.raises(guard.DecompressionBombError, match="same name"):
            guard.validate_zip_container(buf.getvalue(), ".docx")

    def test_directory_entries_are_allowed(self):
        guard = _guard()
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("OEBPS/", b"")
            zf.writestr("OEBPS/c.xhtml", b"<html/>")

        guard.validate_zip_container(buf.getvalue(), ".epub")

    @pytest.mark.parametrize("fmt", ["docx", "pptx", "odt", "epub"])
    def test_pandoc_written_containers_pass(self, fmt, tmp_path):
        pypandoc = pytest.importorskip("pypandoc")
        try:
            pypandoc.get_pandoc_path()
        except OSError:
            pytest.skip("pandoc not available")
        out = tmp_path / f"doc.{fmt}"
        pypandoc.convert_text(
            "# Title\n\nhello\n", fmt, format="md", outputfile=str(out)
        )

        _guard().validate_zip_container(out.read_bytes(), f".{fmt}")


class TestSheetRelationshipsPartsAreUnique:
    def test_sheets_sharing_one_rels_part_are_refused(self):
        """openpyxl finds each sheet's relationships via get_rels_path,
        which can map distinct targets onto one part; the check uses
        that same function. (Pass 1 already refuses the member names
        that could produce this, so it is isolated here.)"""
        pytest.importorskip("openpyxl")
        guard = _guard()
        data = _openpyxl_shaped_xlsx(
            _sheet(1, "rId1") + _sheet(2, "rId2"),
            _ws_rel("rId1", "worksheets/sheet1.xml")
            + _ws_rel("rId2", "worksheets/sheet2.xml"),
        )
        from openpyxl.packaging import relationship

        with patch.object(
            relationship,
            "get_rels_path",
            lambda path: "xl/worksheets/_rels/shared.xml.rels",
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="relationships part"
            ):
                guard.validate_zip_container(data, ".xlsx")


class TestMediaCeilingIsOoxmlOnly:
    def test_epub_entry_with_media_head_stays_a_parsed_part(self):
        """pandoc reads an EPUB chapter by its manifest media type, not
        its bytes, so no EPUB/ODT entry earns the media ceiling."""
        guard = _guard()
        chapter = b"\x00\x00\x00\x18ftypisom" + b"\x00" * 3_000_000
        data = _zip_bytes(
            {"mimetype": b"application/epub+zip", "c.bin": chapter}
        )

        with pytest.raises(guard.DecompressionBombError, match="ceiling"):
            guard.validate_zip_container(data, ".epub", **_SCALED)

    def test_same_entry_in_ooxml_is_media(self):
        guard = _guard()
        chapter = b"\x00\x00\x00\x18ftypisom" + b"\x00" * 3_000_000
        data = _zip_bytes(
            _ooxml_package({"word/media/v.mp4": chapter}, ".docx")
        )

        guard.validate_zip_container(data, ".docx", **_SCALED)


class TestXlsxSkeletonCeiling:
    @pytest.mark.parametrize(
        ("part", "unparsed_step"),
        [
            ("[Content_Types].xml", "read_manifest"),
            ("xl/workbook.xml", "read_workbook"),
            ("xl/_rels/workbook.xml.rels", "read_workbook"),
        ],
    )
    def test_each_oversized_skeleton_part_is_refused_before_parsing(
        self, part, unparsed_step
    ):
        """Only *part* is padded past the ceiling (the other two stay
        under it), and the openpyxl step that would parse it never runs."""
        pytest.importorskip("openpyxl")
        from openpyxl.reader.excel import ExcelReader

        guard = _guard()
        base = _openpyxl_shaped_xlsx(
            _sheet(1, "rId1"), _ws_rel("rId1", "worksheets/sheet1.xml")
        )
        padding = b"<!--" + b"x" * 10_000 + b"-->"
        src = zipfile.ZipFile(io.BytesIO(base))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
            for info in src.infolist():
                body = src.read(info)
                if info.filename == part:
                    body += padding
                out.writestr(info.filename, body)
        data = buf.getvalue()
        sizes = {
            i.filename: i.file_size
            for i in zipfile.ZipFile(io.BytesIO(data)).infolist()
        }
        ceiling = 5_000
        skeleton = [
            "[Content_Types].xml",
            "xl/workbook.xml",
            "xl/_rels/workbook.xml.rels",
        ]
        assert sizes[part] > ceiling
        assert all(
            sizes[other] < ceiling for other in skeleton if other != part
        )

        calls: list = []
        real_step = getattr(ExcelReader, unparsed_step)

        def recording(self, *args, **kwargs):
            calls.append(unparsed_step)
            return real_step(self, *args, **kwargs)

        with (
            patch.object(guard, "MAX_XLSX_SKELETON_PART_BYTES", ceiling),
            patch.object(ExcelReader, unparsed_step, recording),
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="larger than"
            ):
                guard.validate_zip_container(data, ".xlsx")

        assert calls == [], f"{unparsed_step} parsed the oversized {part}"

    def test_skeleton_ceiling_is_generous_for_real_workbooks(self):
        guard = _guard()

        assert 1024 * 1024 <= guard.MAX_XLSX_SKELETON_PART_BYTES
        assert guard.MAX_XLSX_SKELETON_PART_BYTES <= guard.MAX_ENTRY_BYTES // 4


class TestFanOutCheckImportPlacement:
    def test_missing_openpyxl_helper_refuses_rather_than_skips(
        self, monkeypatch
    ):
        """The helper imports sit inside the fail-closed block: if an
        openpyxl upgrade drops one, .xlsx uploads are refused as
        malformed instead of the fan-out check silently switching off."""
        pytest.importorskip("openpyxl")
        from openpyxl.packaging import relationship

        guard = _guard()
        data = _openpyxl_shaped_xlsx(
            _sheet(1, "rId1") + _sheet(2, "rId2"),
            _ws_rel("rId1", "worksheets/sheet1.xml")
            + _ws_rel("rId2", "worksheets/sheet2.xml"),
        )
        guard.validate_zip_container(data, ".xlsx")  # control: accepted

        monkeypatch.delattr(relationship, "get_rels_path")
        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(data, ".xlsx")


# --- Local header vs central directory ------------------------------------
#
# pandoc's zip reader takes an entry's method and compressed size from
# the LOCAL header, zipfile from the central directory. An entry whose
# two headers disagree is checked here as one stream and inflated by
# pandoc as another, so the headers must agree and the DEFLATE stream
# must end exactly at its declared compressed size.

#: Local file header field offsets (APPNOTE 4.3.7).
_LH_FLAGS, _LH_METHOD, _LH_CRC, _LH_CSIZE, _LH_USIZE = 6, 8, 14, 18, 22
#: Central directory record field offsets (APPNOTE 4.3.12).
_CD_CRC, _CD_CSIZE, _CD_USIZE = 16, 20, 24


def _noisy_xml(n: int = 4000) -> bytes:
    """Poorly compressible markup (KB-scale), so a prefix of its
    DEFLATE stream is a genuinely truncated stream."""
    rng = random.Random(6515)
    return b"<p>" + bytes(rng.randrange(97, 123) for _ in range(n)) + b"</p>"


def _single_entry(name: str, payload: bytes) -> bytearray:
    return bytearray(_zip_bytes({name: payload}))


def _inflated_len(stream: bytes) -> int:
    return len(zlib.decompressobj(-zlib.MAX_WBITS).decompress(stream))


class TestLocalHeaderAgreesWithCentralDirectory:
    def _data_start(self, data: bytes) -> int:
        name_len, extra_len = struct.unpack_from("<HH", data, 26)
        return 30 + name_len + extra_len

    def test_central_prefix_of_a_longer_local_stream_is_refused(self):
        """The reviewer's shape: an honest-looking central record that
        covers only a prefix of the stream the local header declares."""
        guard = _guard()
        data = _single_entry("content.xml", _noisy_xml())
        cd = _central_dir_offset(data)
        start = self._data_start(data)
        prefix = 400
        produced = _inflated_len(bytes(data[start : start + prefix]))
        struct.pack_into("<I", data, cd + _CD_CSIZE, prefix)
        struct.pack_into("<I", data, cd + _CD_USIZE, produced)

        with pytest.raises(guard.DecompressionBombError, match="disagrees"):
            guard.validate_zip_container(bytes(data), ".odt")

    def test_stream_truncated_inside_its_range_is_refused(self):
        """Both headers agree on a range that cuts the stream short:
        verification must require the end of the DEFLATE stream."""
        guard = _guard()
        data = _single_entry("content.xml", _noisy_xml())
        cd = _central_dir_offset(data)
        start = self._data_start(data)
        prefix = 400
        produced = _inflated_len(bytes(data[start : start + prefix]))
        for base, csize, usize in (
            (0, _LH_CSIZE, _LH_USIZE),
            (cd, _CD_CSIZE, _CD_USIZE),
        ):
            struct.pack_into("<I", data, base + csize, prefix)
            struct.pack_into("<I", data, base + usize, produced)

        with pytest.raises(guard.DecompressionBombError, match="truncated"):
            guard.validate_zip_container(bytes(data), ".odt")

    def test_trailing_bytes_inside_the_range_are_refused(self):
        guard = _guard()
        data = _single_entry("content.xml", b"<p>hello</p>")
        cd = _central_dir_offset(data)
        for base, csize in ((0, _LH_CSIZE), (cd, _CD_CSIZE)):
            old = struct.unpack_from("<I", data, base + csize)[0]
            struct.pack_into("<I", data, base + csize, old + 2)

        with pytest.raises(guard.DecompressionBombError, match="ends before"):
            guard.validate_zip_container(bytes(data), ".odt")

    def test_short_stream_is_refused(self):
        """Fewer bytes than declared means the declaration is a lie."""
        guard = _guard()
        data = _single_entry("content.xml", b"<p>hello</p>")
        cd = _central_dir_offset(data)
        for base, usize in ((0, _LH_USIZE), (cd, _CD_USIZE)):
            old = struct.unpack_from("<I", data, base + usize)[0]
            struct.pack_into("<I", data, base + usize, old + 10)

        with pytest.raises(guard.DecompressionBombError, match="less than"):
            guard.validate_zip_container(bytes(data), ".odt")

    @pytest.mark.parametrize(
        "offset, fmt, value",
        [
            pytest.param(_LH_METHOD, "<H", zipfile.ZIP_STORED, id="method"),
            pytest.param(_LH_FLAGS, "<H", 0x0800, id="flags"),
            pytest.param(_LH_CRC, "<I", 0x12345678, id="crc"),
            pytest.param(_LH_USIZE, "<I", 9_999_999, id="uncompressed"),
            pytest.param(30, "<B", ord("X"), id="name"),
        ],
    )
    def test_local_field_disagreeing_with_central_is_refused(
        self, offset, fmt, value
    ):
        guard = _guard()
        data = _single_entry("content.xml", b"<p>hello</p>")
        struct.pack_into(fmt, data, offset, value)

        with pytest.raises(guard.DecompressionBombError, match="disagrees"):
            guard.validate_zip_container(bytes(data), ".odt")

    def test_crc_mismatch_in_both_headers_is_refused(self):
        guard = _guard()
        data = _single_entry("content.xml", b"<p>hello</p>")
        cd = _central_dir_offset(data)
        struct.pack_into("<I", data, _LH_CRC, 1)
        struct.pack_into("<I", data, cd + _CD_CRC, 1)

        with pytest.raises(guard.DecompressionBombError, match="CRC"):
            guard.validate_zip_container(bytes(data), ".odt")


_MIB = 1024 * 1024


class TestStreamEndAtOutputStepBoundary:
    """Verification inflates in 1 MiB output steps. A step that hits
    that cap can consume the last compressed byte while zlib still
    holds output (and the end-of-stream marker), so the verifier must
    drain it rather than call a complete stream truncated. Payloads are
    zeros, so each archive stays a few KiB."""

    @pytest.mark.parametrize("level", [1, 6, 9])
    @pytest.mark.parametrize(
        "size",
        [
            pytest.param(base * _MIB + delta, id=f"{base}MiB{delta:+d}")
            for base in (1, 2, 3)
            for delta in (-1, 0, 1, 50, 100, 200, 400)
        ],
    )
    def test_complete_stream_near_an_output_step_is_accepted(self, size, level):
        guard = _guard()
        buf = io.BytesIO()
        with zipfile.ZipFile(
            buf, "w", zipfile.ZIP_DEFLATED, compresslevel=level
        ) as zf:
            entries = _ooxml_package(
                {"word/media/image1.bmp": b"BM" + bytes(size - 2)}, ".docx"
            )
            for name, payload in entries.items():
                zf.writestr(name, payload)

        guard.validate_zip_container(buf.getvalue(), ".docx")


def _streamed_zip(entries: dict[str, bytes], **open_kwargs) -> bytes:
    """A zip written to a non-seekable stream: zipfile then sets flag
    bit 3 and writes data descriptors (as LibreOffice and Java do)."""

    class _Sink(io.RawIOBase):
        def __init__(self):
            self.data = bytearray()

        def writable(self):
            return True

        def write(self, chunk):
            self.data += chunk
            return len(chunk)

    sink = _Sink()
    with zipfile.ZipFile(sink, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in entries.items():
            with zf.open(name, "w", **open_kwargs) as fh:
                fh.write(payload)
    return bytes(sink.data)


class TestDataDescriptors:
    def test_streamed_entries_with_descriptors_are_accepted(self):
        data = _streamed_zip({"content.xml": b"<p>hello</p>", "empty": b""})
        assert struct.unpack_from("<H", data, _LH_FLAGS)[0] & 0x8

        _guard().validate_zip_container(data, ".odt")

    def test_zip64_descriptors_are_accepted(self):
        data = _streamed_zip({"content.xml": b"<p>hello</p>"}, force_zip64=True)

        _guard().validate_zip_container(data, ".odt")

    def test_descriptor_disagreeing_with_central_is_refused(self):
        guard = _guard()
        data = bytearray(_streamed_zip({"content.xml": b"<p>hello</p>"}))
        descriptor = data.index(b"PK\x07\x08")
        size = struct.unpack_from("<I", data, descriptor + 12)[0]
        struct.pack_into("<I", data, descriptor + 12, size + 1_000_000)

        with pytest.raises(guard.DecompressionBombError, match="descriptor"):
            guard.validate_zip_container(bytes(data), ".odt")

    def test_nonzero_local_size_under_bit_3_must_match(self):
        guard = _guard()
        data = bytearray(_streamed_zip({"content.xml": b"<p>hello</p>"}))
        struct.pack_into("<I", data, _LH_USIZE, 9_999_999)

        with pytest.raises(guard.DecompressionBombError, match="disagrees"):
            guard.validate_zip_container(bytes(data), ".odt")


# --- Reference fan-out in .pptx / .docx -----------------------------------


def _rewrite_zip(data: bytes, changes=None, add=None) -> bytes:
    changes, add = changes or {}, add or {}
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as zin,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for info in zin.infolist():
            payload = zin.read(info)
            if info.filename in changes:
                payload = changes[info.filename](payload)
            zout.writestr(info.filename, payload)
        for name, payload in add.items():
            zout.writestr(name, payload)
    return out.getvalue()


def _deck(slides: int = 1) -> bytes:
    pptx = pytest.importorskip("pptx")
    deck = pptx.Presentation()
    for n in range(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"slide {n}"
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


_SLD_ID = re.compile(rb'<p:sldId id="(\d+)" r:id="([^"]+)"/>')
_SLIDE_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
)


def _repeat_first_slide(data: bytes, copies: int, *, distinct_rels: bool):
    """List the first slide part *copies* times in sldIdLst, through one
    rId or through *copies* distinct relationships to the same part."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        presentation = zf.read("ppt/presentation.xml")
    rel_id = _SLD_ID.search(presentation).group(2).decode()

    def sld_ids(xml: bytes) -> bytes:
        ids = b"".join(
            b'<p:sldId id="%d" r:id="%s"/>'
            % (300 + n, (f"rIdFan{n}" if distinct_rels else rel_id).encode())
            for n in range(copies)
        )
        return _SLD_ID.sub(b"", xml).replace(
            b"<p:sldIdLst>", b"<p:sldIdLst>" + ids, 1
        )

    def rels(xml: bytes) -> bytes:
        if not distinct_rels:
            return xml
        extra = "".join(
            f'<Relationship Id="rIdFan{n}" Type="{_SLIDE_REL}" '
            'Target="slides/slide1.xml"/>'
            for n in range(copies)
        )
        return xml.replace(
            b"</Relationships>", extra.encode() + b"</Relationships>"
        )

    return _rewrite_zip(
        data,
        {
            "ppt/presentation.xml": sld_ids,
            "ppt/_rels/presentation.xml.rels": rels,
        },
    )


class TestPptxSlideFanOut:
    def test_premise_python_pptx_yields_a_shared_slide_per_entry(self):
        pptx = pytest.importorskip("pptx")
        data = _repeat_first_slide(_deck(), 5, distinct_rels=False)

        slides = list(pptx.Presentation(io.BytesIO(data)).slides)
        assert len(slides) == 5
        assert len({id(slide.part) for slide in slides}) == 1

    @pytest.mark.parametrize("distinct_rels", [False, True])
    def test_one_slide_listed_repeatedly_is_refused(self, distinct_rels):
        guard = _guard()
        data = _repeat_first_slide(_deck(), 5, distinct_rels=distinct_rels)

        with pytest.raises(guard.DecompressionBombError, match="same slide"):
            guard.validate_zip_container(data, ".pptx")

    def test_rid_redefined_to_a_missing_part_is_resolved_like_pptx(self):
        """python-pptx drops a relationship whose target is missing
        before later duplicates win, so rIdA -> slide1 followed by
        rIdA -> missing still resolves to slide1."""
        pptx = pytest.importorskip("pptx")
        guard = _guard()
        data = _repeat_first_slide(_deck(), 3, distinct_rels=True)

        def shadow(xml: bytes) -> bytes:
            extra = "".join(
                f'<Relationship Id="rIdFan{n}" Type="{_SLIDE_REL}" '
                f'Target="slides/missing{n}.xml"/>'
                for n in range(3)
            )
            return xml.replace(
                b"</Relationships>", extra.encode() + b"</Relationships>"
            )

        data = _rewrite_zip(data, {"ppt/_rels/presentation.xml.rels": shadow})
        # Premise: python-pptx still yields the shared slide three times.
        assert len(list(pptx.Presentation(io.BytesIO(data)).slides)) == 3

        with pytest.raises(guard.DecompressionBombError, match="same slide"):
            guard.validate_zip_container(data, ".pptx")

    def test_unresolvable_slide_entry_is_refused(self):
        guard = _guard()

        def dangling(xml: bytes) -> bytes:
            return xml.replace(
                b"</p:sldIdLst>",
                b'<p:sldId id="999" r:id="rIdNowhere"/></p:sldIdLst>',
                1,
            )

        data = _rewrite_zip(_deck(), {"ppt/presentation.xml": dangling})
        with pytest.raises(
            guard.DecompressionBombError, match="cannot resolve"
        ):
            guard.validate_zip_container(data, ".pptx")

    def test_distinct_slides_are_accepted(self):
        _guard().validate_zip_container(_deck(slides=3), ".pptx")

    def test_upload_path_refuses_the_fan_out(self):
        from local_deep_research.document_loaders.bytes_loader import (
            load_from_bytes,
        )

        guard = _guard()
        data = _repeat_first_slide(_deck(), 5, distinct_rels=False)
        with pytest.raises(guard.DecompressionBombError):
            load_from_bytes(data, ".pptx", "fan.pptx")


_W_MAIN = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_HEADER_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header"
)
_HEADER_CT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.header+xml"
)


def _docx_with_sections(
    sections: int, *, header: bytes | None = b"HEADER", rel_id: str = "rIdHdr"
) -> bytes:
    """A python-docx document whose body holds *sections* one-paragraph
    sections, each (with *header*) referencing one shared header part."""
    docx = pytest.importorskip("docx")
    buf = io.BytesIO()
    docx.Document().save(buf)
    ref = (
        f'<w:headerReference w:type="default" r:id="{rel_id}"/>'.encode()
        if header is not None
        else b""
    )
    body = b"".join(
        b"<w:p><w:pPr><w:sectPr>%s</w:sectPr></w:pPr>"
        b"<w:r><w:t>p%d</w:t></w:r></w:p>" % (ref, n)
        for n in range(sections)
    )
    changes = {
        "word/document.xml": lambda xml: xml.replace(
            b"<w:body>", b"<w:body>" + body, 1
        )
    }
    add = {}
    if header is not None:
        changes["word/_rels/document.xml.rels"] = lambda xml: xml.replace(
            b"</Relationships>",
            f'<Relationship Id="rIdHdr" Type="{_HEADER_REL}" '
            'Target="header1.xml"/></Relationships>'.encode(),
        )
        changes["[Content_Types].xml"] = lambda xml: xml.replace(
            b"</Types>",
            f'<Override PartName="/word/header1.xml" '
            f'ContentType="{_HEADER_CT}"/></Types>'.encode(),
        )
        add["word/header1.xml"] = (
            f'<w:hdr xmlns:w="{_W_MAIN}"><w:p><w:r><w:t>'.encode()
            + header
            + b"</w:t></w:r></w:p></w:hdr>"
        )
    return _rewrite_zip(buf.getvalue(), changes, add)


class TestDocxSectionFanOut:
    def test_premise_python_docx_repeats_a_shared_header_per_section(self):
        docx = pytest.importorskip("docx")
        data = _docx_with_sections(4)

        sections = list(docx.Document(io.BytesIO(data)).sections)
        assert len(sections) == 5  # four paragraph sections + the body's
        headers = [
            s.header for s in sections if not s.header.is_linked_to_previous
        ]
        assert len(headers) == 4
        assert {h.part.partname for h in headers} == {"/word/header1.xml"}

    def test_shared_header_references_over_the_budget_are_refused(self):
        guard = _guard()
        data = _docx_with_sections(20, header=b"H" * 1000)

        with patch.object(guard, "MAX_DOCX_HEADER_FOOTER_BYTES", 10_000):
            with pytest.raises(guard.DecompressionBombError, match="header"):
                guard.validate_zip_container(data, ".docx")
        with patch.object(guard, "MAX_DOCX_HEADER_FOOTER_BYTES", 100_000):
            guard.validate_zip_container(data, ".docx")  # control

    def test_section_work_over_the_ceiling_is_refused(self):
        guard = _guard()
        data = _docx_with_sections(30, header=None)

        with patch.object(guard, "MAX_DOCX_SECTION_WORK", 500):
            with pytest.raises(guard.DecompressionBombError, match="sections"):
                guard.validate_zip_container(data, ".docx")
        with patch.object(guard, "MAX_DOCX_SECTION_WORK", 5_000):
            guard.validate_zip_container(data, ".docx")  # control

    @pytest.mark.parametrize("node", [b"<w:r/>", b"<!---->", b"<?x?>"])
    def test_section_xpath_nodes_count_toward_the_work(self, node):
        # python-docx recomputes ``_sectPrs`` (an absolute XPath that
        # visits every child of every body-level w:p) once per section,
        # so a few sections over one paragraph of many runs is as costly
        # as many body elements: the body-element product alone missed it.
        guard = _guard()
        sections = 10
        data = _docx_with_sections(sections, header=None)
        wide = _rewrite_zip(
            data,
            {
                "word/document.xml": lambda xml: xml.replace(
                    b"<w:body>", b"<w:body><w:p>" + node * 5_000 + b"</w:p>", 1
                )
            },
        )

        with zipfile.ZipFile(io.BytesIO(wide)) as archive:
            from docx.opc.packuri import PackURI

            stats = guard._docx_section_stats(archive, PackURI)
        counted = stats.sections
        assert counted == sections + 1  # plus the body's own w:sectPr
        assert stats.xpath_nodes >= 5_000
        assert counted * stats.body_nodes < 500  # what the old metric saw
        with patch.object(guard, "MAX_DOCX_SECTION_WORK", 500):
            with pytest.raises(guard.DecompressionBombError, match="sections"):
                guard.validate_zip_container(wide, ".docx")
            guard.validate_zip_container(data, ".docx")  # control

    def test_ordinary_multi_section_document_stays_well_under_the_ceiling(
        self,
    ):
        # Styled, multi-run paragraphs as python-docx writes them,
        # measured small and extrapolated to 20 equal sections over
        # 6,000 such paragraphs.
        docx = pytest.importorskip("docx")
        from docx.opc.packuri import PackURI

        guard = _guard()
        document = docx.Document()
        paragraphs_per_section, sections = 20, 5
        for n in range(sections):
            for m in range(paragraphs_per_section):
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.add_run(f"section {n} ")
                paragraph.add_run(f"paragraph {m}").bold = True
                paragraph.add_run(" text")
            document.add_section()
        buf = io.BytesIO()
        document.save(buf)
        data = buf.getvalue()

        guard.validate_zip_container(data, ".docx")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            stats = guard._docx_section_stats(archive, PackURI)
        assert stats.sections == sections + 1
        assert stats.union_pairs == sections  # N1 = sections, N2 = 1
        nodes_per_paragraph = stats.xpath_nodes / (
            paragraphs_per_section * sections
        )
        big_sections, big_paragraphs = 20, 6_000
        # Section k ends after k/20 of the paragraphs.
        big_block_work = (
            sum(k * big_paragraphs // big_sections for k in range(1, 21))
            * big_paragraphs
        )
        big_work = (
            big_sections * big_paragraphs
            + big_sections
            * int(big_paragraphs * nodes_per_paragraph)
            // guard.DOCX_SECTION_XPATH_NODES_PER_UNIT
            + big_sections
            * big_sections
            // guard.DOCX_SECTION_UNION_PAIRS_PER_UNIT
        )
        assert big_block_work <= guard.MAX_DOCX_SECTION_BLOCK_WORK
        assert big_work <= guard.MAX_DOCX_SECTION_WORK

    def test_section_union_pairs_count_toward_the_work(self):
        # The ``_sectPrs`` XPath is a union, and libxml2 drops duplicates
        # by comparing every w:sectPr inside paragraph properties with
        # every body-level w:sectPr, once per section: one w:pPr holding
        # thousands of w:sectPr beside hundreds of body-level ones took
        # minutes from a 36 KB file while the body-element and
        # visited-node terms stayed under the ceiling.
        pytest.importorskip("docx")
        from docx.opc.packuri import PackURI

        guard = _guard()
        in_ppr, at_body = 600, 300  # under MAX_DOCX_SECTIONS together
        data = _docx_with_sections(1, header=None)
        packed = _rewrite_zip(
            data,
            {
                "word/document.xml": lambda xml: xml.replace(
                    b"<w:body>",
                    b"<w:body><w:p><w:pPr>"
                    + b"<w:sectPr/>" * in_ppr
                    + b"</w:pPr></w:p>"
                    + b"<w:sectPr/>" * at_body,
                    1,
                )
            },
        )

        with zipfile.ZipFile(io.BytesIO(packed)) as archive:
            stats = guard._docx_section_stats(archive, PackURI)
        counted = stats.sections
        # plus the fixture's own paragraph and body-level w:sectPr
        assert counted == (in_ppr + 1) + (at_body + 1)
        assert stats.union_pairs == (in_ppr + 1) * (at_body + 1)
        old_work = (
            counted * stats.body_nodes
            + counted
            * stats.xpath_nodes
            // guard.DOCX_SECTION_XPATH_NODES_PER_UNIT
        )
        # A ceiling the metric without the pair term just meets.
        with patch.object(guard, "MAX_DOCX_SECTION_WORK", old_work):
            with pytest.raises(guard.DecompressionBombError, match="pairs"):
                guard.validate_zip_container(packed, ".docx")
            guard.validate_zip_container(data, ".docx")  # control

    def test_many_ordinary_sections_add_few_union_pairs(self):
        # Word-shaped: one body-level w:sectPr, one paragraph w:sectPr
        # per section break, so the pair term grows only linearly.
        pytest.importorskip("docx")
        from docx.opc.packuri import PackURI

        guard = _guard()
        sections = 300
        data = _docx_with_sections(sections, header=None)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            stats = guard._docx_section_stats(archive, PackURI)
        assert stats.sections == sections + 1
        assert stats.union_pairs == sections
        guard.validate_zip_container(data, ".docx")

    def test_unresolvable_header_reference_is_refused(self):
        guard = _guard()
        data = _docx_with_sections(2, rel_id="rIdNowhere")

        with pytest.raises(
            guard.DecompressionBombError, match="cannot resolve"
        ):
            guard.validate_zip_container(data, ".docx")

    def test_real_multi_section_document_is_accepted(self):
        docx = pytest.importorskip("docx")
        document = docx.Document()
        for n in range(5):
            document.add_paragraph(f"section {n}")
            document.add_section()
        document.sections[0].header.paragraphs[0].text = "running head"
        buf = io.BytesIO()
        document.save(buf)

        _guard().validate_zip_container(buf.getvalue(), ".docx")

    def test_defaults_accept_ordinary_documents_of_a_few_thousand_paragraphs(
        self,
    ):
        # The ceilings refuse long documents, not only hostile ones (see
        # "Known gaps" in the guard's docstring): measured on documents
        # python-docx builds, a single section of more than about 11,300
        # one-run paragraphs, or a Word-like document (300 styles, five
        # sections, a table) of more than about 5,250-5,850 paragraphs
        # of 5-10 formatted runs, is refused. One of 5,000 such
        # paragraphs is accepted and one of 7,000 refused.
        guard = _guard()
        guard.validate_zip_container(_ordinary_document(5_000), ".docx")
        with pytest.raises(guard.DecompressionBombError, match="budget"):
            guard.validate_zip_container(_ordinary_document(7_000), ".docx")

    def test_defaults_accept_many_sections_rels_settings_and_styles(self):
        # A few hundred sections, thousands of images or hyperlinks, and
        # settings and styles parts of a few hundred entries, each on
        # its own, are far from their caps.
        guard = _guard()
        assert guard.MAX_DOCX_SECTIONS >= 1_000
        assert guard.MAX_PART_RELATIONSHIPS >= 10_000
        assert guard.MAX_DOCX_SETTINGS_CHILDREN >= 5_000
        assert guard.MAX_DOCX_STYLES >= 5_000
        assert guard.MAX_DOCX_HEADER_FOOTER_BYTES >= 8 * 1024 * 1024
        data = _docx_with_sections(300, header=b"running head")
        data = _pad_rels(data, "word/_rels/document.xml.rels", 2_000)
        data = _pad_styles(data, 400)
        guard.validate_zip_container(data, ".docx")


def _ordinary_document(paragraphs: int) -> bytes:
    """A Word-like document python-docx builds: 300 styles, five
    sections, five formatted runs per paragraph, and a 10 x 4 table."""
    docx = pytest.importorskip("docx")
    from docx.enum.style import WD_STYLE_TYPE

    document = docx.Document()
    names = []
    while len(document.styles.element) < 300:
        style = document.styles.add_style(
            f"Custom {len(names)}", WD_STYLE_TYPE.PARAGRAPH
        )
        style.base_style = document.styles["Normal"]
        names.append(style.name)
    sections = 5
    for section in range(sections):
        for n in range(paragraphs // sections):
            paragraph = document.add_paragraph(style=names[n % len(names)])
            for r in range(5):
                run = paragraph.add_run(f"word{r} text ")
                run.bold = r % 2 == 0
                run.italic = r % 3 == 0
        if section == 1:
            table = document.add_table(rows=10, cols=4)
            for row in table.rows:
                for cell in row.cells:
                    cell.text = "cell"
        if section < sections - 1:
            document.add_section()
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


_HYPERLINK_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    "hyperlink"
)
_SETTINGS_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    "settings"
)


def _count_children(data: bytes, member: str) -> int:
    from lxml import etree

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        return len(etree.fromstring(zf.read(member)))


def _pad_rels(data: bytes, member: str, total: int) -> bytes:
    """Add external hyperlink relationships to *member* until it has
    *total* entries."""
    extra = total - _count_children(data, member)
    assert extra >= 0
    rels = "".join(
        f'<Relationship Id="rPad{n}" Type="{_HYPERLINK_REL}" '
        f'Target="https://example.com/{n}" TargetMode="External"/>'
        for n in range(extra)
    ).encode()
    return _rewrite_zip(
        data,
        {
            member: lambda xml: xml.replace(
                b"</Relationships>", rels + b"</Relationships>"
            )
        },
    )


def _pad_settings(data: bytes, total: int, member: str = "word/settings.xml"):
    """Add empty children to the settings root until it has *total*."""
    extra = total - _count_children(data, member)
    assert extra >= 0
    return _rewrite_zip(
        data,
        {
            member: lambda xml: re.sub(
                rb"(<w:settings\b[^>]*>)",
                lambda m: m.group(1) + b"<w:pad/>" * extra,
                xml,
                count=1,
            )
        },
    )


def _with_body_prefix(data: bytes, prefix: bytes) -> bytes:
    return _rewrite_zip(
        data,
        {
            "word/document.xml": lambda xml: xml.replace(
                b"<w:body>", b"<w:body>" + prefix, 1
            )
        },
    )


_DOCX_WORK_CEILINGS = (
    "MAX_DOCX_SECTION_BLOCK_WORK",
    "MAX_DOCX_SECTION_WORK",
    "MAX_DOCX_HEADER_FOOTER_BYTES",
    "MAX_DOCX_STYLE_WORK",
    "MAX_DOCX_PAGE_BREAK_WORK",
    "MAX_DOCX_RUN_WORK",
    "MAX_DOCX_COPIED_BYTES",
)


@contextmanager
def _isolated(guard, keep: str | None):
    """Lift every .docx work ceiling but *keep* (all of them for None),
    so that a document at *keep*'s ceiling is not refused for the small
    shares of the others in the combined budget."""
    with contextlib.ExitStack() as stack:
        for name in _DOCX_WORK_CEILINGS:
            if name != keep:
                stack.enter_context(patch.object(guard, name, 10**30))
        yield


def _body_paragraphs(paragraphs: int) -> bytes:
    """python-docx's blank document (one body-level w:sectPr) with
    *paragraphs* empty body-level paragraphs."""
    return _with_body_prefix(
        _docx_with_sections(0, header=None), b"<w:p/>" * paragraphs
    )


class TestDocxPerSectionCaps:
    # python-docx and unstructured repeat work once per section that
    # the section-work count did not see (unstructured reads
    # ``settings.odd_and_even_pages_header_footer`` twice per section: a
    # Python scan of every document relationship, then a search of the
    # settings root's children; each ``_sectPrs`` evaluation returns
    # every w:sectPr). Rather than model each, the section count, the
    # relationships per part and the settings root are capped.

    def test_section_cap_accepts_the_limit_and_refuses_one_more(self):
        guard = _guard()
        limit = guard.MAX_DOCX_SECTIONS
        # n paragraph sections plus the body's own w:sectPr
        with _isolated(guard, None):
            guard.validate_zip_container(
                _docx_with_sections(limit - 1, header=None), ".docx"
            )
            with pytest.raises(guard.DecompressionBombError, match="sections"):
                guard.validate_zip_container(
                    _docx_with_sections(limit, header=None), ".docx"
                )

    def test_every_sectpr_counts_toward_the_section_cap(self):
        # One nested in a table cell's paragraph: python-docx does not
        # enumerate it as a section, but the cap counts it anyway.
        guard = _guard()
        limit = guard.MAX_DOCX_SECTIONS
        nested = (
            b"<w:tbl><w:tr><w:tc><w:p><w:pPr><w:sectPr/></w:pPr></w:p>"
            b"</w:tc></w:tr></w:tbl>"
        )
        # body-level w:sectPr, so that no section has paragraphs
        at_limit = _with_body_prefix(
            _docx_with_sections(0, header=None),
            b"<w:sectPr/>" * (limit - 2) + nested,
        )
        guard.validate_zip_container(at_limit, ".docx")
        with pytest.raises(guard.DecompressionBombError, match="sections"):
            guard.validate_zip_container(
                _with_body_prefix(at_limit, b"<w:sectPr/>"), ".docx"
            )

    @pytest.mark.parametrize(
        "in_ppr, at_body", [(14_600, 421), (3_000, 150), (1_001, 0)]
    )
    def test_sectpr_packed_into_one_paragraph_is_refused(self, in_ppr, at_body):
        # The reviewers' shapes: thousands of w:sectPr in one w:pPr,
        # beside body-level ones (14,600 beside 421, a 36 KB file, took
        # ~190 s in partition_docx).
        guard = _guard()
        data = _with_body_prefix(
            _docx_with_sections(0, header=None),
            b"<w:p><w:pPr>"
            + b"<w:sectPr/>" * in_ppr
            + b"</w:pPr></w:p>"
            + b"<w:sectPr/>" * at_body,
        )
        with pytest.raises(guard.DecompressionBombError, match="sections"):
            guard.validate_zip_container(data, ".docx")

    @pytest.mark.parametrize(
        "member", ["word/_rels/document.xml.rels", "_rels/.rels"]
    )
    def test_relationships_cap_accepts_the_limit_and_refuses_one_more(
        self, member
    ):
        guard = _guard()
        limit = guard.MAX_PART_RELATIONSHIPS
        data = _docx_with_sections(3)

        guard.validate_zip_container(_pad_rels(data, member, limit), ".docx")
        with pytest.raises(guard.DecompressionBombError, match="relationships"):
            guard.validate_zip_container(
                _pad_rels(data, member, limit + 1), ".docx"
            )

    def test_presentation_relationships_are_capped_too(self):
        guard = _guard()
        member = "ppt/_rels/presentation.xml.rels"
        guard.validate_zip_container(
            _pad_rels(_deck(1), member, guard.MAX_PART_RELATIONSHIPS), ".pptx"
        )
        with pytest.raises(guard.DecompressionBombError, match="relationships"):
            guard.validate_zip_container(
                _pad_rels(_deck(1), member, guard.MAX_PART_RELATIONSHIPS + 1),
                ".pptx",
            )

    def test_settings_cap_accepts_the_limit_and_refuses_one_more(self):
        guard = _guard()
        limit = guard.MAX_DOCX_SETTINGS_CHILDREN
        data = _docx_with_sections(3)

        guard.validate_zip_container(_pad_settings(data, limit), ".docx")
        with pytest.raises(guard.DecompressionBombError, match="settings"):
            guard.validate_zip_container(
                _pad_settings(data, limit + 1), ".docx"
            )

    def test_settings_part_is_found_through_its_relationship(self):
        # python-docx resolves settings by relationship type, not by
        # name: a padded settings part under another name is refused.
        docx = pytest.importorskip("docx")
        guard = _guard()
        data = _rewrite_zip(
            _docx_with_sections(3),
            {
                "word/_rels/document.xml.rels": lambda xml: xml.replace(
                    b'Target="settings.xml"', b'Target="custom/prefs.xml"'
                ),
                "[Content_Types].xml": lambda xml: xml.replace(
                    b'PartName="/word/settings.xml"',
                    b'PartName="/word/custom/prefs.xml"',
                ),
            },
        )
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            moved = {
                (
                    "word/custom/prefs.xml"
                    if i.filename == "word/settings.xml"
                    else i.filename
                ): zf.read(i)
                for i in zf.infolist()
            }
        data = _zip_bytes(moved)
        document_part = docx.Document(io.BytesIO(data)).part
        # premise: python-docx follows the relationship to the moved part
        assert document_part._settings_part.partname == "/word/custom/prefs.xml"

        guard.validate_zip_container(data, ".docx")  # control
        padded = _pad_settings(
            data,
            guard.MAX_DOCX_SETTINGS_CHILDREN + 1,
            member="word/custom/prefs.xml",
        )
        with pytest.raises(guard.DecompressionBombError, match="settings"):
            guard.validate_zip_container(padded, ".docx")

    def test_settings_part_gets_the_skeleton_treatment(self):
        guard = _guard()
        data = _docx_with_sections(3)
        size = len(zipfile.ZipFile(io.BytesIO(data)).read("word/settings.xml"))
        with patch.object(guard, "MAX_OPC_SKELETON_PART_BYTES", size - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="larger than"
            ):
                guard.validate_zip_container(data, ".docx")

        with_dtd = _rewrite_zip(
            data,
            {
                "word/settings.xml": _with_doctype(
                    b"<!DOCTYPE w:settings>", b"<w:settings"
                )
            },
        )
        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(with_dtd, ".docx")

    def test_every_count_cap_at_its_limit_together_is_accepted(self):
        # Sections at the limit as body-level w:sectPr (python-docx
        # enumerates each as a section), since a thousand sections each
        # holding a paragraph exceed MAX_DOCX_SECTION_BLOCK_WORK.
        guard = _guard()
        data = _with_body_prefix(
            _docx_with_sections(0, header=None),
            b"<w:p><w:r><w:t>text</w:t></w:r></w:p>"
            + b"<w:sectPr/>" * (guard.MAX_DOCX_SECTIONS - 1),
        )
        data = _pad_rels(
            data,
            "word/_rels/document.xml.rels",
            guard.MAX_PART_RELATIONSHIPS,
        )
        data = _pad_settings(data, guard.MAX_DOCX_SETTINGS_CHILDREN)
        data = _pad_styles(data, guard.MAX_DOCX_STYLES)
        stats = _docx_stats(data)
        assert stats.sect_prs == stats.sections == guard.MAX_DOCX_SECTIONS
        guard.validate_zip_container(data, ".docx")


class TestDocxSectionBlockWork:
    # python-docx finds each section's paragraphs and tables with
    # ``preceding-sibling::*[self::w:p | self::w:tbl]``; libxml2 sorts
    # that reverse-axis result with comparisons that walk sibling links
    # to the end of the body, so one section costs about (blocks it
    # returns) x (body nodes): a single-section body of 100,000 empty
    # paragraphs took ~66 s, and 100 sections after 18,000 paragraphs
    # ~86 s, while the body-element count saw 100 x 18,000 units.

    def test_single_section_body_at_the_limit_is_accepted(self):
        pytest.importorskip("docx")
        from docx.opc.packuri import PackURI

        guard = _guard()
        limit = guard.MAX_DOCX_SECTION_BLOCK_WORK
        # blocks = n, body nodes = n + 1 (the body's w:sectPr)
        n = int((limit) ** 0.5)
        while n * (n + 1) > limit:
            n -= 1
        while (n + 1) * (n + 2) <= limit:
            n += 1
        data = _body_paragraphs(n)
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            stats = guard._docx_section_stats(archive, PackURI)
        assert stats.block_work == n * (n + 1) <= limit

        with _isolated(guard, "MAX_DOCX_SECTION_BLOCK_WORK"):
            guard.validate_zip_container(data, ".docx")
            with pytest.raises(guard.DecompressionBombError, match="block"):
                guard.validate_zip_container(_body_paragraphs(n + 1), ".docx")

    def test_sections_after_many_paragraphs_are_refused(self):
        # 100 one-paragraph sections after 18,000 paragraphs: within the
        # section cap and (at 100 x 18,101 units) the section-work
        # ceiling, but ~86 s in partition_docx.
        guard = _guard()
        data = _with_body_prefix(
            _docx_with_sections(100, header=None), b"<w:p/>" * 18_000
        )
        with pytest.raises(guard.DecompressionBombError, match="block"):
            guard.validate_zip_container(data, ".docx")
        guard.validate_zip_container(  # control
            _docx_with_sections(100, header=None), ".docx"
        )

    def test_body_nodes_after_the_sections_count(self):
        # Each comparison walks to the end of the body, so nodes after
        # a section's blocks cost as much as the blocks themselves: 100
        # sections after 1,000 paragraphs, then 50,000 bookmarks before
        # the body's w:sectPr: a 39 KB file that took ~62 s in
        # partition_docx. The squared-blocks sum and the section-work
        # ceiling both miss it.
        pytest.importorskip("docx")
        from docx.opc.packuri import PackURI

        guard = _guard()
        data = _with_body_prefix(
            _docx_with_sections(100, header=None), b"<w:p/>" * 1_000
        )

        def pad_before_body_sectpr(xml: bytes) -> bytes:
            at = xml.rindex(b"<w:sectPr")  # the body-level one
            return xml[:at] + b"<w:bookmarkStart/>" * 50_000 + xml[at:]

        trailing = _rewrite_zip(
            data, {"word/document.xml": pad_before_body_sectpr}
        )
        with zipfile.ZipFile(io.BytesIO(trailing)) as archive:
            stats = guard._docx_section_stats(archive, PackURI)
        assert stats.sections == 101
        prefixes = [1_000 + k for k in range(1, 101)] + [1_100]
        assert sum(p * p for p in prefixes) < guard.MAX_DOCX_SECTION_BLOCK_WORK
        assert (
            stats.sections * stats.body_nodes
            + stats.sections
            * stats.xpath_nodes
            // guard.DOCX_SECTION_XPATH_NODES_PER_UNIT
            < guard.MAX_DOCX_SECTION_WORK
        )
        assert stats.block_work == sum(prefixes) * stats.body_nodes
        with pytest.raises(guard.DecompressionBombError, match="block"):
            guard.validate_zip_container(trailing, ".docx")
        guard.validate_zip_container(data, ".docx")  # control

    @pytest.mark.parametrize(
        "body",
        [
            b"<w:p/>" * 5,
            b"x<w:p/>" * 5 + b"y",
            b"<!---->z" * 3 + b"<w:p/>",
            b"<?pi?><w:p/>",
            b"lead<w:tbl/>",
        ],
    )
    def test_body_nodes_match_python_docx_tree(self, body):
        # Non-whitespace text between body-level elements is a node
        # python-docx keeps and the comparisons walk.
        docx = pytest.importorskip("docx")
        from docx.opc.packuri import PackURI

        guard = _guard()
        data = _with_body_prefix(_docx_with_sections(0, header=None), body)
        tree = docx.Document(io.BytesIO(data)).element.body
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            stats = guard._docx_section_stats(archive, PackURI)
        assert stats.body_nodes == len(tree.xpath("node()"))

    def test_attribute_heavy_body_is_priced_at_its_worst_layout(self):
        # The sibling walk costs more per step when the nodes walked
        # carry attributes (about 16 ns for a bare element, ~48 ns with
        # three attributes, ~100 ns with forty): 1,240 paragraphs before
        # 2 million three-attribute bookmarks, a 234 KiB file at 0.993
        # of the old 2.5 billion ceiling, took ~121 s in partition_docx.
        # The ceiling now prices every step at the worst layout, so the
        # same shape at a fifth of the size is refused.
        pytest.importorskip("docx")
        from docx.opc.packuri import PackURI

        guard = _guard()
        bookmark = b'<w:bookmarkEnd w:id="1" w:a="2" w:b="3"/>'
        data = _with_body_prefix(
            _docx_with_sections(0, header=None),
            b"<w:p/>" * 1_240 + bookmark * 400_000,
        )
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            stats = guard._docx_section_stats(archive, PackURI)
        # the blank document's body ends in its own w:sectPr
        assert stats.block_work == 1_240 * (1_240 + 400_000 + 1)
        assert stats.block_work < 2_500_000_000  # the old ceiling
        with pytest.raises(guard.DecompressionBombError, match="block"):
            guard.validate_zip_container(data, ".docx")


def _pad_styles(data: bytes, total: int, member: str = "word/styles.xml"):
    """Add paragraph styles to the styles root until it has *total*
    children."""
    extra = total - _count_children(data, member)
    assert extra >= 0
    styles = b"".join(
        b'<w:style w:type="paragraph" w:styleId="pad%d"/>' % n
        for n in range(extra)
    )
    return _rewrite_zip(
        data,
        {
            member: lambda xml: xml.replace(
                b"</w:styles>", styles + b"</w:styles>"
            )
        },
    )


def _text_paragraphs(paragraphs: int, run: bytes = b"<w:t>text</w:t>"):
    """python-docx's blank document with *paragraphs* body-level
    paragraphs of one run each (holding *run*)."""
    return _with_body_prefix(
        _docx_with_sections(0, header=None),
        (b"<w:p><w:r>" + run + b"</w:r></w:p>") * paragraphs,
    )


def _docx_stats(data: bytes):
    pytest.importorskip("docx")
    from docx.opc.packuri import PackURI

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return _guard()._docx_section_stats(archive, PackURI)


def _style_work(stats) -> int:
    guard = _guard()
    return stats.styled_paragraphs * (
        stats.style_entries
        + stats.style_attributes // guard.DOCX_STYLE_ATTRIBUTES_PER_ENTRY
        + stats.relationships // guard.DOCX_RELATIONSHIPS_PER_STYLE_ENTRY
    )


class TestDocxStyleLookupWork:
    # unstructured reads ``paragraph.style`` about four times per
    # non-empty paragraph (and per fragment of one split at a rendered
    # page break). Each read scans every relationship of the document
    # part in Python (``part_related_by``) and, for a paragraph with no
    # style of its own (or one that does not resolve), loops over every
    # ``w:style`` in Python (``default_for``): 50 paragraphs beside
    # 300,000 styles, an 805 KB file, took ~132 s in partition_docx.

    def test_premise_unstructured_reads_the_style_per_paragraph(self):
        docx = pytest.importorskip("docx")
        udocx = pytest.importorskip("unstructured.partition.docx")
        from docx.text.paragraph import Paragraph

        reads = []
        style = Paragraph.style

        def counting(paragraph):
            reads.append(paragraph)
            return style.fget(paragraph)

        data = _text_paragraphs(5)
        assert len(docx.Document(io.BytesIO(data)).paragraphs) == 5
        with patch.object(Paragraph, "style", property(counting)):
            udocx.partition_docx(file=io.BytesIO(data))
        assert len(reads) == 4 * 5
        assert _docx_stats(data).styled_paragraphs == 5

    def test_a_missing_styles_part_counts_python_docx_default(self):
        # Without a styles relationship python-docx creates its default
        # styles part, which the style lookup then scans.
        docx = pytest.importorskip("docx")
        guard = _guard()
        data = _rewrite_zip(
            _text_paragraphs(3),
            {
                "word/_rels/document.xml.rels": lambda xml: re.sub(
                    rb"<Relationship [^>]*relationships/styles\"[^>]*/>",
                    b"",
                    xml,
                )
            },
        )
        styles = docx.Document(io.BytesIO(data)).styles.element
        assert len(styles) == guard.DOCX_DEFAULT_STYLES_ENTRIES
        assert _docx_stats(data).style_entries == len(styles)

    def test_styles_cap_accepts_the_limit_and_refuses_one_more(self):
        # No text paragraphs, so the style work stays at zero.
        guard = _guard()
        limit = guard.MAX_DOCX_STYLES
        blank = _docx_with_sections(0, header=None)

        guard.validate_zip_container(_pad_styles(blank, limit), ".docx")
        with pytest.raises(guard.DecompressionBombError, match="styles"):
            guard.validate_zip_container(_pad_styles(blank, limit + 1), ".docx")

    def test_styles_part_is_found_through_its_relationship(self):
        docx = pytest.importorskip("docx")
        guard = _guard()
        data = _rewrite_zip(
            _docx_with_sections(0, header=None),
            {
                "word/_rels/document.xml.rels": lambda xml: xml.replace(
                    b'Target="styles.xml"', b'Target="custom/look.xml"'
                ),
                "[Content_Types].xml": lambda xml: xml.replace(
                    b'PartName="/word/styles.xml"',
                    b'PartName="/word/custom/look.xml"',
                ),
            },
        )
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            moved = {
                (
                    "word/custom/look.xml"
                    if i.filename == "word/styles.xml"
                    else i.filename
                ): zf.read(i)
                for i in zf.infolist()
            }
        data = _zip_bytes(moved)
        document_part = docx.Document(io.BytesIO(data)).part
        # premise: python-docx follows the relationship to the moved part
        assert document_part._styles_part.partname == "/word/custom/look.xml"

        guard.validate_zip_container(data, ".docx")  # control
        padded = _pad_styles(
            data, guard.MAX_DOCX_STYLES + 1, member="word/custom/look.xml"
        )
        with pytest.raises(guard.DecompressionBombError, match="styles"):
            guard.validate_zip_container(padded, ".docx")

    def test_styles_part_gets_the_skeleton_treatment(self):
        guard = _guard()
        data = _docx_with_sections(0, header=None)
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            sizes = {i.filename: i.file_size for i in zf.infolist()}
        # premise: the styles part is the largest skeleton part read
        assert sizes["word/styles.xml"] > max(
            size
            for name, size in sizes.items()
            if name.endswith(".rels") or name == "word/settings.xml"
        )
        with patch.object(
            guard, "MAX_OPC_SKELETON_PART_BYTES", sizes["word/styles.xml"] - 1
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="larger than"
            ):
                guard.validate_zip_container(data, ".docx")

        with_dtd = _rewrite_zip(
            data,
            {
                "word/styles.xml": _with_doctype(
                    b"<!DOCTYPE w:styles>", b"<w:styles"
                )
            },
        )
        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(with_dtd, ".docx")

    def test_style_work_accepts_the_limit_and_refuses_one_more(self):
        guard = _guard()
        limit = guard.MAX_DOCX_STYLE_WORK
        styled = _pad_styles(_text_paragraphs(1), guard.MAX_DOCX_STYLES)
        per_paragraph = _style_work(_docx_stats(styled))
        n = limit // per_paragraph

        def document(paragraphs: int) -> bytes:
            return _pad_styles(
                _text_paragraphs(paragraphs), guard.MAX_DOCX_STYLES
            )

        at_limit = document(n)
        assert _style_work(_docx_stats(at_limit)) == n * per_paragraph
        with _isolated(guard, "MAX_DOCX_STYLE_WORK"):
            guard.validate_zip_container(at_limit, ".docx")
            with pytest.raises(guard.DecompressionBombError, match="style"):
                guard.validate_zip_container(document(n + 1), ".docx")

    def test_relationships_count_toward_the_style_work(self):
        # Each style read also scans every relationship: paragraphs
        # times 10,000 relationships cost about as much as paragraphs
        # times 312 styles part entries.
        guard = _guard()
        paragraphs = 8_000
        data = _text_paragraphs(paragraphs)
        stats = _docx_stats(data)
        assert _style_work(stats) <= guard.MAX_DOCX_STYLE_WORK
        guard.validate_zip_container(data, ".docx")  # control

        padded = _pad_rels(
            data, "word/_rels/document.xml.rels", guard.MAX_PART_RELATIONSHIPS
        )
        assert _docx_stats(padded).relationships == (
            guard.MAX_PART_RELATIONSHIPS
        )
        with pytest.raises(guard.DecompressionBombError, match="style"):
            guard.validate_zip_container(padded, ".docx")

    @pytest.mark.parametrize(
        "paragraph, styled",
        [
            (b"<w:p/>", 0),
            (b"<w:p><w:pPr/></w:p>", 0),
            (b"<w:p><w:r/></w:p>", 1),
            (b"<w:p><w:hyperlink/></w:p>", 1),
            # each rendered page break can split off one more fragment
            (
                b"<w:p><w:r><w:t>a</w:t><w:lastRenderedPageBreak/>"
                b"<w:t>b</w:t><w:lastRenderedPageBreak/><w:t>c</w:t>"
                b"</w:r></w:p>",
                3,
            ),
            # paragraphs inside a table are not looked up
            (b"<w:tbl><w:tr><w:tc><w:p><w:r/></w:p></w:tc></w:tr></w:tbl>", 0),
        ],
    )
    def test_styled_paragraphs_counts_fragments_that_can_have_text(
        self, paragraph, styled
    ):
        data = _with_body_prefix(_docx_with_sections(0, header=None), paragraph)
        assert _docx_stats(data).styled_paragraphs == styled


def _page_break_paragraph(breaks: int, gap: int = 1) -> bytes:
    """One body-level paragraph of one run holding *breaks* rendered
    page breaks, with *gap* text elements before, between and after
    them."""
    text = b"<w:t>a</w:t>" * gap
    return (
        b"<w:p><w:r>"
        + text
        + (b"<w:lastRenderedPageBreak/>" + text) * breaks
        + b"</w:r></w:p>"
    )


def _blank_with(body: bytes) -> bytes:
    return _with_body_prefix(_docx_with_sections(0, header=None), body)


class TestDocxPageBreakWork:
    # unstructured splits a paragraph at each ``w:lastRenderedPageBreak``
    # (python-docx's ``preceding/following_paragraph_fragment``): every
    # split deep-copies the paragraph and evaluates sibling-axis XPath
    # (``precedes_all_content`` and the fragment builders) whose
    # reverse-axis results libxml2 sorts with sibling walks, once per
    # break, recursing into the rest. One paragraph of one run with 600
    # breaks (a 37 KB file) ran for over 180 s; one break after 55,000
    # elements of one run took ~34-61 s.

    def test_premise_unstructured_splits_a_paragraph_per_page_break(self):
        udocx = pytest.importorskip("unstructured.partition.docx")
        data = _blank_with(_page_break_paragraph(3))
        elements = udocx.partition_docx(file=io.BytesIO(data))
        assert [type(e).__name__ for e in elements].count("PageBreak") == 3

    @pytest.mark.parametrize(
        "paragraph",
        [
            _page_break_paragraph(2),
            _page_break_paragraph(3, gap=4),
            b"<w:p><w:pPr><w:jc w:val='left'/></w:pPr>"
            b"<w:hyperlink><w:r><w:t>a</w:t><w:lastRenderedPageBreak/>"
            b"<w:t>b</w:t></w:r></w:hyperlink><!--c-->x<w:r/></w:p>",
        ],
    )
    def test_work_matches_python_docx_tree(self, paragraph):
        # S and N are counted from the stream as python-docx's tree has
        # them (elements, comments, processing instructions and text).
        docx = pytest.importorskip("docx")
        guard = _guard()
        data = _blank_with(paragraph)
        p = docx.Document(io.BytesIO(data)).element.body[0]
        breaks = len(p.xpath(".//w:lastRenderedPageBreak"))
        nodes = len(p.xpath(".//node()"))
        widest = max(len(e.xpath("node()")) for e in p.iter())
        runs = p.xpath(".//w:r")
        # The paragraph's w:r | w:hyperlink union, counted when both are
        # present (the runs here hold one text-branch kind each, so their
        # own unions count nothing).
        branches = [p.xpath("w:r"), p.xpath("w:hyperlink")]
        union = 0
        if all(branches):
            union = -(
                -sum(map(len, branches))
                * len(p.xpath("node()"))
                * guard.DOCX_UNION_READS
                // guard.DOCX_UNION_STEPS_PER_RUN_WORK_UNIT
            )
        stats = _docx_stats(data)
        assert stats.page_break_work == breaks * (breaks + 3) * widest**2
        assert (
            stats.run_work
            == (breaks + 1)
            * (
                guard.DOCX_RUN_WORK_ITEM_WEIGHT * len(runs)
                + sum(len(r.xpath("node()")) for r in runs)
                + union
            )
            + breaks * nodes // guard.DOCX_FRAGMENT_NODES_PER_RUN_WORK_UNIT
        )

    def test_page_break_work_accepts_the_limit_and_refuses_one_more(self):
        # One break after n elements of one run: K = 1 and S = n + 1.
        guard = _guard()

        def document(n: int) -> bytes:
            return _blank_with(
                b"<w:p><w:r>"
                + b"<w:t>a</w:t>" * n
                + b"<w:lastRenderedPageBreak/></w:r></w:p>"
            )

        assert _docx_stats(document(10)).page_break_work == 4 * 11**2
        n = math.isqrt(guard.MAX_DOCX_PAGE_BREAK_WORK // 4) - 1
        with _isolated(guard, "MAX_DOCX_PAGE_BREAK_WORK"):
            guard.validate_zip_container(document(n), ".docx")
            with pytest.raises(
                guard.DecompressionBombError, match="page breaks"
            ):
                guard.validate_zip_container(document(n + 1), ".docx")

    def test_paragraphs_without_page_breaks_cost_no_page_break_work(self):
        data = _blank_with(_page_break_paragraph(0, 5_000))
        assert _docx_stats(data).page_break_work == 0


class TestDocxRunWork:
    # unstructured reads each fragment's text several times (about five
    # ``run.text`` evaluations per run, each an XPath plus a Python loop
    # over the run's children) and walks table rows and cells: 120
    # paragraphs of 2,479 one-character runs, a 76 KB file, took ~40 s,
    # and the cost was bounded only by the 128 MiB part ceiling.

    def test_run_work_accepts_the_limit_and_refuses_one_more(self):
        # One paragraph of n one-child runs: (weight + 1) units a run.
        guard = _guard()
        per_run = guard.DOCX_RUN_WORK_ITEM_WEIGHT + 1

        def document(runs: int) -> bytes:
            return _blank_with(
                b"<w:p>" + b"<w:r><w:t>a</w:t></w:r>" * runs + b"</w:p>"
            )

        assert _docx_stats(document(10)).run_work == 10 * per_run
        n = guard.MAX_DOCX_RUN_WORK // per_run
        with _isolated(guard, "MAX_DOCX_RUN_WORK"):
            guard.validate_zip_container(document(n), ".docx")
            with pytest.raises(guard.DecompressionBombError, match="runs"):
                guard.validate_zip_container(document(n + 1), ".docx")

    def test_table_rows_and_cells_count(self):
        guard = _guard()
        cell = b"<w:tc><w:p><w:r><w:t>a</w:t></w:r></w:p></w:tc>"
        data = _blank_with(
            b"<w:tbl>" + (b"<w:tr>" + cell * 3 + b"</w:tr>") * 2 + b"</w:tbl>"
        )
        # 2 rows, 6 cells and 6 runs, each run with one child, and the
        # cells' 6 paragraphs
        assert _docx_stats(data).run_work == (
            guard.DOCX_RUN_WORK_ITEM_WEIGHT * (2 + 6 + 6)
            + 6
            + 6 * guard.DOCX_CELL_BLOCK_ITEM_WEIGHT
        )

    def test_page_breaks_multiply_the_run_work(self):
        # Each fragment's text is read again, and each split copies the
        # paragraph.
        one = _docx_stats(_blank_with(_page_break_paragraph(0, 10))).run_work
        three = _docx_stats(_blank_with(_page_break_paragraph(2, 10)))
        assert three.run_work > 3 * one

    def test_runs_outside_body_blocks_are_not_counted(self):
        # python-docx reads only body-level paragraphs and tables.
        data = _blank_with(
            b"<w:sdt><w:sdtContent><w:p><w:r><w:t>a</w:t></w:r></w:p>"
            b"</w:sdtContent></w:sdt>"
        )
        assert _docx_stats(data).run_work == 0


class TestDocxWorkBudget:
    def test_work_kinds_share_one_budget(self):
        # Two kinds each at 60% of their ceilings: either alone is
        # accepted, together they are refused.
        guard = _guard()
        styled = _text_paragraphs(100)
        stats = _docx_stats(styled)
        style_ceiling = _style_work(stats) * 10 // 6
        run_ceiling = stats.run_work * 10 // 6
        with patch.object(guard, "MAX_DOCX_STYLE_WORK", style_ceiling):
            with pytest.raises(guard.DecompressionBombError, match="budget"):
                with patch.object(guard, "MAX_DOCX_RUN_WORK", run_ceiling):
                    guard.validate_zip_container(styled, ".docx")
            with patch.object(guard, "MAX_DOCX_RUN_WORK", 10**12):
                guard.validate_zip_container(styled, ".docx")
        with patch.object(guard, "MAX_DOCX_RUN_WORK", run_ceiling):
            with patch.object(guard, "MAX_DOCX_STYLE_WORK", 10**12):
                guard.validate_zip_container(styled, ".docx")

    def test_header_references_are_priced_by_default(self):
        # Header text is read once per reference, at up to ~3.5 us per
        # byte (empty runs, paragraphs or cells): 20 references to a
        # 1.2 MB header of empty runs took ~73 s, well under the old
        # 128 MiB budget.
        guard = _guard()
        hostile = _docx_with_sections(20, header=b"a" * 1_200_000)
        with pytest.raises(guard.DecompressionBombError, match="header"):
            guard.validate_zip_container(hostile, ".docx")
        guard.validate_zip_container(  # control
            _docx_with_sections(20, header=b"a" * 200_000), ".docx"
        )


def _spanned_table(span: int, text: bytes = b"x", rows: int = 1) -> bytes:
    """A body-level table of *rows* rows, each one cell spanning *span*
    grid columns and holding one run of *text*."""
    cell = (
        b"<w:tc><w:tcPr><w:gridSpan w:val='%d'/></w:tcPr>"
        b"<w:p><w:r><w:t>%s</w:t></w:r></w:p></w:tc>" % (span, text)
    )
    return b"<w:tbl>" + (b"<w:tr>" + cell + b"</w:tr>") * rows + b"</w:tbl>"


def _grid_value_table(tag: bytes, value: bytes) -> bytes:
    """A one-cell table whose cell (gridSpan) or row (gridBefore,
    gridAfter) carries *value*."""
    if tag == b"gridSpan":
        return (
            b"<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val='%s'/></w:tcPr>"
            b"<w:p/></w:tc></w:tr></w:tbl>" % value
        )
    return (
        b"<w:tbl><w:tr><w:trPr><w:%s w:val='%s'/></w:trPr>"
        b"<w:tc><w:p/></w:tc></w:tr></w:tbl>" % (tag, value)
    )


class TestDocxGridSpans:
    # python-docx's ``_Row.cells`` yields a cell once per grid column it
    # spans and unstructured's ``_convert_table_to_html`` re-reads the
    # cell's text for each yield (and pads a row with one empty string
    # per gridBefore/gridAfter column), so the span multiplies the
    # cell's work and the text it keeps: a 950-byte file with one cell
    # spanning 400,000 columns took 53 s and was accepted, and at
    # 2**31 - 1 python-docx would build a ~17 GB tuple.

    def test_premise_python_docx_yields_a_cell_per_spanned_column(self):
        docx = pytest.importorskip("docx")
        data = _blank_with(_spanned_table(7))
        table = docx.Document(io.BytesIO(data)).tables[0]
        assert len(table.rows[0].cells) == 7

    def test_premise_unstructured_rereads_the_cell_per_yield(self):
        pytest.importorskip("docx")
        udocx = pytest.importorskip("unstructured.partition.docx")
        from docx.table import _Cell

        reads = []
        inner = _Cell.iter_inner_content

        def counting(cell):
            reads.append(cell)
            return inner(cell)

        data = _blank_with(_spanned_table(7))
        with patch.object(_Cell, "iter_inner_content", counting):
            udocx.partition_docx(file=io.BytesIO(data))
        assert len(reads) >= 7

    def test_hostile_span_is_refused_by_default(self):
        guard = _guard()
        with pytest.raises(guard.DecompressionBombError, match="grid"):
            guard.validate_zip_container(
                _blank_with(_spanned_table(400_000)), ".docx"
            )

    @pytest.mark.parametrize("tag", [b"gridSpan", b"gridBefore", b"gridAfter"])
    def test_grid_values_at_the_cap_are_accepted_and_one_more_refused(
        self, tag
    ):
        guard = _guard()
        cap = guard.MAX_DOCX_GRID_SPAN
        guard.validate_zip_container(
            _blank_with(_grid_value_table(tag, b"%d" % cap)), ".docx"
        )
        with pytest.raises(guard.DecompressionBombError, match="grid"):
            guard.validate_zip_container(
                _blank_with(_grid_value_table(tag, b"%d" % (cap + 1))),
                ".docx",
            )

    def test_cap_is_well_above_real_table_widths(self):
        # Word's tables are at most 63 columns wide.
        assert _guard().MAX_DOCX_GRID_SPAN >= 15 * 63

    @pytest.mark.parametrize(
        "tag,value",
        [
            (tag, value)
            for tag in (b"gridSpan", b"gridBefore", b"gridAfter")
            for value in (b"x", b"", b"1.5", b"0x10", b"7 7")
        ],
    )
    def test_values_python_docx_cannot_read_are_refused(self, tag, value):
        guard = _guard()
        with pytest.raises(guard.DecompressionBombError, match="cannot read"):
            guard.validate_zip_container(
                _blank_with(_grid_value_table(tag, value)), ".docx"
            )

    def test_premise_python_docx_cannot_read_those_values(self):
        docx = pytest.importorskip("docx")
        for value in (b"x", b"", b"1.5", b"0x10", b"7 7"):
            data = _blank_with(_grid_value_table(b"gridSpan", value))
            table = docx.Document(io.BytesIO(data)).tables[0]
            with pytest.raises(ValueError):
                table.rows[0].cells  # noqa: B018

    @pytest.mark.parametrize("value", [b"9" * 40, b"0" * 40 + b"5"])
    def test_overlong_values_are_refused_unparsed(self, value):
        # Word writes at most a few digits; int() of a long digit string
        # is not linear, so the guard does not parse one.
        guard = _guard()
        with pytest.raises(guard.DecompressionBombError, match="characters"):
            guard.validate_zip_container(
                _blank_with(_grid_value_table(b"gridSpan", value)), ".docx"
            )

    def test_missing_value_is_refused(self):
        guard = _guard()
        data = _blank_with(
            b"<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan/></w:tcPr>"
            b"<w:p/></w:tc></w:tr></w:tbl>"
        )
        with pytest.raises(guard.DecompressionBombError, match="cannot read"):
            guard.validate_zip_container(data, ".docx")

    @pytest.mark.parametrize(
        "value",
        [b" 7 ", b"+7", b"0_7", b"\xd9\xa7"],  # Arabic-Indic 7
    )
    def test_values_are_read_as_python_docx_reads_them(self, value):
        # python-docx converts the attribute with int(), which accepts
        # surrounding whitespace, a sign, underscores and Unicode digits:
        # a span the guard read differently would escape its weight.
        docx = pytest.importorskip("docx")
        data = _blank_with(_grid_value_table(b"gridSpan", value))
        table = docx.Document(io.BytesIO(data)).tables[0]
        assert len(table.rows[0].cells) == 7
        plain = _blank_with(_grid_value_table(b"gridSpan", b"7"))
        assert _docx_stats(data).run_work == _docx_stats(plain).run_work

    def test_negative_values_span_and_pad_nothing(self):
        guard = _guard()
        for tag in (b"gridSpan", b"gridBefore", b"gridAfter"):
            data = _blank_with(_grid_value_table(tag, b"-5"))
            guard.validate_zip_container(data, ".docx")
        assert (
            _docx_stats(
                _blank_with(_grid_value_table(b"gridSpan", b"-5"))
            ).run_work
            == _docx_stats(
                _blank_with(_grid_value_table(b"gridSpan", b"1"))
            ).run_work
        )

    @pytest.mark.parametrize("span", [1, 2, 40, 1_000])
    def test_a_cell_counts_once_per_spanned_column(self, span):
        # A row (one item), and per yield of the cell: the cell, its
        # paragraph, its run and the run's one child.
        weight = _guard().DOCX_RUN_WORK_ITEM_WEIGHT
        item = _guard().DOCX_CELL_BLOCK_ITEM_WEIGHT
        stats = _docx_stats(_blank_with(_spanned_table(span)))
        assert stats.run_work == weight + span * (2 * weight + item + 1)

    def test_nested_spans_multiply(self):
        weight = _guard().DOCX_RUN_WORK_ITEM_WEIGHT
        inner = _spanned_table(4)
        data = _blank_with(
            b"<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val='3'/></w:tcPr>"
            + inner
            + b"<w:p/></w:tc></w:tr></w:tbl>"
        )
        guard = _guard()
        item = guard.DOCX_CELL_BLOCK_ITEM_WEIGHT
        inner_work = weight + 4 * (2 * weight + item + 1)
        # The outer cell holds a table and a paragraph (two items) among
        # three child nodes, so its union is counted: 2 x 3 steps.
        union = -(
            -2
            * 3
            * guard.DOCX_UNION_READS
            // guard.DOCX_UNION_STEPS_PER_RUN_WORK_UNIT
        )
        assert _docx_stats(data).run_work == weight + 3 * (
            weight + 2 * item + union + inner_work
        )

    def test_row_padding_counts(self):
        guard = _guard()
        weight = guard.DOCX_RUN_WORK_ITEM_WEIGHT
        per_unit = guard.DOCX_GRID_PADDING_PER_RUN_WORK_UNIT
        data = _blank_with(
            b"<w:tbl><w:tr><w:trPr><w:gridBefore w:val='%d'/>"
            b"<w:gridAfter w:val='%d'/></w:trPr><w:tc><w:p/></w:tc>"
            b"</w:tr></w:tbl>" % (3 * per_unit, 2 * per_unit)
        )
        assert (
            _docx_stats(data).run_work
            == 2 * weight + 5 + guard.DOCX_CELL_BLOCK_ITEM_WEIGHT
        )

    def test_many_rows_of_capped_spans_are_refused(self):
        # Each value under the cap, but the spans add up.
        guard = _guard()
        cap = guard.MAX_DOCX_GRID_SPAN
        weight = guard.DOCX_RUN_WORK_ITEM_WEIGHT
        rows = guard.MAX_DOCX_RUN_WORK // (cap * (2 * weight + 1)) + 1
        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(
                _blank_with(_spanned_table(cap, rows=rows)), ".docx"
            )

    def test_spanned_cell_text_counts_as_copies(self):
        text = b"t" * 10_000
        for span in (1, 50):
            stats = _docx_stats(_blank_with(_spanned_table(span, text)))
            assert stats.copied_bytes == (span - 1) * _cell_copy(text)

    def test_ordinary_merged_cells_are_accepted(self):
        docx = pytest.importorskip("docx")
        document = docx.Document()
        table = document.add_table(rows=6, cols=5)
        table.cell(0, 0).merge(table.cell(0, 4))  # a header across
        table.cell(1, 0).merge(table.cell(5, 0))  # a column down
        for row in table.rows[1:]:
            for cell in row.cells[1:]:
                cell.text = "value"
        buf = io.BytesIO()
        document.save(buf)
        _guard().validate_zip_container(buf.getvalue(), ".docx")


def _merged_column(root_text: bytes, continuing: int) -> bytes:
    """A one-column table: a vertically merged cell holding
    *root_text*, continued by *continuing* rows."""
    return (
        b"<w:tbl><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc><w:tcPr>"
        b"<w:vMerge w:val='restart'/></w:tcPr><w:p><w:r><w:t>"
        + root_text
        + b"</w:t></w:r></w:p></w:tc></w:tr>"
        + b"<w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr>"
        * continuing
        + b"</w:tbl>"
    )


#: tcPr encodings of a cell, and what python-docx reads as its vMerge.
_VMERGE_ENCODINGS = {
    "none": None,
    "val_x": "x",
    "restart_then_bare": "restart",
    "second_tcpr": "restart",
    "first_tcpr_without_vmerge": None,
    "bare": "continue",
    "continue": "continue",
    "continue_then_restart": "continue",
}


def _vmerge_table(
    encoding: str,
    span: int,
    *,
    text: bytes = b"x",
    rows: int = 1,
    restart_first: bool = False,
) -> bytes:
    """A one-column table of *rows* rows, each one cell spanning *span*
    columns, holding *text* and carrying the vMerge *encoding* (but for
    a first row that restarts a merge, with *restart_first*)."""
    grid = b"<w:gridSpan w:val='%d'/>" % span
    tc_pr = {
        "none": b"<w:tcPr>" + grid + b"</w:tcPr>",
        "val_x": b"<w:tcPr>" + grid + b"<w:vMerge w:val='x'/></w:tcPr>",
        "restart_then_bare": b"<w:tcPr>"
        + grid
        + b"<w:vMerge w:val='restart'/><w:vMerge/></w:tcPr>",
        "second_tcpr": b"<w:tcPr>"
        + grid
        + b"<w:vMerge w:val='restart'/></w:tcPr><w:tcPr><w:vMerge/></w:tcPr>",
        "first_tcpr_without_vmerge": b"<w:tcPr>"
        + grid
        + b"</w:tcPr><w:tcPr><w:vMerge/></w:tcPr>",
        "bare": b"<w:tcPr>" + grid + b"<w:vMerge/></w:tcPr>",
        "continue": b"<w:tcPr>"
        + grid
        + b"<w:vMerge w:val='continue'/></w:tcPr>",
        "continue_then_restart": b"<w:tcPr>"
        + grid
        + b"<w:vMerge w:val='continue'/><w:vMerge w:val='restart'/></w:tcPr>",
    }[encoding]
    content = b"<w:p><w:r><w:t>" + text + b"</w:t></w:r></w:p>"
    restart = (
        b"<w:tr><w:tc><w:tcPr>"
        + grid
        + b"<w:vMerge w:val='restart'/></w:tcPr>"
        + content
        + b"</w:tc></w:tr>"
    )
    row = b"<w:tr><w:tc>" + tc_pr + content + b"</w:tc></w:tr>"
    rows_xml = restart + row * (rows - 1) if restart_first else row * rows
    return (
        b"<w:tbl><w:tblGrid><w:gridCol/></w:tblGrid>" + rows_xml + b"</w:tbl>"
    )


def _declaration_bytes(prefix: str, uri: str) -> int:
    """What one copied namespace declaration counts."""
    return _guard().DOCX_COPIED_BYTES_PER_NODE + len(prefix) + len(uri)


def _cell_copy(text: bytes) -> int:
    """What one extra copy of a table cell holding *text* counts."""
    return _guard().DOCX_TABLE_COPY_BYTES_PER_TEXT_BYTE * len(text)


class TestDocxCopiedBytes:
    # Fan-outs that turn one bounded paragraph or cell into many copies
    # of its bytes: unstructured keeps two deep copies of the rest of a
    # paragraph per rendered page break, namespace declarations and the
    # text after the paragraph included (80 breaks before 3 MB of text,
    # a 40 KB file, kept ~230 MB more than without them), re-extracts a
    # spanned cell's text per spanned column, and a continuing merged
    # cell's merge-start text per row (300 rows under a 1 MB cell, a
    # 38 KB file, kept ~930 MB).

    def test_premise_each_page_break_deep_copies_the_paragraph(self):
        pytest.importorskip("docx")
        udocx = pytest.importorskip("unstructured.partition.docx")
        from docx.oxml.ns import qn
        from docx.oxml.text import pagebreak

        copies = []
        deepcopy = pagebreak.copy.deepcopy

        def counting(element, *args):
            if getattr(element, "tag", None) == qn("w:p"):
                copies.append(element)
            return deepcopy(element, *args)

        data = _blank_with(_page_break_paragraph(5))
        with patch.object(pagebreak.copy, "deepcopy", counting):
            udocx.partition_docx(file=io.BytesIO(data))
        assert len(copies) >= 5

    def test_premise_a_continuing_merged_cell_yields_the_merge_start(self):
        docx = pytest.importorskip("docx")
        data = _blank_with(_merged_column(b"root", 3))
        table = docx.Document(io.BytesIO(data)).tables[0]
        assert [row.cells[0].text for row in table.rows] == ["root"] * 4

    def test_page_breaks_count_two_copies_of_the_paragraph_per_break(self):
        guard = _guard()
        text = b"t" * 100_000
        for breaks in (1, 7):
            data = _blank_with(
                b"<w:p><w:r>"
                + b"<w:lastRenderedPageBreak/>" * breaks
                + b"</w:r><w:r><w:t>"
                + text
                + b"</w:t></w:r></w:p>"
            )
            copied = _docx_stats(data).copied_bytes
            # Per copy: the text, 4 + breaks nodes (w:p, two w:r, w:t,
            # the breaks), and the w: declaration of w:document, which
            # libxml2 re-declares on the copy.
            nodes = 4 + breaks
            assert copied == 2 * breaks * (
                len(text)
                + nodes * guard.DOCX_COPIED_BYTES_PER_NODE
                + _declaration_bytes("w", _W_MAIN)
            )

    def test_attribute_values_and_nodes_count(self):
        guard = _guard()
        name = b"n" * 50_000
        data = _blank_with(
            b"<w:p><w:r><w:lastRenderedPageBreak/></w:r>"
            b"<w:bookmarkStart w:id='1' w:name='"
            + name
            + b"'/>"
            + b"<w:bookmarkEnd w:id='1'/>" * 1_000
            + b"<w:r><w:t>x</w:t></w:r></w:p>"
        )
        copied = _docx_stats(data).copied_bytes
        assert copied >= len(name) + 1_000 * guard.DOCX_COPIED_BYTES_PER_NODE

    def test_paragraphs_and_tables_without_fan_out_copy_nothing(self):
        data = _blank_with(
            _text_paragraphs_body(50) + _spanned_table(1, b"cell", rows=5)
        )
        assert _docx_stats(data).copied_bytes == 0

    def test_continuing_merged_cells_count_the_largest_earlier_cell(self):
        text = b"r" * 10_000
        stats = _docx_stats(_blank_with(_merged_column(text, 12)))
        assert stats.copied_bytes == 12 * _cell_copy(text)

    def test_hostile_page_break_paragraph_is_refused_by_default(self):
        # 80 breaks before 3 MB of text: ~230 MB kept, from 40 KB.
        guard = _guard()
        data = _blank_with(
            b"<w:p><w:r>"
            + b"<w:lastRenderedPageBreak/>" * 80
            + b"</w:r>"
            + b"<w:r><w:t>"
            + b"ab " * 1_000_000
            + b"</w:t></w:r></w:p>"
        )
        with pytest.raises(guard.DecompressionBombError, match="copy"):
            guard.validate_zip_container(data, ".docx")

    def test_hostile_merged_column_is_refused_by_default(self):
        guard = _guard()
        data = _blank_with(_merged_column(b"ab " * 350_000, 300))
        with pytest.raises(guard.DecompressionBombError, match="copy"):
            guard.validate_zip_container(data, ".docx")

    def test_copied_bytes_accept_the_limit_and_refuse_one_more(self):
        guard = _guard()
        data = _blank_with(_page_break_paragraph(3, gap=50))
        copied = _docx_stats(data).copied_bytes
        assert copied > 0
        with _isolated(guard, "MAX_DOCX_COPIED_BYTES"):
            with patch.object(guard, "MAX_DOCX_COPIED_BYTES", copied):
                guard.validate_zip_container(data, ".docx")
            with patch.object(guard, "MAX_DOCX_COPIED_BYTES", copied - 1):
                with pytest.raises(guard.DecompressionBombError, match="copy"):
                    guard.validate_zip_container(data, ".docx")

    def test_copied_bytes_share_the_budget(self):
        guard = _guard()
        data = _blank_with(_page_break_paragraph(3, gap=50))
        stats = _docx_stats(data)
        copied_ceiling = stats.copied_bytes * 10 // 6
        run_ceiling = stats.run_work * 10 // 6
        with _isolated(guard, None):
            with patch.object(guard, "MAX_DOCX_COPIED_BYTES", copied_ceiling):
                guard.validate_zip_container(data, ".docx")
                with patch.object(guard, "MAX_DOCX_RUN_WORK", run_ceiling):
                    with pytest.raises(
                        guard.DecompressionBombError, match="budget"
                    ):
                        guard.validate_zip_container(data, ".docx")

    # python-docx's and unstructured's ``tc.vMerge == "continue"`` reads
    # the first w:vMerge of the cell's first w:tcPr, and a missing w:val
    # means "continue"; any other cell is an ordinary one, yielded once
    # per spanned column.

    def test_premise_python_docx_reads_the_first_vmerge_of_the_first_tcpr(
        self,
    ):
        docx = pytest.importorskip("docx")
        for encoding, value in _VMERGE_ENCODINGS.items():
            data = _blank_with(
                _vmerge_table(encoding, 7, rows=2, restart_first=True)
            )
            table = docx.Document(io.BytesIO(data)).tables[0]
            tc = table._tbl.tr_lst[1].tc_lst[0]
            assert tc.vMerge == value, encoding
            if value != "continue":
                assert len(table.rows[1].cells) == 7, encoding

    @pytest.mark.parametrize(
        "encoding",
        [e for e, value in _VMERGE_ENCODINGS.items() if value != "continue"],
    )
    def test_a_cell_python_docx_does_not_continue_counts_per_column(
        self, encoding
    ):
        text = b"c" * 10_000
        stats = _docx_stats(_blank_with(_vmerge_table(encoding, 7, text=text)))
        plain = _docx_stats(_blank_with(_vmerge_table("none", 7, text=text)))
        assert stats.copied_bytes == plain.copied_bytes == 6 * _cell_copy(text)
        assert stats.run_work == plain.run_work

    @pytest.mark.parametrize(
        "encoding",
        [e for e, value in _VMERGE_ENCODINGS.items() if value == "continue"],
    )
    def test_a_cell_python_docx_continues_counts_the_largest_earlier_cell(
        self, encoding
    ):
        text = b"c" * 10_000
        data = _blank_with(
            _vmerge_table(encoding, 7, text=text, rows=2, restart_first=True)
        )
        # The restart row: 6 extra copies; the continuing row: the
        # largest earlier cell (its text times its 7 columns).
        assert _docx_stats(data).copied_bytes == (6 + 7) * _cell_copy(text)

    @pytest.mark.parametrize(
        "encoding",
        [e for e, value in _VMERGE_ENCODINGS.items() if value != "continue"],
    )
    def test_hostile_spanned_cell_with_a_vmerge_python_docx_ignores(
        self, encoding
    ):
        # Two rows of one 100 KB cell spanning 1,000 columns with a
        # w:vMerge python-docx does not read as "continue": it yields
        # each cell 1,000 times (~200 MB of text kept, ~13 GB at 1 MB
        # cells), and the guard used to count no copies at all.
        guard = _guard()
        data = _blank_with(
            _vmerge_table(encoding, 1_000, text=b"ab " * 33_000, rows=2)
        )
        with pytest.raises(guard.DecompressionBombError, match="copy"):
            guard.validate_zip_container(data, ".docx")

    def test_premise_a_deep_copy_carries_declarations_and_tail(self):
        from lxml import etree

        root = etree.fromstring(
            b"<d xmlns:w='urn:w' xmlns:u='urn:used' xmlns:x='urn:unused'>"
            b"<w:body><w:p xmlns:p='urn:own'><w:r u:a='1'/></w:p>TAIL"
            b"</w:body></d>"
        )
        copied = copy.deepcopy(root[0][0])
        assert copied.tail == "TAIL"
        assert copied.nsmap == {"p": "urn:own", "w": "urn:w", "u": "urn:used"}

    def test_namespace_declarations_in_a_paragraph_count(self):
        uri = "urn:" + "u" * (_guard().MAX_DOCX_NAMESPACE_URI_CHARS - 4)
        base = _docx_stats(_blank_with(_page_break_paragraph(3))).copied_bytes
        for declared in (
            _page_break_paragraph(3).replace(
                b"<w:r>", b"<w:r xmlns:z='%s'>" % uri.encode(), 1
            ),
            _page_break_paragraph(3).replace(
                b"<w:p>", b"<w:p xmlns:z='%s'>" % uri.encode(), 1
            ),
        ):
            copied = _docx_stats(_blank_with(declared)).copied_bytes
            assert copied - base == 2 * 3 * _declaration_bytes("z", uri)

    def test_used_ancestor_namespace_declarations_count(self):
        uri = "urn:" + "u" * (_guard().MAX_DOCX_NAMESPACE_URI_CHARS - 4)
        root = b"<w:document xmlns:z='%s' " % uri.encode()

        def stats(body: bytes):
            data = _rewrite_zip(
                _blank_with(body),
                {
                    "word/document.xml": lambda xml: xml.replace(
                        b"<w:document ", root, 1
                    )
                },
            )
            return _docx_stats(data).copied_bytes

        unused = stats(_page_break_paragraph(3))
        used = stats(
            _page_break_paragraph(3).replace(b"<w:r>", b"<w:r z:a='1'>", 1)
        )
        # The attribute (its node and value) and the declaration, per
        # copy.
        assert used - unused == 2 * 3 * (
            _guard().DOCX_COPIED_BYTES_PER_ATTRIBUTE
            + 1
            + _declaration_bytes("z", uri)
        )
        assert (
            unused
            == _docx_stats(_blank_with(_page_break_paragraph(3))).copied_bytes
        )

    def test_the_text_after_a_page_break_paragraph_counts(self):
        tail = b"t" * 50_000
        base = _docx_stats(_blank_with(_page_break_paragraph(3))).copied_bytes
        for after in (b"", b"<!--c-->", b"<w:p/>"):
            body = _page_break_paragraph(3) + tail + after
            copied = _docx_stats(_blank_with(body)).copied_bytes
            assert copied - base == 2 * 3 * len(tail)
        # Not after a paragraph without breaks.
        body = _text_paragraphs_body(1) + tail
        assert _docx_stats(_blank_with(body)).copied_bytes == 0

    @pytest.mark.parametrize("where", ["w:p", "w:r", "tail", "ancestor"])
    def test_hostile_declaration_heavy_page_break_paragraph_is_refused(
        self, where
    ):
        # 80 breaks before 1 MB in a namespace URI or after the
        # paragraph (a ~37 KB file): every copy keeps it, ~150 MB in
        # all (~1.5 GB at 9.5 MB), and the guard used to count none.
        # The text is refused as copies; a URI that long is now refused
        # outright (MAX_DOCX_NAMESPACE_URI_CHARS).
        guard = _guard()
        uri = b"urn:" + b"u" * 1_000_000
        paragraph = _page_break_paragraph(80, gap=0)
        data = None
        if where == "w:p":
            paragraph = paragraph.replace(b"<w:p>", b"<w:p xmlns:z='%s'>" % uri)
        elif where == "w:r":
            paragraph = paragraph.replace(b"<w:r>", b"<w:r xmlns:z='%s'>" % uri)
        elif where == "tail":
            paragraph += uri
        else:
            paragraph = paragraph.replace(b"<w:r>", b"<w:r z:a='1'>")
            data = _rewrite_zip(
                _blank_with(paragraph),
                {
                    "word/document.xml": lambda xml: xml.replace(
                        b"<w:document ", b"<w:document xmlns:z='%s' " % uri, 1
                    )
                },
            )
        if data is None:
            data = _blank_with(paragraph)
        match = "copy" if where == "tail" else "namespace URI"
        with pytest.raises(guard.DecompressionBombError, match=match):
            guard.validate_zip_container(data, ".docx")

    # Each attribute is copied as two libxml2 nodes (~275 bytes per
    # empty attribute per copy measured), which an empty value used to
    # leave uncounted.

    @pytest.mark.parametrize("attributes", [10, 100])
    def test_each_attribute_counts_a_node_and_its_value(self, attributes):
        # 10 are read with attrib.values(), 100 with the @* XPath.
        guard = _guard()
        base = _docx_stats(_blank_with(_page_break_paragraph(3))).copied_bytes
        for value in (b"", b"vv"):
            marker = (
                b"<w:bookmarkStart "
                + b" ".join(b"a%d='%s'" % (i, value) for i in range(attributes))
                + b"/>"
            )
            paragraph = _page_break_paragraph(3).replace(
                b"<w:p>", b"<w:p>" + marker, 1
            )
            copied = _docx_stats(_blank_with(paragraph)).copied_bytes
            assert copied - base == 2 * 3 * (
                guard.DOCX_COPIED_BYTES_PER_NODE
                + attributes
                * (guard.DOCX_COPIED_BYTES_PER_ATTRIBUTE + len(value))
            )

    def test_hostile_attribute_heavy_page_break_paragraph_is_refused(self):
        # 20 rendered page breaks, then 200 elements of 256 empty
        # attributes (a ~40 KB file): the fragments unstructured keeps
        # held ~310 MB more (~14 MB of attribute nodes per copy), and
        # the guard used to count ~1 MB.
        guard = _guard()
        element = (
            b"<w:bookmarkEnd "
            + b" ".join(b"a%d=''" % i for i in range(256))
            + b"/>"
        )
        paragraph = (
            b"<w:p>"
            + b"<w:r><w:t>a</w:t><w:lastRenderedPageBreak/><w:t>a</w:t></w:r>"
            * 20
            + element * 200
            + b"<w:r><w:t>z</w:t></w:r></w:p>"
        )
        with pytest.raises(guard.DecompressionBombError, match="copy"):
            guard.validate_zip_container(_blank_with(paragraph), ".docx")

    # unstructured's HTML keeps every extra copy of a cell as a string,
    # then HTML-escaped in its row's and in its table's string: a '"'
    # becomes '&quot;', and one astral character makes a string four
    # bytes per character.

    def test_premise_table_html_escapes_every_copy_of_a_spanned_cell(self):
        pytest.importorskip("docx")
        udocx = pytest.importorskip("unstructured.partition.docx")
        text = '"q"\U0001f600'
        data = _blank_with(_spanned_table(5, text.encode()))
        (table,) = [
            e
            for e in udocx.partition_docx(file=io.BytesIO(data))
            if e.category == "Table"
        ]
        html = table.metadata.text_as_html
        assert html.count("&quot;q&quot;\U0001f600") == 5

    def test_spanned_cell_copies_count_each_text_byte_at_the_html_factor(
        self,
    ):
        text = ('"' * 1_000 + "\U0001f600").encode()
        stats = _docx_stats(_blank_with(_spanned_table(21, text)))
        assert stats.copied_bytes == 20 * _cell_copy(text)
        assert _cell_copy(text) == (
            _guard().DOCX_TABLE_COPY_BYTES_PER_TEXT_BYTE * (1_000 + 4)
        )

    @pytest.mark.parametrize("shape", ["spanned", "merged"])
    def test_hostile_quoted_astral_cell_copies_are_refused(self, shape):
        # 20 extra copies of a 1 MB cell of '"' and one emoji (a ~38 KB
        # file) kept ~1 GB in unstructured's table HTML (~49 bytes per
        # byte per copy); the guard counted 20 MB and accepted it.
        guard = _guard()
        text = b'"' * 1_000_000 + "\U0001f600".encode()
        if shape == "spanned":
            body = _spanned_table(21, text)
        else:
            body = _merged_column(text, 20)
        with pytest.raises(guard.DecompressionBombError, match="copy"):
            guard.validate_zip_container(_blank_with(body), ".docx")

    def test_ordinary_spanned_and_merged_text_is_accepted(self):
        # A 2 KB header cell across 20 columns and a 2 KB cell merged
        # down 100 rows: ~15 MB counted.
        guard = _guard()
        text = b"Quarterly results " * 110
        body = _spanned_table(20, text) + _merged_column(text, 100)
        guard.validate_zip_container(_blank_with(body), ".docx")


_W_DECL = b"xmlns:w='%s'" % _W_MAIN.encode()
_PLAIN_CELL = b"<w:tc><w:p/></w:tc>"
_CONTINUING_CELL = b"<w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc>"


def _merged_rows(width: int, rows: int) -> bytes:
    """A table of *rows* rows of *width* cells: a plain first row, and
    every cell of the others continuing a vertical merge."""
    return (
        b"<w:tbl><w:tr>"
        + _PLAIN_CELL * width
        + b"</w:tr>"
        + (b"<w:tr>" + _CONTINUING_CELL * width + b"</w:tr>") * (rows - 1)
        + b"</w:tbl>"
    )


def _random_tc(rng: random.Random, span: int, row: int) -> bytes:
    """A w:tc of *span* columns in table row *row* with a random vMerge,
    in encodings python-docx reads one way and a careless reader
    another (a second w:tcPr, a w:tcPr after the content, an omitted,
    zero or negative gridSpan)."""
    merge = rng.choices(
        (
            None,
            b"<w:vMerge/>",
            b"<w:vMerge w:val='continue'/>",
            b"<w:vMerge w:val='restart'/>",
            b"<w:vMerge w:val='x'/>",
        ),
        weights=(35, 25, 20, 15, 5) if row else (70, 2, 3, 20, 5),
    )[0]
    if rng.random() < 0.08:
        span = rng.choice((0, -1))
    parts = (
        []
        if span == 1 and rng.random() < 0.5
        else [b"<w:gridSpan w:val='%d'/>" % span]
    )
    if merge is not None:
        parts.insert(rng.randint(0, len(parts)), merge)
    tc_pr = b"<w:tcPr>" + b"".join(parts) + b"</w:tcPr>"
    pieces = [tc_pr, b"<w:p/>"] if rng.random() < 0.9 else [b"<w:p/>", tc_pr]
    if rng.random() < 0.1:
        pieces.append(
            rng.choice(
                (
                    b"<w:tcPr><w:gridSpan w:val='2'/></w:tcPr>",
                    b"<w:tcPr><w:vMerge/></w:tcPr>",
                    b"<w:tcPr><w:vMerge w:val='restart'/></w:tcPr>",
                )
            )
        )
    return b"<w:tc>" + b"".join(pieces) + b"</w:tc>"


def _random_merged_table(rng: random.Random) -> bytes:
    """A table of a few rows that mostly share one column layout, with
    random merges, gridBefore values and other children between the
    cells."""
    layout = [rng.choice((1, 1, 1, 2, 3)) for _ in range(rng.randint(1, 5))]
    rows = []
    for index in range(rng.randint(2, 7)):
        spans = (
            layout
            if rng.random() < 0.85
            else [rng.choice((1, 2)) for _ in range(rng.randint(1, 5))]
        )
        parts = []
        if rng.random() < 0.15:
            parts.append(
                b"<w:trPr><w:gridBefore w:val='%d'/></w:trPr>"
                % rng.choice((0, 1, 2, -1))
            )
            if rng.random() < 0.3:
                parts.append(b"<w:trPr><w:gridBefore w:val='1'/></w:trPr>")
        for span in spans:
            if rng.random() < 0.1:
                parts.append(b"<w:bookmarkStart w:id='0' w:name='b'/>")
            parts.append(_random_tc(rng, span, index))
        rows.append(b"<w:tr>" + b"".join(parts) + b"</w:tr>")
    return b"<w:tbl>" + b"".join(rows) + b"</w:tbl>"


class TestDocxMergedRows:
    # python-docx resolves a cell continuing a vertical merge
    # (``row.cells``) by computing its grid offset, which reads the
    # span of every earlier cell in its row, and finding the cell at
    # that offset in the row above, recursively up the merge;
    # unstructured reads ``row.cells`` twice per table. Each step costs
    # time proportional to the rows' width, and a merge D rows deep
    # takes D steps per cell: two rows of 1,000 cells, the second all
    # continuing (a 37 KB file), took 6-7 s per read, and every column
    # merged down 50 rows of 64 cells 42 s per read, while the guard
    # counted 15,000 and 32,000 run-work units (of 1.5 million). Rows are now
    # capped at 256 cells and 2,048 grid columns, and every step is
    # priced in MAX_DOCX_RUN_WORK.

    def test_premise_python_docx_steps_once_per_merged_row_above(self):
        docx = pytest.importorskip("docx")
        from docx.oxml.table import CT_Tc

        steps = []
        above = CT_Tc._tc_above

        def counting(tc):
            steps.append(tc)
            return above.fget(tc)

        table = docx.Document(
            io.BytesIO(_blank_with(_merged_rows(3, 5)))
        ).tables[0]
        with patch.object(CT_Tc, "_tc_above", property(counting)):
            for row in table.rows:
                assert len(row.cells) == 3
        # Row r's cells each step up r rows.
        assert len(steps) == 3 * (1 + 2 + 3 + 4)

    def test_wide_merged_row_is_refused(self):
        # 37 KB; 6-7 s per read of its cells, ~15 run-work units a cell
        # before.
        guard = _guard()
        with pytest.raises(guard.DecompressionBombError, match="cells"):
            guard.validate_zip_container(
                _blank_with(_merged_rows(1_000, 2)), ".docx"
            )

    def test_row_cell_cap_accepts_the_limit_and_refuses_one_more(self):
        guard = _guard()
        cap = guard.MAX_DOCX_ROW_CELLS
        row = b"<w:tbl><w:tr>%s</w:tr></w:tbl>"
        guard.validate_zip_container(
            _blank_with(row % (_PLAIN_CELL * cap)), ".docx"
        )
        with pytest.raises(guard.DecompressionBombError, match="cells"):
            guard.validate_zip_container(
                _blank_with(row % (_PLAIN_CELL * (cap + 1))), ".docx"
            )

    def test_cells_of_nested_rows_count_per_row(self):
        # Only a row's own w:tc children count towards its cells.
        guard = _guard()
        half = guard.MAX_DOCX_ROW_CELLS // 2 + 1
        inner = b"<w:tbl><w:tr>" + _PLAIN_CELL * half + b"</w:tr></w:tbl>"
        outer = (
            b"<w:tbl><w:tr>"
            + _PLAIN_CELL * half
            + b"<w:tc>"
            + inner
            + b"<w:p/></w:tc></w:tr></w:tbl>"
        )
        guard.validate_zip_container(_blank_with(outer), ".docx")

    def test_row_grid_column_cap_accepts_the_limit_and_refuses_one_more(self):
        guard = _guard()
        cap = guard.MAX_DOCX_ROW_GRID_COLUMNS
        span = guard.MAX_DOCX_GRID_SPAN

        def row(columns: int) -> bytes:
            # gridBefore, cells of at most the span cap, gridAfter.
            cells = []
            rest = columns - 2
            while rest:
                width = min(rest, span)
                cells.append(
                    b"<w:tc><w:tcPr><w:gridSpan w:val='%d'/></w:tcPr>"
                    b"<w:p/></w:tc>" % width
                )
                rest -= width
            return (
                b"<w:tbl><w:tr><w:trPr><w:gridBefore w:val='1'/>"
                b"<w:gridAfter w:val='1'/></w:trPr>"
                + b"".join(cells)
                + b"</w:tr></w:tbl>"
            )

        with _isolated(guard, None):
            guard.validate_zip_container(_blank_with(row(cap)), ".docx")
            with pytest.raises(guard.DecompressionBombError, match="columns"):
                guard.validate_zip_container(_blank_with(row(cap + 1)), ".docx")

    def test_cap_is_well_above_real_table_widths(self):
        # Word's tables are at most 63 columns wide.
        guard = _guard()
        assert guard.MAX_DOCX_ROW_CELLS >= 4 * 63
        assert guard.MAX_DOCX_ROW_GRID_COLUMNS >= 2 * guard.MAX_DOCX_GRID_SPAN

    def test_merge_steps_grow_with_the_merge_depth(self):
        # One column merged down R rows takes R * (R - 1) / 2 steps.
        guard = _guard()
        weight = guard.DOCX_MERGE_STEP_WEIGHT
        for rows in (2, 10, 40):
            work = _docx_stats(_blank_with(_merged_rows(1, rows))).run_work
            assert work >= weight * rows * (rows - 1) // 2

    def test_long_merged_column_is_refused(self):
        # 600 rows of one merged column, a ~37 KB file: ~180,000 steps
        # of ~0.1 ms per read (one 300-row column took 8 s to partition),
        # counted ~9,000 units before.
        guard = _guard()
        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(
                _blank_with(_merged_rows(1, 600)), ".docx"
            )

    def test_wide_deep_merge_is_refused(self):
        # 64 columns merged down 30 rows: ~28,000 steps at ~0.5 ms per
        # read each.
        guard = _guard()
        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(
                _blank_with(_merged_rows(64, 30)), ".docx"
            )

    def test_steps_follow_python_docx_on_random_tables(self):
        # The guard walks the merges as python-docx does (spans and
        # gridBefore as python-docx reads them, tc_at_grid_offset's
        # search): never fewer steps, and the same number whenever
        # python-docx resolves every cell.
        pytest.importorskip("docx")
        from docx.oxml.table import CT_Tc

        guard = _guard()
        steps = []
        above = CT_Tc._tc_above

        def counting(tc):
            steps.append(tc)
            return above.fget(tc)

        from docx.oxml import parse_xml
        from docx.table import Table

        rng = random.Random(6515)
        exact = deep = 0
        for _ in range(150):
            xml = _random_merged_table(rng)
            data = _blank_with(xml)
            with (
                patch.object(guard, "DOCX_MERGE_STEP_WEIGHT", 10**9),
                patch.object(guard, "DOCX_MERGE_CELLS_PER_UNIT", 10**12),
                patch.object(guard, "DOCX_MERGE_NODES_PER_UNIT", 10**12),
            ):
                counted = _docx_stats(data).run_work // 10**9
            del steps[:]
            failed = False
            table = Table(
                parse_xml(xml.replace(b"<w:tbl>", b"<w:tbl %s>" % _W_DECL, 1)),
                None,  # type: ignore[arg-type]
            )
            with patch.object(CT_Tc, "_tc_above", property(counting)):
                for row in table.rows:
                    try:
                        row.cells  # noqa: B018
                    except ValueError:
                        failed = True
            assert counted >= len(steps)
            if not failed:
                assert counted == len(steps)
                exact += 1
                deep += len(steps) > 1
        assert exact > 40 and deep > 25

    def test_real_wide_merged_tables_are_accepted(self):
        # A 63-column table (Word's widest) with a merged header and its
        # first column merged down 30 rows, as python-docx writes it.
        docx = pytest.importorskip("docx")
        document = docx.Document()
        table = document.add_table(rows=31, cols=63)
        table.cell(0, 0).merge(table.cell(0, 62))
        table.cell(1, 0).merge(table.cell(30, 0))
        for row in table.rows[1:]:
            row.cells[1].text = "value"
        buf = io.BytesIO()
        document.save(buf)
        guard = _guard()
        guard.validate_zip_container(buf.getvalue(), ".docx")
        stats = _docx_stats(buf.getvalue())
        assert stats.run_work < guard.MAX_DOCX_RUN_WORK // 4


_MARKER = b"<w:bookmarkEnd w:id='0'/>"


def _table_with_siblings(
    width: int, rows: int, between: bytes = b"", trailing: bytes = b""
) -> bytes:
    """_merged_rows(width, rows) with *between* before every row after
    the first and *trailing* after the cells of every row."""
    return (
        b"<w:tbl><w:tr>"
        + _PLAIN_CELL * width
        + trailing
        + b"</w:tr>"
        + (
            between
            + b"<w:tr>"
            + _CONTINUING_CELL * width
            + trailing
            + b"</w:tr>"
        )
        * (rows - 1)
        + b"</w:tbl>"
    )


class TestDocxTableSiblings:
    # Every step python-docx takes up a vertical merge reads the
    # continuing cell's row's grid_before (an lxml find over all the
    # row's children when it has no w:trPr) and evaluates _tr_above
    # (ancestor::w:tr[1]/preceding-sibling::w:tr[1], which visits each
    # sibling between that row and the one above). Neither was priced:
    # a row padded with 200,000 empty elements after its cell (a 1.6 KB
    # file, at 0.09 of the run-work ceiling) took 7.6 s per read of the
    # cells, and 64 merged columns with 200,000 elements between each
    # two of 26 rows (81 KB, at 0.90) 58 s per read. Such siblings are
    # now capped and priced into every step.

    def test_wrappers_do_not_add_direct_row_or_table_siblings(self):
        """The guard prices direct siblings, not descendants of a wrapper.

        This checks the shape python-docx actually parses. A wall-clock
        ratio is too sensitive to CI contention to establish that premise.
        """
        docx = pytest.importorskip("docx")
        loose = _MARKER * 100
        wrapped = b"<w:sdt><w:sdtContent>" + loose + b"</w:sdtContent></w:sdt>"

        def table(between: bytes, trailing: bytes):
            return docx.Document(
                io.BytesIO(
                    _blank_with(_table_with_siblings(1, 2, between, trailing))
                )
            ).tables[0]

        loose_between = table(loose, b"")
        wrapped_between = table(wrapped, b"")
        loose_trailing = table(b"", loose)
        wrapped_trailing = table(b"", wrapped)

        assert len(loose_between._tbl) == len(wrapped_between._tbl) + 99
        assert (
            len(loose_trailing.rows[1]._tr)
            == len(wrapped_trailing.rows[1]._tr) + 99
        )
        for parsed in (
            loose_between,
            wrapped_between,
            loose_trailing,
            wrapped_trailing,
        ):
            assert len(parsed.rows[1].cells) == 1

    def test_caps_are_at_least_a_wrapped_word_row(self):
        # Word's widest row (63 cells), each in a content control, with
        # w:trPr, w:tblPrEx and range markers.
        guard = _guard()
        assert guard.MAX_DOCX_ROW_OTHER_CHILDREN >= 63 + 2 + 32
        assert guard.MAX_DOCX_TABLE_ROW_GAP >= 2 + 32

    @pytest.mark.parametrize("node", [_MARKER, b"<!--c-->", b"<?pi x?>"])
    def test_row_children_cap_accepts_the_limit_and_refuses_one_more(
        self, node
    ):
        guard = _guard()
        cap = guard.MAX_DOCX_ROW_OTHER_CHILDREN
        guard.validate_zip_container(
            _blank_with(_table_with_siblings(2, 2, trailing=node * cap)),
            ".docx",
        )
        with pytest.raises(guard.DecompressionBombError, match="not cells"):
            guard.validate_zip_container(
                _blank_with(
                    _table_with_siblings(2, 2, trailing=node * (cap + 1))
                ),
                ".docx",
            )

    def test_row_children_count_wherever_they_sit_in_the_row(self):
        # grid_before's find walks all of them when there is no w:trPr.
        guard = _guard()
        cap = guard.MAX_DOCX_ROW_OTHER_CHILDREN
        half = _MARKER * (cap // 2 + 1)
        row = b"<w:tbl><w:tr>" + half + _PLAIN_CELL + half + b"</w:tr></w:tbl>"
        with pytest.raises(guard.DecompressionBombError, match="not cells"):
            guard.validate_zip_container(_blank_with(row), ".docx")

    @pytest.mark.parametrize("node", [_MARKER, b"<!--c-->", b"<?pi x?>"])
    def test_children_between_rows_cap_accepts_the_limit_and_refuses_one_more(
        self, node
    ):
        guard = _guard()
        cap = guard.MAX_DOCX_TABLE_ROW_GAP
        guard.validate_zip_container(
            _blank_with(_table_with_siblings(2, 3, between=node * cap)),
            ".docx",
        )
        with pytest.raises(guard.DecompressionBombError, match="between"):
            guard.validate_zip_container(
                _blank_with(
                    _table_with_siblings(2, 3, between=node * (cap + 1))
                ),
                ".docx",
            )

    def test_children_before_the_first_row_count_with_tblpr_and_tblgrid(self):
        # A continuing cell in the first row walks them all (and then
        # python-docx raises).
        guard = _guard()
        cap = guard.MAX_DOCX_TABLE_ROW_GAP

        def table(markers: int) -> bytes:
            return (
                b"<w:tbl><w:tblPr/><w:tblGrid/>"
                + _MARKER * markers
                + b"<w:tr>"
                + _CONTINUING_CELL
                + b"</w:tr></w:tbl>"
            )

        guard.validate_zip_container(_blank_with(table(cap - 2)), ".docx")
        with pytest.raises(guard.DecompressionBombError, match="between"):
            guard.validate_zip_container(_blank_with(table(cap - 1)), ".docx")

    def test_wrapped_content_and_children_after_the_last_row_are_free(self):
        # No step walks inside a w:sdt or past the last row: a row group
        # or cell in a content control holds any number of elements, and
        # a table may keep all its rows in w:sdt wrappers.
        guard = _guard()
        wrapped = (
            b"<w:sdt><w:sdtContent>"
            + _MARKER * 5_000
            + (b"</w:sdtContent></w:sdt>")
        )
        sdt_rows = (
            b"<w:sdt><w:sdtContent><w:tr>"
            + _PLAIN_CELL
            + b"</w:tr></w:sdtContent></w:sdt>"
        )
        for body in (
            _table_with_siblings(4, 5, between=wrapped, trailing=wrapped),
            b"<w:tbl><w:tblPr/><w:tblGrid/>" + sdt_rows * 1_000 + b"</w:tbl>",
            _table_with_siblings(4, 5)[: -len(b"</w:tbl>")]
            + _MARKER * 5_000
            + b"</w:tbl>",
            # Rows of a nested table inside a wrapper are not table rows
            # python-docx reads.
            b"<w:tbl><w:tr><w:tc><w:sdt><w:sdtContent><w:tr>"
            + _MARKER * 1_000
            + _PLAIN_CELL
            + b"</w:tr></w:sdtContent></w:sdt><w:p/></w:tc></w:tr></w:tbl>",
        ):
            guard.validate_zip_container(_blank_with(body), ".docx")

    def test_hostile_padded_rows_are_refused(self):
        # The gate's shapes: 200,000 elements after one row's cell, and
        # between the rows of 64 columns merged down 26 rows.
        guard = _guard()
        padding = b"<w:x/>" * 200_000
        for body in (
            _table_with_siblings(1, 2, trailing=padding),
            _table_with_siblings(64, 26, between=padding),
        ):
            with pytest.raises(guard.DecompressionBombError):
                guard.validate_zip_container(_blank_with(body), ".docx")

    def test_every_step_is_priced_at_the_rows_it_walks_between(self):
        # One column merged down 40 rows takes 780 steps; each walks the
        # siblings before its row once more.
        guard = _guard()
        cap = guard.MAX_DOCX_TABLE_ROW_GAP
        plain = _docx_stats(_blank_with(_table_with_siblings(1, 40))).run_work
        padded = _docx_stats(
            _blank_with(_table_with_siblings(1, 40, between=_MARKER * cap))
        ).run_work
        steps = 40 * 39 // 2
        assert padded - plain >= steps * (
            cap // guard.DOCX_MERGE_NODES_PER_UNIT
        )

    def test_every_step_is_priced_at_its_rows_other_children(self):
        # 64 cells continuing a plain row: 64 steps, each reading the
        # continuing row's grid_before over all its children (the
        # siblings after the last cell were not counted).
        guard = _guard()
        cap = guard.MAX_DOCX_ROW_OTHER_CHILDREN

        def table(markers: int) -> bytes:
            return (
                b"<w:tbl><w:tr>"
                + _PLAIN_CELL * 64
                + b"</w:tr><w:tr>"
                + _CONTINUING_CELL * 64
                + _MARKER * markers
                + b"</w:tr></w:tbl>"
            )

        plain = _docx_stats(_blank_with(table(0))).run_work
        padded = _docx_stats(_blank_with(table(cap))).run_work
        assert padded - plain >= 64 * (cap // guard.DOCX_MERGE_NODES_PER_UNIT)

    def test_every_step_is_priced_at_its_rows_first_trpr(self):
        # grid_before looks in the row's first w:trPr wherever it sits:
        # one after 256 continuing cells holding 100,000 children (a
        # 38 KB file) added ~0.28 s per read at 0.04 of the ceiling.
        guard = _guard()

        def table(markers: int) -> bytes:
            return (
                b"<w:tbl><w:tr>"
                + _PLAIN_CELL * 64
                + b"</w:tr><w:tr>"
                + _CONTINUING_CELL * 64
                + b"<w:trPr>"
                + _MARKER * markers
                + b"</w:trPr></w:tr></w:tbl>"
            )

        plain = _docx_stats(_blank_with(table(0))).run_work
        padded = _docx_stats(_blank_with(table(4_096))).run_work
        assert padded - plain >= 64 * (4_096 // guard.DOCX_MERGE_NODES_PER_UNIT)

    def test_real_tables_with_markers_and_wrappers_are_accepted(self):
        # A Word-like table: w:tblPrEx and w:trPr on every row, bookmark,
        # comment, permission and proofing markers between and inside
        # rows, every cell of a 63-column row in a content control, a
        # group of rows in one, and the first column merged down.
        guard = _guard()
        markers = (
            b"<w:bookmarkStart w:id='1' w:name='r'/><w:bookmarkEnd w:id='1'/>"
            b"<w:commentRangeStart w:id='2'/><w:commentRangeEnd w:id='2'/>"
            b"<w:permStart w:id='3' w:edGrp='everyone'/><w:permEnd w:id='3'/>"
            b"<w:proofErr w:type='spellStart'/><w:proofErr w:type='spellEnd'/>"
        )
        row_pr = (
            b"<w:tblPrEx><w:tblBorders/></w:tblPrEx>"
            b"<w:trPr><w:trHeight w:val='300'/></w:trPr>"
        )
        cell = b"<w:tc><w:p><w:r><w:t>value</w:t></w:r></w:p></w:tc>"
        wrapped_cell = (
            b"<w:sdt><w:sdtPr/><w:sdtContent>%s</w:sdtContent></w:sdt>"
        )
        rows = [
            b"<w:tr>" + row_pr + markers + wrapped_cell % cell * 63 + b"</w:tr>"
        ]
        for _ in range(40):
            rows.append(
                b"<w:tr>"
                + row_pr
                + markers
                + _CONTINUING_CELL
                + cell * 62
                + markers
                + b"</w:tr>"
            )
        group = (
            b"<w:sdt><w:sdtPr/><w:sdtContent>"
            + (b"<w:tr>" + row_pr + cell * 63 + b"</w:tr>") * 3
            + b"</w:sdtContent></w:sdt>"
        )
        body = (
            b"<w:tbl><w:tblPr/><w:tblGrid/>"
            + markers
            + (markers + group).join(rows)
            + markers
            + b"</w:tbl>"
        )
        data = _blank_with(body)
        guard.validate_zip_container(data, ".docx")
        assert _docx_stats(data).run_work < guard.MAX_DOCX_RUN_WORK // 4

    def test_python_docx_table_with_bookmarks_is_accepted(self):
        # python-docx's own table (63 columns, a merged header and first
        # column), with a bookmark around every row, as Word keeps one
        # spanning rows.
        docx = pytest.importorskip("docx")
        from docx.oxml import parse_xml

        document = docx.Document()
        table = document.add_table(rows=31, cols=63)
        table.cell(0, 0).merge(table.cell(0, 62))
        table.cell(1, 0).merge(table.cell(30, 0))
        for index, row in enumerate(table.rows):
            tr = row._tr
            for tag, position in (("bookmarkStart", 0), ("bookmarkEnd", 1)):
                marker = parse_xml(
                    b"<w:%s %s w:id='%d' w:name='b%d'/>"
                    % (tag.encode(), _W_DECL, index, index)
                )
                if position:
                    tr.addnext(marker)
                else:
                    tr.addprevious(marker)
        buf = io.BytesIO()
        document.save(buf)
        guard = _guard()
        guard.validate_zip_container(buf.getvalue(), ".docx")
        stats = _docx_stats(buf.getvalue())
        assert stats.run_work < guard.MAX_DOCX_RUN_WORK // 4


class TestDocxCellContentPerYield:
    # unstructured reads a cell's content once per grid column it is
    # yielded for (iter_inner_content, paragraphs, Paragraph.text), so
    # its nodes other than runs are walked span times too: a cell
    # spanning 1,000 columns with 60,000 empty elements after its
    # paragraph took ~1.7 s to partition and counted 5,005 units; 30
    # such rows (a 54 KB file) were accepted at 0.10 of the ceiling.

    def test_cell_nodes_count_once_per_spanned_column(self):
        guard = _guard()
        nodes = 4 * guard.DOCX_CELL_NODES_PER_UNIT

        def table(markers: int) -> bytes:
            return (
                b"<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val='100'/>"
                b"</w:tcPr><w:p/>"
                + _MARKER * markers
                + b"</w:tc></w:tr></w:tbl>"
            )

        extra = (
            _docx_stats(_blank_with(table(nodes))).run_work
            - _docx_stats(_blank_with(table(0))).run_work
        )
        assert extra >= 100 * 4

    def test_hostile_padded_spanned_cells_are_refused(self):
        # Seven rows of one cell spanning 1,000 columns with 60,000
        # elements after its paragraph (~12 s to partition): 35,035
        # units counted before.
        guard = _guard()
        cell = (
            b"<w:tc><w:tcPr><w:gridSpan w:val='1000'/></w:tcPr><w:p/>"
            + _MARKER * 60_000
            + b"</w:tc>"
        )
        body = b"<w:tbl>" + (b"<w:tr>" + cell + b"</w:tr>") * 7 + b"</w:tbl>"
        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(_blank_with(body), ".docx")

    def test_ordinary_cells_count_no_node_units(self):
        # A cell of a few paragraphs and runs holds well under the unit.
        guard = _guard()
        cell = (
            b"<w:tc><w:tcPr><w:tcW w:w='100' w:type='dxa'/></w:tcPr>"
            + b"<w:p><w:pPr/><w:r><w:rPr><w:b/></w:rPr><w:t>v</w:t></w:r></w:p>"
            * 8
            + b"</w:tc>"
        )
        assert guard.DOCX_CELL_NODES_PER_UNIT >= 64
        weight = guard.DOCX_RUN_WORK_ITEM_WEIGHT
        stats = _docx_stats(
            _blank_with(b"<w:tbl><w:tr>" + cell + b"</w:tr></w:tbl>")
        )
        assert stats.run_work == weight * 2 + 8 * (
            weight + 2 + guard.DOCX_CELL_BLOCK_ITEM_WEIGHT
        )


#: Namespace declarations on the root of the documents _blank_with
#: builds (python-docx's template).
_TEMPLATE_ROOT_DECLARATIONS = 17


def _with_attributes(count: int, tag: bytes = b"w:p") -> bytes:
    return (
        b"<%s " % tag + b" ".join(b"a%d=''" % i for i in range(count)) + b"/>"
    )


def _declarations(count: int, start: int = 0) -> bytes:
    return b" ".join(
        b"xmlns:n%d='urn:n%d'" % (i, i) for i in range(start, start + count)
    )


class TestDocxAbnormalMarkupIsRefused:
    # No writer puts more than a few dozen attributes on an element or
    # declares more than a few dozen namespaces in scope (python-docx's
    # template: 17 on w:document; Word ~35), and both have costs that
    # grow faster than their size: lxml reads each attribute value by
    # rescanning the element's attribute list (one paragraph of 50,000
    # empty attributes, a 145 KB file, kept the guard itself busy for
    # 40 s), and libxml2 resolves each copied element's namespace by
    # walking the declarations in scope. Both are refused outright.

    def test_premise_template_root_declarations(self):
        data = _blank_with(b"")
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("word/document.xml")
        root = xml[: xml.index(b"<w:body>")]
        assert root.count(b"xmlns:") == _TEMPLATE_ROOT_DECLARATIONS

    @pytest.mark.parametrize("tag", [b"w:p", b"w:bookmarkStart"])
    def test_attribute_cap_accepts_the_limit_and_refuses_one_more(self, tag):
        guard = _guard()
        cap = guard.MAX_DOCX_ELEMENT_ATTRIBUTES
        body = _with_attributes(cap, tag)
        if tag != b"w:p":
            body = b"<w:p>" + body + b"</w:p>"
        guard.validate_zip_container(_blank_with(body), ".docx")
        body = body.replace(b"/>", b" extra=''/>", 1)
        with pytest.raises(guard.DecompressionBombError, match="attributes"):
            guard.validate_zip_container(_blank_with(body), ".docx")

    def test_attribute_cap_applies_outside_paragraphs_and_tables(self):
        # The w:body itself, which the guard prices nothing for.
        guard = _guard()
        cap = guard.MAX_DOCX_ELEMENT_ATTRIBUTES
        attrs = b" ".join(b"a%d=''" % i for i in range(cap + 1))
        data = _rewrite_zip(
            _blank_with(b""),
            {
                "word/document.xml": lambda xml: xml.replace(
                    b"<w:body>", b"<w:body " + attrs + b">", 1
                )
            },
        )
        with pytest.raises(guard.DecompressionBombError, match="attributes"):
            guard.validate_zip_container(data, ".docx")

    def test_attribute_flood_is_refused_before_any_value_is_read(self):
        # Before, 30,000 attributes took ~6 s to price (and were
        # accepted); the count is checked before any value is read.
        guard = _guard()
        data = _blank_with(_with_attributes(30_000))
        reads = []
        text_bytes = guard._text_bytes

        def counting(text):
            reads.append(text)
            return text_bytes(text)

        started = time.perf_counter()
        with patch.object(guard, "_text_bytes", counting):
            with pytest.raises(
                guard.DecompressionBombError, match="attributes"
            ):
                guard.validate_zip_container(data, ".docx")
        assert len(reads) < 100
        # Generous: ~0.03 s measured.
        assert time.perf_counter() - started < 3

    def test_namespace_cap_accepts_the_limit_and_refuses_one_more(self):
        guard = _guard()
        room = guard.MAX_DOCX_NAMESPACES_IN_SCOPE - _TEMPLATE_ROOT_DECLARATIONS
        for declared, refused in ((room, False), (room + 1, True)):
            data = _blank_with(b"<w:p " + _declarations(declared) + b"/>")
            if refused:
                with pytest.raises(
                    guard.DecompressionBombError, match="namespace"
                ):
                    guard.validate_zip_container(data, ".docx")
            else:
                guard.validate_zip_container(data, ".docx")

    def test_namespace_cap_counts_declarations_in_scope(self):
        guard = _guard()
        half = (
            guard.MAX_DOCX_NAMESPACES_IN_SCOPE - _TEMPLATE_ROOT_DECLARATIONS
        ) // 2 + 1
        # Siblings: never in scope together.
        siblings = (
            (b"<w:p " + _declarations(half) + b"/>")
            + b"<w:p "
            + _declarations(half, half)
            + b"/>"
        )
        guard.validate_zip_container(_blank_with(siblings), ".docx")
        # Nested: the run sees its own and the paragraph's.
        nested = (
            b"<w:p "
            + _declarations(half)
            + b"><w:r "
            + _declarations(half, half)
            + b"/></w:p>"
        )
        with pytest.raises(guard.DecompressionBombError, match="namespace"):
            guard.validate_zip_container(_blank_with(nested), ".docx")

    # lxml hands back every element and attribute name as
    # "{URI}local", building the string again on each read, so one long
    # URI declared once and used through a short prefix multiplies
    # through every name read under it: an 8 MB URI used by 3,000
    # elements, a 44 KB file, exhausted 3 GB in the guard, and a 1 MB
    # URI on 1,024 attribute names grew it by ~0.5 GB. URIs (and
    # prefixes) are capped where they are declared, before any name is
    # read.

    @pytest.mark.parametrize("where", ["w:p", "w:r", "default", "root"])
    def test_namespace_uri_cap_accepts_the_limit_and_refuses_one_more(
        self, where
    ):
        guard = _guard()
        cap = guard.MAX_DOCX_NAMESPACE_URI_CHARS

        def document(length: int) -> bytes:
            uri = b"urn:" + b"u" * (length - 4)
            if where == "root":
                return _rewrite_zip(
                    _blank_with(b"<w:p/>"),
                    {
                        "word/document.xml": lambda xml: xml.replace(
                            b"<w:document ",
                            b"<w:document xmlns:z='%s' " % uri,
                            1,
                        )
                    },
                )
            declaration = (
                b"xmlns='%s'" % uri
                if where == "default"
                else b"xmlns:z='%s'" % uri
            )
            if where == "w:r":
                return _blank_with(
                    b"<w:p><w:r %s><w:t>a</w:t></w:r></w:p>" % declaration
                )
            return _blank_with(b"<w:p %s><x/></w:p>" % declaration)

        guard.validate_zip_container(document(cap), ".docx")
        with pytest.raises(guard.DecompressionBombError, match="namespace URI"):
            guard.validate_zip_container(document(cap + 1), ".docx")

    def test_namespace_prefix_cap_accepts_the_limit_and_refuses_one_more(self):
        guard = _guard()
        cap = guard.MAX_DOCX_NAMESPACE_PREFIX_CHARS
        for length, refused in ((cap, False), (cap + 1, True)):
            prefix = b"p" * length
            data = _blank_with(
                b"<w:p xmlns:%s='urn:x'><%s:x/></w:p>" % (prefix, prefix)
            )
            if refused:
                with pytest.raises(
                    guard.DecompressionBombError, match="prefix"
                ):
                    guard.validate_zip_container(data, ".docx")
            else:
                guard.validate_zip_container(data, ".docx")

    def test_long_uri_is_refused_before_any_name_is_read(self):
        # A 2 MB URI used by 3,000 elements: refused when declared, so
        # the guard reads no element under it (counted through the
        # attribute values and texts it sizes).
        guard = _guard()
        uri = b"urn:" + b"u" * 2_000_000
        data = _blank_with(
            b"<w:p xmlns:u='%s'>" % uri + b"<u:x a='1'/>" * 3_000 + b"</w:p>"
        )
        reads = []
        text_bytes = guard._text_bytes

        def counting(text):
            reads.append(text)
            return text_bytes(text)

        started = time.perf_counter()
        with patch.object(guard, "_text_bytes", counting):
            with pytest.raises(
                guard.DecompressionBombError, match="namespace URI"
            ):
                guard.validate_zip_container(data, ".docx")
        assert len(reads) < 100
        # Generous: ~0.05 s measured (~11 s, and accepted, before).
        assert time.perf_counter() - started < 3

    def test_guard_memory_stays_small_under_many_names_at_the_uri_cap(self):
        # 20,480 distinct attribute names under a URI at the cap,
        # declared on w:document: the guard used to keep every name
        # (~24 MB of Python objects measured); it keeps the URIs.
        import tracemalloc

        guard = _guard()
        uri = b"urn:" + b"u" * (guard.MAX_DOCX_NAMESPACE_URI_CHARS - 4)
        elements = b"".join(
            b"<w:bookmarkStart "
            + b" ".join(b"z:a%d_%d=''" % (e, i) for i in range(256))
            + b"/>"
            for e in range(80)
        )
        data = _rewrite_zip(
            _blank_with(b"<w:p>" + elements + b"</w:p>"),
            {
                "word/document.xml": lambda xml: xml.replace(
                    b"<w:document ", b"<w:document xmlns:z='%s' " % uri, 1
                )
            },
        )
        tracemalloc.start()
        try:
            _docx_stats(data)
            _current, peak = tracemalloc.get_traced_memory()
        finally:
            tracemalloc.stop()
        # ~0.8 MB measured.
        assert peak < 6 * 1024 * 1024

    def test_guard_time_at_the_uri_cap_stays_linear(self):
        # 100,000 short elements named through a URI at the cap cost
        # the guard about a third more than through Word's own (~0.5 s
        # against ~0.4 s measured): each name is read a bounded number
        # of times.
        guard = _guard()

        def elapsed(length: int) -> float:
            uri = b"urn:" + b"u" * (length - 4)
            data = _blank_with(
                b"<w:p xmlns:u='%s'>" % uri + b"<u:x/>" * 100_000 + b"</w:p>"
            )
            started = time.perf_counter()
            guard.validate_zip_container(data, ".docx")
            return time.perf_counter() - started

        word = elapsed(71)
        assert elapsed(guard.MAX_DOCX_NAMESPACE_URI_CHARS) < 3 * word + 1

    # libxml2 re-homes a removed subtree's namespace declarations through
    # a cache it searches linearly, so lxml's remove(), which
    # python-docx calls on a paragraph copy per rendered page break,
    # costs the square of the declarations under the removed node
    # (40,000 one-declaration elements: 0.4-0.8 s per removal; six page
    # breaks before them, a 142 KB file, took 5 s to partition and were
    # accepted). A body-level paragraph may hold at most
    # MAX_DOCX_PARAGRAPH_NAMESPACES declarations.

    def test_paragraph_namespace_cap_accepts_the_limit_and_refuses_one_more(
        self,
    ):
        guard = _guard()
        cap = guard.MAX_DOCX_PARAGRAPH_NAMESPACES
        # 128 per run (in scope with the template's 17), one on the w:p.
        per_run = 128

        def paragraph(declared: int) -> bytes:
            runs = []
            rest = declared - 1
            start = 1
            while rest:
                count = min(rest, per_run)
                runs.append(b"<w:r %s/>" % _declarations(count, start))
                start += count
                rest -= count
            return b"<w:p %s>" % _declarations(1) + b"".join(runs) + b"</w:p>"

        guard.validate_zip_container(_blank_with(paragraph(cap)), ".docx")
        with pytest.raises(guard.DecompressionBombError, match="namespace"):
            guard.validate_zip_container(
                _blank_with(paragraph(cap + 1)), ".docx"
            )
        # The cap is per paragraph.
        guard.validate_zip_container(
            _blank_with(paragraph(cap) + paragraph(cap)), ".docx"
        )

    def test_hostile_declaration_flood_with_page_breaks_is_refused(self):
        # 200 groups of 200 one-declaration elements in one run, after
        # six page breaks: 5 s to partition, accepted before.
        guard = _guard()
        groups = b"".join(
            b"<w:x>"
            + b"".join(
                b"<w:y xmlns:n='urn:%d'/>" % (g * 200 + i) for i in range(200)
            )
            + b"</w:x>"
            for g in range(200)
        )
        body = (
            b"<w:p><w:r><w:t>a</w:t>"
            + b"<w:lastRenderedPageBreak/><w:t>a</w:t>" * 6
            + b"</w:r><w:r>"
            + groups
            + b"</w:r></w:p>"
        )
        with pytest.raises(guard.DecompressionBombError, match="namespace"):
            guard.validate_zip_container(_blank_with(body), ".docx")

    def test_paragraphs_of_many_pictures_are_accepted(self):
        # Pictures declare namespaces inside the paragraph: python-docx
        # two per inline picture (on wp:inline), Word up to five
        # (a:graphicFrameLocks, a:graphic, pic:pic, a14:useLocalDpi and
        # a16:creationId each declare their own).
        docx = pytest.importorskip("docx")
        document = docx.Document()
        run = document.add_paragraph().add_run()
        image = _small_png()
        for _ in range(100):
            run.add_picture(io.BytesIO(image))
        buf = io.BytesIO()
        document.save(buf)
        data = buf.getvalue()
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("word/document.xml")
        body = xml[xml.index(b"<w:body>") :]
        assert body.count(b"xmlns:") == 200
        guard = _guard()
        guard.validate_zip_container(data, ".docx")
        # 200 Word-style pictures in one paragraph.
        word_picture = (
            b"<w:r><w:drawing><w:x>"
            b"<a:graphicFrameLocks xmlns:a='%(a)s'/>"
            b"<a:graphic xmlns:a='%(a)s'><pic:pic xmlns:pic='%(pic)s'>"
            b"<a14:useLocalDpi xmlns:a14='%(a14)s'/>"
            b"<a16:creationId xmlns:a16='%(a16)s'/>"
            b"</pic:pic></a:graphic></w:x></w:drawing></w:r>"
            % {
                b"a": b"http://schemas.openxmlformats.org/drawingml/2006/main",
                b"pic": b"http://schemas.openxmlformats.org/drawingml/2006/"
                b"picture",
                b"a14": b"http://schemas.microsoft.com/office/drawing/2010/"
                b"main",
                b"a16": b"http://schemas.microsoft.com/office/drawing/2014/"
                b"main",
            }
        )
        guard.validate_zip_container(
            _blank_with(b"<w:p>" + word_picture * 200 + b"</w:p>"), ".docx"
        )

    def test_documents_python_docx_writes_stay_far_under_both_caps(self):
        docx = pytest.importorskip("docx")
        document = docx.Document()
        document.add_heading("Title", 0)
        paragraph = document.add_paragraph("Some ")
        paragraph.add_run("bold").bold = True
        table = document.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 2))
        document.add_page_break()
        buf = io.BytesIO()
        document.save(buf)
        data = buf.getvalue()
        guard = _guard()
        guard.validate_zip_container(data, ".docx")
        from lxml import etree

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
        assert max(len(e.attrib) for e in root.iter()) * 8 < (
            guard.MAX_DOCX_ELEMENT_ATTRIBUTES
        )
        assert max(len(e.nsmap) for e in root.iter()) * 8 < (
            guard.MAX_DOCX_NAMESPACES_IN_SCOPE
        )
        uris = set().union(*(e.nsmap.values() for e in root.iter()))
        assert max(len(uri) for uri in uris) * 8 < (
            guard.MAX_DOCX_NAMESPACE_URI_CHARS
        )
        prefixes = set().union(*(e.nsmap.keys() for e in root.iter()))
        assert max(len(p or "") for p in prefixes) * 6 < (
            guard.MAX_DOCX_NAMESPACE_PREFIX_CHARS
        )


def _text_paragraphs_body(paragraphs: int) -> bytes:
    return b"<w:p><w:r><w:t>text</w:t></w:r></w:p>" * paragraphs


def _with_doctype(doctype: bytes, root_tag: bytes):
    """A rewrite that inserts *doctype* before the part's root element."""

    def rewrite(xml: bytes) -> bytes:
        assert root_tag in xml
        return xml.replace(root_tag, doctype + root_tag, 1)

    return rewrite


_INTERNAL_DTD = b'<!DOCTYPE w:document [<!ENTITY e "">]>'
_EXTERNAL_DTD = b'<!DOCTYPE w:document SYSTEM "x.dtd">'


class TestPackagePartsWithADtdAreRefused:
    # python-docx parses with resolve_entities=False, which keeps each
    # entity reference as an _Entity node that its per-section XPath
    # walks, while iterparse emits no event for it: the section-work
    # count missed them. Package parts never carry a DTD (OPC forbids
    # it), and a DTD is the only way to declare an entity.

    def _docx_with_entity_refs(self, doctype: bytes, refs: int = 50):
        data = _docx_with_sections(3, header=None)
        return _rewrite_zip(
            data,
            {
                "word/document.xml": lambda xml: _with_doctype(
                    doctype, b"<w:document"
                )(
                    # standalone="yes" would make an undeclared
                    # entity (the external-subset case) a parse error.
                    re.sub(
                        rb" standalone=['\"]yes['\"]", b"", xml, count=1
                    ).replace(
                        b"<w:body>",
                        b"<w:body><w:p>" + b"&e;" * refs + b"</w:p>",
                        1,
                    )
                )
            },
        )

    @pytest.mark.parametrize(
        "doctype", [_INTERNAL_DTD, _EXTERNAL_DTD], ids=["internal", "system"]
    )
    def test_premise_entity_references_are_nodes_the_count_misses(
        self, doctype
    ):
        docx = pytest.importorskip("docx")
        from lxml import etree

        data = self._docx_with_entity_refs(doctype)
        body = docx.Document(io.BytesIO(data)).element.body
        assert sum(isinstance(n, etree._Entity) for n in body[0]) == 50

    @pytest.mark.parametrize(
        "doctype", [_INTERNAL_DTD, _EXTERNAL_DTD], ids=["internal", "system"]
    )
    def test_document_xml_with_a_dtd_is_refused(self, doctype):
        pytest.importorskip("docx")
        guard = _guard()
        data = self._docx_with_entity_refs(doctype)

        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(data, ".docx")
        guard.validate_zip_container(  # control: same document, no DTD
            _docx_with_sections(3, header=None), ".docx"
        )

    def test_relationships_part_with_a_dtd_is_refused(self):
        pytest.importorskip("docx")
        guard = _guard()
        data = _rewrite_zip(
            _docx_with_sections(1),
            {
                "word/_rels/document.xml.rels": _with_doctype(
                    b"<!DOCTYPE Relationships>", b"<Relationships"
                )
            },
        )

        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(data, ".docx")

    def test_presentation_part_with_a_dtd_is_refused(self):
        guard = _guard()
        data = _rewrite_zip(
            _deck(1),
            {
                "ppt/presentation.xml": _with_doctype(
                    b"<!DOCTYPE p:presentation>", b"<p:presentation"
                )
            },
        )

        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(data, ".pptx")
        guard.validate_zip_container(_deck(1), ".pptx")  # control

    def test_workbook_part_with_a_dtd_is_refused(self):
        openpyxl = pytest.importorskip("openpyxl")
        guard = _guard()
        buf = io.BytesIO()
        openpyxl.Workbook().save(buf)
        data = _rewrite_zip(
            buf.getvalue(),
            {
                "xl/workbook.xml": _with_doctype(
                    b"<!DOCTYPE workbook>", b"<workbook"
                )
            },
        )

        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(data, ".xlsx")
        guard.validate_zip_container(buf.getvalue(), ".xlsx")  # control

    def test_header_part_with_a_dtd_is_refused(self):
        # Header/footer parts are streamed for their unions now, so a
        # DTD there is refused as in document.xml.
        guard = _guard()
        data = _rewrite_zip(
            _docx_with_sections(1),
            {"word/header1.xml": _with_doctype(b"<!DOCTYPE w:hdr>", b"<w:hdr")},
        )

        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(data, ".docx")
        guard.validate_zip_container(_docx_with_sections(1), ".docx")


def _real_package(ext: str) -> bytes:
    """A package the format's own library writes."""
    buf = io.BytesIO()
    if ext == ".docx":
        pytest.importorskip("docx").Document().save(buf)
    elif ext == ".pptx":
        return _deck(1)
    else:
        pytest.importorskip("openpyxl").Workbook().save(buf)
    return buf.getvalue()


def _detected_type(tmp_path, data: bytes, ext: str) -> str:
    """What unstructured's ``detect_filetype`` calls *data* saved as
    *ext* (the call langchain's Word and PowerPoint loaders make)."""
    filetype = pytest.importorskip("unstructured.file_utils.filetype")
    path = tmp_path / f"probe{ext}"
    path.write_bytes(data)
    return filetype.detect_filetype(str(path)).name


def _renamed_members(data: bytes, renames: dict[str, str]) -> bytes:
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as zin,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout,
    ):
        for info in zin.infolist():
            payload = zin.read(info)
            for old, new in renames.items():
                payload = payload.replace(old.encode(), new.encode())
            zout.writestr(renames.get(info.filename, info.filename), payload)
    return out.getvalue()


class TestPackageTypeMatchesTheExtension:
    # langchain's Word and PowerPoint loaders ask unstructured's
    # detect_filetype (whenever python-magic imports) and send what it
    # calls DOC or PPT to partition_doc/partition_ppt: soffice with no
    # timeout, then partition_docx on its output without this guard. A
    # zip whose member names match none of its OOXML patterns is typed
    # from a root ``mimetype`` member, which the guard accepted before.

    @pytest.mark.parametrize(
        ("ext", "mimetype", "detected"),
        [
            (".docx", "application/msword", "DOC"),
            (".pptx", "application/vnd.ms-powerpoint", "PPT"),
        ],
    )
    def test_mimetype_naming_a_legacy_format_is_refused(
        self, tmp_path, ext, mimetype, detected
    ):
        guard = _guard()
        data = _zip_bytes({"mimetype": mimetype.encode()})
        assert len(data) < 200
        assert _detected_type(tmp_path, data, ext) == detected

        with pytest.raises(guard.DecompressionBombError, match="mimetype"):
            guard.validate_zip_container(data, ext)

    @pytest.mark.parametrize("ext", [".docx", ".pptx", ".xlsx"])
    def test_any_mimetype_member_of_an_office_package_is_refused(self, ext):
        # Even beside the main part, which makes the sniffer type the
        # package correctly: OPC packages have no such member.
        guard = _guard()
        data = _rewrite_zip(
            _real_package(ext), add={"mimetype": b"application/msword"}
        )

        with pytest.raises(guard.DecompressionBombError, match="mimetype"):
            guard.validate_zip_container(data, ext)

    @pytest.mark.parametrize("ext", [".docx", ".pptx", ".xlsx"])
    def test_real_office_packages_are_accepted(self, tmp_path, ext):
        data = _real_package(ext)
        assert _detected_type(tmp_path, data, ext) == ext[1:].upper()

        _guard().validate_zip_container(data, ext)

    def test_docx_without_its_main_part_is_refused(self):
        # word/document.xml is there (the sniffer says DOCX), but no
        # package relationship names it: python-docx cannot open this.
        guard = _guard()
        data = _zip_bytes(
            {
                "word/document.xml": (
                    f"<w:document xmlns:w='{_W_MAIN}'><w:body/></w:document>"
                ).encode()
            }
        )

        with pytest.raises(guard.DecompressionBombError, match="main document"):
            guard.validate_zip_container(data, ".docx")

    def test_pptx_without_its_main_part_is_refused(self):
        guard = _guard()
        data = _rewrite_zip(
            _deck(1),
            {
                "_rels/.rels": lambda xml: xml.replace(
                    b"/officeDocument", b"/officeDocumentX"
                )
            },
        )

        with pytest.raises(guard.DecompressionBombError, match="main document"):
            guard.validate_zip_container(data, ".pptx")

    @pytest.mark.parametrize(
        ("written", "ext"),
        [
            (".docx", ".pptx"),
            (".pptx", ".docx"),
            (".xlsx", ".docx"),
            (".docx", ".xlsx"),
            (".xlsx", ".pptx"),
        ],
    )
    def test_package_of_another_kind_is_refused(self, written, ext):
        guard = _guard()

        with pytest.raises(guard.DecompressionBombError, match="part names"):
            guard.validate_zip_container(_real_package(written), ext)

    def test_relocated_main_part_is_refused(self):
        # python-docx would read word/main.xml through the relationship,
        # but no member name tells the sniffer this is a .docx, so a
        # mimetype member would decide (or it calls it a plain zip).
        guard = _guard()
        data = _renamed_members(
            _real_package(".docx"), {"word/document.xml": "word/main.xml"}
        )
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            assert b"word/main.xml" in archive.read("_rels/.rels")

        with pytest.raises(guard.DecompressionBombError, match="part names"):
            guard.validate_zip_container(data, ".docx")

    def test_xlsx_must_hold_xl_workbook_xml(self):
        # The sniffer's pattern matches xl/workbook2.xml, but pandas picks
        # openpyxl only for xl/workbook.xml.
        guard = _guard()
        data = _xlsx_with_sheets(
            ["rId1"],
            [("rId1", "worksheets/sheet1.xml")],
            workbook="xl/workbook2.xml",
        )

        with pytest.raises(
            guard.DecompressionBombError, match="xl/workbook.xml"
        ):
            guard.validate_zip_container(data, ".xlsx")

    @pytest.mark.parametrize(
        ("ext", "mimetype", "other"),
        [
            (
                ".odt",
                b"application/vnd.oasis.opendocument.text",
                b"application/epub+zip",
            ),
            (
                ".epub",
                b"application/epub+zip",
                b"application/vnd.oasis.opendocument.text",
            ),
        ],
    )
    def test_odt_and_epub_mimetype_must_name_the_extension(
        self, ext, mimetype, other
    ):
        guard = _guard()
        guard.validate_zip_container(
            _zip_bytes({"mimetype": mimetype, "content.xml": b"<a/>"}), ext
        )
        guard.validate_zip_container(_zip_bytes({"content.xml": b"<a/>"}), ext)

        for wrong in (other, b"application/msword", b"\xff\xfe", b"x" * 300):
            with pytest.raises(guard.DecompressionBombError, match="mimetype"):
                guard.validate_zip_container(
                    _zip_bytes({"mimetype": wrong, "content.xml": b"<a/>"}),
                    ext,
                )

    def test_upload_path_never_reaches_the_legacy_partitioner(self):
        doc_partitioner = pytest.importorskip("unstructured.partition.doc")
        from local_deep_research.document_loaders.bytes_loader import (
            load_from_bytes,
        )

        guard = _guard()
        data = _zip_bytes({"mimetype": b"application/msword"})
        with patch.object(
            doc_partitioner,
            "partition_doc",
            side_effect=AssertionError("routed to soffice"),
        ):
            with pytest.raises(guard.DecompressionBombError):
                load_from_bytes(data, ".docx", "sniffed.docx")


def _one_cell_table(span: int, content: bytes) -> bytes:
    return (
        b"<w:tbl><w:tr><w:tc><w:tcPr><w:gridSpan w:val='%d'/></w:tcPr>" % span
        + content
        + b"</w:tc></w:tr></w:tbl>"
    )


def _union_units(steps: int, reads: int | None = None) -> int:
    guard = _guard()
    reads = guard.DOCX_UNION_READS if reads is None else reads
    return -(-steps * reads // guard.DOCX_UNION_STEPS_PER_RUN_WORK_UNIT)


def _without_sections(data: bytes) -> bytes:
    return _rewrite_zip(
        data,
        {
            "word/document.xml": lambda xml: re.sub(
                rb"<w:sectPr.*?</w:sectPr>", b"", xml, flags=re.S
            )
        },
    )


class TestDocxNodeSetUnions:
    # python-docx selects a cell's, a header's (and, without sections,
    # the body's) paragraphs and tables with ``./w:p | ./w:tbl``, a
    # paragraph's runs and hyperlinks with ``w:r | w:hyperlink`` and a
    # run's text with a six-branch union; libxml2 merges the branches by
    # comparing every node with every node, then sorts the result with
    # comparisons that walk sibling links, so one evaluation costs up to
    # (selected nodes) x (child nodes) steps once a later branch's node
    # precedes an earlier one's. None of this was counted.

    @pytest.mark.parametrize("tables_first", [False, True])
    def test_cell_of_tables_and_paragraphs_is_refused(self, tables_first):
        # One cell spanning 1,000 columns holding 16,000 empty paragraphs
        # and 16,000 empty tables, a 37 KB file: accepted at 0.09 of the
        # ceiling before, ~1.5 s per yield (~25 minutes).
        guard = _guard()
        blocks = [b"<w:p/>" * 16_000, b"<w:tbl/>" * 16_000]
        if tables_first:
            blocks.reverse()
        data = _blank_with(_one_cell_table(1000, b"".join(blocks) + b"<w:p/>"))
        assert len(data) < 40_000

        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(data, ".docx")

    def test_cell_paragraphs_count_once_per_yield(self):
        guard = _guard()
        item = guard.DOCX_CELL_BLOCK_ITEM_WEIGHT

        def run_work(paragraphs: int) -> int:
            return _docx_stats(
                _blank_with(_one_cell_table(10, b"<w:p/>" * paragraphs))
            ).run_work

        assert run_work(101) - run_work(1) == 10 * 100 * item
        # 32,000 empty paragraphs in a cell spanning 1,000 columns: about
        # 20 us per paragraph and yield, ~10 minutes; 125 units per
        # yield (one per 256 nodes) before.
        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(
                _blank_with(_one_cell_table(1000, b"<w:p/>" * 32_000)),
                ".docx",
            )

    def test_cell_union_counts_selected_nodes_times_children(self):
        guard = _guard()
        weight = guard.DOCX_RUN_WORK_ITEM_WEIGHT
        item = guard.DOCX_CELL_BLOCK_ITEM_WEIGHT
        content = b"<w:tbl/>" * 3 + b"<w:p/>" * 40 + _MARKER * 200
        stats = _docx_stats(_blank_with(_one_cell_table(7, content)))
        # The cell's children: w:tcPr, 43 blocks and 200 markers.
        children = 1 + 43 + 200
        assert stats.run_work == weight + 7 * (
            weight + 43 * item + _union_units(43 * children)
        )

    def test_paragraph_union_is_counted(self):
        # A hyperlink before 8,000 runs and 40,000 empty elements in one
        # paragraph, a 38 KB file: accepted at 0.03 of the ceiling
        # before and took 15 s to partition (4,000 runs beside 20,000:
        # 2.8 s; the cost quadruples per doubling).
        guard = _guard()
        run = b"<w:r><w:t xml:space='preserve'>a </w:t></w:r>"
        link = b"<w:hyperlink><w:r><w:t>h</w:t></w:r></w:hyperlink>"
        data = _blank_with(
            b"<w:p>" + link + run * 8_000 + b"<w:x/>" * 40_000 + b"</w:p>"
        )

        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(data, ".docx")
        plain = _docx_stats(_blank_with(b"<w:p>" + run * 3 + b"</w:p>"))
        linked = _docx_stats(_blank_with(b"<w:p>" + link + run * 3 + b"</w:p>"))
        weight = guard.DOCX_RUN_WORK_ITEM_WEIGHT
        assert plain.run_work == 3 * (weight + 1)
        assert linked.run_work == 4 * (weight + 1) + _union_units(4 * 4)

    def test_run_union_is_counted(self):
        # 8,000 w:t, then a w:br, then 40,000 empty elements in one run,
        # a 38 KB file: accepted at 0.03 of the ceiling before, 14 s to
        # partition (4,000 and 20,000: 2.5 s).
        guard = _guard()
        text = b"<w:t xml:space='preserve'>a </w:t>"
        data = _blank_with(
            b"<w:p><w:r>"
            + text * 8_000
            + b"<w:br/>"
            + b"<w:x/>" * 40_000
            + b"</w:r></w:p>"
        )

        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(data, ".docx")
        weight = guard.DOCX_RUN_WORK_ITEM_WEIGHT
        one_kind = _docx_stats(
            _blank_with(b"<w:p><w:r>" + text * 5 + b"</w:r></w:p>")
        )
        assert one_kind.run_work == weight + 5
        # Three kinds: two merges, of 7 selected nodes among 7 children.
        mixed = _docx_stats(
            _blank_with(
                b"<w:p><w:r>" + text * 5 + b"<w:tab/><w:br/></w:r></w:p>"
            )
        )
        assert mixed.run_work == weight + 7 + _union_units(2 * 7 * 7)

    def test_body_union_is_counted_without_sections(self):
        paragraph = b"<w:p><w:r><w:t>a</w:t></w:r></w:p>"

        def stats(prefix: bytes):
            return _docx_stats(
                _without_sections(_blank_with(prefix + paragraph * 50))
            )

        plain, tabled = stats(b""), stats(b"<w:tbl/>")
        assert plain.sections == tabled.sections == 0
        # The table's own work is a table item; the union is read once.
        assert tabled.run_work - plain.run_work == _union_units(
            51 * tabled.body_nodes, reads=1
        )
        # With a section the body union is not read (the section lookup
        # is priced instead).
        with_sections = _docx_stats(_blank_with(b"<w:tbl/>" + paragraph * 50))
        without_table = _docx_stats(_blank_with(paragraph * 50))
        assert with_sections.run_work == without_table.run_work

    def test_header_unions_count_per_reference(self):
        # A header holding a table before 8,000 paragraphs and 40,000
        # empty elements, referenced by ten sections, a 37 KB file:
        # accepted at 0.23 of the header ceiling before and took 35 s
        # (two references: 8 s).
        guard = _guard()
        header = (
            b"h</w:t></w:r></w:p><w:tbl/>"
            + b"<w:p/>" * 8_000
            + b"<w:x/>" * 40_000
            + b"<w:p><w:r><w:t>x"
        )

        with pytest.raises(guard.DecompressionBombError, match="header/footer"):
            guard.validate_zip_container(
                _docx_with_sections(10, header=header), ".docx"
            )
        # Without a table first, the root's union has one branch.
        plain = _docx_with_sections(3, header=b"plain")
        with zipfile.ZipFile(io.BytesIO(plain)) as archive:
            size = archive.getinfo("word/header1.xml").file_size
        assert _docx_stats(plain).referenced_bytes == 3 * size
        tabled = _docx_with_sections(
            3,
            header=b"h</w:t></w:r></w:p><w:tbl/><w:p/>"
            + b"<w:x/>" * 1_000
            + b"<w:p><w:r><w:t>x",
        )
        with zipfile.ZipFile(io.BytesIO(tabled)) as archive:
            size = archive.getinfo("word/header1.xml").file_size
        # The root's union: 4 selected nodes among 1,004 children, read
        # once per reference.
        assert _docx_stats(tabled).referenced_bytes == 3 * size + (
            3 * 4 * 1_004 // guard.DOCX_UNION_STEPS_PER_HEADER_BYTE
        )

    def test_word_like_document_with_links_and_nested_tables_is_accepted(self):
        # 2,000 paragraphs of eight runs around a hyperlink, and a 20 x 5
        # table whose cells each hold a nested 2 x 2 table and a
        # paragraph: the unions add little to the text extraction.
        guard = _guard()
        run = b"<w:r><w:rPr><w:b/></w:rPr><w:t xml:space='preserve'>word </w:t></w:r>"
        link = b"<w:hyperlink><w:r><w:t>link</w:t></w:r></w:hyperlink>"
        paragraph = b"<w:p>" + run * 4 + link + run * 4 + b"</w:p>"
        nested = (
            b"<w:tbl>"
            + (
                b"<w:tr>"
                + (b"<w:tc><w:p>" + run + b"</w:p></w:tc>") * 2
                + b"</w:tr>"
            )
            * 2
            + b"</w:tbl>"
        )
        cell = b"<w:tc>" + nested + b"<w:p>" + run + b"</w:p></w:tc>"
        table = (
            b"<w:tbl>" + (b"<w:tr>" + cell * 5 + b"</w:tr>") * 20 + b"</w:tbl>"
        )
        data = _blank_with(paragraph * 2_000 + table)

        guard.validate_zip_container(data, ".docx")
        stats = _docx_stats(data)
        assert stats.run_work < guard.MAX_DOCX_RUN_WORK // 2


_BREAK_IN_RUN = b"<w:p><w:r><w:lastRenderedPageBreak/></w:r></w:p>"
_BREAK_IN_LINK = (
    b"<w:p><w:hyperlink><w:r><w:lastRenderedPageBreak/></w:r></w:hyperlink>"
    b"</w:p>"
)
# The same paragraph with its break outside every branch of the union.
_BREAK_IN_SMART_TAG = (
    b"<w:p><w:smartTag><w:r><w:lastRenderedPageBreak/></w:r></w:smartTag></w:p>"
)
_PAGE_BREAK_BRANCHES = (
    "./w:body/w:p/w:r/w:lastRenderedPageBreak",
    "./w:body/w:p/w:hyperlink/w:r/w:lastRenderedPageBreak",
    "./w:body/w:tbl/w:tr/w:tc/w:p/w:r/w:lastRenderedPageBreak",
    "./w:body/w:tbl/w:tr/w:tc/w:p/w:hyperlink/w:r/w:lastRenderedPageBreak",
)


def _cell_of(content: bytes) -> bytes:
    return b"<w:tbl><w:tr><w:tc>" + content + b"</w:tc></w:tr></w:tbl>"


def _run_work_with_pair_price(data: bytes, pairs_per_unit: int) -> int:
    with patch.object(
        _guard(), "DOCX_PAGE_BREAK_UNION_PAIRS_PER_UNIT", pairs_per_unit
    ):
        return _docx_stats(data).run_work


def _run_work_with_sort_price(data: bytes, steps_per_unit: int) -> int:
    with patch.object(
        _guard(), "DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT", steps_per_unit
    ):
        return _docx_stats(data).run_work


def _expected_sort_steps(data: bytes) -> int:
    """The sum over the elements of ``document.xml`` of the union's
    breaks at or below each, B, weighted for the sort's binary insertion
    (B + 4 * min(B, 64)), times its child nodes as python-docx's tree
    has them (text included), computed from the parsed tree."""
    from lxml import etree

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        root = etree.fromstring(zf.read("word/document.xml"))
    ns = {"w": _W_MAIN}
    breaks = set()
    for branch in _PAGE_BREAK_BRANCHES:
        breaks.update(root.xpath(branch, namespaces=ns))
    below: dict = {}
    steps = 0
    for element in reversed(list(root.iter())):
        count = (element in breaks) + sum(below.get(c, 0) for c in element)
        below[element] = count
        nodes = len(element) + bool(element.text)
        nodes += sum(bool(child.tail) for child in element)
        steps += (count + 4 * min(count, 64)) * nodes
    return steps


class TestDocxPageBreakUnion:
    # unstructured asks once per document whether it has rendered page
    # breaks, with a four-branch union over the whole body; libxml2
    # compares each branch's nodes with every node merged before them,
    # so 20,000 breaks in one table cell after 20,000 in body paragraphs
    # took 0.8 s, quadrupling per doubling, while a break in a cell cost
    # about one unit of MAX_DOCX_RUN_WORK: 96,000 in each of the two
    # cell branches, ~0.9 of that ceiling, were accepted (~18 s).

    def test_premise_unstructured_unions_four_branches_over_the_body(self):
        import inspect

        docx = pytest.importorskip("docx")
        udocx = pytest.importorskip("unstructured.partition.docx")
        source = inspect.getsource(udocx)
        for branch in _PAGE_BREAK_BRANCHES:
            assert branch in source
        data = _blank_with(
            _BREAK_IN_RUN * 2
            + _BREAK_IN_LINK * 3
            + _cell_of(_BREAK_IN_RUN * 5 + _BREAK_IN_LINK * 7)
            + _BREAK_IN_SMART_TAG
        )
        element = docx.Document(io.BytesIO(data)).element
        assert [len(element.xpath(b)) for b in _PAGE_BREAK_BRANCHES] == [
            2,
            3,
            5,
            7,
        ]

    def test_breaks_over_the_ceiling_are_refused(self):
        guard = _guard()
        limit = guard.MAX_DOCX_PAGE_BREAKS
        refusal = f"more than {limit} rendered page breaks"
        # One branch: no merge and no sort to price (two branches of
        # half each in one cell are refused by their sort's price).
        at_ceiling = _blank_with(_cell_of(_BREAK_IN_RUN * limit))
        guard.validate_zip_container(at_ceiling, ".docx")
        over = _blank_with(_cell_of(_BREAK_IN_RUN * limit + _BREAK_IN_LINK))
        with pytest.raises(guard.DecompressionBombError, match=refusal):
            guard.validate_zip_container(over, ".docx")
        # Breaks outside the union's branches count toward it too.
        elsewhere = _blank_with(
            _cell_of(_BREAK_IN_SMART_TAG * (guard.MAX_DOCX_PAGE_BREAKS + 1))
        )
        with pytest.raises(guard.DecompressionBombError, match=refusal):
            guard.validate_zip_container(elsewhere, ".docx")

    def test_hostile_cross_branch_document_is_refused_by_default(self):
        # Accepted at ~0.85 of MAX_DOCX_RUN_WORK before; the union alone
        # would take ~16 s.
        guard = _guard()
        data = _blank_with(
            _cell_of(_BREAK_IN_RUN * 90_000 + _BREAK_IN_LINK * 90_000)
        )
        assert len(data) < 200_000
        with pytest.raises(
            guard.DecompressionBombError, match="rendered page breaks"
        ):
            guard.validate_zip_container(data, ".docx")

    def test_cross_branch_pairs_are_priced(self):
        # Each branch is merged against the branches before it, so the
        # pairs are the sum over pairs of branches of their products.
        a, b, c, d = 3, 5, 7, 11
        data = _blank_with(
            _BREAK_IN_RUN * a
            + _BREAK_IN_LINK * b
            + _cell_of(_BREAK_IN_RUN * c + _BREAK_IN_LINK * d)
        )
        pairs = a * b + a * c + a * d + b * c + b * d + c * d
        # Priced at one unit per pair, or rounded up to one unit.
        assert (
            _run_work_with_pair_price(data, 1)
            - _run_work_with_pair_price(data, 10**30)
            == pairs - 1
        )
        # One branch alone, or breaks outside the branches: no merge.
        for content in (
            _BREAK_IN_RUN * 50,
            _BREAK_IN_RUN * 7 + _BREAK_IN_SMART_TAG * 11,
        ):
            single = _blank_with(_cell_of(content))
            assert _run_work_with_pair_price(
                single, 1
            ) == _run_work_with_pair_price(single, 10**30)

    @pytest.mark.parametrize(
        "body",
        [
            # Rows of hyperlink breaks before rows of plain ones, then
            # empty table children.
            b"<w:tbl>"
            + (b"<w:tr><w:tc>" + _BREAK_IN_LINK + b"</w:tc></w:tr>") * 3
            + (b"<w:tr><w:tc>" + _BREAK_IN_RUN + b"</w:tc></w:tr>") * 2
            + b"<w:bookmarkEnd/>" * 50
            + b"</w:tbl>",
            # Body paragraphs and a cell, text between them.
            _BREAK_IN_LINK * 3
            + b"text"
            + _cell_of(_BREAK_IN_RUN * 4 + b"<w:x/>" * 9)
            + _BREAK_IN_RUN * 2,
        ],
    )
    def test_sort_steps_are_priced(self, body):
        # Each element counts the union's breaks below it times its
        # child nodes, priced at one unit per step here.
        data = _blank_with(body)
        steps = _expected_sort_steps(data)
        assert steps > 0
        assert (
            _run_work_with_sort_price(data, 1)
            - _run_work_with_sort_price(data, 10**30)
            == steps - 1
        )

    @pytest.mark.parametrize(
        "content",
        [
            _BREAK_IN_RUN * 50 + b"<w:x/>" * 500,
            _BREAK_IN_LINK * 7 + _BREAK_IN_SMART_TAG * 11 + b"<w:x/>" * 500,
        ],
    )
    def test_sort_of_one_branch_is_not_priced(self, content):
        # One non-empty branch is already in document order.
        data = _blank_with(_cell_of(content))
        assert _run_work_with_sort_price(data, 1) == _run_work_with_sort_price(
            data, 10**30
        )

    def test_hostile_sort_filler_is_refused_by_default(self):
        # 1,000 rows of hyperlink breaks before 1,000 rows of plain
        # ones, then empty table children the sort walks per break:
        # 160,000 took 2.2-2.6 s per evaluation and 640,000 10.5 s, all
        # counted at 0.02 of MAX_DOCX_RUN_WORK before.
        guard = _guard()

        def row(paragraph: bytes) -> bytes:
            return b"<w:tr><w:tc>" + paragraph + b"</w:tc></w:tr>"

        data = _blank_with(
            b"<w:tbl>"
            + row(_BREAK_IN_LINK) * 1_000
            + row(_BREAK_IN_RUN) * 1_000
            + b"<w:bookmarkEnd/>" * 1_000_000
            + b"</w:tbl>"
        )
        assert len(data) < 80_000
        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(data, ".docx")

    def test_sort_price_covers_siblings_far_apart_in_memory(self):
        # Filler elements with 16 empty attributes each lie a cache miss
        # apart: 1,000 rows of each branch before 200,000 of them took
        # 16 s per evaluation (40 ns per step), counted at 0.29 of
        # MAX_DOCX_RUN_WORK at 30 ns per step before.
        guard = _guard()
        filler = (
            b"<w:bookmarkEnd"
            + b"".join(b' w:a%d=""' % n for n in range(16))
            + b"/>"
        )

        def row(paragraph: bytes) -> bytes:
            return b"<w:tr><w:tc>" + paragraph + b"</w:tc></w:tr>"

        data = _blank_with(
            b"<w:tbl>"
            + row(_BREAK_IN_LINK) * 1_000
            + row(_BREAK_IN_RUN) * 1_000
            + filler * 200_000
            + b"</w:tbl>"
        )
        assert len(data) < 200_000
        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(data, ".docx")

    def test_sort_price_covers_binary_insertion(self):
        # Below 64 nodes libxml2 sorts by binary insertion: 62 breaks of
        # one branch before one of another walked the filler after them
        # about 12 times per break (1.1 s per evaluation over 200,000
        # plain filler elements, 5.4 s over elements with 16
        # attributes). Unweighted, this shape counts at 0.25 of
        # MAX_DOCX_RUN_WORK; weighted, it is refused.
        guard = _guard()

        def row(paragraph: bytes) -> bytes:
            return b"<w:tr><w:tc>" + paragraph + b"</w:tc></w:tr>"

        data = _blank_with(
            b"<w:tbl>"
            + row(_BREAK_IN_LINK) * 62
            + row(_BREAK_IN_RUN)
            + b"<w:bookmarkEnd/>" * 1_200_000
            + b"</w:tbl>"
        )
        assert len(data) < 100_000
        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(data, ".docx")

    def test_sort_price_defaults(self):
        guard = _guard()
        assert guard.DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT <= 200
        assert guard.DOCX_PAGE_BREAK_SORT_INSERTION_WALKS >= 4
        assert guard.DOCX_PAGE_BREAK_SORT_INSERTION_RUN >= 64

    def test_pairs_refuse_below_the_count_ceiling(self):
        # 40,000 breaks in each cell branch: 1.6e9 pairs, ~3.2-3.6 s for
        # the union, counted at 0.37 of MAX_DOCX_RUN_WORK before.
        guard = _guard()
        data = _blank_with(
            _cell_of(_BREAK_IN_RUN * 40_000 + _BREAK_IN_LINK * 40_000)
        )
        with patch.object(guard, "MAX_DOCX_PAGE_BREAKS", 10**9):
            with pytest.raises(
                guard.DecompressionBombError, match="cells cost"
            ):
                guard.validate_zip_container(data, ".docx")

    def test_word_like_long_document_is_accepted(self):
        # 280 pages of ten paragraphs, each page's first paragraph
        # starting with a rendered break (some inside a hyperlink), then
        # a four-column table of 600 rows over 20 pages, each page
        # boundary crossing a row and putting a break in each of its
        # cells: Word's one break per page and flow.
        guard = _guard()
        text = b"<w:r><w:t xml:space='preserve'>Some words on the page. </w:t></w:r>"
        page = b"".join(
            (
                b"<w:p><w:r><w:lastRenderedPageBreak/><w:t>Start</w:t></w:r>"
                if n % 7
                else b"<w:p><w:hyperlink><w:r><w:lastRenderedPageBreak/>"
                b"<w:t>Link</w:t></w:r></w:hyperlink>"
            )
            + text * 4
            + b"</w:p>"
            + (b"<w:p>" + text * 6 + b"</w:p>") * 9
            for n in range(280)
        )

        def row(n: int) -> bytes:
            mark = b"<w:lastRenderedPageBreak/>" if n % 30 == 29 else b""
            cell = (
                b"<w:tc><w:p><w:r>"
                + mark
                + b"<w:t>cell</w:t></w:r></w:p></w:tc>"
            )
            return b"<w:tr>" + cell * 4 + b"</w:tr>"

        table = b"<w:tbl>" + b"".join(row(n) for n in range(600)) + b"</w:tbl>"
        data = _blank_with(page + table)

        guard.validate_zip_container(data, ".docx")
        assert (
            _run_work_with_pair_price(data, 1)
            - _run_work_with_pair_price(data, 10**30)
            == 240 * 40 + 40 * 80 + 240 * 80 - 1
        )
        # Its sort steps are a small share of the ceiling.
        steps = _expected_sort_steps(data)
        assert (
            0
            < steps
            < (
                guard.MAX_DOCX_RUN_WORK
                * guard.DOCX_PAGE_BREAK_SORT_STEPS_PER_UNIT
            )
            // 100
        )


def _inline_drawing(runs: bytes, inline: bytes = b"wp:inline") -> bytes:
    return (
        b"<w:r><w:drawing><"
        + inline
        + b">"
        + runs
        + b"</"
        + inline
        + b"></w:drawing></w:r>"
    )


_PLAIN_RUN = b"<w:r><w:t>x</w:t></w:r>"


class TestDocxInlineDrawingRuns:
    # unstructured reads a body-level paragraph's text with
    # ``w:r | w:hyperlink | w:r/descendant::wp:inline[ancestor::w:drawing][1]//w:r``;
    # the third branch's runs are not the paragraph's children, so the
    # union's merge and sort are not bounded by the paragraph's own
    # child count, and a paragraph with a single w:r branch counted
    # nothing: one run whose inline drawing holds a run, before 10,000
    # plain runs (37 KB), took 0.37 s per evaluation at 0.04 of the
    # ceiling (40,000: 8.9 s), ten runs whose drawings hold 5,000 runs
    # each 2.9 s at 0.17.

    def test_premise_unstructured_reads_runs_in_inline_drawings(self):
        import inspect

        docx = pytest.importorskip("docx")
        udocx = pytest.importorskip("unstructured.partition.docx")
        union = (
            "w:r | w:hyperlink | "
            "w:r/descendant::wp:inline[ancestor::w:drawing][1]//w:r"
        )
        assert union in inspect.getsource(udocx)
        data = _blank_with(
            b"<w:p>"
            + _inline_drawing(b"<w:r/>" * 4)
            + _PLAIN_RUN * 3
            + b"</w:p>"
        )
        p = docx.Document(io.BytesIO(data)).paragraphs[0]._p
        assert len(p.xpath(union)) == 1 + 3 + 4

    @pytest.mark.parametrize("nested, plain", [(50, 30), (1, 300)])
    def test_inline_runs_are_a_third_branch(self, nested, plain):
        def stats(inline: bytes):
            return _docx_stats(
                _blank_with(
                    b"<w:p>"
                    + _inline_drawing(b"<w:r/>" * nested, inline)
                    + _PLAIN_RUN * plain
                    + b"</w:p>"
                )
            )

        # Under wp:anchor the runs are not selected, and nothing else
        # differs. The inline's own child nodes (the nested runs, which
        # have none) are walked by the sort as well.
        selected = 1 + plain + nested
        children = 1 + plain
        assert stats(b"wp:inline").run_work - stats(
            b"wp:anchor"
        ).run_work == _union_units(
            selected * max(children, selected) + selected * nested
        )

    @pytest.mark.parametrize("nested, filler", [(3, 0), (40, 500), (1, 2_000)])
    def test_nodes_in_inline_drawings_are_priced(self, nested, filler):
        # A run whose drawing holds runs before filler, then a
        # hyperlink: sorting the drawing's runs walks the filler after
        # them, so the drawing's nodes count once per node selected.
        def stats(inline: bytes):
            return _docx_stats(
                _blank_with(
                    b"<w:p>"
                    + _inline_drawing(
                        b"<w:c>"
                        + b"<w:r/>" * nested
                        + b"<w:x/>" * filler
                        + b"</w:c>",
                        inline,
                    )
                    + b"<w:hyperlink/></w:p>"
                )
            )

        selected = 1 + 1 + nested
        children = 2
        # The inline's one child, and w:c's runs and filler.
        drawing_nodes = 1 + nested + filler
        # Three non-empty branches under wp:inline, two under wp:anchor.
        with_drawing = _union_units(
            2 * selected * max(children, selected) + selected * drawing_nodes
        )
        without = _union_units(children * children)
        assert (
            stats(b"wp:inline").run_work - stats(b"wp:anchor").run_work
            == with_drawing - without
        )

    def test_hostile_drawing_filler_is_refused(self):
        # One run with a hyperlink after it, its drawing holding 1,000
        # runs before 100,000 empty elements: 0.31-0.39 s per evaluation
        # of the paragraph union, counted at 0.04 of MAX_DOCX_RUN_WORK
        # before.
        guard = _guard()
        drawing = _inline_drawing(
            b"<w:c>" + b"<w:r/>" * 1_000 + b"<w:x/>" * 100_000 + b"</w:c>"
        )
        data = _blank_with(b"<w:p>" + drawing + b"<w:hyperlink/></w:p>")
        assert len(data) < 40_000
        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(data, ".docx")

    @pytest.mark.parametrize(
        "paragraph",
        [
            # An inline outside a w:drawing.
            b"<w:p><w:r><w:pict><wp:inline><w:r/><w:r/></wp:inline></w:pict>"
            b"</w:r>" + _PLAIN_RUN * 5 + b"</w:p>",
            # A drawing in a hyperlink's run, not the paragraph's.
            b"<w:p><w:hyperlink>"
            + _inline_drawing(b"<w:r/><w:r/>")
            + b"</w:hyperlink></w:p>",
        ],
    )
    def test_runs_the_third_branch_does_not_select_are_not_counted(
        self, paragraph
    ):
        moved = paragraph.replace(b"wp:inline", b"wp:anchor")
        assert (
            _docx_stats(_blank_with(paragraph)).run_work
            == _docx_stats(_blank_with(moved)).run_work
        )

    @pytest.mark.parametrize(
        "paragraph",
        [
            _inline_drawing(b"<w:r/>") + _PLAIN_RUN * 10_000,
            _inline_drawing(b"<w:r/>" * 5_000) * 10,
        ],
        ids=["one-before-many", "ten-of-many"],
    )
    def test_hostile_inline_runs_are_refused(self, paragraph):
        guard = _guard()
        data = _blank_with(b"<w:p>" + paragraph + b"</w:p>")
        assert len(data) < 40_000
        with pytest.raises(guard.DecompressionBombError, match="cells cost"):
            guard.validate_zip_container(data, ".docx")

    def test_paragraphs_with_inline_text_boxes_are_accepted(self):
        # 500 paragraphs of 20 runs, each with an inline drawing holding
        # a text box of two ten-run paragraphs.
        guard = _guard()
        box = (
            b"<wps:txbx><w:txbxContent>"
            + (b"<w:p>" + _PLAIN_RUN * 10 + b"</w:p>") * 2
            + b"</w:txbxContent></wps:txbx>"
        )
        paragraph = (
            b"<w:p>"
            + _PLAIN_RUN * 10
            + _inline_drawing(box)
            + _PLAIN_RUN * 10
            + b"</w:p>"
        )
        data = _blank_with(paragraph * 500)
        guard.validate_zip_container(data, ".docx")
        assert _docx_stats(data).run_work < guard.MAX_DOCX_RUN_WORK // 4


_PPTX_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_PPTX_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_PPTX_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _placeholder_sp(
    shape_id: int, idx: int, ph_type: str = "body", junk: int = 0
):
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="s{shape_id}"/>'
        f"<p:cNvSpPr/>{'<p:x/>' * junk}<p:nvPr>"
        f'<p:ph type="{ph_type}" idx="{idx}"/></p:nvPr></p:nvSpPr><p:spPr/>'
        "<p:txBody><a:bodyPr/><a:p><a:r><a:t>t</a:t></a:r></a:p></p:txBody>"
        "</p:sp>"
    ).encode()


def _deck_with_placeholders(
    on_slide: int,
    on_layout: int,
    *,
    slides: int = 1,
    junk: int = 0,
    links: int = 0,
) -> bytes:
    """Blank-layout slides each holding *on_slide* placeholders without
    a position, all of the idx of the layout's last placeholder, after
    *on_layout* other placeholders without one (each with *junk* empty
    elements in its non-visual properties); with *links* external
    hyperlink relationships on each slide."""
    pptx = pytest.importorskip("pptx")
    deck = pptx.Presentation()
    layout = deck.slide_layouts[6]
    for _ in range(slides):
        deck.slides.add_slide(layout)
    layout_name = layout.part.partname[1:]
    buf = io.BytesIO()
    deck.save(buf)
    slide_shapes = b"".join(
        _placeholder_sp(100 + n, 7777) for n in range(on_slide)
    )
    layout_shapes = b"".join(
        _placeholder_sp(100 + n, 5000 + n, junk=junk) for n in range(on_layout)
    ) + _placeholder_sp(99_999, 7777)
    link_rels = "".join(
        f'<Relationship Id="rIdLink{n}" Type="{_PPTX_R}/hyperlink" '
        f'Target="https://example.com/{n}" TargetMode="External"/>'
        for n in range(links)
    ).encode()
    changes = {
        layout_name: lambda xml: xml.replace(
            b"</p:spTree>", layout_shapes + b"</p:spTree>", 1
        )
    }
    for n in range(1, slides + 1):
        changes[f"ppt/slides/slide{n}.xml"] = lambda xml: xml.replace(
            b"</p:spTree>", slide_shapes + b"</p:spTree>", 1
        )
        changes[f"ppt/slides/_rels/slide{n}.xml.rels"] = lambda xml: (
            xml.replace(b"</Relationships>", link_rels + b"</Relationships>", 1)
        )
    return _rewrite_zip(buf.getvalue(), changes)


def _placeholder_work(data: bytes) -> int:
    from pptx.opc.packuri import PackURI

    guard = _guard()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        slides = guard._pptx_slide_targets(archive, PackURI)
        return guard._pptx_placeholder_work(archive, PackURI, slides, 10**30)


def _expected_read_units(data: bytes, slide: int = 1) -> int:
    """One inherited read of slide *slide*, counted independently from
    the parsed parts python-pptx resolves."""
    pptx = pytest.importorskip("pptx")
    guard = _guard()
    part = pptx.Presentation(io.BytesIO(data)).slides[slide - 1].part
    layout = part.slide_layout.part
    master = layout.slide_master.part

    def shape_tree(element):
        tree = element.find(f"{{{_PPTX_P}}}cSld/{{{_PPTX_P}}}spTree")
        below = tree.xpath(
            "*/node()[not(self::text())]"
            " | */*/node()[not(self::text())]"
            " | */*/*/node()[not(self::text())]"
        )
        return len(tree), 2 * len(below)

    layout_shapes, layout_nodes = shape_tree(layout._element)
    master_shapes, master_nodes = shape_tree(master._element)
    relationships = len(part.rels) + len(layout.rels)
    return (
        layout_shapes
        + master_shapes
        - (
            -(layout_nodes + master_nodes)
            // guard.PPTX_PLACEHOLDER_NODES_PER_UNIT
        )
        - (-relationships // guard.PPTX_PLACEHOLDER_RELATIONSHIPS_PER_UNIT)
    )


class TestPptxPlaceholderWork:
    # unstructured reads top and left of every slide shape up to three
    # times each; a placeholder without a position of its own inherits
    # it through python-pptx's uncached layout (and master) lookups,
    # which build a proxy for each layout placeholder per read: 100
    # placeholders over a 100-placeholder layout, a 30 KB file, took
    # 5.2 s, 200 over 200, 18.6 s, and nothing counted them.

    def test_premise_each_inherited_read_scans_the_layout(self):
        pptx = pytest.importorskip("pptx")
        from pptx.shapes import shapetree

        data = _deck_with_placeholders(3, 5)
        layout = pptx.Presentation(io.BytesIO(data)).slide_layouts[6]
        # The template's three, the five added and the matching one.
        on_layout = len(list(layout.placeholders))
        assert on_layout == 9
        calls = []
        original = shapetree.LayoutPlaceholders._shape_factory

        def counting(self, element):
            calls.append(element)
            return original(self, element)

        with patch.object(
            shapetree.LayoutPlaceholders, "_shape_factory", counting
        ):
            slide = pptx.Presentation(io.BytesIO(data)).slides[0]
            placeholders = list(slide.placeholders)
            for shape in placeholders:
                assert shape.top is not None
            assert len(calls) == 3 * on_layout
            for shape in placeholders:
                assert shape.left is not None
            assert len(calls) == 2 * 3 * on_layout

    def test_premise_unstructured_reads_positions_of_every_shape(self):
        import inspect

        upptx = pytest.importorskip("unstructured.partition.pptx")
        source = inspect.getsource(upptx)
        assert "return shape.top or 0, shape.left or 0" in source
        assert (
            "(shape.top and shape.left) and (shape.top < 0 or shape.left < 0)"
            in source
        )

    def test_work_is_placeholders_times_one_read(self):
        data = _deck_with_placeholders(4, 6, slides=3)
        read = _expected_read_units(data)
        assert read > 6
        assert _placeholder_work(data) == 3 * 4 * read

    def test_layout_nodes_and_slide_relationships_are_counted(self):
        plain = _deck_with_placeholders(2, 3)
        junk = _deck_with_placeholders(2, 3, junk=5_000)
        linked = _deck_with_placeholders(2, 3, links=1_000)
        for data in (plain, junk, linked):
            assert _placeholder_work(data) == 2 * _expected_read_units(data)
        assert _placeholder_work(junk) - _placeholder_work(plain) >= (
            2 * 3 * 2 * 5_000 // _guard().PPTX_PLACEHOLDER_NODES_PER_UNIT
        )
        assert _placeholder_work(linked) - _placeholder_work(plain) >= (
            2 * 1_000 // _guard().PPTX_PLACEHOLDER_RELATIONSHIPS_PER_UNIT
        )

    def test_ceiling_accepts_the_limit_and_refuses_one_more(self):
        guard = _guard()
        data = _deck_with_placeholders(5, 8, slides=2)
        work = _placeholder_work(data)
        with patch.object(guard, "MAX_PPTX_PLACEHOLDER_WORK", work):
            guard.validate_zip_container(data, ".pptx")
        with patch.object(guard, "MAX_PPTX_PLACEHOLDER_WORK", work - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="placeholders cost"
            ):
                guard.validate_zip_container(data, ".pptx")

    def test_hostile_deck_is_refused_by_default(self):
        # Three slides of 200 placeholders over a 200-placeholder layout:
        # ~56 s to partition (one such slide: 18.6 s), accepted before.
        guard = _guard()
        data = _deck_with_placeholders(200, 200, slides=3)
        assert len(data) < 60_000
        with pytest.raises(
            guard.DecompressionBombError, match="placeholders cost"
        ):
            guard.validate_zip_container(data, ".pptx")

    def test_upload_path_refuses_the_fan_out(self):
        from local_deep_research.document_loaders.bytes_loader import (
            load_from_bytes,
        )

        guard = _guard()
        data = _deck_with_placeholders(200, 200, slides=3)
        with pytest.raises(guard.DecompressionBombError):
            load_from_bytes(data, ".pptx", "placeholders.pptx")

    def test_slide_relationships_over_the_cap_are_refused(self):
        guard = _guard()
        data = _deck_with_placeholders(1, 0, links=guard.MAX_PART_RELATIONSHIPS)
        with pytest.raises(
            guard.DecompressionBombError, match="relationships part"
        ):
            guard.validate_zip_container(data, ".pptx")

    @pytest.mark.parametrize("part", ["slide", "layout", "master"])
    def test_slide_layout_and_master_with_a_dtd_are_refused(self, part):
        pptx = pytest.importorskip("pptx")
        guard = _guard()
        data = _deck(1)
        slide = pptx.Presentation(io.BytesIO(data)).slides[0].part
        layout = slide.slide_layout.part
        names = {
            "slide": (slide.partname[1:], b"<p:sld"),
            "layout": (layout.partname[1:], b"<p:sldLayout"),
            "master": (layout.slide_master.part.partname[1:], b"<p:sldMaster"),
        }
        name, root = names[part]
        hostile = _rewrite_zip(
            data, {name: _with_doctype(b"<!DOCTYPE " + root[1:] + b">", root)}
        )
        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(hostile, ".pptx")
        guard.validate_zip_container(data, ".pptx")  # control

    def test_real_decks_are_accepted(self):
        pptx = pytest.importorskip("pptx")
        guard = _guard()
        # Every layout of the default template, each placeholder filled.
        deck = pptx.Presentation()
        for layout in deck.slide_layouts:
            slide = deck.slides.add_slide(layout)
            for shape in slide.placeholders:
                if shape.has_text_frame:
                    shape.text_frame.text = "text"
        buf = io.BytesIO()
        deck.save(buf)
        guard.validate_zip_container(buf.getvalue(), ".pptx")
        assert _placeholder_work(buf.getvalue()) < 600
        # 300 title-and-content slides: 34 units each.
        deck = pptx.Presentation()
        for n in range(300):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = f"Slide {n}"
            slide.placeholders[1].text = "first point\nsecond point"
        buf = io.BytesIO()
        deck.save(buf)
        guard.validate_zip_container(buf.getvalue(), ".pptx")
        assert _placeholder_work(buf.getvalue()) <= 40 * 300


_CT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_CT_OFFICE = "application/vnd.openxmlformats-officedocument."
_CT_SLIDE = _CT_OFFICE + "presentationml.slide+xml"
_CT_DOCUMENT = _CT_OFFICE + "wordprocessingml.document.main+xml"
_STYLES_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"
)


def _retyped(data: bytes, partname: str, content_type: str) -> bytes:
    """*data* with an ``Override`` giving *partname* *content_type*
    appended to ``[Content_Types].xml`` (a later entry wins)."""
    override = (
        f'<Override PartName="{partname}" ContentType="{content_type}"/>'
    ).encode()
    return _rewrite_zip(
        data,
        {
            "[Content_Types].xml": lambda xml: xml.replace(
                b"</Types>", override + b"</Types>", 1
            )
        },
    )


def _empty_slide() -> bytes:
    return (
        f'<p:sld xmlns:a="{_PPTX_A}" xmlns:r="{_PPTX_R}" xmlns:p="{_PPTX_P}">'
        "<p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id='1' name=''/>"
        "<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/></p:spTree>"
        "</p:cSld></p:sld>"
    ).encode()


def _chained_layout_deck(on_slide: int, on_layout: int) -> bytes:
    """``_deck_with_placeholders(on_slide, on_layout)``, but the slide's
    layout relationship targets an empty slide (typed as a slide, not
    listed in the presentation), whose own layout relationship targets
    the layout of *on_layout* placeholders."""
    direct = _deck_with_placeholders(on_slide, on_layout)
    with zipfile.ZipFile(io.BytesIO(direct)) as zf:
        slide_rels = zf.read("ppt/slides/_rels/slide1.xml.rels")
    layout_rel = re.compile(rb'(Type="[^"]*/slideLayout"\s+Target=")[^"]+"')
    assert layout_rel.search(slide_rels)
    chained = _rewrite_zip(
        direct,
        {
            "ppt/slides/_rels/slide1.xml.rels": lambda xml: layout_rel.sub(
                rb'\g<1>slide99.xml"', xml
            )
        },
        add={
            "ppt/slides/slide99.xml": _empty_slide(),
            "ppt/slides/_rels/slide99.xml.rels": slide_rels,
        },
    )
    return _retyped(chained, "/ppt/slides/slide99.xml", _CT_SLIDE)


def _chained_docx_part(reltype: str, member: str, payload: bytes) -> bytes:
    """python-docx's blank document whose *reltype* relationship
    targets a second document part (typed as one), whose own *reltype*
    relationship targets *member* holding *payload*, typed as the part
    *reltype* names."""
    docx = pytest.importorskip("docx")
    buf = io.BytesIO()
    docx.Document().save(buf)
    target = re.compile(
        rb'(<Relationship [^>]*Type="'
        + re.escape(reltype.encode())
        + rb'"[^>]*Target=")[^"]+"'
        + rb"|"
        + rb'(<Relationship [^>]*Target=")[^"]+("[^>]*Type="'
        + re.escape(reltype.encode())
        + rb'")'
    )

    def retarget(xml: bytes) -> bytes:
        def sub(match):
            if match.group(1):
                return match.group(1) + b'document2.xml"'
            return match.group(2) + b"document2.xml" + match.group(3)

        out, count = target.subn(sub, xml)
        assert count == 1
        return out

    chained = _rewrite_zip(
        buf.getvalue(),
        {"word/_rels/document.xml.rels": retarget},
        add={
            "word/document2.xml": (
                f'<w:document xmlns:w="{_W_MAIN}"><w:body/></w:document>'
            ).encode(),
            "word/_rels/document2.xml.rels": (
                f'<Relationships xmlns="{_PKG_RELS}"><Relationship Id="rId1" '
                f'Type="{reltype}" Target="{member.split("/")[-1]}"/>'
                "</Relationships>"
            ).encode(),
            member: payload,
        },
    )
    kind = reltype.rsplit("/", 1)[1]
    chained = _retyped(
        chained, f"/{member}", f"{_CT_OFFICE}wordprocessingml.{kind}+xml"
    )
    return _retyped(chained, "/word/document2.xml", _CT_DOCUMENT)


def _many_styles(count: int) -> bytes:
    return (
        f'<w:styles xmlns:w="{_W_MAIN}">'.encode()
        + b"".join(
            b'<w:style w:type="paragraph" w:styleId="S%d"/>' % n
            for n in range(count)
        )
        + b"</w:styles>"
    )


class TestRelationshipTargetContentTypes:
    # python-pptx and python-docx build each part as the class its
    # content type names, whatever relationship reaches it, and a slide
    # (or document) part answers ``slide_layout`` (or ``styles`` and
    # ``settings``) by following its own relationship of that type. A
    # slide's layout relationship to another slide therefore makes
    # python-pptx read that slide's layout, which the guard never
    # priced: 300 placeholders over a chained 300-placeholder layout
    # were counted at 1,200 units, the same layout reached directly at
    # about 90,000.

    def test_premise_python_pptx_follows_a_chained_layout(self):
        pptx = pytest.importorskip("pptx")
        data = _chained_layout_deck(3, 5)
        slide = pptx.Presentation(io.BytesIO(data)).slides[0]
        assert (
            slide.part.part_related_by(
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/slideLayout"
            ).partname
            == "/ppt/slides/slide99.xml"
        )
        # The template's three placeholders, the five added and the
        # matching one: the layout two relationships away.
        assert len(list(slide.slide_layout.placeholders)) == 9

    def test_chained_layout_is_refused(self):
        guard = _guard()
        direct = _deck_with_placeholders(300, 300)
        with pytest.raises(
            guard.DecompressionBombError, match="placeholders cost"
        ):
            guard.validate_zip_container(direct, ".pptx")
        chained = _chained_layout_deck(300, 300)
        assert len(chained) < 80_000
        with pytest.raises(
            guard.DecompressionBombError, match="slide layout part has the"
        ):
            guard.validate_zip_container(chained, ".pptx")

    @pytest.mark.parametrize(
        ("reltype", "member", "attribute"),
        [
            (_STYLES_REL, "word/styles2.xml", "styles"),
            (_SETTINGS_REL, "word/settings2.xml", "settings"),
        ],
    )
    def test_premise_python_docx_follows_a_chained_part(
        self, reltype, member, attribute
    ):
        docx = pytest.importorskip("docx")
        root = "w:styles" if attribute == "styles" else "w:settings"
        payload = (
            f'<{root} xmlns:w="{_W_MAIN}">'.encode()
            + b"<w:pad/>" * 7
            + f"</{root}>".encode()
        )
        data = _chained_docx_part(reltype, member, payload)
        document = docx.Document(io.BytesIO(data))
        assert len(getattr(document, attribute).element) == 7

    @pytest.mark.parametrize(
        ("reltype", "member", "payload"),
        [
            (
                _STYLES_REL,
                "word/styles2.xml",
                lambda guard: _many_styles(guard.MAX_DOCX_STYLES * 4),
            ),
            (
                _SETTINGS_REL,
                "word/settings2.xml",
                lambda guard: (
                    f'<w:settings xmlns:w="{_W_MAIN}">'.encode()
                    + b"<w:pad/>" * (guard.MAX_DOCX_SETTINGS_CHILDREN * 4)
                    + b"</w:settings>"
                ),
            ),
        ],
        ids=["styles", "settings"],
    )
    def test_chained_styles_and_settings_are_refused(
        self, reltype, member, payload
    ):
        # Four times the cap behind a second document part: the guard
        # counted that part's one child before.
        guard = _guard()
        data = _chained_docx_part(reltype, member, payload(guard))
        with pytest.raises(guard.DecompressionBombError, match="content type"):
            guard.validate_zip_container(data, ".docx")

    @pytest.mark.parametrize(
        ("ext", "part", "content_type", "what"),
        [
            (
                ".pptx",
                "main",
                _CT_OFFICE + "presentationml.slideshow.main+xml",
                "main",
            ),
            (
                ".pptx",
                "slide",
                _CT_OFFICE + "presentationml.slideLayout+xml",
                "slide",
            ),
            (".pptx", "layout", _CT_SLIDE, "slide layout"),
            (
                ".pptx",
                "master",
                _CT_OFFICE + "presentationml.slideLayout+xml",
                "slide master",
            ),
            (
                ".docx",
                "main",
                "application/vnd.ms-word.document.macroEnabled.main+xml",
                "main",
            ),
            (".docx", "styles", "application/xml", "styles"),
            (".docx", "settings", _CT_DOCUMENT, "settings"),
            (
                ".docx",
                "header",
                _CT_OFFICE + "wordprocessingml.footer+xml",
                "header",
            ),
            (
                ".docx",
                "footer",
                _CT_OFFICE + "wordprocessingml.header+xml",
                "footer",
            ),
        ],
    )
    def test_mistyped_target_is_refused(self, ext, part, content_type, what):
        guard = _guard()
        data, partnames = self._real(ext)
        guard.validate_zip_container(data, ext)  # control
        hostile = _retyped(data, partnames[part], content_type)
        with pytest.raises(
            guard.DecompressionBombError, match=f"{what} part has the wrong"
        ):
            guard.validate_zip_container(hostile, ext)

    @staticmethod
    def _real(ext: str) -> tuple[bytes, dict[str, str]]:
        """A package python-pptx or python-docx writes (a deck whose
        slide has placeholders, a document with a header and a footer),
        and the partnames of its parts."""
        if ext == ".pptx":
            pptx = pytest.importorskip("pptx")
            data = _deck_with_placeholders(2, 2)
            slide = pptx.Presentation(io.BytesIO(data)).slides[0]
            layout = slide.slide_layout
            return data, {
                "main": "/ppt/presentation.xml",
                "slide": str(slide.part.partname),
                "layout": str(layout.part.partname),
                "master": str(layout.slide_master.part.partname),
            }
        docx = pytest.importorskip("docx")
        document = docx.Document()
        document.add_paragraph("body")
        section = document.sections[0]
        section.header.paragraphs[0].text = "header"
        section.footer.paragraphs[0].text = "footer"
        buf = io.BytesIO()
        document.save(buf)
        data = buf.getvalue()
        loaded = docx.Document(io.BytesIO(data))
        part = loaded.part
        return data, {
            "main": str(part.partname),
            "styles": str(part._styles_part.partname),
            "settings": str(part._settings_part.partname),
            "header": str(loaded.sections[0].header.part.partname),
            "footer": str(loaded.sections[0].footer.part.partname),
        }

    @pytest.mark.parametrize("ext", [".pptx", ".docx"])
    def test_real_packages_are_accepted(self, ext):
        guard = _guard()
        data, _ = self._real(ext)
        guard.validate_zip_container(data, ext)
        # And python-pptx's and python-docx's untouched defaults.
        guard.validate_zip_container(_real_package(ext), ext)

    def test_content_types_resolve_as_the_libraries_resolve_them(self):
        docx = pytest.importorskip("docx")
        guard = _guard()
        data, partnames = self._real(".docx")
        styles = partnames["styles"]
        styles_ct = _CT_OFFICE + "wordprocessingml.styles+xml"
        # Part names match without regard to case...
        upper = _retyped(data, styles.upper(), "application/xml")
        with pytest.raises(guard.DecompressionBombError, match="styles part"):
            guard.validate_zip_container(upper, ".docx")
        # ...and a later entry replaces an earlier one, either way.
        fixed = _retyped(upper, styles, styles_ct)
        guard.validate_zip_container(fixed, ".docx")
        assert docx.Document(io.BytesIO(fixed)).styles is not None
        # Without an Override, the Default of the extension applies.
        override = re.compile(
            rb'<Override PartName="' + re.escape(styles.encode()) + rb'"[^>]*/>'
        )
        bare = _rewrite_zip(
            data,
            {"[Content_Types].xml": lambda xml: override.sub(b"", xml)},
        )
        with pytest.raises(guard.DecompressionBombError, match="styles part"):
            guard.validate_zip_container(bare, ".docx")
        defaulted = _rewrite_zip(
            bare,
            {
                "[Content_Types].xml": lambda xml: xml.replace(
                    b"</Types>",
                    f'<Default Extension="XML" ContentType="{styles_ct}"/>'
                    "</Types>".encode(),
                )
            },
        )
        assert docx.Document(io.BytesIO(defaulted)).styles is not None
        guard.validate_zip_container(defaulted, ".docx")

    @pytest.mark.parametrize("ext", [".pptx", ".docx"])
    def test_package_without_content_types_is_refused(self, ext):
        guard = _guard()
        data = _real_package(ext)
        with zipfile.ZipFile(io.BytesIO(data)) as zin:
            entries = {
                info.filename: zin.read(info)
                for info in zin.infolist()
                if info.filename != "[Content_Types].xml"
            }
        with pytest.raises(guard.DecompressionBombError, match="content types"):
            guard.validate_zip_container(_zip_bytes(entries), ext)


_MAIN_RELS = {
    ".docx": "word/_rels/document.xml.rels",
    ".pptx": "ppt/_rels/presentation.xml.rels",
}


def _relationships(count: int, target: str, external: bool = False) -> bytes:
    """A relationships part of *count* entries all targeting *target*."""
    mode = ' TargetMode="External"' if external else ""
    body = "".join(
        f'<Relationship Id="h{n}" Type="{_REL_NS}/customXml" '
        f'Target="{target}"{mode}/>'
        for n in range(count)
    )
    return f'<Relationships xmlns="{_PKG_RELS}">{body}</Relationships>'.encode()


def _package_with_parts(
    ext: str,
    parts: int,
    hubs: dict[int, int],
    *,
    external: bool = False,
    name: str = "item",
) -> bytes:
    """A real package whose main part relates to *parts* tiny parts
    ``cx/<name>NNNNN.xml``; part *i* of *hubs* gets a relationships part
    of ``hubs[i]`` entries, all to the last of them (or external)."""
    names = [f"{name}{n:05d}.xml" for n in range(parts)]
    links = "".join(
        f'<Relationship Id="rX{n}" Type="{_REL_NS}/customXml" '
        f'Target="../cx/{part}"/>'
        for n, part in enumerate(names)
    ).encode()
    add = {f"cx/{part}": b"<a/>" for part in names}
    for index, count in hubs.items():
        add[f"cx/_rels/{names[index]}.rels"] = _relationships(
            count,
            "https://example.com/" if external else names[-1],
            external,
        )
    return _rewrite_zip(
        _real_package(ext),
        {
            _MAIN_RELS[ext]: lambda xml: xml.replace(
                b"</Relationships>", links + b"</Relationships>"
            )
        },
        add,
    )


def _name_bytes(name: str) -> int:
    """Bytes CPython stores *name* in: 1, 2 or 4 per character by its
    widest code point (PEP 393)."""
    widest = max(map(ord, name), default=0)
    return len(name) * (1 if widest < 256 else 2 if widest < 65_536 else 4)


def _walk_steps(data: bytes) -> int:
    """Internal relationships of every relationships part times the
    members, each member one step plus one per 256 bytes of its name as
    CPython stores it."""
    from lxml import etree

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        internal = sum(
            1
            for member in names
            if member.endswith(".rels")
            for rel in etree.fromstring(zf.read(member))
            if rel.get("TargetMode") != "External"
        )
    return internal * (len(names) + sum(map(_name_bytes, names)) // 256)


def _orphan_rels(ext: str, body: str, root_attrs: str = "") -> bytes:
    """A real package plus a relationships part no part owns, holding
    *body* under its root."""
    return _rewrite_zip(
        _real_package(ext),
        add={
            "cx/_rels/none.xml.rels": (
                f'<Relationships xmlns="{_PKG_RELS}"{root_attrs}>{body}'
                "</Relationships>"
            ).encode()
        },
    )


class TestPackageRelationships:
    # python-docx and python-pptx parse the relationships part of every
    # part they reach, not just those the guard reads; python-docx also
    # scans a list of the parts it has visited once per internal
    # relationship (PackageReader._walk_phys_parts, OpcPackage.iter_rels):
    # 9,000 parts and one unread part with 200,000 relationships to the
    # last of them, a 1.66 MB file, took 41 s to open.

    def test_premise_python_docx_scans_a_list_of_visited_parts(self):
        import inspect

        pytest.importorskip("docx")
        from docx.opc import package, pkgreader

        walk = inspect.getsource(pkgreader.PackageReader._walk_phys_parts)
        assert "visited_partnames = []" in walk
        assert "partname in visited_partnames" in walk
        iter_rels = inspect.getsource(package.OpcPackage.iter_rels)
        assert "visited = [] if visited is None" in iter_rels

    @pytest.mark.parametrize("ext", [".docx", ".pptx"])
    def test_premise_unread_relationships_parts_are_parsed(self, ext):
        # The library reaches the part and reads its relationships.
        data = _package_with_parts(ext, 3, {1: 2})
        if ext == ".docx":
            opened = pytest.importorskip("docx").Document(io.BytesIO(data))
        else:
            opened = pytest.importorskip("pptx").Presentation(io.BytesIO(data))
        parts = {
            str(rel.target_part.partname): rel.target_part
            for rel in opened.part.rels.values()
            if not rel.is_external
        }
        assert len(parts["/cx/item00001.xml"].rels) == 2

    @pytest.mark.parametrize("ext", [".docx", ".pptx"])
    def test_every_relationships_part_is_capped(self, ext):
        guard = _guard()
        limit = guard.MAX_PART_RELATIONSHIPS
        guard.validate_zip_container(
            _package_with_parts(ext, 2, {0: limit}), ext
        )
        with pytest.raises(
            guard.DecompressionBombError,
            match=f"more than {limit} entries",
        ):
            guard.validate_zip_container(
                _package_with_parts(ext, 2, {0: limit + 1}), ext
            )
        # Reached or not: a relationships part of no part counts too.
        orphan = _rewrite_zip(
            _real_package(ext),
            add={"cx/_rels/none.xml.rels": _relationships(limit + 1, "x")},
        )
        with pytest.raises(
            guard.DecompressionBombError,
            match=f"more than {limit} entries",
        ):
            guard.validate_zip_container(orphan, ext)

    @pytest.mark.parametrize("ext", [".docx", ".pptx"])
    def test_relationships_are_capped_across_the_package(self, ext):
        guard = _guard()
        data = _package_with_parts(
            ext, 6, {n: 9_000 for n in range(6)}, external=True
        )
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            total = sum(
                _count_children(data, member)
                for member in zf.namelist()
                if member.endswith(".rels")
            )
        assert guard.MAX_PACKAGE_RELATIONSHIPS < total
        with pytest.raises(
            guard.DecompressionBombError, match="entries together"
        ):
            guard.validate_zip_container(data, ext)
        with patch.object(guard, "MAX_PACKAGE_RELATIONSHIPS", total):
            guard.validate_zip_container(data, ext)
        with patch.object(guard, "MAX_PACKAGE_RELATIONSHIPS", total - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="entries together"
            ):
                guard.validate_zip_container(data, ext)

    @pytest.mark.parametrize("ext", [".docx", ".pptx"])
    def test_unread_relationships_parts_get_the_skeleton_treatment(self, ext):
        guard = _guard()
        data = _package_with_parts(ext, 2, {0: 50})
        member = "cx/_rels/item00000.xml.rels"
        with_dtd = _rewrite_zip(
            data,
            {
                member: _with_doctype(
                    b"<!DOCTYPE Relationships>", b"<Relationships"
                )
            },
        )
        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(with_dtd, ext)
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            size = zf.getinfo(member).file_size
        with patch.object(guard, "MAX_OPC_SKELETON_PART_BYTES", size - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="larger than"
            ):
                guard.validate_zip_container(data, ext)

    def test_relationship_walk_is_priced_exactly(self):
        guard = _guard()
        data = _package_with_parts(".docx", 40, {3: 70, 39: 11})
        steps = _walk_steps(data)
        with patch.object(guard, "MAX_DOCX_RELATIONSHIP_WALK", steps):
            guard.validate_zip_container(data, ".docx")
        with patch.object(guard, "MAX_DOCX_RELATIONSHIP_WALK", steps - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="relationship walk"
            ):
                guard.validate_zip_container(data, ".docx")
        # External relationships are never looked up in the list.
        external = _package_with_parts(".docx", 40, {3: 500}, external=True)
        assert _walk_steps(external) < steps
        with patch.object(guard, "MAX_DOCX_RELATIONSHIP_WALK", steps):
            guard.validate_zip_container(external, ".docx")
        # Long names cost their bytes.
        long = _package_with_parts(
            ".docx", 40, {3: 70, 39: 11}, name="n" * 2_000
        )
        assert _walk_steps(long) > steps * 3
        with patch.object(guard, "MAX_DOCX_RELATIONSHIP_WALK", steps):
            with pytest.raises(
                guard.DecompressionBombError, match="relationship walk"
            ):
                guard.validate_zip_container(long, ".docx")

    def test_hostile_relationship_walk_is_refused_by_default(self):
        # Every relationships part under MAX_PART_RELATIONSHIPS; 9,000
        # parts and 27,000 internal relationships cost python-docx about
        # 2.4e8 list comparisons on open.
        guard = _guard()
        data = _package_with_parts(".docx", 9_000, {0: 9_000, 1: 9_000})
        assert len(data) < 2_000_000
        with pytest.raises(
            guard.DecompressionBombError, match="relationship walk"
        ):
            guard.validate_zip_container(data, ".docx")
        # python-pptx keeps sets: no walk to price.
        guard.validate_zip_container(
            _package_with_parts(".pptx", 9_000, {0: 9_000, 1: 9_000}), ".pptx"
        )

    def test_document_with_many_images_is_accepted(self):
        # One part and one internal relationship per image: 5,000 of
        # them count about an eighth of the ceiling.
        guard = _guard()
        data = _package_with_parts(".docx", 5_000, {})
        assert _walk_steps(data) < guard.MAX_DOCX_RELATIONSHIP_WALK // 4
        guard.validate_zip_container(data, ".docx")

    @pytest.mark.parametrize("ext", [".docx", ".pptx"])
    def test_an_entry_with_child_elements_is_refused(self, ext):
        # The OPC schema gives Relationship text content only. One entry
        # holding millions of children (16 MiB, ~16 KB deflated, in a
        # part no library reads) cost the guard 3.1 s and ~550 MB, since
        # an entry was only dropped at its end.
        guard = _guard()
        rel = f'Id="r1" Type="{_REL_NS}/customXml" Target="x.xml"'
        guard.validate_zip_container(
            _orphan_rels(ext, f"<Relationship {rel}>text</Relationship>"), ext
        )
        for children in (
            "<a/>",
            "<a/>" * 200_000,
            "<x:a xmlns:x='u'><b/></x:a>",
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="child elements"
            ):
                guard.validate_zip_container(
                    _orphan_rels(
                        ext, f"<Relationship {rel}>{children}</Relationship>"
                    ),
                    ext,
                )
        # Any element below the root's children, whatever it is called.
        with pytest.raises(
            guard.DecompressionBombError, match="child elements"
        ):
            guard.validate_zip_container(
                _orphan_rels(ext, "<other><a/></other>"), ext
            )

    @pytest.mark.parametrize("ext", [".docx", ".pptx"])
    def test_comments_and_instructions_count_as_entries(self, ext):
        # Each stays in the streamed tree until the next entry ends.
        guard = _guard()
        limit = 50
        rel = f'<Relationship Id="r1" Type="{_REL_NS}/customXml" Target="x"/>'
        with patch.object(guard, "MAX_PART_RELATIONSHIPS", limit):
            guard.validate_zip_container(
                _orphan_rels(ext, rel + "<!---->" * 24 + "<?p?>" * 25), ext
            )
            for body in (
                rel + "<!---->" * 25 + "<?p?>" * 25,
                f'<Relationship Id="r1" Type="t" Target="x">{"<!---->" * limit}'
                "</Relationship>",
            ):
                with pytest.raises(
                    guard.DecompressionBombError,
                    match=f"more than {limit} entries",
                ):
                    guard.validate_zip_container(_orphan_rels(ext, body), ext)
        # At the default cap, a part of nothing but comments is refused.
        with pytest.raises(
            guard.DecompressionBombError, match="more than .* entries"
        ):
            guard.validate_zip_container(
                _orphan_rels(
                    ext, "<!---->" * (guard.MAX_PART_RELATIONSHIPS + 1)
                ),
                ext,
            )

    @pytest.mark.parametrize("ext", [".docx", ".pptx"])
    def test_elements_with_many_attributes_are_refused(self, ext):
        guard = _guard()
        limit = guard.MAX_RELS_ELEMENT_ATTRIBUTES
        assert 4 <= limit <= 64

        def attributes(count: int) -> str:
            return "".join(f' a{n}=""' for n in range(count))

        guard.validate_zip_container(
            _orphan_rels(ext, f"<Relationship{attributes(limit)}/>"), ext
        )
        guard.validate_zip_container(
            _orphan_rels(ext, "", root_attrs=attributes(limit)), ext
        )
        for body, root in (
            (f"<Relationship{attributes(limit + 1)}/>", ""),
            ("", attributes(limit + 1)),
        ):
            with pytest.raises(
                guard.DecompressionBombError,
                match=f"more than {limit} attributes",
            ):
                guard.validate_zip_container(
                    _orphan_rels(ext, body, root_attrs=root), ext
                )

    def test_ordinary_relationships_parts_still_pass(self):
        # Text content and namespace declarations are allowed.
        guard = _guard()
        body = (
            f'<Relationship xmlns:x="urn:x" Id="r1" Type="{_REL_NS}/customXml"'
            ' Target="x" TargetMode="External">  </Relationship>'
        )
        for ext in (".docx", ".pptx"):
            guard.validate_zip_container(_orphan_rels(ext, body), ext)

    @pytest.mark.parametrize("char", ["\u00e9", "\u30a2", "\U0001f600"])
    def test_relationship_walk_prices_names_as_stored(self, char):
        # CPython compares equal-length names over len * width bytes:
        # names of four-byte characters measured two to four times the
        # cost per character of ASCII ones.
        guard = _guard()
        ascii_data = _package_with_parts(
            ".docx", 40, {3: 70, 39: 11}, name="n" * 1_000
        )
        data = _package_with_parts(
            ".docx", 40, {3: 70, 39: 11}, name=char * 1_000
        )
        steps = _walk_steps(data)
        width = {"\u00e9": 1, "\u30a2": 2, "\U0001f600": 4}[char]
        if width == 1:
            assert steps == _walk_steps(ascii_data)
        else:
            # Every name stored in *width* bytes per character (less
            # the fixed step per member and the ASCII member names).
            assert steps > _walk_steps(ascii_data) * (width - 1)
        with patch.object(guard, "MAX_DOCX_RELATIONSHIP_WALK", steps):
            guard.validate_zip_container(data, ".docx")
        with patch.object(guard, "MAX_DOCX_RELATIONSHIP_WALK", steps - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="relationship walk"
            ):
                guard.validate_zip_container(data, ".docx")

    def test_relationship_defaults(self):
        guard = _guard()
        assert guard.MAX_PART_RELATIONSHIPS >= 10_000
        assert guard.MAX_PACKAGE_RELATIONSHIPS >= 50_000
        assert guard.MAX_DOCX_RELATIONSHIP_WALK <= 200_000_000
        assert guard.DOCX_RELATIONSHIP_WALK_NAME_BYTES_PER_STEP <= 256


class TestXlsxSheetCount:
    # openpyxl checks each sheet's target against a list of the members
    # and pandas looks every sheet up by name in lists of all of them:
    # 10,000 empty sheets took 40 s in pandas.read_excel, and 100,000
    # sheets without parts next to 9,000 members 7.6 s in openpyxl alone.

    def test_premise_lookups_are_linear_scans(self):
        import inspect

        pytest.importorskip("openpyxl")
        pytest.importorskip("pandas")
        from openpyxl.reader.excel import ExcelReader
        from pandas.io.excel._base import BaseExcelReader

        source = inspect.getsource(ExcelReader.read_worksheets)
        assert "rel.target not in self.valid_files" in source
        data = _xlsx_with_sheets(["rId1"], [("rId1", "worksheets/sheet1.xml")])
        assert isinstance(ExcelReader(io.BytesIO(data)).valid_files, list)
        assert "not in self.sheet_names" in inspect.getsource(
            BaseExcelReader.raise_if_bad_sheet_by_name
        )

    def test_sheet_cap_accepts_the_limit_and_refuses_one_more(self):
        guard = _guard()
        limit = guard.MAX_XLSX_SHEETS

        def workbook(sheets: int) -> bytes:
            return _xlsx_with_sheets(
                [f"r{n}" for n in range(sheets)],
                [(f"r{n}", f"worksheets/s{n}.xml") for n in range(sheets)],
            )

        guard.validate_zip_container(workbook(limit), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError, match=f"more than {limit} sheets"
        ):
            guard.validate_zip_container(workbook(limit + 1), ".xlsx")

    def test_sheet_defaults(self):
        guard = _guard()
        assert 500 <= guard.MAX_XLSX_SHEETS <= 2_000


_SML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _stylesheet(
    formats: list[tuple[int, str]] = (),
    cell_formats: list[int] = (0,),
    named: int = 1,
    named_formats: int = 1,
    prefix: str = "",
    namespace: str = _SML_NS,
    style_ids: list[int] | None = None,
    fonts: str = "<font/>",
    fills: str = "<fill><patternFill/></fill>",
    borders: str = "<border/>",
    style_format: str = "",
) -> bytes:
    """An ``xl/styles.xml`` openpyxl loads: *formats* as (numFmtId,
    formatCode), one cellXfs ``<xf>`` per id in *cell_formats*, *named*
    named styles over *named_formats* cellStyleXfs entries (all with a
    custom format id, so openpyxl rebuilds its format dict per style),
    under *prefix* bound to *namespace*. *style_ids* overrides the
    named styles' ``xfId`` values; *fonts*, *fills* and *borders* are
    the lists' unprefixed contents (one entry each by default), and
    *style_format* the children of every cellStyleXfs ``<xf>``."""
    p = f"{prefix}:" if prefix else ""
    ns = f' xmlns{":" + prefix if prefix else ""}="{namespace}"'
    numfmts = "".join(
        f'<{p}numFmt numFmtId="{n}" formatCode="{code}"/>'
        for n, code in formats
    )
    xfs = "".join(
        f'<{p}xf numFmtId="{n}" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        for n in cell_formats
    )
    style_xf = '<{p}xf numFmtId="164" fontId="0" fillId="0" borderId="0"'
    style_xfs = (
        style_xf.format(p=p)
        + (f">{style_format}</{p}xf>" if style_format else "/>")
    ) * named_formats
    if style_ids is None:
        style_ids = [n % named_formats for n in range(named)]
    styles = "".join(
        f'<{p}cellStyle name="s{n}" xfId="{xf_id}"/>'
        for n, xf_id in enumerate(style_ids)
    )
    if prefix:
        fonts, fills, borders = (
            re.sub(r"<(/?)(?=\w)", rf"<\1{p}", xml)
            for xml in (fonts, fills, borders)
        )
    return (
        f"<{p}styleSheet{ns}><{p}numFmts>{numfmts}</{p}numFmts>"
        f"<{p}fonts>{fonts}</{p}fonts>"
        f"<{p}fills>{fills}</{p}fills>"
        f"<{p}borders>{borders}</{p}borders>"
        f"<{p}cellStyleXfs>{style_xfs}</{p}cellStyleXfs>"
        f"<{p}cellXfs>{xfs}</{p}cellXfs>"
        f"<{p}cellStyles>{styles}</{p}cellStyles></{p}styleSheet>"
    ).encode()


def _xlsx_with_styles(styles: bytes) -> bytes:
    """A workbook openpyxl writes, with *styles* as ``xl/styles.xml``."""
    return _rewrite_zip(
        _real_package(".xlsx"), {"xl/styles.xml": lambda _xml: styles}
    )


def _load_xlsx(data: bytes):
    """Open *data* the way pandas does."""
    openpyxl = pytest.importorskip("openpyxl")
    return openpyxl.load_workbook(
        io.BytesIO(data), read_only=True, data_only=True, keep_links=False
    )


def _gradient_fill(stops: int, tag: str = "gradientFill") -> str:
    """A ``<fill>`` of *stops* stops, each 13 characters of attribute
    values, under a first child named *tag*."""
    return (
        f"<fill><{tag}>"
        + "".join(
            f'<stop position="{n / 1_000:.3f}"><color rgb="FF000000"/></stop>'
            for n in range(stops)
        )
        + f"</{tag}></fill>"
    )


def _format_units(guard, code: str) -> int:
    return len(code) * (guard.XLSX_NUMBER_FORMAT_SCAN_WEIGHT + code.count("["))


class TestXlsxStyles:
    # openpyxl loads xl/styles.xml on every read (pandas' read-only load
    # too) and repeats work per cell format and per named style: one
    # 1,000,000-character format code under 2,000 cell formats, a 3.3 KB
    # file, took 53 s to load.

    def test_premise_openpyxl_repeats_work_per_format(self):
        import inspect

        pytest.importorskip("openpyxl")
        from openpyxl.reader.excel import ExcelReader
        from openpyxl.styles import stylesheet
        from openpyxl.xml.constants import ARC_STYLE

        assert ARC_STYLE == "xl/styles.xml"
        assert "apply_stylesheet(" in inspect.getsource(ExcelReader.read)
        normalise = inspect.getsource(stylesheet.Stylesheet._normalise_numbers)
        assert "for idx, style in enumerate(self.cell_styles)" in normalise
        assert "is_date_format(fmt)" in normalise
        assert "is_timedelta_format(fmt)" in normalise
        expand = inspect.getsource(stylesheet.Stylesheet._expand_named_style)
        assert "self.custom_formats" in expand
        assert isinstance(
            inspect.getattr_static(stylesheet.Stylesheet, "custom_formats"),
            property,
        )
        # Children are matched by local name, in any namespace.
        data = _xlsx_with_styles(
            _stylesheet(
                [(164, "0.0")], [164, 164], prefix="x", namespace="urn:x"
            )
        )
        workbook = _load_xlsx(data)
        assert len(workbook._cell_styles) == 2
        workbook.close()

    def test_long_format_codes_are_refused(self):
        guard = _guard()
        limit = guard.MAX_XLSX_NUMBER_FORMAT_CHARS
        assert 255 <= limit <= 1_024
        guard.validate_zip_container(
            _xlsx_with_styles(_stylesheet([(164, "0" * limit)], [164])),
            ".xlsx",
        )
        for styles in (
            _stylesheet([(164, "0" * (limit + 1))], [164]),
            # Referenced or not, in any namespace.
            _stylesheet([(170, "0" * (limit + 1))], [0], prefix="s"),
            _stylesheet([(164, "0" * (limit + 1))], [164], namespace="urn:x"),
            # The reported shape: one 1,000,000-character code, 2,000 xfs.
            _stylesheet([(164, "0" * 1_000_000)], [164] * 2_000),
        ):
            with pytest.raises(
                guard.DecompressionBombError,
                match=f"longer than {limit} characters",
            ):
                guard.validate_zip_container(_xlsx_with_styles(styles), ".xlsx")

    @pytest.mark.parametrize(
        "styles",
        [
            lambda n: _stylesheet([], [0] * n),
            lambda n: _stylesheet([], [0], named=n, named_formats=1),
            lambda n: _stylesheet([], [0], named=1, named_formats=n),
            # Repeated lists are counted together.
            lambda n: _stylesheet([], [0] * (n // 2)).replace(
                b"</styleSheet>",
                b"<cellXfs>"
                + b'<xf numFmtId="0"/>' * (n - n // 2)
                + b"</cellXfs>"
                b"</styleSheet>",
            ),
        ],
    )
    def test_cell_format_lists_are_capped(self, styles):
        guard = _guard()
        assert guard.MAX_XLSX_CELL_FORMATS >= 65_490
        with patch.object(guard, "MAX_XLSX_CELL_FORMATS", 40):
            guard.validate_zip_container(_xlsx_with_styles(styles(39)), ".xlsx")
            with pytest.raises(
                guard.DecompressionBombError, match="more than 40 cell"
            ):
                guard.validate_zip_container(
                    _xlsx_with_styles(styles(41)), ".xlsx"
                )

    def test_number_format_work_is_priced_exactly(self):
        guard = _guard()
        codes = [
            (164, "[" * 300),
            (165, "[$-409]d/m/yy h:mm"),
            (165, "0.00"),
            (166, "x"),
        ]
        cell_formats = [164] * 7 + [165] * 3 + [0, 14, 200]
        data = _xlsx_with_styles(_stylesheet(codes, cell_formats))
        units = (
            7 * _format_units(guard, "[" * 300)
            # A code given twice counts at the costlier entry.
            + 3 * _format_units(guard, "[$-409]d/m/yy h:mm")
            # Built-in or unknown formats count at 64 characters.
            + 3 * 64 * guard.XLSX_NUMBER_FORMAT_SCAN_WEIGHT
        )
        with patch.object(guard, "MAX_XLSX_NUMBER_FORMAT_WORK", units):
            guard.validate_zip_container(data, ".xlsx")
        with patch.object(guard, "MAX_XLSX_NUMBER_FORMAT_WORK", units - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="number formats cost"
            ):
                guard.validate_zip_container(data, ".xlsx")

    def test_bracket_codes_are_refused_by_default(self):
        # One scan restarts at every unclosed "[": 4,000 cell formats of
        # a 1,024-bracket code take about 2.3 s to load.
        guard = _guard()
        code = "[" * guard.MAX_XLSX_NUMBER_FORMAT_CHARS
        assert (
            4_000 * _format_units(guard, code)
            > guard.MAX_XLSX_NUMBER_FORMAT_WORK
        )
        with pytest.raises(
            guard.DecompressionBombError, match="number formats cost"
        ):
            guard.validate_zip_container(
                _xlsx_with_styles(_stylesheet([(164, code)], [164] * 4_000)),
                ".xlsx",
            )

    def test_named_style_work_is_priced(self):
        guard = _guard()
        formats = [(164 + n, "0.0") for n in range(30)]
        data = _xlsx_with_styles(
            _stylesheet(formats, [0], named=25, named_formats=20)
        )
        # 20 distinct xfIds, each a style of the fixed work, the 30
        # formats and the default font, fill and border (1, 2 and 1
        # elements).
        units = 20 * (
            guard.XLSX_NAMED_STYLE_BASE_UNITS
            + 30
            + 4 * guard.XLSX_STYLE_NODE_UNITS
        )
        with patch.object(guard, "MAX_XLSX_NAMED_STYLE_WORK", units):
            guard.validate_zip_container(data, ".xlsx")
        with patch.object(guard, "MAX_XLSX_NAMED_STYLE_WORK", units - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="named styles cost"
            ):
                guard.validate_zip_container(data, ".xlsx")
        # 5,000 named styles over 5,000 custom formats: about 3.5 s.
        many = [(164 + n, "0") for n in range(5_000)]
        with pytest.raises(
            guard.DecompressionBombError, match="named styles cost"
        ):
            guard.validate_zip_container(
                _xlsx_with_styles(
                    _stylesheet(many, [0], named=5_000, named_formats=5_000)
                ),
                ".xlsx",
            )

    def test_negative_xf_ids_count_as_named_styles(self):
        # openpyxl drops a named style whose xfId repeats but indexes a
        # negative xfId from the end, so n cellStyleXfs entries carry 2n
        # named styles: 8,800 over 4,400 formats and 4,400 custom
        # formats took 5.3 s to load, and were counted as 4,400.
        guard = _guard()
        n = 6
        formats = [(164 + i, "0.0") for i in range(10)]
        ids = list(range(-n, n))
        workbook = _load_xlsx(
            _xlsx_with_styles(
                _stylesheet(formats, [0], named_formats=n, style_ids=ids)
            )
        )
        assert len(workbook._named_styles) == 2 * n
        workbook.close()
        # Repeated ids are dropped, as openpyxl drops them.
        data = _xlsx_with_styles(
            _stylesheet(
                formats, [0], named_formats=n, style_ids=ids + [0, -1, 5]
            )
        )
        units = (
            2
            * n
            * (
                guard.XLSX_NAMED_STYLE_BASE_UNITS
                + len(formats)
                + 4 * guard.XLSX_STYLE_NODE_UNITS
            )
        )
        with patch.object(guard, "MAX_XLSX_NAMED_STYLE_WORK", units):
            guard.validate_zip_container(data, ".xlsx")
        with patch.object(guard, "MAX_XLSX_NAMED_STYLE_WORK", units - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="named styles cost"
            ):
                guard.validate_zip_container(data, ".xlsx")
        # The reported shape.
        many = [(164 + i, "0") for i in range(4_400)]
        with pytest.raises(
            guard.DecompressionBombError, match="named styles cost"
        ):
            guard.validate_zip_container(
                _xlsx_with_styles(
                    _stylesheet(
                        many,
                        [0],
                        named_formats=4_400,
                        style_ids=list(range(-4_400, 4_400)),
                    )
                ),
                ".xlsx",
            )

    def test_named_styles_price_the_objects_they_bind(self):
        # apply_stylesheet binds every named style, hashing its font,
        # fill and border (every stop of a gradient fill) and comparing
        # them with equal listed entries (every character of a font
        # name): 1,000 named styles of a 1,000-stop gradient fill, an
        # 13 KB file, took about 3 s to load, 2,000 of 2,000 stops 12.4 s.
        guard = _guard()
        node = guard.XLSX_STYLE_NODE_UNITS
        # Text counts too, though openpyxl reads these from attributes.
        fonts = (
            '<font><name val="{}"/><sz val="11">{}</sz></font><font/>'.format(
                "n" * 1_000, "t" * 512
            )
        )
        border = (
            '<border><left style="thin"><color indexed="64"/></left></border>'
        )
        data = _xlsx_with_styles(
            _stylesheet(
                [(164, "0" * 600), (165, "0")],
                named=7,
                named_formats=7,
                fonts=fonts,
                fills=_gradient_fill(10) + "<fill><patternFill/></fill>",
                borders="<border/>" + border,
            )
        )
        # The 2 custom formats and the longest code's 600 characters;
        # the largest font: 3 elements and 1,514 characters; fill: 2
        # elements and 10 stops of 2 elements and 13 characters each;
        # border: 3 elements and 6 characters.
        units = 7 * (
            guard.XLSX_NAMED_STYLE_BASE_UNITS
            + 2
            + 600 // guard.XLSX_STYLE_CHARS_PER_UNIT
            + (3 * node + 1_514 // guard.XLSX_STYLE_CHARS_PER_UNIT)
            + (22 * node + 130 // guard.XLSX_STYLE_CHARS_PER_UNIT)
            + (3 * node + 6 // guard.XLSX_STYLE_CHARS_PER_UNIT)
        )
        with patch.object(guard, "MAX_XLSX_NAMED_STYLE_WORK", units):
            guard.validate_zip_container(data, ".xlsx")
        with patch.object(guard, "MAX_XLSX_NAMED_STYLE_WORK", units - 1):
            with pytest.raises(
                guard.DecompressionBombError, match="named styles cost"
            ):
                guard.validate_zip_container(data, ".xlsx")
        # The reported shape, refused by its stops alone and by its
        # named styles with the stop and entry ceilings lifted.
        reported = _xlsx_with_styles(
            _stylesheet(
                named=1_000,
                named_formats=1_000,
                fills=_gradient_fill(1_000),
            )
        )
        with pytest.raises(guard.DecompressionBombError, match="stops"):
            guard.validate_zip_container(reported, ".xlsx")
        with (
            patch.object(guard, "MAX_XLSX_GRADIENT_STOPS", 1_000),
            patch.object(guard, "MAX_XLSX_STYLE_ENTRY_ELEMENTS", 2_002),
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="named styles cost"
            ):
                guard.validate_zip_container(reported, ".xlsx")
        # 2,000 named styles of a fill at the stop ceiling load in
        # about 1.6 s.
        guard.validate_zip_container(
            _xlsx_with_styles(
                _stylesheet(
                    named=2_000,
                    named_formats=2_000,
                    fills=_gradient_fill(guard.MAX_XLSX_GRADIENT_STOPS),
                )
            ),
            ".xlsx",
        )

    def test_gradient_stops_are_capped(self):
        guard = _guard()
        limit = guard.MAX_XLSX_GRADIENT_STOPS
        assert 3 <= limit <= 256
        # openpyxl reads a fill's first child as a gradient whatever
        # its tag.
        workbook = _load_xlsx(
            _xlsx_with_styles(_stylesheet(fills=_gradient_fill(3, "x")))
        )
        assert type(workbook._fills[0]).__name__ == "GradientFill"
        assert len(workbook._fills[0].stop) == 3
        workbook.close()
        guard.validate_zip_container(
            _xlsx_with_styles(
                _stylesheet(fills=_gradient_fill(limit) * 2, prefix="s")
            ),
            ".xlsx",
        )
        for fills in (
            _gradient_fill(limit + 1),
            _gradient_fill(limit + 1, "x"),
            _gradient_fill(3) + _gradient_fill(limit + 1),
        ):
            with pytest.raises(
                guard.DecompressionBombError,
                match=f"more than {limit} stops",
            ):
                guard.validate_zip_container(
                    _xlsx_with_styles(_stylesheet(fills=fills)), ".xlsx"
                )

    def test_premise_from_tree_reads_fields_from_child_elements(self):
        pytest.importorskip("openpyxl")
        from openpyxl.styles.cell_style import CellStyle
        from openpyxl.styles.named_styles import _NamedCellStyle

        guard = _guard()
        base = _stylesheet([(164, "0")], [0])
        workbook = _load_xlsx(
            _xlsx_with_styles(
                base.replace(
                    b'formatCode="0"/>',
                    b'formatCode="0"><formatCode>0.000</formatCode></numFmt>',
                ).replace(
                    b'xfId="0"/></cellXfs>',
                    b'xfId="0"><numFmtId>164</numFmtId></xf></cellXfs>',
                )
            )
        )
        assert list(workbook._number_formats) == ["0.000"]
        assert workbook._cell_styles[0].numFmtId == 164
        workbook.close()
        assert guard._xlsx_text_fields(CellStyle) == {
            "numFmtId",
            "fontId",
            "fillId",
            "borderId",
            "xfId",
            "quotePrefix",
            "pivotButton",
            "applyNumberFormat",
            "applyFont",
            "applyFill",
            "applyBorder",
        }
        assert guard._xlsx_text_fields(_NamedCellStyle) == {
            "name",
            "xfId",
            "builtinId",
            "iLevel",
            "hidden",
            "customBuiltin",
        }

    @pytest.mark.parametrize(
        "old, new",
        [
            # A format code given as a child's text: one 1,000,000-
            # character code under 2,000 cell formats, a 1,881-byte
            # file, took 52 s to load.
            (
                b'formatCode="0"/>',
                b'formatCode="0"><formatCode>LONG</formatCode></numFmt>',
            ),
            (
                b'formatCode="0"/>',
                b'formatCode="0"><x:formatCode xmlns:x="urn:x">LONG'
                b"</x:formatCode></numFmt>",
            ),
            (
                b'formatCode="0"/>',
                b'formatCode="0"><numFmtId>165</numFmtId></numFmt>',
            ),
            (b'formatCode="0"/>', b'formatCode="0"><foo/></numFmt>'),
            # A cell format's format id, a named style's format's font
            # id, a named style's xfId.
            (
                b'xfId="0"/></cellXfs>',
                b'xfId="0"><numFmtId>164</numFmtId></xf></cellXfs>',
            ),
            (
                b'borderId="0"/></cellStyleXfs>',
                b'borderId="0"><fontId>0</fontId></xf></cellStyleXfs>',
            ),
            (
                b'xfId="0"/></cellStyles>',
                b'xfId="0"><xfId>0</xfId></cellStyle></cellStyles>',
            ),
        ],
    )
    def test_fields_given_as_child_elements_are_refused(self, old, new):
        guard = _guard()
        base = _stylesheet([(164, "0")], [164] * 2_000)
        assert old in base
        styles = base.replace(old, new.replace(b"LONG", b"0" * 1_000_000))
        with pytest.raises(guard.DecompressionBombError, match="child element"):
            guard.validate_zip_container(_xlsx_with_styles(styles), ".xlsx")

    def test_child_elements_excel_writes_pass(self):
        guard = _guard()
        children = (
            b'<alignment horizontal="center" wrapText="1"/>'
            b'<protection locked="0"/><extLst><ext uri="urn:x"/></extLst>'
            b"<foo/>"
        )
        styles = (
            _stylesheet([(164, "0")], [164], style_format=children.decode())
            .replace(
                b'xfId="0"/></cellXfs>',
                b'xfId="0">' + children + b"</xf></cellXfs>",
            )
            .replace(
                b'xfId="0"/></cellStyles>',
                b'xfId="0"><extLst><ext uri="urn:x"/></extLst>'
                b"</cellStyle></cellStyles>",
            )
        )
        data = _xlsx_with_styles(styles)
        workbook = _load_xlsx(data)
        assert workbook._alignments[1].horizontal == "center"
        workbook.close()
        guard.validate_zip_container(data, ".xlsx")

    def test_unloadable_ids_and_dtds_are_refused(self):
        guard = _guard()
        bad = _stylesheet([(164, "0")], [164]).replace(
            b'numFmtId="164" formatCode', b'numFmtId="x" formatCode'
        )
        with pytest.raises((TypeError, ValueError)):
            _load_xlsx(_xlsx_with_styles(bad))
        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(_xlsx_with_styles(bad), ".xlsx")
        with_dtd = b"<!DOCTYPE styleSheet>" + _stylesheet([(164, "0")], [164])
        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(_xlsx_with_styles(with_dtd), ".xlsx")

    def test_ordinary_workbooks_pass(self):
        guard = _guard()
        openpyxl = pytest.importorskip("openpyxl")
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row in range(1, 301):
            cell = sheet.cell(row=row, column=1, value=row)
            cell.number_format = f'[$-409]#,##0.{"0" * (row % 9)}" u{row}"'
        workbook.add_named_style(openpyxl.styles.NamedStyle(name="extra"))
        buf = io.BytesIO()
        workbook.save(buf)
        guard.validate_zip_container(buf.getvalue(), ".xlsx")
        # Excel's limits: 65,490 cell formats of an ordinary date code.
        limit = guard.MAX_XLSX_CELL_FORMATS
        code = "[$-409]d/m/yy\\ h:mm;@"
        guard.validate_zip_container(
            _xlsx_with_styles(_stylesheet([(164, code)], [164] * limit)),
            ".xlsx",
        )
        # No styles part at all is fine too.
        data = _real_package(".xlsx")
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            kept = {
                n: zf.read(n) for n in zf.namelist() if n != "xl/styles.xml"
            }
        guard.validate_zip_container(_zip_bytes(kept), ".xlsx")

    def test_style_defaults(self):
        guard = _guard()
        assert guard.MAX_XLSX_NUMBER_FORMAT_WORK <= 4_000_000_000
        assert guard.MAX_XLSX_NAMED_STYLE_WORK <= 20_000_000
        assert guard.MAX_XLSX_GRADIENT_STOPS <= 256
        assert guard.XLSX_NAMED_STYLE_BASE_UNITS >= 500
        assert guard.XLSX_STYLE_NODE_UNITS >= 10
        assert guard.XLSX_STYLE_CHARS_PER_UNIT <= 1_024
        assert guard.MAX_XLSX_STYLE_LIST_ENTRIES <= 16_384
        assert guard.MAX_XLSX_STYLE_HASH_CHAIN <= 8
        assert guard.XLSX_STYLE_INT_MAX < 2**61 - 1
        assert guard.XLSX_STYLE_INT_MIN > -(2**61 - 1)
        assert guard.MAX_XLSX_STYLE_FLOAT <= 2**53


_HASH_MODULUS = 2**61 - 1


def _one_hash_floats(count: int) -> list[str]:
    """*count* distinct floats of one hash: Python hashes ``x`` and
    ``x * 2**-61`` alike."""
    return [repr(2.0 ** (-61 * n)) for n in range(count)]


def _cell_formats(entries: list[str]) -> bytes:
    """A stylesheet with *entries* appended to its ``cellXfs``."""
    return _stylesheet().replace(
        b"</cellXfs>", "".join(entries).encode() + b"</cellXfs>"
    )


class TestXlsxStyleHashCollisions:
    # openpyxl keeps fonts, fills, borders, alignments, protections and
    # cell formats in IndexedLists (dicts keyed by the style object,
    # whose hash is taken over its field values) and adds each named
    # style's objects to them again. Distinct objects of one hash are
    # compared with each other on every insertion and lookup: 4,000
    # fonts differing only in a charset of 1 + k * (2**61 - 1) took
    # 6.6 s to load (51 KB), 8,000 of them under 8,000 named styles 81 s.

    def test_premise_distinct_style_objects_share_a_hash(self):
        pytest.importorskip("openpyxl")
        import inspect

        from openpyxl.styles import stylesheet
        from openpyxl.styles.cell_style import CellStyleList
        from openpyxl.styles.fonts import Font
        from openpyxl.utils.indexed_list import IndexedList

        fonts = [Font(charset=1 + n * _HASH_MODULUS) for n in range(3)]
        assert len({hash(font) for font in fonts}) == 1
        assert fonts[0] != fonts[1] != fonts[2] != fonts[0]
        floats = [Font(sz=float(x)) for x in _one_hash_floats(3)]
        assert len({hash(font) for font in floats}) == 1
        assert floats[0] != floats[1] != floats[2] != floats[0]
        assert "IndexedList(stylesheet.fonts)" in inspect.getsource(
            stylesheet.apply_stylesheet
        )
        assert "IndexedList(" in inspect.getsource(CellStyleList._to_array)
        assert "self._dict[val] = idx" in inspect.getsource(IndexedList)

    @pytest.mark.parametrize(
        "styles",
        [
            # The reported shape: fonts of one charset hash.
            lambda v: _stylesheet(
                fonts=f'<font><charset val="{1 + _HASH_MODULUS}"/></font>'
            ),
            # Ids the guard and openpyxl key dicts by.
            lambda v: _stylesheet([(164 + _HASH_MODULUS, "0.0")], [164]),
            lambda v: _stylesheet(style_ids=[_HASH_MODULUS]),
            # A child element's text and a float attribute.
            lambda v: _stylesheet(
                fonts=f"<font><color><theme>{2**32}</theme></color></font>"
            ),
            lambda v: _stylesheet(
                fills=f'<fill><gradientFill degree="{v}"/></fill>'
            ),
            lambda v: _stylesheet(fonts=f'<font><sz val="{v}"/></font>'),
            # A child's text is read even when the child has children.
            lambda v: _stylesheet(
                fills="<fill><gradientFill>"
                f"<degree>{v}<x/></degree></gradientFill></fill>"
            ),
        ],
        ids=[
            "charset",
            "numFmtId",
            "xfId",
            "theme-text",
            "degree",
            "sz",
            "degree-text-before-a-child",
        ],
    )
    def test_out_of_range_style_numbers_are_refused(self, styles):
        guard = _guard()
        for value in ("1e300", "nan", "-inf", str(2**33)):
            with pytest.raises(
                guard.DecompressionBombError, match="out of range"
            ):
                guard.validate_zip_container(
                    _xlsx_with_styles(styles(value)), ".xlsx"
                )

    def test_style_numbers_in_range_pass(self):
        guard = _guard()
        styles = _stylesheet(
            fonts=(
                f'<font><charset val="{guard.XLSX_STYLE_INT_MAX}"/>'
                '<sz val="409"/><color indexed="64" tint="-0.249977111117893"/>'
                "</font>"
                f'<font><charset val="{guard.XLSX_STYLE_INT_MIN}"/></font>'
            ),
            fills='<fill><gradientFill degree="90"><stop position="0">'
            '<color theme="1"/></stop></gradientFill></fill>',
        )
        guard.validate_zip_container(_xlsx_with_styles(styles), ".xlsx")
        _load_xlsx(_xlsx_with_styles(styles)).close()

    @pytest.mark.parametrize(
        "what, styles",
        [
            (
                "fonts",
                lambda values: _stylesheet(
                    fonts="".join(
                        f'<font><sz val="{v}"/></font>' for v in values
                    )
                ),
            ),
            (
                "fills",
                lambda values: _stylesheet(
                    fills="".join(
                        f'<fill><gradientFill degree="{v}"/></fill>'
                        for v in values
                    )
                ),
            ),
            (
                "borders",
                lambda values: _stylesheet(
                    borders="".join(
                        f'<border><left style="thin"><color tint="{v}"/>'
                        "</left></border>"
                        for v in values
                    )
                ),
            ),
            (
                "alignments",
                lambda values: _cell_formats(
                    [
                        f'<xf numFmtId="0"><alignment indent="{v}"/></xf>'
                        for v in values
                    ]
                ),
            ),
            # Each named style's format adds its alignment as it binds.
            (
                "alignments",
                lambda values: _stylesheet(
                    named=len(values) + 1,
                    style_ids=list(range(len(values) + 1)),
                ).replace(
                    b"</cellStyleXfs>",
                    "".join(
                        f'<xf numFmtId="0"><alignment indent="{v}"/></xf>'
                        for v in values
                    ).encode()
                    + b"</cellStyleXfs>",
                ),
            ),
            # Any tag in the fonts list is a font, in any namespace.
            (
                "fonts",
                lambda values: _stylesheet(
                    fonts="".join(
                        f'<x:f xmlns:x="urn:x"><x:sz val="{v}"/></x:f>'
                        for v in values
                    )
                ),
            ),
        ],
    )
    def test_distinct_style_objects_of_one_hash_are_refused(self, what, styles):
        guard = _guard()
        limit = guard.MAX_XLSX_STYLE_HASH_CHAIN
        values = _one_hash_floats(limit + 1)
        guard.validate_zip_container(
            _xlsx_with_styles(styles(values[:limit])), ".xlsx"
        )
        # Equal entries, however often repeated or however spelled, are
        # one value.
        respelled = [f"+{v}" for v in values[:limit]]
        guard.validate_zip_container(
            _xlsx_with_styles(styles((values[:limit] + respelled) * 50)),
            ".xlsx",
        )
        with pytest.raises(
            guard.DecompressionBombError, match=f"distinct {what} of one hash"
        ):
            guard.validate_zip_container(
                _xlsx_with_styles(styles(values)), ".xlsx"
            )

    def test_cell_formats_of_one_hash_are_refused(self, monkeypatch):
        # Free 32-bit fields of a cell format's style array can be
        # steered to one tuple hash; a constant hash stands in for that.
        pytest.importorskip("openpyxl")
        from openpyxl.styles.cell_style import StyleArray

        guard = _guard()
        limit = guard.MAX_XLSX_STYLE_HASH_CHAIN
        styles = _cell_formats(
            [f'<xf numFmtId="0" fontId="{n + 1}"/>' for n in range(limit)]
        )
        monkeypatch.setattr(StyleArray, "__hash__", lambda self: 7)
        with pytest.raises(
            guard.DecompressionBombError, match="distinct cell formats"
        ):
            guard.validate_zip_container(_xlsx_with_styles(styles), ".xlsx")
        monkeypatch.undo()
        guard.validate_zip_container(_xlsx_with_styles(styles), ".xlsx")

    @pytest.mark.parametrize(
        "container, entry",
        [
            ("fonts", "<font><sz val='{n}'/></font>"),
            ("fills", "<fill><gradientFill degree='{n}'/></fill>"),
            (
                "borders",
                "<border><left style='thin'><color indexed='{n}'/></left></border>",
            ),
        ],
    )
    def test_style_lists_are_capped(self, container, entry):
        guard = _guard()
        limit = guard.MAX_XLSX_STYLE_LIST_ENTRIES
        assert limit >= 512  # Excel's own font limit per workbook

        def styles(count: int, lists: int = 1) -> bytes:
            per_list = count // lists
            body = "".join(entry.format(n=n) for n in range(per_list))
            return _stylesheet(**{container: body}).replace(
                f"<{container}>".encode(),
                (f"<{container}>{body}</{container}>" * (lists - 1)).encode()
                + f"<{container}>".encode(),
                1,
            )

        guard.validate_zip_container(_xlsx_with_styles(styles(limit)), ".xlsx")
        for over in (styles(limit + 1), styles(limit + 2, lists=2)):
            with pytest.raises(
                guard.DecompressionBombError,
                match=f"more than {limit} {container} entries",
            ):
                guard.validate_zip_container(_xlsx_with_styles(over), ".xlsx")


def _before_root(xml: bytes, nodes: bytes) -> bytes:
    """*xml* with *nodes* inserted before its root element (after its
    XML declaration, if any)."""
    at = xml.index(b"?>") + 2 if xml.startswith(b"<?xml") else 0
    return xml[:at] + nodes + xml[at:]


_XLSX_LIST_SITES = {
    # NestedSequence lists: openpyxl builds an entry per child node.
    "fonts": ("<fonts><font/>", "<fonts><font/>{node}"),
    "fills": (
        "<fills><fill><patternFill/></fill>",
        "<fills><fill><patternFill/></fill>{node}",
    ),
    "borders": ("<borders><border/>", "<borders><border/>{node}"),
    "dxfs": ("</styleSheet>", "<dxfs><dxf/>{node}</dxfs></styleSheet>"),
    "indexedColors": (
        "</styleSheet>",
        "<colors><indexedColors><rgbColor rgb='FF000000'/>{node}"
        "</indexedColors></colors></styleSheet>",
    ),
    "mruColors": (
        "</styleSheet>",
        "<colors><mruColors><color rgb='FF000000'/>{node}</mruColors>"
        "</colors></styleSheet>",
    ),
    # Sequence lists and single entries, which openpyxl skips them in.
    "numFmts": ("<numFmts>", "<numFmts>{node}"),
    "cellStyleXfs": ("<cellStyleXfs>", "<cellStyleXfs>{node}"),
    "cellXfs": ("<cellXfs>", "<cellXfs>{node}"),
    "cellStyles": ("<cellStyles>", "<cellStyles>{node}"),
    "tableStyles": (
        "</styleSheet>",
        "<tableStyles count='0'>{node}</tableStyles></styleSheet>",
    ),
    "font": ("<fonts><font/>", "<fonts><font>{node}</font>"),
    "fill": (
        "<fills><fill><patternFill/></fill>",
        "<fills><fill><patternFill/>{node}</fill>",
    ),
    "gradientFill": (
        "<fills><fill><patternFill/></fill>",
        "<fills><fill><patternFill/></fill><fill><gradientFill>"
        "<stop position='0'><color rgb='FF000000'/></stop>{node}"
        "</gradientFill></fill>",
    ),
    "xf": ('xfId="0"/>', 'xfId="0">{node}</xf>'),
}


class TestXmlCommentsAndProcessingInstructions:
    # openpyxl parses xl/styles.xml and the workbook skeleton with lxml,
    # which keeps comments and processing instructions as nodes, and
    # NestedSequence builds an entry per child node: 2,000,000 empty
    # comments after one font, a 25 KB file the guard accepted (it
    # counted elements), took openpyxl 37 s and 960 MB to load as
    # 2,000,001 fonts.

    def test_premise_openpyxl_builds_an_entry_per_child_node(self):
        import inspect

        pytest.importorskip("lxml")
        pytest.importorskip("openpyxl")
        from openpyxl import xml as openpyxl_xml
        from openpyxl.descriptors.sequence import NestedSequence

        assert openpyxl_xml.LXML
        source = inspect.getsource(NestedSequence.from_tree)
        assert "for el in node" in source
        for node in ("<!---->", "<?p?>"):
            for site, expected in (
                ("fonts", lambda wb: len(wb._fonts)),
                ("borders", lambda wb: len(wb._borders)),
                ("dxfs", lambda wb: len(wb._differential_styles.styles)),
            ):
                old, new = _XLSX_LIST_SITES[site]
                styles = _stylesheet().replace(
                    old.encode(), new.format(node=node * 3).encode(), 1
                )
                workbook = _load_xlsx(_xlsx_with_styles(styles))
                assert expected(workbook) == 4, (site, node)
                workbook.close()
        # Where openpyxl skips them (between the lists), nothing is built.
        workbook = _load_xlsx(
            _xlsx_with_styles(
                _stylesheet().replace(b"<fills>", b"<!----><?p?><fills>", 1)
            )
        )
        assert len(workbook._fonts) == 1
        workbook.close()

    @pytest.mark.parametrize("node", ["<!---->", "<?p?>"])
    @pytest.mark.parametrize("site", sorted(_XLSX_LIST_SITES))
    def test_styles_comment_below_the_top_level_is_refused(self, site, node):
        guard = _guard()
        old, new = _XLSX_LIST_SITES[site]
        base = _stylesheet()
        assert base.count(old.encode()) >= 1, site
        control = base.replace(old.encode(), new.format(node="").encode(), 1)
        guard.validate_zip_container(_xlsx_with_styles(control), ".xlsx")
        styles = base.replace(old.encode(), new.format(node=node).encode(), 1)
        with pytest.raises(
            guard.DecompressionBombError,
            match="comment or processing instruction below its top level",
        ):
            guard.validate_zip_container(_xlsx_with_styles(styles), ".xlsx")

    def test_reported_shape_is_refused_quickly(self):
        guard = _guard()
        styles = _stylesheet(fonts="<font/>" + "<!---->" * 2_000_000)
        data = _xlsx_with_styles(styles)
        assert len(data) < 64 * 1024
        start = time.perf_counter()
        with pytest.raises(
            guard.DecompressionBombError, match="below its top level"
        ):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    @pytest.mark.parametrize("node", ["<!--c-->", "<?p x?>"])
    def test_styles_comment_at_the_top_level_is_accepted(self, node):
        """Between the lists, before and after the root: openpyxl skips
        them, at the cost of a tree node each."""
        guard = _guard()
        styles = _before_root(
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            + _stylesheet()
            .replace(b"<fills>", f"{node}<fills>".encode(), 1)
            .replace(b"</styleSheet>", f"{node}</styleSheet>".encode())
            + node.encode(),
            node.encode(),
        )
        data = _xlsx_with_styles(styles)
        guard.validate_zip_container(data, ".xlsx")
        workbook = _load_xlsx(data)
        assert len(workbook._fonts) == 1
        workbook.close()

    @pytest.mark.parametrize("node", ["<!---->", "<?p?>"])
    @pytest.mark.parametrize(
        "part, old, new",
        [
            # NestedSequence lists of the workbook.
            ("xl/workbook.xml", "<bookViews>", "<bookViews>{node}"),
            ("xl/workbook.xml", "<sheets>", "<sheets>{node}"),
            # Elsewhere below the root's children.
            (
                "xl/workbook.xml",
                "<definedNames/>",
                "<definedNames>{node}</definedNames>",
            ),
            ("xl/workbook.xml", 'r:id="rId1"/>', 'r:id="rId1">{node}</sheet>'),
            ("[Content_Types].xml", "/>", ">{node}</Default>"),
            (
                "xl/_rels/workbook.xml.rels",
                'Id="rId1"/>',
                'Id="rId1">{node}</Relationship>',
            ),
        ],
    )
    def test_skeleton_comment_below_the_top_level_is_refused(
        self, part, old, new, node
    ):
        guard = _guard()
        base = _real_package(".xlsx")

        def rewrite(fill: str) -> bytes:
            def change(xml: bytes) -> bytes:
                assert old.encode() in xml, (part, old)
                return xml.replace(
                    old.encode(), new.format(node=fill).encode(), 1
                )

            return _rewrite_zip(base, {part: change})

        guard.validate_zip_container(rewrite(""), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError,
            match="workbook skeleton part has a comment or processing",
        ):
            guard.validate_zip_container(rewrite(node), ".xlsx")

    def test_reported_workbook_shape_is_refused_quickly(self):
        # 2,390,000 comments among the bookViews of a 16 MiB workbook
        # part took the guard's own openpyxl parse 19.6 s and 1.1 GB.
        guard = _guard()
        data = _rewrite_zip(
            _real_package(".xlsx"),
            {
                "xl/workbook.xml": lambda xml: xml.replace(
                    b"<bookViews>", b"<bookViews>" + b"<!---->" * 200_000, 1
                )
            },
        )
        start = time.perf_counter()
        with pytest.raises(
            guard.DecompressionBombError, match="below its top level"
        ):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    @pytest.mark.parametrize("node", ["<!--c-->", "<?p x?>"])
    def test_skeleton_comment_at_the_top_level_is_accepted(self, node):
        guard = _guard()
        top = node.encode()
        data = _rewrite_zip(
            _real_package(".xlsx"),
            {
                part: lambda xml, root=root: _before_root(
                    xml.replace(root, top + root, 1) + top, top
                )
                for part, root in (
                    ("xl/workbook.xml", b"<sheets>"),
                    ("[Content_Types].xml", b"<Default "),
                )
            },
        )
        guard.validate_zip_container(data, ".xlsx")
        workbook = _load_xlsx(data)
        assert workbook.sheetnames == ["Sheet"]
        workbook.close()
        # The relationships list is read as every child node of the
        # root (ElementList), so there openpyxl itself drops the whole
        # list and cannot find the sheets; the guard, which runs
        # openpyxl's reader, refuses it as malformed.
        rels = _rewrite_zip(
            _real_package(".xlsx"),
            {
                "xl/_rels/workbook.xml.rels": lambda xml: xml.replace(
                    b"<Relationship ", top + b"<Relationship ", 1
                )
            },
        )
        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(rels, ".xlsx")
        with pytest.raises(KeyError):
            _load_xlsx(rels)

    @pytest.mark.parametrize(
        "ext, part",
        [
            (".xlsx", "xl/styles.xml"),
            (".pptx", "ppt/slides/slide1.xml"),
            (".docx", "word/header1.xml"),
            (".docx", "word/document.xml"),
        ],
    )
    def test_nodes_outside_the_root_are_capped(self, ext, part):
        """lxml's iterparse emits the comment event for a node before
        the root in time linear in the nodes before it: 60,000 empty
        comments before a slide's root, a 420 KB part, kept the guard
        busy for 16.7 s."""
        guard = _guard()
        limit = guard.MAX_XML_OUTSIDE_ROOT_NODES
        assert limit <= 1_024
        base = {
            ".xlsx": lambda: _xlsx_with_styles(_stylesheet()),
            ".pptx": lambda: _deck(1),
            ".docx": lambda: _docx_with_sections(1),
        }[ext]()

        def padded(before: int, after: int = 0) -> bytes:
            return _rewrite_zip(
                base,
                {
                    part: lambda xml: (
                        _before_root(xml, b"<!---->" * before)
                        + b"<?p?>" * after
                    )
                },
            )

        guard.validate_zip_container(padded(limit), ext)
        guard.validate_zip_container(padded(limit - 1, 1), ext)
        for over in (padded(limit + 1), padded(limit, 1)):
            with pytest.raises(
                guard.DecompressionBombError,
                match=f"more than {limit} comments or processing",
            ):
                guard.validate_zip_container(over, ext)
        start = time.perf_counter()
        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(padded(200_000), ext)
        assert time.perf_counter() - start < 5


def _xlsx_with_workbook(change) -> bytes:
    """An openpyxl-written workbook whose ``xl/workbook.xml`` is passed
    through *change* (bytes to bytes)."""
    return _rewrite_zip(_real_package(".xlsx"), {"xl/workbook.xml": change})


_R_DECL = (
    b' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
    b'relationships"'
)


def _workbook_lists(xml: bytes, lists: bytes) -> bytes:
    """*xml* (a workbook part) with *lists* before its ``definedNames``
    and the relationships namespace declared on its root."""
    xml = xml.replace(b"<workbook ", b"<workbook" + _R_DECL + b" ", 1)
    return xml.replace(b"<definedNames/>", lists + b"<definedNames/>", 1)


# (constant, list entries as a function of a count, list name)
_STYLE_LIST_CASES = {
    "numFmts": (
        "MAX_XLSX_NUMBER_FORMATS",
        lambda n: _stylesheet([(164 + k, "0") for k in range(n)]),
    ),
    "dxfs": (
        "MAX_XLSX_DIFFERENTIAL_FORMATS",
        lambda n: _stylesheet().replace(
            b"</styleSheet>",
            b"<dxfs>" + b"<dxf/>" * (n - n // 2) + b"</dxfs>"
            b"<dxfs>" + b"<dxf/>" * (n // 2) + b"</dxfs></styleSheet>",
        ),
    ),
    "tableStyles": (
        "MAX_XLSX_TABLE_STYLES",
        lambda n: _stylesheet().replace(
            b"</styleSheet>",
            b"<tableStyles>"
            + b"".join(b"<tableStyle name='t%d'/>" % k for k in range(n))
            + b"</tableStyles></styleSheet>",
        ),
    ),
    "tableStyle": (
        "MAX_XLSX_TABLE_STYLE_ELEMENTS",
        lambda n: _stylesheet().replace(
            b"</styleSheet>",
            b"<tableStyles><tableStyle name='t'>"
            + b"<tableStyleElement type='wholeTable'/>" * n
            + b"</tableStyle></tableStyles></styleSheet>",
        ),
    ),
    "indexedColors": (
        "MAX_XLSX_COLOR_LIST_ENTRIES",
        lambda n: _stylesheet().replace(
            b"</styleSheet>",
            b"<colors><indexedColors>"
            + b"<rgbColor rgb='FF000000'/>" * n
            + b"</indexedColors></colors></styleSheet>",
        ),
    ),
    "mruColors": (
        "MAX_XLSX_COLOR_LIST_ENTRIES",
        lambda n: _stylesheet().replace(
            b"</styleSheet>",
            b"<colors><mruColors>"
            + b"<color rgb='FF000000'/>" * n
            + b"</mruColors></colors></styleSheet>",
        ),
    ),
    "extLst": (
        "MAX_XLSX_EXTENSIONS",
        # Counted over every extension list of the part.
        lambda n: (
            _stylesheet(
                cell_formats=[0, 0],
            )
            .replace(
                b'xfId="0"/>',
                b'xfId="0"><extLst>'
                + b"<ext uri='u'/>" * (n // 2)
                + b"</extLst></xf>",
                1,
            )
            .replace(
                b"</styleSheet>",
                b"<extLst>"
                + b"<ext uri='u'/>" * (n - n // 2)
                + b"</extLst></styleSheet>",
            )
        ),
    ),
    "cellXfs": (
        # Every child of a cell-format list counts, not only <xf>.
        "MAX_XLSX_CELL_FORMATS",
        lambda n: _stylesheet().replace(
            b"<cellXfs>", b"<cellXfs>" + b"<x/>" * (n - 1), 1
        ),
    ),
}


class TestXlsxListCeilings:
    # openpyxl builds an object from every entry of every list of the
    # stylesheet and of the workbook part, and from every child it knows
    # of every other element, keeping only the last of a repeated one:
    # 500,000 <font/> in one dxf, a 10 KB file, took 7.4 s to load, and
    # 1,800,000 <calcPr/> in a 16 MiB workbook part 17 s, in the guard's
    # own parse as long again.

    def test_premise_every_list_openpyxl_builds_is_covered(self):
        """The list-valued fields openpyxl reads from the stylesheet and
        the workbook part, walked from its classes: an upgrade adding
        one fails here until the guard counts it."""
        pytest.importorskip("openpyxl")
        from openpyxl.descriptors import Descriptor
        from openpyxl.descriptors.sequence import Sequence
        from openpyxl.descriptors.serialisable import Serialisable
        from openpyxl.packaging.workbook import WorkbookPackage
        from openpyxl.styles.stylesheet import Stylesheet

        def lists(root) -> set:
            found, seen, pending = set(), set(), [root]
            while pending:
                cls = pending.pop()
                if cls in seen:
                    continue
                seen.add(cls)
                pending.extend(cls.__subclasses__())
                for name in dir(cls):
                    desc = getattr(cls, name, None)
                    if not isinstance(desc, Descriptor):
                        continue
                    if isinstance(desc, Sequence):
                        found.add(f"{cls.__name__}.{name}")
                    kind = getattr(desc, "expected_type", None)
                    if isinstance(kind, type) and issubclass(
                        kind, Serialisable
                    ):
                        pending.append(kind)
            return found

        # Each with the guard's bound: its list (by the container's
        # name), or, for stops, MAX_XLSX_GRADIENT_STOPS(_TOTAL).
        assert lists(Stylesheet) == {
            "Stylesheet.fonts",  # fonts
            "Stylesheet.fills",  # fills
            "Stylesheet.borders",  # borders
            "Stylesheet.dxfs",  # dxfs
            "NumberFormatList.numFmt",  # numFmts
            "CellStyleList.xf",  # cellXfs, cellStyleXfs
            # (their other children; openpyxl fails to load these)
            "CellStyleList.alignment",
            "CellStyleList.protection",
            "_NamedCellStyleList.cellStyle",  # cellStyles
            "TableStyleList.tableStyle",  # tableStyles
            "TableStyle.tableStyleElement",  # tableStyle
            "ColorList.indexedColors",  # indexedColors
            "ColorList.mruColors",  # mruColors
            "GradientFill.stop",  # gradient stops
            "ExtensionList.ext",  # extLst
        }
        guard = _guard()
        covered = set(guard._XLSX_WORKBOOK_LISTS)
        assert lists(WorkbookPackage) == {
            "WorkbookPackage.bookViews",
            "WorkbookPackage.sheets",
            "WorkbookPackage.customWorkbookViews",
            "WorkbookPackage.externalReferences",
            "WorkbookPackage.pivotCaches",
            "SmartTagList.smartTagType",
            "FunctionGroupList.functionGroup",
            "WebPublishObjectList.webPublishObject",
            # (count: the part's 16 MiB ceiling; each value, and the
            # print titles and areas: _check_xlsx_defined_names)
            "DefinedNameList.definedName",
            "ExtensionList.ext",
        }
        assert covered == {
            "bookViews",
            "sheets",
            "customWorkbookViews",
            "externalReferences",
            "pivotCaches",
            "smartTagTypes",
            "functionGroups",
            "webPublishObjects",
        }
        assert set(guard._XLSX_STYLE_LISTS) | set(guard._XLSX_COLOR_LISTS) == {
            "numFmts",
            "fonts",
            "fills",
            "borders",
            "cellStyleXfs",
            "cellXfs",
            "cellStyles",
            "dxfs",
            "tableStyles",
            "indexedColors",
            "mruColors",
        }

    def test_premise_openpyxl_builds_every_entry_and_repeat(self):
        pytest.importorskip("openpyxl")
        styles = _stylesheet().replace(
            b"</styleSheet>",
            b"<dxfs>"
            + b"<dxf><font/><font><b/></font></dxf>" * 3
            + b"</dxfs></styleSheet>",
        )
        workbook = _load_xlsx(_xlsx_with_styles(styles))
        dxfs = workbook._differential_styles.styles
        assert len(dxfs) == 3
        # The repeated font was built and replaced by the last one.
        assert all(dxf.font.b for dxf in dxfs)
        workbook.close()

    @pytest.mark.parametrize("name", sorted(_STYLE_LIST_CASES))
    def test_style_list_accepts_its_cap_and_refuses_one_more(
        self, name, monkeypatch
    ):
        guard = _guard()
        constant, styles = _STYLE_LIST_CASES[name]
        monkeypatch.setattr(guard, constant, 6)
        guard.validate_zip_container(_xlsx_with_styles(styles(6)), ".xlsx")
        _load_xlsx(_xlsx_with_styles(styles(6))).close()
        with pytest.raises(
            guard.DecompressionBombError, match=f"more than 6 {name} entries"
        ):
            guard.validate_zip_container(_xlsx_with_styles(styles(7)), ".xlsx")

    def test_differential_format_elements_are_capped(self, monkeypatch):
        guard = _guard()
        monkeypatch.setattr(guard, "MAX_XLSX_DIFFERENTIAL_FORMAT_ELEMENTS", 9)

        def styles(dxfs: int) -> bytes:
            # Each <dxf> is three elements.
            return _stylesheet().replace(
                b"</styleSheet>",
                b"<dxfs>"
                + b"<dxf><font><b/></font></dxf>" * dxfs
                + b"</dxfs></styleSheet>",
            )

        guard.validate_zip_container(_xlsx_with_styles(styles(3)), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError, match="more than 9 elements"
        ):
            guard.validate_zip_container(_xlsx_with_styles(styles(4)), ".xlsx")

    def test_gradient_stops_are_capped_together(self, monkeypatch):
        guard = _guard()
        monkeypatch.setattr(guard, "MAX_XLSX_GRADIENT_STOPS_TOTAL", 10)

        def styles(in_fills: int, in_dxf: int) -> bytes:
            return _stylesheet(
                fills=_gradient_fill(in_fills - in_fills // 2)
                + _gradient_fill(in_fills // 2)
            ).replace(
                b"</styleSheet>",
                b"<dxfs><dxf>"
                + _gradient_fill(in_dxf).encode()
                + b"</dxf></dxfs></styleSheet>",
            )

        guard.validate_zip_container(_xlsx_with_styles(styles(6, 4)), ".xlsx")
        _load_xlsx(_xlsx_with_styles(styles(6, 4))).close()
        for over in (styles(6, 5), styles(11, 0)):
            with pytest.raises(
                guard.DecompressionBombError,
                match="more than 10 stops together",
            ):
                guard.validate_zip_container(_xlsx_with_styles(over), ".xlsx")
        # One gradient in a differential format is held like one in fills.
        monkeypatch.undo()
        limit = guard.MAX_XLSX_GRADIENT_STOPS
        with pytest.raises(
            guard.DecompressionBombError, match=f"more than {limit} stops"
        ):
            guard.validate_zip_container(
                _xlsx_with_styles(styles(0, limit + 1)), ".xlsx"
            )

    def test_top_level_elements_are_capped(self, monkeypatch):
        from lxml import etree

        guard = _guard()
        base = _stylesheet()
        top = len(etree.fromstring(base))
        monkeypatch.setattr(guard, "MAX_XLSX_TOP_LEVEL_ELEMENTS", top + 3)

        def styles(extra: int) -> bytes:
            return base.replace(
                b"</styleSheet>", b"<colors/>" * extra + b"</styleSheet>"
            )

        guard.validate_zip_container(_xlsx_with_styles(styles(3)), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError,
            match=f"styles part has more than {top + 3} top-level",
        ):
            guard.validate_zip_container(_xlsx_with_styles(styles(4)), ".xlsx")

    @pytest.mark.parametrize(
        "old, new",
        [
            ("<fonts><font/>", "<fonts><font><b/><b/></font>"),
            (
                "<fonts><font/>",
                "<fonts><font><color rgb='FF000000'/><color/></font>",
            ),
            (
                'xfId="0"/>',
                'xfId="0"><alignment/><alignment horizontal="left"/></xf>',
            ),
            (
                "<fills><fill><patternFill/></fill>",
                "<fills><fill><patternFill><fgColor/><fgColor/></patternFill></fill>",
            ),
            ("<borders><border/>", "<borders><border><left/><left/></border>"),
            (
                "</styleSheet>",
                "<dxfs><dxf><font/><font/></dxf></dxfs></styleSheet>",
            ),
            (
                "</styleSheet>",
                "<colors><mruColors/><mruColors/></colors></styleSheet>",
            ),
            (
                "</styleSheet>",
                "<dxfs><dxf><fill><gradientFill><stop position='0'/>"
                "<bottom/><bottom/></gradientFill></fill></dxf></dxfs>"
                "</styleSheet>",
            ),
        ],
    )
    def test_repeated_child_of_an_entry_is_refused(self, old, new):
        guard = _guard()
        base = _stylesheet()
        assert old.encode() in base
        styles = base.replace(old.encode(), new.encode(), 1)
        with pytest.raises(
            guard.DecompressionBombError, match="repeats its .* child"
        ):
            guard.validate_zip_container(_xlsx_with_styles(styles), ".xlsx")

    def test_reported_repeat_shapes_are_refused_quickly(self):
        guard = _guard()
        for styles in (
            _stylesheet().replace(
                b"</styleSheet>",
                b"<dxfs><dxf>"
                + b"<font/>" * 500_000
                + b"</dxf></dxfs></styleSheet>",
            ),
            _stylesheet().replace(
                b"</styleSheet>", b"<colors/>" * 500_000 + b"</styleSheet>"
            ),
        ):
            start = time.perf_counter()
            with pytest.raises(guard.DecompressionBombError):
                guard.validate_zip_container(_xlsx_with_styles(styles), ".xlsx")
            assert time.perf_counter() - start < 5

    def test_excel_like_lists_and_extensions_pass(self):
        """Excel's own shapes: repeated children inside an extension
        (which openpyxl does not read), several stops in a gradient, a
        table style of several elements and differential formats with a
        font, fill and border each."""
        guard = _guard()
        styles = _stylesheet(
            fills="<fill><patternFill/></fill>" + _gradient_fill(3)
        ).replace(
            b"</styleSheet>",
            b"<dxfs><dxf><font><b/><color rgb='FF9C0006'/></font><fill>"
            b"<patternFill><bgColor rgb='FFFFC7CE'/></patternFill></fill>"
            b"<border><left/><right/></border></dxf><dxf><font><i/></font>"
            b"</dxf></dxfs><tableStyles count='1'><tableStyle name='t'>"
            b"<tableStyleElement type='wholeTable' dxfId='0'/>"
            b"<tableStyleElement type='headerRow' dxfId='1'/></tableStyle>"
            b"</tableStyles><colors><mruColors><color rgb='FF00B050'/>"
            b"<color rgb='FF0070C0'/></mruColors></colors><extLst>"
            b"<ext uri='{EB79DEF2-80B8-43e5-95BD-54CBDDF9F1E1}' "
            b"xmlns:x14='urn:x14'><x14:slicerStyles><x14:slicerStyle "
            b"name='a'/><x14:slicerStyle name='b'/></x14:slicerStyles>"
            b"</ext><ext uri='u2'/></extLst></styleSheet>",
        )
        data = _xlsx_with_styles(styles)
        guard.validate_zip_container(data, ".xlsx")
        workbook = _load_xlsx(data)
        assert len(workbook._differential_styles.styles) == 2
        workbook.close()

    @pytest.mark.parametrize(
        "name, entry",
        [
            ("bookViews", "<workbookView/>"),
            ("customWorkbookViews", "<customWorkbookView/>"),
            ("externalReferences", "<externalReference/>"),
            ("pivotCaches", "<pivotCache/>"),
            ("smartTagTypes", "<smartTagType/>"),
            ("functionGroups", "<functionGroup/>"),
            ("webPublishObjects", "<webPublishObject/>"),
            ("extLst", "<ext uri='u'/>"),
        ],
    )
    def test_workbook_list_accepts_its_cap_and_refuses_one_more(
        self, name, entry, monkeypatch
    ):
        guard = _guard()
        constant = (
            "MAX_XLSX_EXTENSIONS"
            if name == "extLst"
            else "MAX_XLSX_WORKBOOK_LIST_ENTRIES"
        )
        monkeypatch.setattr(guard, constant, 6)
        from lxml import etree

        with zipfile.ZipFile(io.BytesIO(_real_package(".xlsx"))) as archive:
            root = etree.fromstring(archive.read("xl/workbook.xml"))
        listed = sum(len(c) for c in root if etree.QName(c).localname == name)

        def workbook(count: int) -> bytes:
            # Over two lists of that name, beside those openpyxl writes.
            count -= listed
            body = "".join(
                f"<{name}>{entry * n}</{name}>"
                for n in (count - count // 2, count // 2)
            ).encode()
            return _xlsx_with_workbook(lambda xml: _workbook_lists(xml, body))

        # The guard's own workbook checks (openpyxl may fail to load
        # these placeholder entries later; that is refused as malformed).
        with pytest.raises(
            guard.DecompressionBombError, match=f"more than 6 {name} entries"
        ):
            guard.validate_zip_container(workbook(7), ".xlsx")
        try:
            guard.validate_zip_container(workbook(6), ".xlsx")
        except guard.DecompressionBombError as exc:
            assert str(exc) == "malformed spreadsheet package", exc

    def test_sheets_without_ids_count_towards_the_sheet_cap(self, monkeypatch):
        guard = _guard()
        monkeypatch.setattr(guard, "MAX_XLSX_SHEETS", 5)

        def workbook(extra: int) -> bytes:
            return _xlsx_with_workbook(
                lambda xml: xml.replace(
                    b"</sheets>",
                    b"".join(
                        b"<sheet name='n%d' sheetId='%d'/>" % (k, k + 2)
                        for k in range(extra)
                    )
                    + b"</sheets>",
                    1,
                )
            )

        guard.validate_zip_container(workbook(4), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError, match="more than 5 sheets"
        ):
            guard.validate_zip_container(workbook(5), ".xlsx")

    def test_workbook_top_level_and_repeats(self, monkeypatch):
        guard = _guard()
        base = _real_package(".xlsx")
        with zipfile.ZipFile(io.BytesIO(base)) as archive:
            xml = archive.read("xl/workbook.xml")
        from lxml import etree

        with zipfile.ZipFile(io.BytesIO(base)) as archive:
            styles_top = len(etree.fromstring(archive.read("xl/styles.xml")))
        top = len(etree.fromstring(xml))
        # The cap is shared with the styles part, which must stay under it.
        cap = max(top, styles_top) + 2
        monkeypatch.setattr(guard, "MAX_XLSX_TOP_LEVEL_ELEMENTS", cap)

        def extra(count: int) -> bytes:
            return _xlsx_with_workbook(
                lambda x: x.replace(
                    b"</workbook>", b"<calcPr/>" * count + b"</workbook>"
                )
            )

        # Repeated top-level elements are allowed up to the cap.
        guard.validate_zip_container(extra(cap - top), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError,
            match=f"workbook part has more than {cap} top-level",
        ):
            guard.validate_zip_container(extra(cap - top + 1), ".xlsx")
        monkeypatch.undo()
        for change in (
            # A repeated child of a list entry, or of a top-level element.
            lambda x: x.replace(
                b"<workbookView ",
                b"<workbookView><extLst/><extLst/></workbookView><workbookView ",
                1,
            ),
            lambda x: x.replace(
                b"<workbookPr/>", b"<workbookPr><a/><a/></workbookPr>", 1
            ),
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="repeats its .* child"
            ):
                guard.validate_zip_container(
                    _xlsx_with_workbook(change), ".xlsx"
                )
        start = time.perf_counter()
        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(extra(1_000_000), ".xlsx")
        assert time.perf_counter() - start < 5

    def test_excel_like_workbook_passes(self):
        """Excel writes repeated children inside its calculation-feature
        extension, which openpyxl does not read."""
        guard = _guard()
        data = _xlsx_with_workbook(
            lambda xml: xml.replace(
                b"</workbook>",
                b"<extLst><ext uri='{140A7094-0E35-4892-8432-C4D2E57EDEB5}'"
                b" xmlns:x15='urn:x15'><x15:workbookPr chartTrackingRefBase='1'/>"
                b"</ext><ext uri='{B58B0392-4F1F-4190-BB64-5DF3571DCE5F}' "
                b"xmlns:xcalcf='urn:xcalcf'><xcalcf:calcFeatures>"
                b"<xcalcf:feature name='microsoft.com:RD'/>"
                b"<xcalcf:feature name='microsoft.com:Single'/>"
                b"</xcalcf:calcFeatures></ext></extLst></workbook>",
            )
        )
        guard.validate_zip_container(data, ".xlsx")
        workbook = _load_xlsx(data)
        assert workbook.sheetnames == ["Sheet"]
        workbook.close()

    def test_new_ceilings_are_generous_and_pinned(self):
        guard = _guard()
        assert guard.MAX_XLSX_NUMBER_FORMATS >= 250  # Excel's own limit
        assert (
            guard.MAX_XLSX_DIFFERENTIAL_FORMATS == guard.MAX_XLSX_CELL_FORMATS
        )
        assert guard.MAX_XLSX_COLOR_LIST_ENTRIES >= 64  # Excel's palette
        assert guard.MAX_XLSX_TABLE_STYLE_ELEMENTS >= 28  # table parts
        assert guard.MAX_XLSX_TABLE_STYLES >= 1_024
        assert guard.MAX_XLSX_EXTENSIONS >= 2 * guard.MAX_XLSX_CELL_FORMATS
        assert guard.MAX_XLSX_WORKBOOK_LIST_ENTRIES >= 1_024
        assert guard.MAX_XLSX_TOP_LEVEL_ELEMENTS >= 64
        assert guard.MAX_XLSX_GRADIENT_STOPS_TOTAL >= 256 * 256
        assert guard.MAX_XLSX_DIFFERENTIAL_FORMAT_ELEMENTS <= 500_000

    def test_document_with_a_comment_before_its_root_is_read(self):
        """One comment before document.xml's root (a generator's note)
        is outside python-docx's tree, and the guard accepts it."""
        docx = pytest.importorskip("docx")
        guard = _guard()
        base = docx.Document()
        base.add_paragraph("hello")
        buf = io.BytesIO()
        base.save(buf)
        data = _rewrite_zip(
            buf.getvalue(),
            {"word/document.xml": lambda xml: _before_root(xml, b"<!--g-->")},
        )
        guard.validate_zip_container(data, ".docx")
        assert [p.text for p in docx.Document(io.BytesIO(data)).paragraphs] == [
            "hello"
        ]


def _with_defined_names(entries: str) -> bytes:
    """An openpyxl-written workbook with *entries* as its defined names."""
    return _xlsx_with_workbook(
        lambda xml: xml.replace(
            b"<definedNames/>",
            b"<definedNames>" + entries.encode() + b"</definedNames>",
            1,
        )
    )


def _defined_name(name: str, value: str, sheet: int | None = 0) -> str:
    local = "" if sheet is None else f' localSheetId="{sheet}"'
    return f'<definedName name="{name}"{local}>{value}</definedName>'


def _stops(count: int) -> str:
    return "".join(
        f'<stop position="{n / max(count, 1):.6f}"><color rgb="FF000000"/></stop>'
        for n in range(count)
    )


def _with_sheet_names(*names: str) -> bytes:
    """An openpyxl-written workbook of one sheet per name (each with its
    part), the names written into the workbook part as given."""
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    workbook.active.title = "s0"
    for n in range(1, len(names)):
        workbook.create_sheet(f"s{n}")
    buf = io.BytesIO()
    workbook.save(buf)

    def rename(xml: bytes) -> bytes:
        for n, name in enumerate(names):
            xml = xml.replace(
                f'name="s{n}"'.encode(), f'name="{name}"'.encode(), 1
            )
        return xml

    return _rewrite_zip(buf.getvalue(), {"xl/workbook.xml": rename})


def _with_sheet_root(change) -> bytes:
    """An openpyxl-written workbook whose first worksheet's root start
    tag (bytes up to its ``>``) is passed through *change*."""

    def sheet(xml: bytes) -> bytes:
        start = xml.index(b"<worksheet")
        end = xml.index(b">", start)
        return xml[:start] + change(xml[start:end]) + xml[end:]

    return _rewrite_zip(
        _real_package(".xlsx"), {"xl/worksheets/sheet1.xml": sheet}
    )


class TestXlsxReadPathStringWork:
    # Text openpyxl's read path (pandas' read-only load) processes in
    # time worse than linear in its length: the print titles and print
    # areas among the defined names (regular expressions, quadratic),
    # sheet names (compared with every other name), and namespace URIs
    # (copied into every element name lxml hands back). Every other
    # regular expression on that path is anchored or bounded and runs in
    # linear time.

    def test_premise_print_names_are_parsed_on_every_load(self):
        import inspect

        pytest.importorskip("openpyxl")
        from openpyxl.reader.excel import ExcelReader
        from openpyxl.reader.workbook import WorkbookParser
        from openpyxl.workbook.defined_name import DefinedName
        from openpyxl.worksheet.print_settings import PrintArea

        assert "assign_names()" in inspect.getsource(ExcelReader.read)
        source = inspect.getsource(WorkbookParser.assign_names)
        assert "PrintTitles.from_string(defn.value)" in source
        assert "PrintArea.from_string(defn.value)" in source
        # The reserved names are matched as a prefix.
        assert DefinedName(name="_xlnm.Print_AreaZ").is_reserved == "Print_Area"
        # So one sheet's distinct spellings are each parsed.
        data = _with_defined_names(
            _defined_name("_xlnm.Print_Area", "s!$A$1:$B$2")
            + _defined_name("_xlnm.Print_AreaZ", "s!$A$1:$B$3")
        )
        calls: list = []
        real = PrintArea.from_string.__func__

        def spy(cls, value):
            calls.append(value)
            return real(cls, value)

        with patch.object(PrintArea, "from_string", classmethod(spy)):
            _load_xlsx(data).close()
        assert sorted(calls) == ["s!$A$1:$B$2", "s!$A$1:$B$3"]

    def test_premise_print_name_parse_is_quadratic(self):
        pytest.importorskip("openpyxl")
        from openpyxl.worksheet.print_settings import PRINT_AREA_RE

        def cost(length: int) -> float:
            best = math.inf
            for _ in range(3):
                start = time.perf_counter()
                list(PRINT_AREA_RE.finditer("'" * length))
                best = min(best, time.perf_counter() - start)
            return best

        # Four times the length, about sixteen times the time.
        assert cost(4_000) > 6 * cost(1_000)

    def test_defined_name_value_is_capped(self):
        guard = _guard()
        guard.validate_zip_container(
            _with_defined_names(_defined_name("n", "1" * 8_192, None)),
            ".xlsx",
        )
        for entry in (
            _defined_name("n", "1" * 8_193, None),
            # A child's text, which openpyxl reads in place of a field.
            f'<definedName name="n"><comment>{"1" * 8_193}</comment>'
            "</definedName>",
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="longer than 8192"
            ):
                guard.validate_zip_container(
                    _with_defined_names(entry), ".xlsx"
                )

    def test_premise_from_tree_reads_the_value_from_an_attribute(self):
        pytest.importorskip("openpyxl")
        from lxml import etree
        from openpyxl.workbook.defined_name import DefinedName

        for attribute in ("attr_text", "attr-text"):
            node = etree.fromstring(
                f'<definedName name="n" {attribute}="Sheet!$A$1"/>'
            )
            assert DefinedName.from_tree(node).value == "Sheet!$A$1"
        # Text, when there is any, replaces the attribute.
        node = etree.fromstring(
            '<definedName name="n" attr_text="a">Sheet!$A$1</definedName>'
        )
        assert DefinedName.from_tree(node).value == "Sheet!$A$1"

    @pytest.mark.parametrize("attribute", ["attr_text", "attr-text"])
    @pytest.mark.parametrize(
        "name", ["_xlnm.Print_Area", "_xlnm.Print_Titles", "n"]
    )
    def test_defined_name_value_in_an_attribute_is_measured(
        self, name, attribute
    ):
        guard = _guard()

        def entry(length: int) -> str:
            return (
                f'<definedName name="{name}" localSheetId="0" '
                f'{attribute}="{"#" * length}"/>'
            )

        guard.validate_zip_container(_with_defined_names(entry(100)), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError, match="longer than 8192"
        ):
            guard.validate_zip_container(
                _with_defined_names(entry(8_193)), ".xlsx"
            )

    def test_print_names_in_attributes_are_priced(self, monkeypatch):
        guard = _guard()
        base = guard.XLSX_PRINT_NAME_BASE_UNITS
        monkeypatch.setattr(
            guard, "MAX_XLSX_PRINT_NAME_WORK", 2 * (100**2 + base)
        )
        value = "#" * 100
        names = (
            f'<definedName name="_xlnm.Print_Area" localSheetId="0" '
            f'attr_text="{value}"/>'
        )
        guard.validate_zip_container(
            _with_defined_names(names + names.replace("Area", "Titles")),
            ".xlsx",
        )
        with pytest.raises(
            guard.DecompressionBombError, match="print areas and titles"
        ):
            guard.validate_zip_container(
                _with_defined_names(
                    names
                    + names.replace("Area", "Titles")
                    + names.replace("Area", "Area2")
                ),
                ".xlsx",
            )

    def test_reported_attribute_print_name_shape_is_refused_quickly(self):
        guard = _guard()
        for name in ("_xlnm.Print_Titles", "_xlnm.Print_Area"):
            data = _with_defined_names(
                f'<definedName name="{name}" localSheetId="0" '
                f'attr_text="{"#" * 40_000}"/>'
            )
            assert len(data) < 8_000
            start = time.perf_counter()
            with pytest.raises(guard.DecompressionBombError):
                guard.validate_zip_container(data, ".xlsx")
            assert time.perf_counter() - start < 5

    def test_reported_print_name_shapes_are_refused_quickly(self):
        guard = _guard()
        for name in ("_xlnm.Print_Titles", "_xlnm.Print_Area"):
            data = _with_defined_names(_defined_name(name, "#" * 40_000))
            assert len(data) < 8_000
            start = time.perf_counter()
            with pytest.raises(guard.DecompressionBombError):
                guard.validate_zip_container(data, ".xlsx")
            assert time.perf_counter() - start < 5

    def test_print_names_are_priced_together(self, monkeypatch):
        guard = _guard()
        base = guard.XLSX_PRINT_NAME_BASE_UNITS
        monkeypatch.setattr(
            guard, "MAX_XLSX_PRINT_NAME_WORK", 3 * (100**2 + base)
        )
        value = "#" * 100
        names = (
            _defined_name("_xlnm.Print_Area", value)
            + _defined_name("_xlnm.Print_Titles", value, 1)
            # Spelled through a child element, which openpyxl reads in
            # place of the attribute.
            + f'<definedName localSheetId="0">{value}'
            "<name>_xlnm.Print_AreaZ</name></definedName>"
            # Other names are not parsed and not counted.
             + _defined_name("n", "#" * 8_000) * 20
        )
        guard.validate_zip_container(_with_defined_names(names), ".xlsx")
        for extra in (
            _defined_name("_xlnm.Print_Area2", value),
            _defined_name("_xlnm.Print_Titles", "#", 3),
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="print areas and titles"
            ):
                guard.validate_zip_container(
                    _with_defined_names(names + extra), ".xlsx"
                )

    def test_print_name_at_the_length_ceiling_is_accepted(self):
        guard = _guard()
        limit = guard.MAX_XLSX_DEFINED_NAME_CHARS
        assert guard.MAX_XLSX_PRINT_NAME_WORK == (
            limit**2 + guard.XLSX_PRINT_NAME_BASE_UNITS
        )
        guard.validate_zip_container(
            _with_defined_names(_defined_name("_xlnm.Print_Area", "'" * limit)),
            ".xlsx",
        )
        with pytest.raises(
            guard.DecompressionBombError, match="print areas and titles"
        ):
            guard.validate_zip_container(
                _with_defined_names(
                    _defined_name("_xlnm.Print_Area", "'" * limit)
                    + _defined_name("_xlnm.Print_Titles", "")
                ),
                ".xlsx",
            )

    def test_excel_like_print_names_pass(self):
        guard = _guard()
        data = _with_defined_names(
            _defined_name("_xlnm.Print_Area", "Sheet!$A$1:$H$50")
            + _defined_name("_xlnm.Print_Titles", "Sheet!$1:$2,Sheet!$A:$A")
            + _defined_name("Rates", "Sheet!$B$2:$B$9", None)
        )
        guard.validate_zip_container(data, ".xlsx")
        workbook = _load_xlsx(data)
        assert workbook.sheetnames == ["Sheet"]
        workbook.close()

    def test_premise_chartsheets_are_renamed_by_a_search_of_every_name(self):
        import inspect

        pytest.importorskip("openpyxl")
        from openpyxl.workbook.child import avoid_duplicate_name

        source = inspect.getsource(avoid_duplicate_name)
        assert "n.lower() == value.lower()" in source
        assert '",".join(names)' in source and "findall" in source

    def test_sheet_name_length_is_capped(self):
        guard = _guard()
        guard.validate_zip_container(_with_sheet_names("a" * 255), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError, match="longer than 255"
        ):
            guard.validate_zip_container(_with_sheet_names("a" * 256), ".xlsx")

    def test_sheets_of_one_name_are_refused(self):
        guard = _guard()
        guard.validate_zip_container(
            _with_sheet_names("Data", "Data2"), ".xlsx"
        )
        for data in (
            _with_sheet_names("Data", "DATA"),
            _with_sheet_names("Straße", "STRASSE".lower().replace("ss", "ß")),
            # A sheet without a relationship counts too.
            _xlsx_with_workbook(
                lambda xml: xml.replace(
                    b"</sheets>", b'<sheet name="sheet" sheetId="9"/></sheets>'
                )
            ),
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="two sheets of one name"
            ):
                guard.validate_zip_container(data, ".xlsx")

    def test_premise_lxml_copies_the_uri_into_every_name(self):
        from lxml import etree

        uri = "urn:" + "a" * 100
        root = etree.fromstring(f'<r xmlns:q="{uri}"><q:z/><q:z/></r>')
        first, second = (child.tag for child in root)
        assert first == second == "{%s}z" % uri
        assert first is not second  # built anew on every read

    def test_long_namespace_uri_is_refused_in_any_part_openpyxl_parses(self):
        guard = _guard()

        def declared(length: int):
            uri = ("urn:" + "a" * length)[:length]
            return f' xmlns:q="{uri}"'.encode()

        guard.validate_zip_container(
            _with_sheet_root(lambda tag: tag + declared(1_024)), ".xlsx"
        )
        for data in (
            _with_sheet_root(lambda tag: tag + declared(1_025)),
            # Declared deeper, in a part openpyxl reads only on load.
            _rewrite_zip(
                _real_package(".xlsx"),
                {
                    "xl/worksheets/sheet1.xml": lambda xml: xml.replace(
                        b"</worksheet>",
                        b"<q:x" + declared(1_025) + b"/></worksheet>",
                    )
                },
            ),
            # The document properties openpyxl parses with lxml.
            _rewrite_zip(
                _real_package(".xlsx"),
                {
                    "docProps/core.xml": lambda xml: xml.replace(
                        b"<cp:coreProperties ",
                        b"<cp:coreProperties" + declared(1_025) + b" ",
                        1,
                    )
                },
            ),
        ):
            with pytest.raises(
                guard.DecompressionBombError, match="namespace URI"
            ):
                guard.validate_zip_container(data, ".xlsx")
        # docProps/app.xml is never parsed (not by openpyxl, pandas,
        # unstructured or msoffcrypto), so it is not scanned.
        guard.validate_zip_container(
            _rewrite_zip(
                _real_package(".xlsx"),
                {
                    "docProps/app.xml": lambda xml: xml.replace(
                        b"<Properties ",
                        b"<Properties" + declared(1_025) + b" ",
                        1,
                    )
                },
            ),
            ".xlsx",
        )

    def test_reported_namespace_shape_is_refused_quickly(self):
        guard = _guard()
        data = _rewrite_zip(
            _real_package(".xlsx"),
            {
                "xl/worksheets/sheet1.xml": lambda xml: xml.replace(
                    b"<worksheet ",
                    b'<worksheet xmlns:q="urn:' + b"a" * 1_000_000 + b'" ',
                    1,
                ).replace(b"</worksheet>", b"<q:z/>" * 2_000 + b"</worksheet>")
            },
        )
        assert len(data) < 10_000
        start = time.perf_counter()
        with pytest.raises(guard.DecompressionBombError, match="namespace"):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    def test_premise_libxml2_stops_where_expat_reads_on(self):
        from lxml import etree

        # openpyxl's iterparse when defusedxml is installed (expat).
        defused = pytest.importorskip("defusedxml.ElementTree")

        for document in (
            b"<r><" + b"n" * 50_001 + b"/></r>",
            b'<r a="' + b"v" * 10_500_000 + b'"/>',
            b"<r><!--" + b"c" * 10_500_000 + b"--></r>",
        ):
            with pytest.raises(etree.XMLSyntaxError):
                etree.fromstring(document)
            # openpyxl reads worksheets and shared strings this way.
            list(defused.iterparse(io.BytesIO(document)))

    @pytest.mark.parametrize(
        "lead",
        [
            b"<" + b"n" * 50_001 + b"/>",
            b'<lead a="' + b"v" * 10_500_000 + b'"/>',
            b"<!--" + b"c" * 10_500_000 + b"-->",
        ],
        ids=["long-name", "long-attribute", "long-comment"],
    )
    def test_uri_after_what_only_expat_reads_is_refused(self, lead):
        """openpyxl's expat parse of a worksheet reads on past markup
        where libxml2 stops, so a declaration after it is checked."""
        guard = _guard()

        def sheet(uri_chars: int, elements: int):
            uri = b"urn:" + b"a" * (uri_chars - 4)
            tail = (
                lead
                + b'<q:w xmlns:q="'
                + uri
                + b'">'
                + b"<q:z/>" * elements
                + b"</q:w></worksheet>"
            )
            return _rewrite_zip(
                _real_package(".xlsx"),
                {
                    "xl/worksheets/sheet1.xml": lambda xml: xml.replace(
                        b"</worksheet>", tail
                    )
                },
            )

        guard.validate_zip_container(sheet(1_024, 10), ".xlsx")
        with pytest.raises(guard.DecompressionBombError, match="namespace"):
            guard.validate_zip_container(sheet(1_025, 10), ".xlsx")
        # The reported shape: a 1,000,000-character URI used 40,000
        # times.
        data = sheet(1_000_000, 40_000)
        start = time.perf_counter()
        with pytest.raises(guard.DecompressionBombError, match="namespace"):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    def test_dtd_in_a_worksheet_is_refused(self):
        guard = _guard()
        data = _rewrite_zip(
            _real_package(".xlsx"),
            {
                "xl/worksheets/sheet1.xml": lambda xml: xml.replace(
                    b"<worksheet ", b"<!DOCTYPE worksheet><worksheet ", 1
                )
            },
        )
        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(data, ".xlsx")

    def test_namespace_scan_reads_each_parsed_member_once(self):
        """The scan streams each part openpyxl's load parses once, with
        that load's parser family, and nothing else (the theme, whose
        bytes openpyxl keeps unparsed, the package relationships,
        ``docProps/app.xml``, media)."""
        guard = _guard()
        data = _rewrite_zip(
            _real_package(".xlsx"), add={"xl/media/image1.png": _small_png()}
        )
        scanned: list = []
        real_scan = guard._scan_xlsx_part

        def spy(archive, info, families):
            scanned.append((info.filename, families))
            return real_scan(archive, info, families)

        opened: list = []
        real_open = zipfile.ZipFile.open

        def open_spy(self, name, *args, **kwargs):
            opened.append(
                name.filename if isinstance(name, zipfile.ZipInfo) else name
            )
            return real_open(self, name, *args, **kwargs)

        with (
            patch.object(guard, "_scan_xlsx_part", spy),
            patch.object(zipfile.ZipFile, "open", open_spy),
        ):
            guard.validate_zip_container(data, ".xlsx")
        lxml, expat = guard._LXML, guard._EXPAT
        assert sorted(scanned) == sorted(
            [
                ("[Content_Types].xml", frozenset({lxml, expat})),
                ("xl/workbook.xml", frozenset({lxml})),
                ("xl/_rels/workbook.xml.rels", frozenset({lxml})),
                ("docProps/core.xml", frozenset({lxml})),
                ("xl/styles.xml", frozenset({lxml})),
                ("xl/worksheets/sheet1.xml", frozenset({expat})),
            ]
        )
        for name in (
            "xl/theme/theme1.xml",
            "_rels/.rels",
            "docProps/app.xml",
            "xl/media/image1.png",
        ):
            assert name not in opened

    def test_premise_an_extension_list_entry_is_read_as_a_fill(self):
        """openpyxl reads every child of ``fills`` as a fill, and a
        fill's first child as a gradient unless it is a pattern, so an
        ``extLst`` there is not an extension list."""
        pytest.importorskip("openpyxl")
        for fills in (
            "<fill><patternFill/></fill><extLst><ext>"
            + _stops(5)
            + "</ext></extLst>",
            "<fill><patternFill/></fill><fill><extLst>"
            + _stops(5)
            + "</extLst></fill>",
        ):
            workbook = _load_xlsx(_xlsx_with_styles(_stylesheet(fills=fills)))
            assert [len(fill.stop) for fill in workbook._fills[1:]] == [5]
            workbook.close()

    @pytest.mark.parametrize(
        "fills, dxfs",
        [
            ("<fill><patternFill/></fill><extLst><ext>{}</ext></extLst>", ""),
            ("<fill><patternFill/></fill><fill><extLst>{}</extLst></fill>", ""),
            (
                "<fill><patternFill/></fill>",
                "<dxfs><dxf><fill><extLst>{}</extLst></fill></dxf></dxfs>",
            ),
        ],
    )
    def test_gradient_stops_under_an_extension_tag_are_counted(
        self, fills, dxfs
    ):
        guard = _guard()

        def styles(stops: int) -> bytes:
            return _stylesheet(fills=fills.format(_stops(stops))).replace(
                b"</styleSheet>",
                dxfs.format(_stops(stops)).encode() + b"</styleSheet>",
            )

        guard.validate_zip_container(_xlsx_with_styles(styles(256)), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError, match="more than 256 stops"
        ):
            guard.validate_zip_container(
                _xlsx_with_styles(styles(257)), ".xlsx"
            )

    def test_reported_extension_stop_shape_is_refused_quickly(self):
        guard = _guard()
        data = _xlsx_with_styles(
            _stylesheet(
                fills="<fill><patternFill/></fill><extLst><ext>"
                + _stops(40_000)
                + "</ext></extLst>"
            )
        )
        start = time.perf_counter()
        with pytest.raises(guard.DecompressionBombError, match="stops"):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    @pytest.mark.parametrize(
        "old, new, held",
        [
            (
                "<fills><fill><patternFill/></fill>",
                "<fills><fill><patternFill/><extLst><ext>{}</ext></extLst>"
                "</fill>",
                4,
            ),
            (
                'xfId="0"/>',
                'xfId="0"><extLst><ext uri="u">{}</ext></extLst></xf>',
                3,
            ),
            ("<fonts><font/>", "<fonts><font>{}</font>", 1),
        ],
    )
    def test_style_entry_elements_are_capped(self, old, new, held, monkeypatch):
        """*held* elements besides the filler: the entry and the
        elements around the filler."""
        guard = _guard()
        base = _stylesheet()
        assert old.encode() in base

        def styles(elements: int) -> bytes:
            # Distinct names: a repeat would be refused on its own.
            filler = "".join(f"<x{n}/>" for n in range(elements - held))
            return base.replace(old.encode(), new.format(filler).encode(), 1)

        monkeypatch.setattr(guard, "MAX_XLSX_STYLE_ENTRY_ELEMENTS", 12)
        guard.validate_zip_container(_xlsx_with_styles(styles(12)), ".xlsx")
        with pytest.raises(
            guard.DecompressionBombError, match="entry has more than 12"
        ):
            guard.validate_zip_container(_xlsx_with_styles(styles(13)), ".xlsx")

    def test_large_style_entry_is_refused_quickly(self):
        guard = _guard()
        data = _xlsx_with_styles(
            _stylesheet(
                fills="<fill><patternFill/><extLst><ext uri='u'>"
                + "<x/>" * 80_000
                + "</ext></extLst></fill>"
            )
        )
        start = time.perf_counter()
        with pytest.raises(guard.DecompressionBombError, match="elements"):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    def test_excel_cell_format_extensions_pass(self):
        guard = _guard()
        styles = _stylesheet().replace(
            b'xfId="0"/>',
            b'xfId="0"><alignment horizontal="left"/><protection locked="0"/>'
            b'<extLst><ext uri="{C7286773-470A-42A8-94C5-96B5CB345126}" '
            b'xmlns:xfpb="urn:xfpb"><xfpb:xfComplement i="0"/></ext>'
            b"</extLst></xf>",
            1,
        )
        guard.validate_zip_container(_xlsx_with_styles(styles), ".xlsx")

    def test_new_ceilings_are_pinned(self):
        guard = _guard()
        assert guard.MAX_XLSX_DEFINED_NAME_CHARS == 8_192  # Excel's limit
        assert guard.MAX_XLSX_SHEET_NAME_CHARS >= 100  # Google Sheets
        assert guard.MAX_XLSX_NAMESPACE_URI_CHARS == 1_024
        assert (
            guard.MAX_XLSX_STYLE_ENTRY_ELEMENTS
            >= 2 * guard.MAX_XLSX_GRADIENT_STOPS + 2
        )


def _with_part(name: str, change) -> bytes:
    """An openpyxl-written workbook with member *name* passed through
    *change*."""
    return _rewrite_zip(_real_package(".xlsx"), {name: change})


def _reencoded(
    xml: bytes, codec: str, declared: str, bom: bytes = b""
) -> bytes:
    """*xml* (UTF-8, with or without a declaration) re-encoded in
    *codec* under a declaration naming *declared*."""
    text = xml.decode("utf-8")
    if text.startswith("<?xml"):
        text = text[text.index("?>") + 2 :]
    declaration = f'<?xml version="1.0" encoding="{declared}"?>'
    return bom + (declaration + text).encode(codec)


def _without_declaration(xml: bytes) -> bytes:
    """*xml* without its XML declaration, if it has one."""
    if xml.startswith(b"<?xml"):
        xml = xml[xml.index(b"?>") + 2 :]
    return xml.lstrip(b"\r\n")


def _utf16_lt_names(count: int) -> list[str]:
    """*count* distinct attribute names of three CJK ideographs or
    Hangul syllables each, every one with a UTF-16 low byte of ``3C``
    (a ``<`` byte in either byte order)."""
    letters = [
        chr(c)
        for c in [*range(0x4E3C, 0x9FA5, 0x100), *range(0xAC3C, 0xD7A3, 0x100)]
    ]
    names = []
    for first in letters:
        for second in letters:
            for third in letters:
                names.append(first + second + third)
                if len(names) == count:
                    return names
    raise AssertionError("not enough names")


def _utf16_leading_space(xml: str, codec: str) -> bytes:
    """*xml* in UTF-16 (*codec*), no byte-order mark, after a space."""
    return (" " + xml).encode(codec)


def _declared_utf16le(xml: str) -> bytes:
    """An ASCII declaration naming ``UTF-16LE``, then *xml* in UTF-16LE
    (the declaration's ``?>`` included)."""
    return b'<?xml version="1.0" encoding="UTF-16LE"' + ("?>" + xml).encode(
        "utf-16-le"
    )


def _with_shared_strings(xml: bytes) -> bytes:
    """An openpyxl-written workbook with *xml* as its shared strings
    part ``xl/sharedStrings.xml``."""

    def types(ct: bytes) -> bytes:
        return ct.replace(
            b"</Types>",
            b'<Override PartName="/xl/sharedStrings.xml" ContentType='
            b'"application/vnd.openxmlformats-officedocument.'
            b'spreadsheetml.sharedStrings+xml"/></Types>',
        )

    return _rewrite_zip(
        _real_package(".xlsx"),
        {"[Content_Types].xml": types},
        add={"xl/sharedStrings.xml": xml},
    )


def _long_uri_core(elements: int, uri_chars: int = 1_000_000) -> str:
    return (
        "<cp:coreProperties xmlns:cp="
        '"http://schemas.openxmlformats.org/package/2006/metadata/'
        'core-properties" xmlns:q="urn:'
        + "a" * uri_chars
        + '">'
        + "<q:z/>" * elements
        + "</cp:coreProperties>"
    )


def _long_uri_custom(elements: int, uri_chars: int = 1_000_000) -> str:
    return (
        "<Properties xmlns="
        '"http://schemas.openxmlformats.org/officeDocument/2006/'
        'custom-properties" xmlns:q="urn:'
        + "a" * uri_chars
        + '">'
        + "<q:z/>" * elements
        + "</Properties>"
    )


def _with_custom_properties(xml: bytes) -> bytes:
    def types(ct: bytes) -> bytes:
        return ct.replace(
            b"</Types>",
            b'<Override PartName="/docProps/custom.xml" ContentType='
            b'"application/vnd.openxmlformats-officedocument.custom-'
            b'properties+xml"/></Types>',
        )

    return _rewrite_zip(
        _real_package(".xlsx"),
        {"[Content_Types].xml": types},
        add={"docProps/custom.xml": xml},
    )


_SVG_WITH_DOCTYPE = (
    b'<?xml version="1.0" encoding="UTF-8"?>\n'
    b'<!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN" '
    b'"http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd">\n'
    b'<svg xmlns="http://www.w3.org/2000/svg" version="1.1" width="2" '
    b'height="2"><rect width="2" height="2"/></svg>'
)


class TestXlsxScanFollowsTheLoaders:
    """The namespace scan covers exactly the parts openpyxl's read-only
    load (and msoffcrypto) parses, each with the parser that load uses
    for it, and refuses encodings, nesting and start tags the loaders'
    parsers handle differently or without a bound."""

    def test_premise_lxml_reads_utf32_from_a_buffer_only(self):
        """openpyxl parses the document properties with lxml
        ``fromstring``, which decodes UTF-32 with a byte-order mark;
        lxml's push parser and expat reject it at the first byte."""
        from lxml import etree
        from xml.parsers import expat

        document = '<?xml version="1.0" encoding="UTF-32"?><r/>'.encode(
            "utf-32"
        )
        etree.fromstring(document, etree.XMLParser(resolve_entities=False))
        pushed = etree.XMLParser(resolve_entities=False)
        with pytest.raises(etree.XMLSyntaxError):
            pushed.feed(document)
            pushed.close()
        with pytest.raises(expat.ExpatError):
            expat.ParserCreate(namespace_separator="}").Parse(document, True)

    @pytest.mark.parametrize("part", ["core", "custom"])
    def test_utf32_properties_with_a_long_uri_are_refused(self, part):
        guard = _guard()

        def data(elements: int, uri_chars: int = 1_000_000) -> bytes:
            if part == "core":
                xml = _long_uri_core(elements, uri_chars)
            else:
                xml = _long_uri_custom(elements, uri_chars)
            body = ('<?xml version="1.0" encoding="UTF-32"?>' + xml).encode(
                "utf-32"
            )
            if part == "core":
                return _with_part("docProps/core.xml", lambda _xml: body)
            return _with_custom_properties(body)

        # openpyxl's load reads the part (a short URI, to stay quick).
        workbook = _load_xlsx(data(10, uri_chars=10))
        workbook.close()
        # The reported shape: a 1,000,000-character URI over 20,000
        # elements, about 10 KB, which took openpyxl 13 s or more to load.
        reported = data(20_000)
        assert len(reported) < 20_000
        start = time.perf_counter()
        with pytest.raises(
            guard.DecompressionBombError, match="not encoded in UTF-8"
        ):
            guard.validate_zip_container(reported, ".xlsx")
        assert time.perf_counter() - start < 5

    @pytest.mark.parametrize(
        "codec, declared, bom",
        [
            ("utf-8", "ISO-8859-1", b""),
            ("cp1252", "windows-1252", b""),
            ("shift_jis", "Shift_JIS", b""),
            ("utf-8", "US-ASCII", b""),
            ("utf-32-le", "UTF-32", b""),
            ("utf-32-be", "UTF-32BE", b""),
            ("utf-32-le", "UTF-32", b"\xff\xfe\x00\x00"),
            ("cp037", "cp037", b""),
            ("utf-16-le", "ISO-8859-1", b"\xff\xfe"),
            ("utf-8", "UTF-8' encoding='ISO-8859-1", b""),
        ],
        ids=[
            "latin-1",
            "windows-1252",
            "shift-jis",
            "ascii",
            "utf-32le",
            "utf-32be",
            "utf-32-bom",
            "ebcdic",
            "utf-16-bom-latin-1-declared",
            "two-encodings",
        ],
    )
    @pytest.mark.parametrize(
        "member",
        ["xl/worksheets/sheet1.xml", "docProps/core.xml", "xl/styles.xml"],
    )
    def test_other_encodings_are_refused(self, member, codec, declared, bom):
        """A declaration naming two encodings is not one XML 1.0 spells,
        so it is refused as malformed."""
        guard = _guard()
        data = _with_part(
            member, lambda xml: _reencoded(xml, codec, declared, bom)
        )
        with pytest.raises(
            guard.DecompressionBombError,
            match="not encoded in UTF-8|declaration is malformed",
        ):
            guard.validate_zip_container(data, ".xlsx")

    @pytest.mark.parametrize(
        "head",
        [
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n',
            b"\xef\xbb\xbf<?xml version='1.0' encoding='utf-8'?>",
            b'\xef\xbb\xbf<?xml version = "1.0"\r\n\tencoding = "Utf-8" ?>',
            b'<?xml version="1.0" standalone="yes"?>',
            b"",
            b"\xef\xbb\xbf",
        ],
        ids=[
            "utf-8",
            "utf-8-bom",
            "bom-spaced-declaration",
            "no-encoding",
            "no-declaration",
            "bom-only",
        ],
    )
    def test_utf8_parts_pass_and_load(self, head):
        """UTF-8, with or without a byte-order mark, with a declaration
        naming UTF-8 or no encoding, or with no declaration at all."""
        guard = _guard()
        data = _real_package(".xlsx")
        for member in (
            "xl/worksheets/sheet1.xml",
            "docProps/core.xml",
            "xl/styles.xml",
            "xl/workbook.xml",
        ):
            data = _rewrite_zip(
                data, {member: lambda xml: head + _without_declaration(xml)}
            )
        guard.validate_zip_container(data, ".xlsx")
        workbook = _load_xlsx(data)
        list(workbook.worksheets[0].iter_rows())
        workbook.close()

    @pytest.mark.parametrize(
        "codec, bom",
        [("utf-16-le", b"\xff\xfe"), ("utf-16-be", b"\xfe\xff")],
        ids=["utf-16le-bom", "utf-16be-bom"],
    )
    @pytest.mark.parametrize(
        "member",
        [
            "xl/worksheets/sheet1.xml",
            "docProps/core.xml",
            "xl/styles.xml",
            "xl/workbook.xml",
        ],
    )
    def test_utf16_parts_with_a_bom_are_refused(self, member, codec, bom):
        """These passed before, when the guard decoded UTF-16 to count
        the ``<`` and ``=`` of a part; it now counts bytes and accepts
        UTF-8 only (``_require_utf8_xml``), since the parsers detect
        UTF-16 differently and no producer surveyed writes it. openpyxl
        does load such a part, so this is a refusal of a spec-legal
        (OPC allows UTF-16) but unwritten encoding."""
        guard = _guard()
        data = _with_part(
            member, lambda xml: _reencoded(xml, codec, "UTF-16", bom)
        )
        workbook = _load_xlsx(data)  # the premise: openpyxl reads it
        workbook.close()
        with pytest.raises(
            guard.DecompressionBombError, match="not encoded in UTF-8"
        ):
            guard.validate_zip_container(data, ".xlsx")

    def test_premise_expat_reads_a_zero_second_byte_as_utf16(self):
        """expat reads a part whose first or second byte is zero as
        UTF-16 without a byte-order mark (a leading space, ``20 00``,
        then ``<``), where libxml2 sees UTF-8 and fails; in UTF-16 the
        low byte of U+4E3C is ``3C``, a ``<`` to a byte count."""
        from lxml import etree
        from xml.parsers import expat

        names = _utf16_lt_names(3)
        for codec in ("utf-16-le", "utf-16-be"):
            document = _utf16_leading_space(
                "<r " + " ".join(f'{n}="1"' for n in names) + "/>", codec
            )
            assert document.count(b"<") == 1 + 3 * len(names)
            seen = []
            parser = expat.ParserCreate(namespace_separator="}")
            parser.StartElementHandler = lambda _n, attrs: seen.append(attrs)
            parser.Parse(document, True)
            assert list(seen[0]) == names
            with pytest.raises(etree.XMLSyntaxError):
                etree.fromstring(document)

    def test_premise_libxml2_switches_to_a_declared_utf16le(self):
        """After an ASCII declaration naming ``UTF-16LE`` libxml2 reads
        the rest as UTF-16LE (which expat rejects), so a byte count of
        what follows misreads it."""
        from lxml import etree
        from xml.parsers import expat

        document = _declared_utf16le('<r 㸿="1" 丼="2"/>')
        root = etree.fromstring(document)
        assert root.attrib.keys() == ["㸿", "丼"]
        with pytest.raises(expat.ExpatError):
            expat.ParserCreate(namespace_separator="}").Parse(document, True)

    @pytest.mark.parametrize("codec", ["utf-16-le", "utf-16-be"])
    @pytest.mark.parametrize(
        "member", ["xl/worksheets/sheet1.xml", "xl/sharedStrings.xml"]
    )
    def test_bomless_utf16_expat_parts_are_refused(self, member, codec):
        """A worksheet or shared-strings part (openpyxl streams both
        with expat) that starts with a UTF-16 space and no byte-order
        mark is refused before any parser reads it. Before, the guard
        counted it as UTF-8, attribute names of U+4E3C and the like
        reset its ``=`` count, and one start tag of 8.66 million
        attributes took the guard to 2,252 MB before the names ceiling
        refused it (200,000 attributes held about 34 MB here)."""
        guard = _guard()
        attributes = " ".join(
            f'{n}="1"'
            for n in _utf16_lt_names(guard.MAX_XML_TAG_ATTRIBUTE_SIGNS + 1)
        )
        if member == "xl/sharedStrings.xml":
            xml = f'<sst xmlns="{_SHEET_MAIN}" {attributes}><si><t>a</t></si></sst>'
            data = _with_shared_strings(_utf16_leading_space(xml, codec))
        else:
            xml = (
                f'<worksheet xmlns="{_SHEET_MAIN}"><sheetData>'
                f'<row r="1" {attributes}/></sheetData></worksheet>'
            )
            data = _with_part(
                member, lambda _xml: _utf16_leading_space(xml, codec)
            )
        start = time.perf_counter()
        with pytest.raises(
            guard.DecompressionBombError, match="not encoded in UTF-8"
        ):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    @pytest.mark.parametrize(
        "member", ["docProps/core.xml", "xl/styles.xml", "xl/workbook.xml"]
    )
    def test_a_declared_utf16_head_is_refused(self, member):
        """An ASCII declaration naming ``UTF-16LE``, after which libxml2
        (openpyxl's parser for these parts) reads UTF-16LE. Before, the
        guard took the part for UTF-8 once a U+3E3F (bytes ``3F 3E``)
        supplied the ``?>`` it looked for, and counted its ``<`` and
        ``=`` in bytes: 4,000,000 attributes on ``core.xml`` cost about
        988 MB in the guard. A part of this shape that libxml2 reads in
        full is refused now, without attributes."""
        from lxml import etree

        guard = _guard()

        def declared(xml: bytes) -> bytes:
            text = _without_declaration(xml).decode("utf-8")
            end = text.index(">")
            if text[end - 1] == "/":
                end -= 1
            return _declared_utf16le(text[:end] + ' 㸿="1"' + text[end:])

        data = _with_part(member, declared)
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            assert "㸿" in etree.fromstring(zf.read(member)).attrib
        with pytest.raises(
            guard.DecompressionBombError, match="not encoded in UTF-8"
        ):
            guard.validate_zip_container(data, ".xlsx")

    @pytest.mark.parametrize(
        "head",
        [
            b" \x00<\x00",
            b"\x00<",
            b"<\x00?\x00",
            b"\x4c\x6f\xa7\x94",
            b"\xef\xbb\xbf\xff\xfe<\x00",
            b'<?xml version="1.0" encoding="UTF-8"?>\x00',
            b'<?xml version="1.0" encoding="UTF8"?>',
            b'<?xml version="1.0" encoding="UTF-16"?>',
            b'<?xml encoding="UTF-8"?>',
        ],
        ids=[
            "utf16le-leading-space",
            "utf16be",
            "utf16le",
            "ebcdic",
            "utf8-bom-then-utf16-bom",
            "nul-after-declaration",
            "utf8-alias",
            "utf8-bytes-declared-utf16",
            "no-version",
        ],
    )
    def test_heads_both_parsers_may_not_read_as_utf8_are_refused(self, head):
        guard = _guard()
        data = _with_part(
            "xl/worksheets/sheet1.xml",
            lambda xml: head + _without_declaration(xml),
        )
        with pytest.raises(
            guard.DecompressionBombError, match="UTF-8|declaration"
        ):
            guard.validate_zip_container(data, ".xlsx")

    def test_an_overlong_xml_declaration_is_refused(self):
        guard = _guard()
        data = _with_part(
            "xl/worksheets/sheet1.xml",
            lambda xml: (
                b'<?xml version="1.0"'
                + b" " * guard._XML_SCAN_CHUNK
                + b'encoding="UTF-8"?>'
                + xml
            ),
        )
        with pytest.raises(guard.DecompressionBombError, match="declaration"):
            guard.validate_zip_container(data, ".xlsx")

    @pytest.mark.parametrize(
        "member, root_end",
        [
            ("xl/worksheets/sheet1.xml", b"</worksheet>"),
            ("[Content_Types].xml", b"</Types>"),
        ],
    )
    def test_expat_parts_are_held_to_the_depth_ceiling(self, member, root_end):
        """openpyxl streams worksheets with expat, and msoffcrypto reads
        ``[Content_Types].xml`` with minidom; the root is one level."""
        guard = _guard()

        def nested(levels: int) -> bytes:
            return _with_part(
                member,
                lambda xml: xml.replace(
                    root_end, b"<x>" * levels + b"</x>" * levels + root_end
                ),
            )

        limit = guard.MAX_XLSX_XML_DEPTH
        guard.validate_zip_container(nested(limit - 1), ".xlsx")
        with pytest.raises(guard.DecompressionBombError, match="deep"):
            guard.validate_zip_container(nested(limit), ".xlsx")

    def test_lxml_parts_are_held_to_the_depth_ceiling(self):
        """openpyxl's lxml tree parse fails beyond 256 levels, but
        libxml2's push parser driving a target sets no limit."""
        guard = _guard()

        def nested(levels: int) -> bytes:
            return _with_part(
                "docProps/core.xml",
                lambda xml: xml.replace(
                    b"</cp:coreProperties>",
                    b"<x>" * levels
                    + b"</x>" * levels
                    + b"</cp:coreProperties>",
                ),
            )

        guard.validate_zip_container(nested(255), ".xlsx")
        _load_xlsx(nested(255)).close()
        with pytest.raises(guard.DecompressionBombError, match="deep"):
            guard.validate_zip_container(nested(256), ".xlsx")
        from lxml import etree

        with pytest.raises(etree.XMLSyntaxError, match="depth"):
            _load_xlsx(nested(256))

    def test_premise_lxml_push_parser_with_a_target_has_no_depth_limit(self):
        from lxml import etree

        class Target:
            def close(self):
                return None

        document = b"<a>" * 5_000 + b"</a>" * 5_000
        parser = etree.XMLParser(target=Target(), resolve_entities=False)
        parser.feed(document)
        parser.close()
        with pytest.raises(etree.XMLSyntaxError, match="depth"):
            etree.fromstring(document, etree.XMLParser(resolve_entities=False))

    def test_reported_deep_worksheet_is_refused_quickly(self):
        """4 MiB of ``<a>``: expat kept about 145 bytes per open element
        (128 MiB, a 130 KB file, about 6 GB) before the depth ceiling."""
        guard = _guard()
        data = _with_part(
            "xl/worksheets/sheet1.xml",
            lambda xml: xml.replace(
                b"</worksheet>", b"<a>" * (4 * 1024 * 1024 // 3)
            ),
        )
        assert len(data) < 50_000
        start = time.perf_counter()
        with pytest.raises(guard.DecompressionBombError, match="deep"):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    @pytest.mark.parametrize(
        "member, root_end",
        [
            ("xl/worksheets/sheet1.xml", b"</worksheet>"),
            ("docProps/core.xml", b"</cp:coreProperties>"),
        ],
    )
    def test_start_tags_are_held_to_the_attribute_ceiling(
        self, member, root_end
    ):
        """Counted as the ``=`` between two ``<``, before either parser
        stores the tag's attributes (in a value, ``=`` counts too)."""
        guard = _guard()
        limit = guard.MAX_XML_TAG_ATTRIBUTE_SIGNS

        def tag(signs: int) -> bytes:
            return _with_part(
                member,
                lambda xml: xml.replace(
                    root_end,
                    b'<x a0="" a1="' + b"=" * (signs - 2) + b'"/>' + root_end,
                ),
            )

        guard.validate_zip_container(tag(limit), ".xlsx")
        with pytest.raises(guard.DecompressionBombError, match="attributes"):
            guard.validate_zip_container(tag(limit + 1), ".xlsx")

    def test_reported_attribute_shape_is_refused_quickly(self):
        """One start tag of 1,000,000 attributes, 11 MB of a worksheet,
        held about 175 MB in the scan before."""
        guard = _guard()
        data = _with_part(
            "xl/worksheets/sheet1.xml",
            lambda xml: xml.replace(
                b"</worksheet>",
                b"<x "
                + b" ".join(b'a%d=""' % i for i in range(1_000_000))
                + b"/></worksheet>",
            ),
        )
        start = time.perf_counter()
        with pytest.raises(guard.DecompressionBombError, match="attributes"):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    def test_a_cell_of_32767_equals_signs_passes_and_loads(self):
        """Excel holds up to 32,767 characters in a cell."""
        guard = _guard()
        text = "=" * 32_767
        data = _with_part(
            "xl/worksheets/sheet1.xml",
            lambda xml: xml.replace(
                b"<sheetData></sheetData>",
                b'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>'
                + text.encode()
                + b"</t></is></c></row></sheetData>",
            ),
        )
        guard.validate_zip_container(data, ".xlsx")
        workbook = _load_xlsx(data)
        assert next(workbook.worksheets[0].values)[0] == text
        workbook.close()

    def test_a_member_no_loader_parses_is_not_scanned(self):
        """A DTD, a long URI and deep nesting in an unreferenced member
        cost the scan nothing: nothing parses it."""
        guard = _guard()
        junk = (
            b"<!DOCTYPE r><r xmlns:q='urn:"
            + b"a" * 100_000
            + b"'>"
            + b"<a>" * 1_000_000
        )
        data = _rewrite_zip(
            _real_package(".xlsx"),
            add={"junk/junk.xml": junk, "xl/junk.xml": junk},
        )
        opened: list = []
        real_open = zipfile.ZipFile.open

        def spy(self, name, *args, **kwargs):
            opened.append(
                name.filename if isinstance(name, zipfile.ZipInfo) else name
            )
            return real_open(self, name, *args, **kwargs)

        # The guard's own verification inflates entries from the raw
        # bytes; every ZipFile.open is the scan's or openpyxl's.
        with patch.object(zipfile.ZipFile, "open", spy):
            guard.validate_zip_container(data, ".xlsx")
        assert "junk/junk.xml" not in opened
        assert "xl/junk.xml" not in opened

    def test_svg_with_a_doctype_in_media_passes_and_loads(self):
        """Illustrator and icon SVGs carry the SVG 1.1 DOCTYPE; openpyxl
        never parses an image."""
        guard = _guard()

        def types(ct: bytes) -> bytes:
            return ct.replace(
                b"</Types>",
                b'<Default Extension="svg" ContentType="image/svg+xml"/>'
                b"</Types>",
            )

        data = _rewrite_zip(
            _real_package(".xlsx"),
            {"[Content_Types].xml": types},
            add={"xl/media/image1.svg": _SVG_WITH_DOCTYPE},
        )
        guard.validate_zip_container(data, ".xlsx")
        _load_xlsx(data).close()
        pd = pytest.importorskip("pandas")
        pd.read_excel(io.BytesIO(data), sheet_name=None, engine="openpyxl")

    @pytest.mark.parametrize(
        "change",
        [
            b"<!DOCTYPE worksheet>",
            b'<q:w xmlns:q="urn:' + b"a" * 1_025 + b'"/>',
            b"<x>" * 300,
        ],
        ids=["dtd", "long-uri", "deep"],
    )
    def test_a_worksheet_at_a_media_path_is_still_scanned(self, change):
        """openpyxl follows the workbook's relationship wherever it
        points, so the part is scanned as a worksheet."""
        guard = _guard()
        sheet = (
            f'<worksheet xmlns="{_SHEET_MAIN}"><sheetData/></worksheet>'
        ).encode()
        if change.startswith(b"<!DOCTYPE"):
            sheet = change + sheet
        else:
            sheet = sheet.replace(b"</worksheet>", change + b"</worksheet>")
        data = _openpyxl_shaped_xlsx(
            _sheet(1, "rId1"),
            _ws_rel("rId1", "media/image1.svg"),
            extra={"xl/media/image1.svg": sheet},
        )
        with pytest.raises(guard.DecompressionBombError):
            guard.validate_zip_container(data, ".xlsx")

    def test_shared_strings_at_a_media_path_are_still_scanned(self):
        """openpyxl reads the shared strings from the part the manifest
        names, wherever it is."""
        guard = _guard()
        strings = (
            f'<sst xmlns="{_SHEET_MAIN}"><si><t>a</t></si>'
            + "<x>" * 300
            + "</sst>"
        ).encode()

        def types(ct: bytes) -> bytes:
            return ct.replace(
                b"</Types>",
                b'<Override PartName="/xl/media/image1.svg" ContentType='
                b'"application/vnd.openxmlformats-officedocument.'
                b'spreadsheetml.sharedStrings+xml"/></Types>',
            )

        data = _openpyxl_shaped_xlsx(
            _sheet(1, "rId1"),
            _ws_rel("rId1", "worksheets/sheet1.xml"),
            extra={"xl/media/image1.svg": strings},
        )
        data = _rewrite_zip(data, {"[Content_Types].xml": types})
        with pytest.raises(guard.DecompressionBombError, match="deep"):
            guard.validate_zip_container(data, ".xlsx")

    def test_chartsheets_and_sheet_relationships_are_scanned_with_lxml(self):
        guard = _guard()
        data = _xlsx_with_image_references(1, _small_png())
        families: dict = {}
        real_scan = guard._scan_xlsx_part

        def spy(archive, info, scanned_with):
            families[info.filename] = scanned_with
            return real_scan(archive, info, scanned_with)

        with patch.object(guard, "_scan_xlsx_part", spy):
            guard.validate_zip_container(data, ".xlsx")
        lxml = frozenset({guard._LXML})
        assert families["xl/chartsheets/sheet2.xml"] == lxml
        assert families["xl/chartsheets/_rels/sheet2.xml.rels"] == lxml
        assert families["xl/worksheets/sheet1.xml"] == frozenset({guard._EXPAT})
        # The drawing and its image are never read (openpyxl_hardening).
        assert "xl/drawings/drawing1.xml" not in families
        assert "xl/media/image1.png" not in families

    def test_a_part_lxml_rejects_is_refused_where_openpyxl_parses_it(self):
        """libxml2 stops at a name over 50,000 characters, so openpyxl's
        load of such a part fails; expat would read on."""
        guard = _guard()
        data = _with_part(
            "docProps/core.xml",
            lambda xml: xml.replace(
                b"</cp:coreProperties>",
                b"<" + b"n" * 50_001 + b"/></cp:coreProperties>",
            ),
        )
        with pytest.raises(guard.DecompressionBombError, match="malformed"):
            guard.validate_zip_container(data, ".xlsx")

    @pytest.mark.parametrize(
        "member, root_end, name",
        [
            ("xl/worksheets/sheet1.xml", b"</worksheet>", "<n{}/>"),
            ("xl/worksheets/sheet1.xml", b"</worksheet>", '<x a{}=""/>'),
            ("xl/worksheets/sheet1.xml", b"</worksheet>", '<x xmlns:p{}="u"/>'),
            ("docProps/core.xml", b"</cp:coreProperties>", "<n{}/>"),
            ("docProps/core.xml", b"</cp:coreProperties>", '<x a{}=""/>'),
        ],
        ids=[
            "expat-elements",
            "expat-attributes",
            "expat-prefixes",
            "lxml-elements",
            "lxml-attributes",
        ],
    )
    def test_distinct_names_are_capped(self, member, root_end, name):
        """Both parsers keep every distinct name for the whole parse."""
        guard = _guard()

        def names(count: int) -> bytes:
            filler = "".join(name.format(i) for i in range(count)).encode()
            return _with_part(
                member, lambda xml: xml.replace(root_end, filler + root_end)
            )

        guard.validate_zip_container(names(1_000), ".xlsx")
        with pytest.raises(guard.DecompressionBombError, match="distinct"):
            guard.validate_zip_container(
                names(guard.MAX_XLSX_PART_NAMES + 1), ".xlsx"
            )

    def test_reported_distinct_name_shape_is_refused_quickly(self):
        """1,000,000 distinct element names, 8.9 MB of a worksheet, held
        about 140 MB in the scan before the ceiling."""
        guard = _guard()
        data = _with_part(
            "xl/worksheets/sheet1.xml",
            lambda xml: xml.replace(
                b"</worksheet>",
                b"".join(b"<a%d/>" % i for i in range(1_000_000))
                + b"</worksheet>",
            ),
        )
        start = time.perf_counter()
        with pytest.raises(guard.DecompressionBombError, match="distinct"):
            guard.validate_zip_container(data, ".xlsx")
        assert time.perf_counter() - start < 5

    def test_scan_ceilings_are_pinned(self):
        guard = _guard()
        assert guard.MAX_XLSX_PART_NAMES >= 1_000
        assert guard.MAX_XLSX_XML_DEPTH == 256  # libxml2's own ceiling
        assert guard.MAX_XML_TAG_ATTRIBUTE_SIGNS >= 32_767  # Excel's cell
        assert guard._XML_SCAN_CHUNK <= guard.MAX_XML_TAG_ATTRIBUTE_SIGNS


def _with_start_tag(data: bytes, member: str, anchor: bytes, attributes: int):
    """*data* with one ``<x>`` of *attributes* empty attributes inserted
    after *anchor* in *member* (into the first ``Relationship`` start
    tag when *anchor* is empty)."""
    names = b" ".join(b'a%d=""' % i for i in range(attributes))

    def change(xml: bytes) -> bytes:
        if not anchor:
            return xml.replace(
                b"<Relationship ", b"<Relationship " + names + b" ", 1
            )
        assert anchor in xml
        return xml.replace(anchor, anchor + b"<x " + names + b"/>", 1)

    return _rewrite_zip(data, {member: change})


_STREAMED_PARTS = [
    pytest.param(".docx", "word/document.xml", b"<w:body>", id="docx-document"),
    pytest.param(
        ".pptx", "ppt/slides/slide1.xml", b"<p:cSld>", id="pptx-slide"
    ),
    pytest.param(".docx", "word/_rels/document.xml.rels", b"", id="docx-rels"),
    pytest.param(
        ".pptx", "ppt/_rels/presentation.xml.rels", b"", id="pptx-rels"
    ),
]


class TestStreamedPartsAreCountedBeforeParsing:
    """Every ``.docx``/``.pptx`` part the guard streams with lxml's
    ``iterparse`` (and every relationships part) is held to UTF-8 and
    its start tags to ``MAX_XML_TAG_ATTRIBUTE_SIGNS`` before lxml reads
    it: ``iterparse`` builds a start tag's element in full before the
    per-element attribute ceilings see it, about 200 bytes per
    attribute, with no limit of its own."""

    def test_premise_iterparse_builds_a_huge_start_tag(self):
        from lxml import etree

        count = 70_000
        document = (
            b"<r><x "
            + b" ".join(b'a%d=""' % i for i in range(count))
            + b"/></r>"
        )
        for _event, element in etree.iterparse(
            io.BytesIO(document), events=("start",)
        ):
            if element.tag == "x":
                assert len(element.attrib) == count
                break

    @pytest.mark.parametrize("ext, member, anchor", _STREAMED_PARTS)
    def test_a_start_tag_over_the_sign_ceiling_is_refused_before_parsing(
        self, ext, member, anchor
    ):
        """Before, the guard's walk built such an element first (2,000,000
        attributes in ``document.xml``, 23 MB, held 404 MB, so a 128 MiB
        part about 2.4 GB) and a slide's was not refused at all."""
        guard = _guard()
        data = _with_start_tag(
            _real_package(ext),
            member,
            anchor,
            guard.MAX_XML_TAG_ATTRIBUTE_SIGNS + 1,
        )
        with pytest.raises(
            guard.DecompressionBombError, match="may carry more"
        ):
            guard.validate_zip_container(data, ext)

    @pytest.mark.parametrize("ext, member, anchor", _STREAMED_PARTS)
    def test_utf16_streamed_parts_are_refused(self, ext, member, anchor):
        """A UTF-16 part (with a byte-order mark, which python-docx and
        python-pptx would read) passed before; its bytes cannot be
        counted, so it is refused like an ``.xlsx`` part."""
        guard = _guard()
        data = _rewrite_zip(
            _real_package(ext),
            {
                member: lambda xml: _reencoded(
                    xml, "utf-16-le", "UTF-16", b"\xff\xfe"
                )
            },
        )
        with pytest.raises(
            guard.DecompressionBombError, match="not encoded in UTF-8"
        ):
            guard.validate_zip_container(data, ext)

    @pytest.mark.parametrize("ext, member, anchor", _STREAMED_PARTS)
    def test_utf8_streamed_parts_with_a_bom_pass(self, ext, member, anchor):
        guard = _guard()
        data = _rewrite_zip(
            _real_package(ext),
            {member: lambda xml: b"\xef\xbb\xbf" + xml},
        )
        guard.validate_zip_container(data, ext)


@contextmanager
def _parser_input(monkeypatch):
    """Count every byte handed to an lxml or expat parser (``iterparse``
    reads, ``fromstring`` buffers, ``XMLParser.feed`` and expat
    ``Parse``) while the block runs; yields a one-element list."""
    from lxml import etree
    from xml.parsers import expat

    received = [0]
    real_iterparse, real_fromstring = etree.iterparse, etree.fromstring
    real_parser, real_create = etree.XMLParser, expat.ParserCreate

    class Metered:
        def __init__(self, source):
            self._source = source

        def read(self, size=-1):
            chunk = self._source.read(size)
            received[0] += len(chunk)
            return chunk

    class MeteredParser(real_parser):
        def feed(self, data):
            received[0] += len(data)
            return super().feed(data)

    class MeteredExpat:
        def __init__(self, parser):
            object.__setattr__(self, "_parser", parser)

        def __getattr__(self, name):
            return getattr(self._parser, name)

        def __setattr__(self, name, value):
            setattr(self._parser, name, value)

        def Parse(self, data, final=False):
            received[0] += len(data)
            return self._parser.Parse(data, final)

    def fromstring(text, *args, **kwargs):
        received[0] += len(text)
        return real_fromstring(text, *args, **kwargs)

    monkeypatch.setattr(
        etree,
        "iterparse",
        lambda source, *a, **k: real_iterparse(Metered(source), *a, **k),
    )
    monkeypatch.setattr(etree, "fromstring", fromstring)
    monkeypatch.setattr(etree, "XMLParser", MeteredParser)
    monkeypatch.setattr(
        expat,
        "ParserCreate",
        lambda *a, **k: MeteredExpat(real_create(*a, **k)),
    )
    try:
        yield received
    finally:
        monkeypatch.undo()


def _dtd_subset(shape: str, root: bytes, count: int) -> bytes:
    """An internal subset of *count* declarations: attribute-list
    declarations (no ``=`` anywhere), entity declarations, or ``xmlns:``
    attribute defaults, which libxml2 adds to *root*'s start tag."""
    if shape == "attlist":
        return (
            b"<!ATTLIST "
            + root
            + b" "
            + b" ".join(b'a%d CDATA "x"' % i for i in range(count))
            + b">"
        )
    if shape == "entity":
        return b"".join(b'<!ENTITY e%d "x">' % i for i in range(count))
    assert shape == "xmlns"
    return (
        b"<!ATTLIST "
        + root
        + b" "
        + b" ".join(b'xmlns:p%d CDATA "u%d"' % (i, i) for i in range(count))
        + b">"
    )


_ROOT_START = re.compile(rb"<(?![?!])([^\s/>]+)")


def _with_dtd_subset(shape: str, count: int):
    """A rewrite that puts a DTD of *count* *shape* declarations before
    the part's root element (a part of about 2.5-3.5 MB at 150,000)."""

    def rewrite(xml: bytes) -> bytes:
        match = _ROOT_START.search(xml)
        root = match.group(1)
        return (
            xml[: match.start()]
            + b"<!DOCTYPE "
            + root
            + b" ["
            + _dtd_subset(shape, root, count)
            + b"]>"
            + xml[match.start() :]
        )

    return rewrite


def _after_root_start(insert: bytes):
    """A rewrite that puts *insert* right after the root's start tag."""

    def rewrite(xml: bytes) -> bytes:
        match = _ROOT_START.search(xml)
        end = xml.index(b">", match.end()) + 1
        assert xml[end - 2 : end] != b"/>"
        return xml[:end] + insert + xml[end:]

    return rewrite


def _package_for(source: str) -> bytes:
    if source == "docx":
        return _real_package(".docx")
    if source == "docx-header":
        return _docx_with_sections(1)
    assert source == "pptx"
    return _deck(1)


#: (package, member, extension): every kind of site that parses a
#: ``.docx``/``.pptx`` part: the walkers of ``document.xml`` and a slide,
#: the relationships check, the header/footer union walker, and the
#: buffered parses of ``_skeleton_xml``.
_DTD_SITES = [
    pytest.param("docx", "word/document.xml", ".docx", id="docx-document"),
    pytest.param("pptx", "ppt/slides/slide1.xml", ".pptx", id="pptx-slide"),
    pytest.param(
        "docx", "word/_rels/document.xml.rels", ".docx", id="docx-rels"
    ),
    pytest.param("docx-header", "word/header1.xml", ".docx", id="docx-header"),
    pytest.param(
        "docx", "[Content_Types].xml", ".docx", id="docx-content-types"
    ),
    pytest.param("docx", "word/settings.xml", ".docx", id="docx-settings"),
    pytest.param("docx", "word/styles.xml", ".docx", id="docx-styles"),
    pytest.param(
        "pptx", "ppt/presentation.xml", ".pptx", id="pptx-presentation"
    ),
    pytest.param(
        "pptx", "[Content_Types].xml", ".pptx", id="pptx-content-types"
    ),
]

#: The sites ``TestStreamedPartsAreCountedBeforeParsing`` does not
#: cover: the header/footer union walker and ``_skeleton_xml``.
_HEADER_AND_SKELETON_SITES = [
    p
    for p in _DTD_SITES
    if p.id not in ("docx-document", "pptx-slide", "docx-rels")
]


class TestDtdIsRefusedInTheBytes:
    """A document type declaration is refused in a part's bytes
    (``_XmlProlog``, through ``_CountedXmlStream``) before any parser
    the guard drives receives it. Checking at the root's first event,
    as before, came after libxml2 had read and built the whole internal
    subset: attribute-list declarations carry no ``=`` for the
    start-tag count, so a 17.9 MB ``document.xml`` of one million held
    424 MB in the guard, and ``xmlns:`` defaults added namespace
    declarations to the root its bytes did not show."""

    def test_premise_libxml2_reads_the_whole_subset_before_the_root(self):
        from lxml import etree

        xml = _with_dtd_subset("attlist", 20_000)(b"<r/>")
        read = [0]

        class Source:
            def __init__(self):
                self._data = io.BytesIO(xml)

            def read(self, size=-1):
                chunk = self._data.read(size)
                read[0] += len(chunk)
                return chunk

        for _event, element in etree.iterparse(
            Source(), events=("start",), resolve_entities=False
        ):
            assert element.getroottree().docinfo.internalDTD is not None
            break
        assert read[0] == len(xml)

    def test_premise_an_attlist_default_adds_namespaces_to_the_root(self):
        from lxml import etree

        xml = _with_dtd_subset("xmlns", 3)(b"<r/>")
        assert b"=" not in xml.split(b"]>")[-1]
        root = etree.fromstring(xml, etree.XMLParser(resolve_entities=False))
        assert len(root.nsmap) == 3

    @pytest.mark.parametrize("shape", ["attlist", "entity", "xmlns"])
    @pytest.mark.parametrize("source, member, ext", _DTD_SITES)
    def test_a_large_dtd_is_refused_before_a_parser_reads_it(
        self, monkeypatch, source, member, ext, shape
    ):
        """Before, every parser site received the whole part (about
        3 MB here) and libxml2 built the subset first; now no parser
        receives any of it (the other parts of the package total about
        0.4 MB)."""
        guard = _guard()
        data = _rewrite_zip(
            _package_for(source), {member: _with_dtd_subset(shape, 150_000)}
        )
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            assert archive.getinfo(member).file_size > 2_500_000
        with _parser_input(monkeypatch) as received:
            with pytest.raises(guard.DecompressionBombError, match="DTD"):
                guard.validate_zip_container(data, ext)
        assert received[0] < 1_000_000

    @pytest.mark.parametrize(
        "member, families",
        [
            ("xl/worksheets/sheet1.xml", "expat"),
            ("xl/styles.xml", "lxml"),
            ("[Content_Types].xml", "lxml and expat"),
            ("xl/workbook.xml", "lxml"),
        ],
    )
    def test_a_large_dtd_in_a_spreadsheet_part_is_refused_before_a_parser_reads_it(
        self, monkeypatch, member, families
    ):
        """The scan's parsers refused at their doctype callbacks, but
        lxml's push parser had first buffered the whole subset."""
        openpyxl = pytest.importorskip("openpyxl")
        guard = _guard()
        buf = io.BytesIO()
        openpyxl.Workbook().save(buf)
        data = _rewrite_zip(
            buf.getvalue(), {member: _with_dtd_subset("attlist", 150_000)}
        )
        with _parser_input(monkeypatch) as received:
            with pytest.raises(guard.DecompressionBombError, match="DTD"):
                guard.validate_zip_container(data, ".xlsx")
        assert received[0] < 1_000_000

    def test_a_declaration_across_a_chunk_boundary_is_refused(
        self, monkeypatch
    ):
        guard = _guard()

        def rewrite(xml: bytes) -> bytes:
            dtd = _with_dtd_subset("attlist", 150_000)(xml)
            at = dtd.index(b"<!DOCTYPE")
            pad = guard._XML_SCAN_CHUNK - 1 - at - len(b"<!---->")
            assert pad > 0
            # "<" ends the first chunk; "!DOCTYPE" starts the second.
            return dtd[:at] + b"<!--" + b" " * pad + b"-->" + dtd[at:]

        data = _rewrite_zip(
            _real_package(".docx"), {"word/document.xml": rewrite}
        )
        with _parser_input(monkeypatch) as received:
            with pytest.raises(guard.DecompressionBombError, match="DTD"):
                guard.validate_zip_container(data, ".docx")
        assert received[0] < 1_000_000

    @pytest.mark.parametrize(
        "prolog",
        [
            b'<!ENTITY e "x">',
            b'<!ATTLIST w:document a CDATA "x">',
            b"<!ELEMENT w:document ANY>",
        ],
        ids=["entity", "attlist", "element"],
    )
    def test_a_declaration_without_a_doctype_is_refused(self, prolog):
        guard = _guard()
        data = _rewrite_zip(
            _real_package(".docx"),
            {"word/document.xml": _with_doctype(prolog, b"<w:document")},
        )
        with pytest.raises(guard.DecompressionBombError, match="DTD"):
            guard.validate_zip_container(data, ".docx")

    def test_comments_cdata_and_instructions_holding_a_declaration_pass(self):
        """``<!`` is examined only before the root element, outside
        comments and processing instructions: a declaration spelled in
        a comment, a processing instruction or a CDATA section is text."""
        guard = _guard()
        spelled = b'<!DOCTYPE w:document [<!ENTITY e "x">]>'

        def rewrite(xml: bytes) -> bytes:
            xml = _with_doctype(
                b"<!-- " + spelled + b" --><?pi " + spelled + b"?><!---->",
                b"<w:document",
            )(xml)
            return xml.replace(
                b"<w:body>",
                b"<w:body><!-- "
                + spelled
                + b" --><w:p><w:r><w:t><![CDATA["
                + spelled
                + b"]]></w:t></w:r></w:p>",
                1,
            )

        data = _rewrite_zip(
            _real_package(".docx"), {"word/document.xml": rewrite}
        )
        guard.validate_zip_container(data, ".docx")

    def test_a_comment_in_a_spreadsheet_part_holding_a_declaration_passes(self):
        openpyxl = pytest.importorskip("openpyxl")
        guard = _guard()
        buf = io.BytesIO()
        openpyxl.Workbook().save(buf)
        data = _rewrite_zip(
            buf.getvalue(),
            {
                "xl/worksheets/sheet1.xml": _with_doctype(
                    b"<!-- <!DOCTYPE worksheet> -->", b"<worksheet"
                )
            },
        )
        guard.validate_zip_container(data, ".xlsx")

    def test_a_double_hyphen_in_a_comment_before_the_root_is_refused(self):
        """Both parsers reject it; the scan would otherwise have to
        guess where they end the comment."""
        guard = _guard()
        data = _rewrite_zip(
            _real_package(".docx"),
            {
                "word/document.xml": _with_doctype(
                    b"<!-- a -- b -->", b"<w:document"
                )
            },
        )
        with pytest.raises(guard.DecompressionBombError, match="comment"):
            guard.validate_zip_container(data, ".docx")

    @pytest.mark.parametrize(
        "xml, verdict",
        [
            (
                b'<?xml version="1.0"?>\n<!-- <!DOCTYPE r> --><?p <!DOCTYPE?>'
                b"<!----><r><![CDATA[<!DOCTYPE]]><!--<!DOCTYPE--></r>",
                None,
            ),
            (b"<r/><!DOCTYPE r>", None),
            (b"\xef\xbb\xbf <!DOCTYPE r><r/>", "DTD"),
            (b'<?xml version="1.0"?><!ENTITY e "x"><r/>', "DTD"),
            (b'<!-- a --><!ATTLIST r a CDATA "x"><r/>', "DTD"),
            (b"<![CDATA[x]]><r/>", "DTD"),
            (b"<?p ?? > <!-- ?><!doctype r><r/>", "DTD"),
            (b"<!-- a -- b --><r/>", "comment"),
            (b"<!-x--><r/>", "comment"),
            (b"<!-- a ---><r/>", "comment"),
        ],
    )
    def test_prolog_scan_is_the_same_at_every_chunk_boundary(
        self, xml, verdict
    ):
        guard = _guard()

        def scan(cuts) -> str | None:
            prolog = guard._XmlProlog()
            try:
                start = 0
                for cut in [*cuts, len(xml)]:
                    prolog.feed(xml[start:cut])
                    start = cut
            except guard.DecompressionBombError as exc:
                return "DTD" if "DTD" in str(exc) else "comment"
            return None

        assert scan([]) == verdict
        assert scan(range(1, len(xml))) == verdict
        for cut in range(len(xml) + 1):
            assert scan([cut]) == verdict


class TestHeaderAndSkeletonPartsAreCountedBeforeParsing:
    """The header/footer union walker and ``_skeleton_xml`` (the content
    types, ``presentation.xml``, and the settings and styles parts) take
    their bytes through ``_CountedXmlStream`` as the other walkers do."""

    @pytest.mark.parametrize("source, member, ext", _HEADER_AND_SKELETON_SITES)
    def test_a_start_tag_over_the_sign_ceiling_is_refused(
        self, source, member, ext
    ):
        guard = _guard()
        names = b" ".join(
            b'a%d=""' % i for i in range(guard.MAX_XML_TAG_ATTRIBUTE_SIGNS + 1)
        )
        data = _rewrite_zip(
            _package_for(source),
            {member: _after_root_start(b"<x " + names + b"/>")},
        )
        with pytest.raises(
            guard.DecompressionBombError, match="may carry more"
        ):
            guard.validate_zip_container(data, ext)

    @pytest.mark.parametrize("source, member, ext", _HEADER_AND_SKELETON_SITES)
    def test_utf16_parts_are_refused(self, source, member, ext):
        guard = _guard()
        data = _rewrite_zip(
            _package_for(source),
            {
                member: lambda xml: _reencoded(
                    xml, "utf-16-le", "UTF-16", b"\xff\xfe"
                )
            },
        )
        with pytest.raises(
            guard.DecompressionBombError, match="not encoded in UTF-8"
        ):
            guard.validate_zip_container(data, ext)

    @pytest.mark.parametrize("source, member, ext", _HEADER_AND_SKELETON_SITES)
    def test_controls_pass(self, source, member, ext):
        _guard().validate_zip_container(_package_for(source), ext)

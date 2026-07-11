"""Regression tests for office-document upload support (issue #4414).

ODT/DOC/DOCX/PPT/PPTX/XLS/XLSX were advertised as supported but failed at
runtime because their ``unstructured`` parser dependencies (python-docx,
python-pptx, openpyxl) were never installed. These tests assert that:

1. With the parser dependencies installed, the office extensions are
   registered and text actually round-trips through the upload extraction
   path (``extract_text_from_bytes``).
2. The registry is *honest*: a format is advertised only when its real
   runtime dependency is importable, so a missing dependency hides the
   format instead of accepting an upload that later fails silently.
"""

import importlib
import shutil
import tempfile
from pathlib import Path

import pytest

REGISTRY_MODULE = "local_deep_research.document_loaders.loader_registry"


def _reload(module_name):
    return importlib.reload(importlib.import_module(module_name))


class TestOfficeExtensionsRegistered:
    """Office formats are advertised when their parser deps are present."""

    def test_modern_office_formats_registered(self):
        # Modern office formats work with pip-only parser deps (no system
        # binaries), so they are registered whenever those deps are present.
        pytest.importorskip("docx")
        pytest.importorskip("pptx")
        pytest.importorskip("openpyxl")
        pytest.importorskip("msoffcrypto")

        from local_deep_research.document_loaders.loader_registry import (
            LOADER_REGISTRY,
        )

        for ext in (".docx", ".pptx", ".xlsx"):
            assert ext in LOADER_REGISTRY, f"{ext} should be registered"

    def test_legacy_binary_formats_gated_on_libreoffice(self):
        # .doc/.ppt are OLE binary formats that unstructured converts via
        # LibreOffice; they must be advertised only when soffice is present.
        import local_deep_research.document_loaders.loader_registry as mod

        for ext in (".doc", ".ppt"):
            assert (ext in mod.LOADER_REGISTRY) == (
                mod.HAS_LIBREOFFICE
                and (mod.HAS_DOCX_DEP if ext == ".doc" else mod.HAS_PPTX_DEP)
            )

    def test_odt_registered_when_docx_and_pandoc_present(self):
        pytest.importorskip("docx")
        pytest.importorskip("pypandoc")

        from local_deep_research.document_loaders.loader_registry import (
            LOADER_REGISTRY,
        )

        assert ".odt" in LOADER_REGISTRY


class TestOfficeExtractionRoundTrip:
    """Real files extract text through the public upload path."""

    def test_docx_extraction(self):
        pytest.importorskip("docx")
        from docx import Document

        from local_deep_research.document_loaders import (
            extract_text_from_bytes,
        )

        marker = "Hello DOCX world, extraction works."
        with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
            doc = Document()
            doc.add_paragraph(marker)
            doc.save(tmp.name)
            content = Path(tmp.name).read_bytes()

        text = extract_text_from_bytes(content, ".docx", "sample.docx")
        assert text is not None
        assert marker in text

    def test_xls_extraction(self):
        # xlwt (a dev-only writer) builds the legacy fixture; the loader reads
        # it with pandas + xlrd, bypassing the fragile unstructured path.
        xlwt = pytest.importorskip("xlwt")
        pytest.importorskip("xlrd")

        from local_deep_research.document_loaders import (
            extract_text_from_bytes,
        )

        wb = xlwt.Workbook()
        ws = wb.add_sheet("Data")
        for r, (a, b) in enumerate(
            [("Animal", "Count"), ("Otter", "3"), ("Badger", "7")]
        ):
            ws.write(r, 0, a)
            ws.write(r, 1, b)
        with tempfile.NamedTemporaryFile(suffix=".xls", delete=False) as tmp:
            xls_path = tmp.name
        try:
            wb.save(xls_path)
            content = Path(xls_path).read_bytes()
        finally:
            Path(xls_path).unlink(missing_ok=True)

        text = extract_text_from_bytes(content, ".xls", "sample.xls")
        assert text is not None
        assert "Animal" in text and "Otter" in text and "Badger" in text

    def test_odt_extraction(self):
        pypandoc = pytest.importorskip("pypandoc")
        pytest.importorskip("docx")
        if not shutil.which("pandoc"):
            try:
                pypandoc.get_pandoc_version()
            except OSError:
                pytest.skip("pandoc binary not available")

        from local_deep_research.document_loaders import (
            extract_text_from_bytes,
        )

        marker = "Hello ODT world, extraction works."
        with tempfile.NamedTemporaryFile(suffix=".odt", delete=False) as tmp:
            odt_path = tmp.name
        try:
            # Generate a genuine ODT (with styles.xml etc.) via pandoc.
            pypandoc.convert_text(
                marker, "odt", format="md", outputfile=odt_path
            )
            content = Path(odt_path).read_bytes()
        finally:
            Path(odt_path).unlink(missing_ok=True)

        text = extract_text_from_bytes(content, ".odt", "sample.odt")
        assert text is not None
        assert marker in text


class TestXLSLoaderEncryption:
    """Encrypted XLS files produce a clear error, not a generic traceback."""

    def test_encrypted_xls_raises_value_error(self, monkeypatch):
        pytest.importorskip("xlrd")

        from local_deep_research.document_loaders.xls_loader import XLSLoader

        def fake_read_excel(*args, **kwargs):
            raise Exception("Workbook is encrypted")

        monkeypatch.setattr("pandas.read_excel", fake_read_excel)

        loader = XLSLoader("/tmp/fake.xls")
        with pytest.raises(ValueError, match="encrypted"):
            loader.load()

    def test_corrupted_xls_raises_original_error(self, monkeypatch):
        pytest.importorskip("xlrd")

        from local_deep_research.document_loaders.xls_loader import XLSLoader

        def fake_read_excel(*args, **kwargs):
            raise Exception("Unsupported format")

        monkeypatch.setattr("pandas.read_excel", fake_read_excel)

        loader = XLSLoader("/tmp/fake.xls")
        with pytest.raises(Exception, match="Unsupported format"):
            loader.load()


class TestHonestDetection:
    """A missing runtime dependency hides the format from the registry.

    The registry computes its capability flags at import time, so we reload
    the module with ``importlib.util.find_spec`` patched to report specific
    modules as absent. Patching ``_module_available`` directly would not
    survive the reload (it redefines that function from source), so we patch
    the lower-level ``find_spec`` it calls.
    """

    @pytest.fixture
    def restore_module(self):
        # Restore the pre-test module dict instead of reloading: pytest
        # finalizes this fixture before monkeypatch (LIFO, it is listed
        # after monkeypatch in the test signatures), so a teardown reload
        # would re-execute the registry with find_spec still patched and
        # leak the blocked-dependency state to later tests.
        module = importlib.import_module(REGISTRY_MODULE)
        snapshot = dict(module.__dict__)
        yield
        module.__dict__.clear()
        module.__dict__.update(snapshot)

    @staticmethod
    def _reload_without(monkeypatch, blocked):
        import importlib.util as importlib_util

        real_find_spec = importlib_util.find_spec

        def fake_find_spec(name, *args, **kwargs):
            if name in blocked:
                return None
            return real_find_spec(name, *args, **kwargs)

        monkeypatch.setattr(importlib_util, "find_spec", fake_find_spec)
        return _reload(REGISTRY_MODULE)

    def test_missing_python_docx_hides_word_and_odt(
        self, monkeypatch, restore_module
    ):
        reloaded = self._reload_without(monkeypatch, {"docx"})

        assert reloaded.HAS_DOCX_DEP is False
        for ext in (".doc", ".docx", ".odt"):
            assert ext not in reloaded.LOADER_REGISTRY

    def test_missing_pptx_hides_powerpoint(self, monkeypatch, restore_module):
        reloaded = self._reload_without(monkeypatch, {"pptx"})

        assert reloaded.HAS_PPTX_DEP is False
        for ext in (".ppt", ".pptx"):
            assert ext not in reloaded.LOADER_REGISTRY

    def test_missing_ocr_dep_hides_image_formats(
        self, monkeypatch, restore_module
    ):
        reloaded = self._reload_without(
            monkeypatch, {"pytesseract", "unstructured.pytesseract"}
        )

        assert reloaded.HAS_OCR_DEP is False
        for ext in (".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".heic"):
            assert ext not in reloaded.LOADER_REGISTRY
